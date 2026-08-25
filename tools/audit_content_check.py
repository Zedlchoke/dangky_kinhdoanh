# -*- coding: utf-8 -*-
"""Kiểm toán tính đúng đối chiếu nội dung (audit mục 3).

Chọn 4 case tiêu biểu, dựng data với GIÁ TRỊ SENTINEL RIÊNG CHO TỪNG KEY
(dạng <<key>>), render mọi mẫu của case rồi đọc lại bằng python-docx:

1. Sentinel của mọi field phải xuất hiện trong văn bản xuất ra (đúng mẫu).
2. Không còn tag {..} sót.
3. Kiểm tra NGỮ CẢNH cho các trường trọng yếu: sentinel phải nằm TRONG CÙNG
   ô/dòng với nhãn tương ứng (vd ten doanh nghiệp sau nhãn "Tên doanh nghiệp",
   ngày sinh trên dòng "Ngày, tháng, năm sinh").

Chạy:  python -m tools.audit_content_check
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app_tk import engine_io as eio  # noqa: E402
from tools.docx_render import render  # noqa: E402

OUT_DIR = ROOT / "generated_ho_so" / "_audit_content"
CASES = ["thanh_lap_tnhh_1tv", "dang_ky_ho_kinh_doanh_lan_dau",
         "lap_chi_nhanh_vpdd", "giai_the_doanh_nghiep"]

# (case_id, regex không dấu trên nhãn đã fold, key) — kiểm tra cùng ô/dòng
SEMANTIC_SPOTS = {
    "thanh_lap_tnhh_1tv": [
        (r"(ten doanh nghiep|ten cong ty).*in hoa", "dn.ten_tieng_viet"),
        (r"ngay, thang, nam sinh", "ndd.ngay_sinh"),
        (r"dien thoai", "ndd.dien_thoai"),
        (r"kinh gui|co quan dang ky kinh doanh", "co_quan_dkkd"),
    ],
    "dang_ky_ho_kinh_doanh_lan_dau": [
        (r"ten ho kinh doanh", "hkd_ten_tieng_viet"),
        (r"ho, chu dem va ten|ghi ho ten bang chu in hoa", "chu_hkd.ho_ten"),
        (r"ngay, thang, nam sinh|sinh ngay", "ngay_sinh"),
        (r"ma so dinh danh ca nhan|dinh danh", "chu_hkd.so_cccd"),
    ],
    "lap_chi_nhanh_vpdd": [
        (r"ten van phong dai dien|ten chi nhan", "cn.ten_tieng_viet"),
        (r"ten chi nhan chu quan", "cn_chu_quan.ten"),
        (r"dia chi", "cn.dia_chi"),
    ],
    "giai_the_doanh_nghiep": [
        (r"ten doanh nghiep.*in hoa", "dn.ten_tieng_viet"),
        (r"ly do giai the", "ly_do_giai_the"),
        (r"ma so doanh nghiep", "dn.ma_so_dn"),
    ],
}


def fold(text: str) -> str:
    nfkd = unicodedata.normalize("NFD", text.lower())
    out = "".join(c for c in nfkd if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", out.replace("đ", "d")).strip()


def sentinel(key: str) -> str:
    return "<<%s>>" % key


def doc_texts(path: Path):
    """Toàn bộ text của file: [(vị_trí_mô_tả, text)] gồm body, bảng và header/footer."""
    d = docx.Document(str(path))
    items = [("p%d" % i, p.text) for i, p in enumerate(d.paragraphs)]
    for ti, tbl in enumerate(d.tables):
        for ri, row in enumerate(tbl.rows):
            for ci, cell in enumerate(row.cells):
                items.append(("tbl%d[%d][%d]" % (ti, ri, ci), cell.text))
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith(("word/header", "word/footer")):
                xml = z.read(n).decode("utf-8", "ignore")
                plain = re.sub(r"<[^>]+>", "", xml)
                if plain.strip():
                    items.append((n, plain))
    return items


def main() -> int:
    maps = eio.load_form_maps()
    uc = eio.load_use_cases()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    total_errors: list[str] = []

    for cid in CASES:
        case = next((c for c in uc["cases"] if c["id"] == cid), None)
        print("=== case:", cid, "===")
        if case is None:
            total_errors.append("%s: không có trong use_cases" % cid)
            continue
        errors: list[str] = []
        forms = eio.forms_of_case(case, maps)
        for f in forms:
            fmap = f["map"]
            fid = fmap["id"]
            data: dict = {}
            for fld in fmap.get("fields", []):
                val = "X" if fld.get("type") == "chk" else sentinel(fld["key"])
                eio.nested_set(data, fld["key"], val)
            for t in fmap.get("dynamic_tables", []):
                data[t["loop"]] = [
                    {c["tag"]: "%s r%d" % (sentinel(c["tag"]), i) for c in t["columns"]}
                    for i in range(2)
                ]
            out = OUT_DIR / ("%s_%s.docx" % (cid, fid))
            try:
                out.write_bytes(render(eio.APP_DIR / fmap["template"], data))
            except Exception as exc:  # noqa: BLE001
                errors.append("%s: render lỗi %r" % (fid, exc))
                continue
            texts = doc_texts(out)
            joined = "\n".join(t for _, t in texts)

            # prefix loop: các field dạng <loop>.<col> do kiểm tra cột lo
            loop_prefixes = {t["loop"] for t in fmap.get("dynamic_tables", [])}

            def is_loop_field(k):
                return any(k.startswith(lp + ".") for lp in loop_prefixes)

            # 1) mọi sentinel field phải xuất hiện (field chk điền "X", không sentinel;
            #    field cột loop được kiểm tra riêng phía dưới)
            missing_keys = []
            for fld in fmap.get("fields", []):
                if fld.get("type") == "chk" or is_loop_field(fld["key"]):
                    continue
                s = sentinel(fld["key"])
                if s not in joined:
                    missing_keys.append(fld["key"])
            # loop: chỉ yêu cầu tag cột của dòng 0
            for t in fmap.get("dynamic_tables", []):
                for c in t.get("columns", []):
                    if sentinel(c["tag"]) + " r0" not in joined:
                        missing_keys.append("%s(%s)" % (c["tag"], t["loop"]))
            if missing_keys:
                errors.append("%s: %d sentinel KHÔNG xuất hiện: %s"
                              % (fid, len(missing_keys), ", ".join(missing_keys[:8])))

            # 2) tag sót
            leftover = sorted(set(re.findall(r"\{[/#]?[A-Za-z][A-Za-z0-9_.#/]*\}", joined)))
            if leftover:
                errors.append("%s: còn tag sót: %s" % (fid, ", ".join(leftover[:10])))

            # 3) ngữ cảnh: sentinel phải nằm trong cùng ô/dòng với nhãn.
            #    Nhãn có thể trùng ở nhiều chỗ (vd bảng loop cũng có "Ngày, tháng,
            #    năm sinh") — chỉ cần MỘT vị trí nhãn chứa sentinel là đạt.
            spots = [sp for sp in SEMANTIC_SPOTS.get(cid, [])
                     if any(fl["key"] == sp[1] and fl.get("type") != "chk"
                            for fl in fmap["fields"])]
            for pat, key in spots:
                s = sentinel(key)
                hit = False
                where_label = None
                for loc, txt in texts:
                    if re.search(pat, fold(txt)):
                        where_label = (loc, fold(txt)[:80])
                        if s in txt:
                            hit = True
                            break
                if not hit:
                    msg = "%s: '%s' không nằm cùng ô/dòng với nhãn /%s/ (nhãn thấy ở %s)" % (
                        fid, key, pat, where_label)
                    errors.append(msg)
            print("  %s: %d field, %s/%s spot ngữ nghĩa OK" % (
                fid, len(fmap["fields"]),
                sum(1 for sp in spots if not any(e.startswith(fid + ": '" + sp[1]) for e in errors)),
                len(spots)))
        if errors:
            for e in errors:
                print("  FAIL:", e)
            total_errors += ["%s :: %s" % (cid, e) for e in errors]
        else:
            print("  PASS (%d mẫu)" % len(forms))

    print("-" * 60)
    if total_errors:
        print("AUDIT NỘI DUNG: %d lỗi" % len(total_errors))
        for e in total_errors:
            print(" -", e)
        return 1
    print("AUDIT NỘI DUNG: PASS toàn bộ 4 case")
    return 0


if __name__ == "__main__":
    sys.exit(main())
