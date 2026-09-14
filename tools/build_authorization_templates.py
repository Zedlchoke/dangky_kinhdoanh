# -*- coding: utf-8 -*-
"""Sinh template DOCX ủy quyền (mẫu tham chiếu theo Điều 12/93 NĐ 168).

Chạy: python -X utf8 -m tools.build_authorization_templates
"""
from __future__ import annotations

from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "app" / "templates" / "authorization"


def _write_doc(path: Path, title: str, body_paragraphs: list[str]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    run.bold = True
    p2 = doc.add_paragraph("Độc lập – Tự do – Hạnh phúc")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("……………………")
    p3 = doc.add_paragraph(title)
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.runs[0]
    run3.bold = True
    doc.add_paragraph("")
    for line in body_paragraphs:
        doc.add_paragraph(line)
    doc.add_paragraph("")
    doc.add_paragraph("…………, ngày {ngay} tháng {thang} năm {nam}")
    doc.add_paragraph("")
    doc.add_paragraph("BÊN ỦY QUYỀN\t\t\tBÊN ĐƯỢC ỦY QUYỀN")
    doc.add_paragraph("(Ký, ghi rõ họ tên)\t\t\t(Ký, ghi rõ họ tên)")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def build_all(force: bool = False) -> list[Path]:
    dn_body = [
        "BÊN ỦY QUYỀN (Bên A):",
        "Tên doanh nghiệp: {uy_quyen_ten_chu_the}",
        "Mã số doanh nghiệp: {uy_quyen_ma_so}",
        "Địa chỉ trụ sở: {uy_quyen_dia_chi}",
        "Người đại diện theo pháp luật: {uy_quyen_nguoi_ky}",
        "Chức danh: {uy_quyen_chuc_danh}",
        "",
        "BÊN ĐƯỢC ỦY QUYỀN (Bên B):",
        "Họ và tên: {agent_ho_ten}",
        "Ngày sinh: {agent_ngay_sinh}",
        "Số CCCD: {agent_so_cccd}",
        "Ngày cấp: {agent_ngay_cap}    Nơi cấp: {agent_noi_cap}",
        "Địa chỉ liên hệ: {agent_dia_chi}",
        "Điện thoại: {agent_dien_thoai}    Email: {agent_email}",
        "",
        "ĐIỀU 1. NỘI DUNG VÀ PHẠM VI ỦY QUYỀN",
        "Bên A ủy quyền cho Bên B thay mặt doanh nghiệp làm việc với {co_quan_tiep_nhan} "
        "để {pham_vi_uy_quyen}, bao gồm: nộp hồ sơ, bổ sung hồ sơ, nhận kết quả, "
        "ký và gửi các văn bản liên quan đến thủ tục đăng ký doanh nghiệp.",
        "",
        "ĐIỀU 2. THỜI HẠN ỦY QUYỀN",
        "Giấy ủy quyền có hiệu lực đến ngày {thoi_han_uy_quyen}.",
        "",
        "ĐIỀU 3. CAM KẾT",
        "Hai bên chịu trách nhiệm trước pháp luật về tính hợp pháp, trung thực "
        "và chính xác của việc ủy quyền theo Nghị định 168/2025/NĐ-CP.",
    ]
    hkd_body = [
        "BÊN ỦY QUYỀN (Bên A) — Chủ hộ kinh doanh:",
        "Tên hộ kinh doanh: {uy_quyen_ten_chu_the}",
        "Họ và tên chủ hộ: {uy_quyen_nguoi_ky}",
        "Số CCCD chủ hộ: {uy_quyen_so_cccd}",
        "Địa chỉ: {uy_quyen_dia_chi}",
        "",
        "BÊN ĐƯỢC ỦY QUYỀN (Bên B):",
        "Họ và tên: {agent_ho_ten}",
        "Ngày sinh: {agent_ngay_sinh}",
        "Số CCCD: {agent_so_cccd}",
        "Ngày cấp: {agent_ngay_cap}    Nơi cấp: {agent_noi_cap}",
        "Địa chỉ liên hệ: {agent_dia_chi}",
        "Điện thoại: {agent_dien_thoai}    Email: {agent_email}",
        "",
        "ĐIỀU 1. NỘI DUNG VÀ PHẠM VI ỦY QUYỀN",
        "Bên A ủy quyền cho Bên B làm việc với {co_quan_tiep_nhan} để {pham_vi_uy_quyen}, "
        "bao gồm: nộp hồ sơ, bổ sung, nhận kết quả và các việc liên quan đăng ký hộ kinh doanh.",
        "",
        "ĐIỀU 2. THỜI HẠN ỦY QUYỀN",
        "Giấy ủy quyền có hiệu lực đến ngày {thoi_han_uy_quyen}.",
        "",
        "ĐIỀU 3. CAM KẾT",
        "Hai bên chịu trách nhiệm theo Nghị định 168/2025/NĐ-CP Điều 93.",
    ]
    paths = [
        OUT / "AUTH-UQ-DN-01.docx",
        OUT / "AUTH-UQ-HKD-01.docx",
    ]
    written: list[Path] = []
    specs = [
        (paths[0], "GIẤY ỦY QUYỀN\n(Thủ tục đăng ký doanh nghiệp)", dn_body),
        (paths[1], "GIẤY ỦY QUYỀN\n(Thủ tục đăng ký hộ kinh doanh)", hkd_body),
    ]
    for path, title, body in specs:
        if path.is_file() and not force:
            written.append(path)
            continue
        _write_doc(path, title, body)
        written.append(path)
    return written


if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true",
                    help="Ghi đè template đã có (mặc định: giữ mẫu owner đã chỉnh)")
    args = ap.parse_args()
    built = build_all(force=args.force)
    for p in built:
        print("OK", p)
