# -*- coding: utf-8 -*-
"""Chuyển mỗi template app/templates/PLI_Mxx.docx thành app/data/blocks/<id>.json.

Template đã gắn the dạng {key} (key chấm được: dn.ten_tieng_viet), {chk_xxx},
và bảng loop có hàng khuôn chứa {#ten} ở ô đầu và {/ten} ở ô cuối cùng hàng,
các ô giữa là {tag_cot}. Biến đặc biệt {_stt} trong hàng khuôn là STT tự sinh
(1-based theo thứ tự dòng dữ liệu, engine render thay khi nhân bản hàng),
không được tính là cột dữ liệu trong cols.

Schema JSON đầu ra (blocks-v3 — mọi trường mới đều OPTIONAL, chỉ emit khi khác
default; block cũ thiếu role/format/vertAlign vẫn hợp lệ hoàn toàn):
{
  "id": "PLI_M02",
  "ten_mau": "...",            # lấy từ app/data/form_maps.json
  "blocks": [...],
  "v": 3                       # schema version (L1): 1 block mới có thể có
                               # "role"/"format"/"footnotes", run mới có thể có
                               # "vertAlign"/"role":"footnote-ref"
}
Block "p":
{"t":"p", "align":"left|center|right",
 "role": "...",               # optional, tools/layout_roles.py; thiếu = "body"
 "format": {...},             # optional, chỉ khi khác default:
                              # leftIndentPt/rightIndentPt/firstLineIndentPt/
                              # spaceBeforePt/spaceAfterPt (EMU→pt, làm tròn 1),
                              # lineSpacing (bội, vd 1.15) | lineSpacingPt
                              # (exact/at-least), keepWithNext=true
 "footnotes": {"7": "chú thích…"},  # optional, id→text từ word/footnotes.xml
 "runs":[
    {"text":"..."},                          # \n từ <w:br/>/<w:cr/>, \t từ
                                             # <w:tab/> được giữ nguyên ký tự
    {"text":"1","vertAlign":"superscript"},  # hoặc "subscript"
    {"role":"footnote-ref","footnoteId":7},  # <w:footnoteReference w:id="7"/>
    {"input":{"key":"dn.ten_tieng_viet","w":40}},
    {"input":{"key":"chk_x","chk":true}}     # chk lồng trong paragraph dài
 ]}
{"t":"chk", "label":"...", "key":"chk_x"},   # paragraph chỉ gồm chk + nhãn ngắn
{"t":"table", "widths":null, "rows":[
    {"cells":[{"runs":[...], "colspan":1, "rowspan":1}]}
]},
{"t":"loop", "loop":"ds_thanh_vien", "headerRows":[0],
 "cols":[{"col":1,"tag":"ho_ten"}], "blankRows":4, "rows":[...]}

Chạy:
    python -X utf8 tools/docx_to_blocks.py              # chuyển cả 87 file + tự kiểm định
    python -X utf8 tools/docx_to_blocks.py --one PLI_M02  # chuyển 1 file
"""
from __future__ import annotations

import io
import json
import os
import random
import re
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Length
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import layout_roles

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = ROOT / "app" / "templates"
FORM_MAPS_PATH = ROOT / "app" / "data" / "form_maps.json"
OUT_DIR = ROOT / "app" / "data" / "blocks"

# Tag biến: {key} hoặc {a.b.c}; không khớp {#loop}/{/loop}
TAG_RE = re.compile(r"\{([A-Za-z0-9_][A-Za-z0-9_.]*)\}")
LOOP_OPEN_RE = re.compile(r"\{#([A-Za-z0-9_]+)\}")
# Marker vùng điều kiện P1.2 M72 Option D: {?flag} mở / {/flag} đóng.
# Paragraph CHỈ chứa marker không sinh block; mọi block nằm trong vùng
# được gắn "cond": [flag, ...] (renderer UI lọc theo flag, engine DOCX
# tự loại vùng theo cùng flag - xem tools/docx_render.py).
COND_OPEN_FULL_RE = re.compile(r"^\{\?([A-Za-z0-9_]+)\}$")
COND_CLOSE_FULL_RE = re.compile(r"^\{/([A-Za-z0-9_]+)\}$")
CHK_PREFIX = "chk"


def iter_blocks(doc):
    """Duyệt body theo thứ tự: ('p', Paragraph) hoặc ('tbl', Table)."""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield ("p", Paragraph(child, doc))
        elif child.tag == qn("w:tbl"):
            yield ("tbl", Table(child, doc))


# ---------------------------------------------------------------------------
# Paragraph -> runs
# ---------------------------------------------------------------------------

def _para_full_text(p: Paragraph) -> str:
    return "".join(r.text for r in p.runs)


def _run_flags(run):
    return bool(run.bold), bool(run.italic), bool(run.underline)


def _run_vert_align(run):
    """'superscript'|'subscript' từ w:rPr/w:vertAlign, None nếu không có."""
    rPr = run._element.rPr
    if rPr is None:
        return None
    va = rPr.find(qn("w:vertAlign"))
    if va is None:
        return None
    val = va.get(qn("w:val"))
    return val if val in ("superscript", "subscript") else None


# python-docx 1.2.0 (docx/oxml/text/run.py CT_R.text) map các con của w:r
# thành ký tự: w:t -> text; w:br, w:cr -> "\n"; w:tab -> "\t";
# w:noBreakHyphen -> "-"; w:ptab -> "\t". _char_stream dựng lại ĐÚNG chuỗi đó
# kèm metadata từng ký tự và vị trí footnoteReference.
_CHAR_STREAM_TAGS = {
    qn("w:t"),
    qn("w:br"),
    qn("w:cr"),
    qn("w:tab"),
    qn("w:noBreakHyphen"),
    qn("w:ptab"),
}


def _char_stream(p: Paragraph):
    """Dựng dòng ký tự của paragraph với metadata mỗi ký tự.

    Trả về (chars, footnote_refs):
      chars: list[(ch, meta)], meta = (b, i, u, vertAlign)
      footnote_refs: list[(pos, id)] — pos là vị trí chèn trong stream
    """
    chars: list[tuple] = []
    fn_refs: list[tuple] = []
    for r in p.runs:
        b, i_, u = _run_flags(r)
        meta = (b, i_, u, _run_vert_align(r))
        for child in r._element:
            tag = child.tag
            if tag not in _CHAR_STREAM_TAGS:
                if tag == qn("w:footnoteReference"):
                    try:
                        fid = int(child.get(qn("w:id")))
                    except (TypeError, ValueError):
                        continue  # fail-soft: thiếu id thì bỏ qua ref này
                    fn_refs.append((len(chars), fid))
                continue  # rPr & con khác không sinh text (như Run.text)
            if tag == qn("w:t"):
                txt = child.text or ""
                for ch in txt:
                    chars.append((ch, meta))
            elif tag == qn("w:br") or tag == qn("w:cr"):
                chars.append(("\n", meta))
            elif tag == qn("w:noBreakHyphen"):
                chars.append(("-", meta))
            else:  # w:tab / w:ptab
                chars.append(("\t", meta))
    return chars, fn_refs


def build_runs(p: Paragraph) -> list[dict]:
    """Tách paragraph thành runs: gộp text liền nhau cùng metadata (b/i/u +
    vertAlign), tách run tại mỗi the; the chk -> input.chk=true, the thường ->
    input.key + w. <w:footnoteReference> -> run {"role":"footnote-ref",
    "footnoteId":N} giữ đúng vị trí tương đối.

    Chuỗi text đầu ra khớp chính xác "".join(r.text ...) của python-docx
    (hard break \n và tab \t được bảo toàn, KHÔNG collapse).
    """
    chars, fn_refs = _char_stream(p)
    full = "".join(c for c, _meta in chars)
    n = len(chars)
    out: list[dict] = []

    def emit_text(s: int, e: int) -> None:
        while s < e:
            j = s + 1
            while j < e and chars[j][1] == chars[s][1]:
                j += 1
            b, i_, u, va = chars[s][1]
            run = {"text": full[s:j]}
            if b:
                run["b"] = True
            if i_:
                run["i"] = True
            if u:
                run["u"] = True
            if va:
                run["vertAlign"] = va
            out.append(run)
            s = j

    tags = list(TAG_RE.finditer(full))
    fns = sorted(fn_refs)
    ti = fi = 0
    cur = 0
    prev_tag_end = 0
    while True:
        nxt_tag = tags[ti].start() if ti < len(tags) else None
        nxt_fn = fns[fi][0] if fi < len(fns) else None
        if nxt_tag is None and nxt_fn is None:
            emit_text(cur, n)
            break
        # footnote-ref đứng trước tag/text tại cùng vị trí (zero-width)
        if nxt_fn is not None and (nxt_tag is None or nxt_fn <= nxt_tag):
            emit_text(cur, nxt_fn)
            out.append({"role": "footnote-ref", "footnoteId": fns[fi][1]})
            fi += 1
            cur = nxt_fn
            continue
        m = tags[ti]
        emit_text(cur, m.start())
        key = m.group(1)
        # nhãn trước the = khoảng cách từ the trước đến the này (giữ nguyên
        # công thức cũ để output w không đổi với template hiện có)
        label_len = m.start() - prev_tag_end
        w = min(60, max(12, label_len // 2 + 8))
        if key.startswith(CHK_PREFIX):
            out.append({"input": {"key": key, "chk": True}})
        else:
            out.append({"input": {"key": key, "w": w}})
        cur = prev_tag_end = m.end()
        ti += 1
    return out


def para_text_no_tags(p: Paragraph) -> str:
    return TAG_RE.sub("", _para_full_text(p))


def _align_of(p: Paragraph) -> str:
    a = p.alignment
    if a == WD_ALIGN_PARAGRAPH.CENTER:
        return "center"
    if a == WD_ALIGN_PARAGRAPH.RIGHT:
        return "right"
    return "left"


_EMU_PER_PT = 12700


def _len_pt(length) -> float | None:
    """EMU (docx.shared.Length) -> pt làm tròn 1 số thập phân."""
    if length is None:
        return None
    return round(int(length) / _EMU_PER_PT, 1)


def _para_format_dict(p: Paragraph) -> dict | None:
    """Paragraph format khác default mới emit (schema blocks-v3)."""
    pf = p.paragraph_format
    fmt: dict = {}
    for key, val in (
        ("leftIndentPt", pf.left_indent),
        ("rightIndentPt", pf.right_indent),
        ("firstLineIndentPt", pf.first_line_indent),
        ("spaceBeforePt", pf.space_before),
        ("spaceAfterPt", pf.space_after),
    ):
        pt = _len_pt(val)
        if pt:  # None hoặc 0 = mặc định -> bỏ
            fmt[key] = pt
    ls = pf.line_spacing
    if ls is not None:
        if isinstance(ls, Length):  # exact/at-least (twips) -> pt
            pt = _len_pt(ls)
            if pt:
                fmt["lineSpacingPt"] = pt
        elif abs(float(ls) - 1.0) > 1e-9:  # bội dòng, vd 1.15
            fmt["lineSpacing"] = round(float(ls), 2)
    if pf.keep_with_next:
        fmt["keepWithNext"] = True
    return fmt or None


# ---------------------------------------------------------------------------
# Override role tường minh (tools/layout_role_overrides.json)
# ---------------------------------------------------------------------------

OVERRIDES_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "layout_role_overrides.json"
)


def load_role_overrides(path=None) -> dict:
    """Đọc override JSON; file rỗng {} = dùng rule thuần.

    Shape hỗ trợ (tất cả optional):
      {"by_index": {FORM_ID: {"<block_index_cấp_thân>": role}},
       "by_folded_prefix": [{"prefix": "...", "role": "...",
                             "forms": ["PLI_M01"]}]}   # bỏ "forms" = mọi form
    """
    p = path or OVERRIDES_PATH
    try:
        with io.open(p, "r", encoding="utf-8-sig") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}  # fail-soft: override hỏng/thiếu không được phá converter


def apply_role_override(overrides, fid, body_para_index, text_folded, role):
    """Áp override tường minh (ưu tiên index theo form, rồi prefix gập dấu)."""
    form_map = (overrides.get("by_index") or {}).get(fid or "", {})
    hit = form_map.get(str(body_para_index)) if form_map else None
    if isinstance(hit, str) and hit in layout_roles.ROLES:
        return hit
    for entry in overrides.get("by_folded_prefix") or []:
        if not isinstance(entry, dict):
            continue
        forms = entry.get("forms")
        if forms and fid not in forms:
            continue
        prefix = entry.get("prefix")
        if isinstance(prefix, str) and text_folded.startswith(prefix):
            r = entry.get("role")
            if isinstance(r, str) and r in layout_roles.ROLES:
                return r
    return role


# ---------------------------------------------------------------------------
# Body paragraph -> block (p hoặc chk)
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Ranh giới heading đánh số bị dính vào đuôi paragraph trước (bug cấu trúc V2)
#
# Nguồn gốc: một số template scrape lại từ DOCX gốc gộp cả đuôi bảng lưới và
# heading mục tiếp theo vào CÙNG MỘT w:p (không có <w:br/>), heading chỉ là
# một <w:r> riêng ở cuối paragraph — vd PLI_M72 (blk110/112):
#   "(VNĐ … )Tỷ lệ (%)Ghi chú (nếu có)" + <w:r><w:b/>"8. Cổ đông là nhà đầu
#   tư nước ngoài: " + …
# Không sửa template tay: converter TÁCH ranh giới run này thành block "p"
# riêng rồi classify_role() tự xếp nó vào section-heading. Điều kiện tách nằm
# trong split_glued_numbered_headings(); detector Wave L1
# (tools/census/check_layout_l1.py, nhóm rule (g2)) DÙNG CHÍNH XÁC cùng điều
# kiện nên khi blocks còn dính thì đỏ, sau regen thì về 0.
# ---------------------------------------------------------------------------

# "8. Cổ đông…" / "9. Người đại diện…" — số (1-2 chữ số) + dấu chấm +
# khoảng trắng + chữ. "5.1"/"10.1" (heading thập phân con) KHÔNG khớp vì sau
# dấu chấm là chữ số, không phải khoảng trắng.
NUMBERED_HEADING_RUN_RE = re.compile(r"^\d{1,2}\.[ \t]\S")
NUMBERED_HEADING_FIND_RE = re.compile(r"\d{1,2}\.[ \t](?=\S)")
# Guard chống tách nhầm: phần đuôi sau ranh giới phải là 1 dòng heading, không
# phải cả đoạn văn pháp lý dài bắt đầu bằng "N. …".
MAX_HEADING_TAIL_CHARS = 300
# Heading-role mà phần đuôi tách ra PHẢI đạt được để được coi là ranh giới
# heading thật (cùng nhóm heading của detector L1 rule (g1)).
HEADING_ROLES = frozenset(
    {"section-heading", "form-title", "form-subtitle"}
)


def _text_len(runs: list[dict]) -> int:
    return sum(
        len(r.get("text", ""))
        for r in runs
        if isinstance(r, dict) and isinstance(r.get("text"), str)
    )


def _is_heading_tail(text: str) -> bool:
    """Text hậu kiểm có đóng vai heading đánh số (giống tier-11 của
    layout_roles) hay không. Dòng kẻ chấm trống chữ ("1. ......") là bản in
    để điền tay — không phải heading có nội dung."""
    body = (text or "").strip()
    if not body or not any(ch.isalpha() for ch in body):
        return False
    folded = layout_roles.fold(body)
    return (
        layout_roles.classify_role(folded, None, [{"text": body}])
        in HEADING_ROLES
    )


def _join_tail(runs: list[dict], start: int) -> str:
    """Text ghép của các run text từ vị trí `start` đến hết danh sách."""
    return "".join(
        r.get("text", "") for r in runs[start:]
        if isinstance(r, dict) and isinstance(r.get("text"), str)
    )


def _embedded_heading_cuts(full: str) -> list[int]:
    """Các offset (trên chuỗi `full`) cần chẻ vì có heading đánh số dính đuôi.

    Ứng viên: mọi vị trí khớp NUMBERED_HEADING_FIND_RE (số 1-2 chữ số + dấu
    chấm + khoảng trắng + chữ, KHÔNG khớp "5.1"/"10.1" kiểu thập phân con)
    nằm KHÔNG ở đầu đoạn — glue V2 luôn có text đứng trước; trường hợp
    "N. xxx" ngay đầu paragraph là heading tự nó (block riêng), không glue.

    Guard chặn tách nhầm (mọi điều kiện phải đúng, xem _is_glued_heading_cut):
      * ký tự ngay trước số không được là ".", khoảng trắng hay chữ số;
      * đuôi từ offset đến hết <= MAX_HEADING_TAIL_CHARS;
      * đuôi hậu kiểm đạt một role thuộc HEADING_ROLES.
    """
    cuts: list[int] = []
    for m in NUMBERED_HEADING_FIND_RE.finditer(full):
        pos = m.start()
        if pos == 0 or not full[:pos].strip():
            continue
        if _is_glued_heading_cut(full, pos):
            cuts.append(pos)
    return cuts


def _is_glued_heading_cut(full: str, pos: int) -> bool:
    """`pos` có phải ranh giới heading đánh số bị GHÉP vào text trước đó?

    Đây là predicate CHUNG của cả hai dạng glue audit Wave D; detector L1
    (tools/census/check_layout_l1.py, nhóm rule (g2)) dùng cùng bản sao để
    đỏ đúng chỗ blocks còn dính và tự về 0 sau khi Lead regen."""
    prev = full[pos - 1]
    if prev == ".":
        # "…mục 5." + "1. …" / "5." + "1. Địa chỉ…" (PLII_M01) là đánh số nhiều
        # cấp con, chứ không phải heading mới dính đuôi.
        return False
    if prev.isdigit():
        # Số bị cắt đôi giữa các run ("1"+"0. Thông tin đăng ký thuế" = mục
        # "10.", PLI_M02/M06/M08) hoặc heading nằm chung cột đếm scrape
        # ("123456789"+"2. Thông tin…", PLI_M09) — ranh giới thật KHÔNG suy
        # ra được từ text nên converter KHÔNG tự đoán; vùng lỗi này thuộc
        # trách nhiệm lane sửa template và vẫn bị detector L1 bắt qua rule
        # (e) STT-glue + (f) digit-run dài.
        return False
    if prev.isspace():
        # đã ngắt dòng / hard break rõ ràng: ranh giới hiển thị còn nguyên.
        return False
    tail = full[pos:]
    if not tail.strip():
        return False
    if len(tail) > MAX_HEADING_TAIL_CHARS:
        return False
    return _is_heading_tail(tail)


def _split_runs_at_offsets(runs: list[dict], cuts: list[int]) -> list[dict]:
    """Chẻ các run text tại offset tuyệt đối trong chuỗi text ghép.

    Run mới giữ nguyên mọi metadata (b/i/u/vertAlign), chỉ khác khoá "text".
    Run không chứa offset nào được trả nguyên vẹn (cùng tham chiếu)."""
    if not cuts:
        return list(runs)
    cut_set = set(cuts)
    out: list[dict] = []
    off = 0
    for r in runs:
        txt = r.get("text") if isinstance(r, dict) else None
        if isinstance(txt, str):
            end = off + len(txt)
            internal = sorted(p - off for p in cut_set if off < p < end)
            if internal:
                prev_at = 0
                for cp in internal:
                    head = dict(r)
                    head["text"] = txt[prev_at:cp]
                    out.append(head)
                    prev_at = cp
                last = dict(r)
                last["text"] = txt[prev_at:]
                out.append(last)
                off = end
                continue
            off = end
        out.append(r)
    return out


def _groups_from_offsets(
    runs: list[dict], offsets: list[int]
) -> list[list[dict]]:
    """Chia danh sách run thành các nhóm tại offset tuyệt đối (đã tăng dần)."""
    groups: list[list[dict]] = []
    cur: list[dict] = []
    wanted = set(offsets)
    off = 0
    for r in _split_runs_at_offsets(runs, sorted(wanted)):
        txt = r.get("text") if isinstance(r, dict) else None
        if off in wanted and cur:
            groups.append(cur)
            cur = []
        cur.append(r)
        off += len(txt) if isinstance(txt, str) else 0
    if cur:
        groups.append(cur)
    return groups


def split_glued_numbered_headings(runs: list[dict]) -> list[list[dict]]:
    """Tách danh sách run thành các nhóm tại mỗi ranh giới heading đánh số.

    Trả về [runs] nguyên vẹn khi không có gì phải tách (trường hợp thường);
    ngược lại trả về >= 2 nhóm: các run trước ranh giới, và mỗi đoạn bắt đầu
    bằng một heading đánh số riêng.

    Hai dạng glue của audit Wave D cùng đi qua MỘT predicate
    (_is_glued_heading_cut) nên dễ đoán và dễ kiểm chứng:
      (1) heading là run riêng ở cuối w:p — PLI_M72 blk110/112:
          "(VNĐ … )Tỷ lệ (%)Ghi chú (nếu có)" | "8. Cổ đông là nhà đầu tư…";
      (2) heading ghép chung một run với cột đếm scrape (dạng ")Ghi chú…8.
          Cổ đông…" trong MỘT w:t): _split_runs_at_offsets chẻ theo offset
          rồi quy về dạng (1).

    Điều kiện tách tại mọi ứng viên "N. <chữ>" (TẤT CẢ phải đúng):
      * KHÔNG đứng đầu paragraph (heading ở đầu đoạn tự nó là block riêng);
      * trước đó vẫn còn text khác trong cùng paragraph;
      * ký tự ngay trước số không là ".", khoảng trắng hay chữ số
        ("1"+"0." = mục 10.; "5."+"1." = mục con; cột đếm scrape như
        "1234567892." thì ranh giới thật không suy ra được từ text nên
        converter chủ động KHÔNG tự đoán — lane sửa template xử lý, vùng lỗi
        vẫn bị detector L1 bắt qua rule (e)/(f));
      * đuôi từ ranh giới đến hết <= MAX_HEADING_TAIL_CHARS;
      * đuôi hậu kiểm (_is_heading_tail) đạt một role thuộc HEADING_ROLES
        (tier-11 của tools/layout_roles.py).
    """
    full = _join_tail(runs, 0)
    if not full.strip():
        return [list(runs)]
    cuts = _embedded_heading_cuts(full)
    if not cuts:
        return [list(runs)]
    return _groups_from_offsets(runs, cuts)


def para_block(p: Paragraph, prev_block_ref: dict, ctx: dict | None = None) -> list[dict]:
    """Trả về list block cho 1 paragraph ở thân tài liệu.
    prev_block_ref: {'block': dict|None} để lấy nhãn cho chk đứng một mình.
    ctx (optional): {'form_id', 'body_para_index', 'footnotes', 'overrides'}
    — dùng để gắn role/format/footnotes cho block "p" (schema blocks-v3).

    Một paragraph có thể trả về NHIỀU block "p" khi template gộp heading
    đánh số vào đuôi paragraph (split_glued_numbered_headings); cả NHIỀU block
    đều kế thừa align/format/cond của paragraph nguồn."""
    runs = build_runs(p)
    keys = [r["input"]["key"] for r in runs if "input" in r]
    rest = para_text_no_tags(p).strip()

    if not rest and not keys:
        return []  # paragraph rỗng, bỏ

    def is_chk(k: str) -> bool:
        return k.startswith(CHK_PREFIX)

    # Paragraph chỉ gồm the chk + nhãn ngắn (hoặc chk đứng một mình) -> block chk
    if keys and all(is_chk(k) for k in keys) and len(rest) <= 80:
        blocks = []
        if not rest and keys:
            # chk đứng một mình: mượn nhãn từ paragraph văn bản ngắn ngay trước
            pb = prev_block_ref.get("block")
            if (
                pb is not None
                and pb.get("t") == "p"
                and not prev_block_ref.get("consumed")
                and len("".join(r.get("text", "") for r in pb.get("runs", [])).strip()) <= 80
            ):
                label = "".join(
                    r.get("text", "") for r in pb["runs"] if "text" in r
                ).strip()
                blocks.append({"t": "chk", "label": label, "key": keys[0]})
                prev_block_ref["consumed"] = True
                # xoá block p đã bị hút vào nhãn khỏi danh sách đã emit
                emitted = prev_block_ref.get("emitted")
                if emitted is not None and pb in emitted:
                    emitted.remove(pb)
                keys = keys[1:]
                if not keys:
                    return blocks
        # nhiều chk kèm nhãn ngắn trong cùng paragraph
        if len(keys) == 1:
            blocks.append({"t": "chk", "label": rest, "key": keys[0]})
        else:
            # chia nhãn theo đoạn text quanh từng the
            full = _para_full_text(p)
            pos = 0
            segs_before = []
            for m in TAG_RE.finditer(full):
                segs_before.append(full[pos:m.start()].strip())
                pos = m.end()
            for idx, k in enumerate(keys):
                label = segs_before[idx] if idx < len(segs_before) else ""
                if not label and idx > 0:
                    label = rest
                blocks.append({"t": "chk", "label": label, "key": k})
        return blocks

    # V2: heading đánh số dính đuôi paragraph -> tách thành block "p" riêng
    # (mỗi block classify role độc lập; cùng align/format/cond nguồn).
    blocks_out: list[dict] = []
    for run_group in split_glued_numbered_headings(runs):
        blk = {"t": "p", "align": _align_of(p), "runs": run_group}
        _attach_layout_metadata(blk, p, run_group, ctx)
        blocks_out.append(blk)
    return blocks_out


def _attach_layout_metadata(blk: dict, p: Paragraph, runs: list, ctx: dict | None):
    """Gắn role/format/footnotes (blocks-v3) — chỉ khi khác default."""
    ctx = ctx or {}
    fmt = _para_format_dict(p)
    if fmt:
        blk["format"] = fmt

    raw_full = "".join(r.get("text", "") for r in runs if "text" in r)
    folded = layout_roles.fold(raw_full)
    role = layout_roles.classify_role(folded, fmt, runs)
    role = apply_role_override(
        ctx.get("overrides") or {}, ctx.get("form_id"),
        ctx.get("body_para_index"), folded, role,
    )
    # blocks-v3: MỌI block "p" đều có role tường minh (kể cả "body") để L2
    # không phải regex lại văn bản; block cũ thiếu role vẫn hợp lệ (= body).
    blk["role"] = role

    fn_texts = ctx.get("footnotes")
    if fn_texts:
        ids = sorted(
            {
                r["footnoteId"]
                for r in runs
                if r.get("role") == "footnote-ref" and "footnoteId" in r
            }
        )
        present = {str(i): fn_texts[i] for i in ids if i in fn_texts}
        if present:  # id không tìm thấy text vẫn giữ ở run (fail-soft)
            blk["footnotes"] = present


# ---------------------------------------------------------------------------
# Table -> block table / loop
# ---------------------------------------------------------------------------

def _cell_paragraphs(tc_elem, doc):
    """Các Paragraph trực tiếp của một ô (w:tc)."""
    return [
        Paragraph(child, doc)
        for child in tc_elem.iterchildren()
        if child.tag == qn("w:p")
    ]


def parse_table(tbl: Table):
    """Phân rã bảng thành rows/cells với dedup ô merge, colspan/rowspan.

    Trả về list hàng; mỗi hàng là list ô:
    {"_tc": element, "col": c0, "colspan": n, "rowspan": n, "runs": [...]}
    """
    rows_out = []
    active = {}  # cột -> record ô đang mở dọc (cho rowspan)
    for tr in tbl._tbl.tr_lst:
        row_cells = []
        col = 0
        # Duyệt con trực tiếp w:tc của w:tr (tc_iter): mỗi ô xuất hiện đúng
        # một lần, không cần dedup — và tránh id() trên proxy lxml/python-docx
        # (proxy tạo mới mỗi lần truy cập nên id() không ổn định).
        for tc in tc_iter(tr):
            tc_pr = tc.find(qn("w:tcPr"))
            span = 1
            vmerge_val = None
            if tc_pr is not None:
                gs = tc_pr.find(qn("w:gridSpan"))
                if gs is not None:
                    span = int(gs.get(qn("w:val")))
                vm = tc_pr.find(qn("w:vMerge"))
                if vm is not None:
                    vmerge_val = vm.get(qn("w:val")) or "continue"
            if vmerge_val == "continue":
                rec = None
                for c in range(col, col + span):
                    if c in active:
                        rec = active[c]
                        break
                if rec is not None:
                    rec["rowspan"] += 1
                for c in range(col, col + span):
                    active[c] = rec
                col += span
                continue
            paras = _cell_paragraphs(tc, tbl.part.document)
            runs = []
            for pp in paras:
                runs.extend(build_runs(pp))
            rec = {
                "_tc": tc,
                "col": col,
                "colspan": span,
                "rowspan": 1,
                "runs": runs,
            }
            row_cells.append(rec)
            for c in range(col, col + span):
                active[c] = rec
            col += span
        rows_out.append(row_cells)
    return rows_out


def tc_iter(tr):
    for child in tr.iterchildren():
        if child.tag == qn("w:tc"):
            yield child


def cell_record_to_json(rec: dict) -> dict:
    return {
        "runs": rec["runs"],
        "colspan": rec["colspan"],
        "rowspan": rec["rowspan"],
    }


def cell_full_text(rec: dict) -> str:
    """Text đầy đủ của ô, the thường/chk được dựng lại dạng {key}."""
    parts = []
    for r in rec["runs"]:
        if "text" in r:
            parts.append(r["text"])
        elif "input" in r:
            parts.append("{%s}" % r["input"]["key"])
    return "".join(parts)


def row_visible_text(row_cells: list[dict]) -> str:
    txts = []
    for rec in row_cells:
        txts.append(cell_full_text(rec))
    return ANY_TAG_RE.sub("", "".join(txts)).strip()


HEADER_BAND_FIRST_CELL = ("TÊN DOANH NGHIỆP", "TÊN HỘ KINH DOANH")
HEADER_BAND_MOTTO = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"

# [WAVE v7 HEADER-NAME-SYNC] registry band mirror (nguồn chân lý duy nhất
# do generator ghi/đọc; converter chỉ ĐỌC để biến the {key} ở ô (0,0) band
# thành run "mirror" — KHÔNG phải input thứ hai theo quyết định owner).
_HEADER_NAME_BINDING_PATH = ROOT / "tools" / "header_name_binding.json"
_hdr_name_keys: set | None = None
_hdr_name_forms: dict | None = None


def _header_name_binding_forms() -> dict:
    global _hdr_name_forms
    if _hdr_name_forms is None:
        try:
            data = json.loads(
                _HEADER_NAME_BINDING_PATH.read_text(encoding="utf-8"))
            _hdr_name_forms = data.get("forms", {})
        except (OSError, ValueError, KeyError):
            _hdr_name_forms = {}
    return _hdr_name_forms


def _band_mode_for(form_id: str, key: str) -> str:
    """Mode binding theo (form, key): 'mirror' (mặc định, band read-only) hoặc
    'band_input' (N-v1/N-v2: band LÀ ô nhập duy nhất của form — run input
    giữ nguyên để UI vẽ input editable đúng một lần tại band)."""
    entry = _header_name_binding_forms().get(form_id) or {}
    if entry.get("key") == key:
        return entry.get("mode", "mirror")
    return "mirror"


def is_header_band_table(grid: list[list[dict]]) -> bool:
    """Bảng header 2 band (hdr2col): ô đầu trái là ô TÊN …, bảng chứa khối
    quốc hiệu — chỉ các bảng header band của Việc 2 khớp, bảng dữ liệu thường
    không bao giờ chứa cả hai."""
    if not grid or not grid[0]:
        return False
    first = ANY_TAG_RE.sub("", cell_full_text(grid[0][0]))
    if not any(k in first for k in HEADER_BAND_FIRST_CELL):
        return False
    all_text = ANY_TAG_RE.sub(
        "", "".join(cell_full_text(rec) for row in grid for rec in row))
    return HEADER_BAND_MOTTO in all_text


def _band_first_cell_to_mirror(block: dict, form_id: str = "") -> None:
    """[WAVE v7] ô (0,0) band: run input với key thuộc registry mirror ->
    run {"mirror": {"key": …}} — UI vẽ text read-only đồng bộ store, KHÔNG
    tạo input/tab-stop/validation thứ hai (quyết định owner
    HEADER-NAME-SYNC 2026-09-05). Run không nằm registry (band chưa bind)
    giữ nguyên để gate bắt.
    [N-v1/N-v2 2026-09-06] mode='band_input' (PLI_M33/PLII_M06): giữ run
    input — band là ô NHẬP duy nhất của form, không mirror."""
    if not block.get("rows"):
        return
    mirror_keys = _header_name_binding_forms()
    if not mirror_keys:
        return
    cells = block["rows"][0].get("cells") or []
    if not cells:
        return
    for run in cells[0].get("runs") or []:
        inp = run.get("input")
        if not inp:
            continue
        key = inp.get("key")
        if any(e.get("key") == key for e in mirror_keys.values()) and \
                _band_mode_for(form_id, key) == "mirror":
            run.pop("input")
            run["mirror"] = {"key": key}


# ---------------------------------------------------------------------------
# [N-v3 OUTPUT-MIRROR 2026-09-06] mirror pass TỔNG QUÁT cho occurrence output
# (M48 p17 / M49 p9 nhánh DN / M70 p14): run input BẤT KỲ trong form F có key
# nằm trong registry output_mirrors[F] (section "output_mirrors" của
# tools/header_name_binding.json) -> run {"mirror": {"key": …}} — áp cho cả
# paragraph lẫn table cell, chạy SAU KHI parse toàn tài liệu. READ-ONLY theo
# quyết định owner: KHÔNG tạo input/tab-stop/draft key thứ hai; draft thiếu
# canonical -> render RỖNG đúng pháp định. Ô (0,0) band header KHÔNG đụng —
# đã xử lý riêng bằng _band_first_cell_to_mirror (mirror lẫn band_input đều
# bỏ qua). Block "loop" giữ nguyên: the trong hàng khuôn là cột dữ liệu loop,
# không phải occurrence output.
# ---------------------------------------------------------------------------

_output_mirror_forms: dict | None = None


def _output_mirror_registry() -> dict:
    global _output_mirror_forms
    if _output_mirror_forms is None:
        try:
            data = json.loads(
                _HEADER_NAME_BINDING_PATH.read_text(encoding="utf-8"))
            _output_mirror_forms = data.get("output_mirrors", {})
        except (OSError, ValueError):
            _output_mirror_forms = {}
    return _output_mirror_forms


def _mirror_runs_inplace(runs: list[dict], key: str) -> int:
    """Biến run input đúng key thành run mirror (trừ run đã là mirror)."""
    n = 0
    for run in runs:
        inp = run.get("input")
        if inp and inp.get("key") == key and "mirror" not in run:
            run.pop("input")
            run["mirror"] = {"key": key}
            n += 1
    return n


def _output_mirrors_to_mirror(blocks: list[dict], form_id: str) -> int:
    """Pass tổng quát: input run key ∈ output_mirrors[form_id] -> mirror."""
    entry = _output_mirror_registry().get(form_id)
    if not entry or not entry.get("key"):
        return 0
    key = entry["key"]
    n = 0
    for blk in blocks:
        t = blk.get("t")
        if t == "p":
            n += _mirror_runs_inplace(blk.get("runs") or [], key)
        elif t == "table":
            band_header = blk.get("band") == "header"
            for ri, row in enumerate(blk.get("rows") or []):
                for ci, cell in enumerate(row.get("cells") or []):
                    if band_header and ri == 0 and ci == 0:
                        continue  # ô (0,0) band do pass riêng xử lý
                    n += _mirror_runs_inplace(cell.get("runs") or [], key)
    return n


def table_block(tbl: Table, fid: str = "") -> dict:
    grid = parse_table(tbl)

    # phát hiện hàng khuôn loop: một hàng chứa cả {#ten} và {/ten}
    loop_name = None
    template_idx = None
    for ri, row_cells in enumerate(grid):
        texts = "".join(cell_full_text(rec) for rec in row_cells)
        opens = LOOP_OPEN_RE.findall(texts)
        for name in opens:
            if "{/%s}" % name in texts:
                loop_name, template_idx = name, ri
                break
        if loop_name:
            break

    if loop_name is None:
        block = {
            "t": "table",
            "widths": None,
            "rows": [
                {"cells": [cell_record_to_json(rec) for rec in row]}
                for row in grid
            ],
        }
        if is_header_band_table(grid):
            block["band"] = "header"
            _band_first_cell_to_mirror(block, fid)
        return block

    # bảng loop: headerRows = hàng tiêu đề (không trắng) trước hàng khuôn,
    # blankRows = số hàng trắng giữa header và khuôn (bỏ khỏi rows),
    # hàng khuôn bỏ khỏi rows.
    blank_idx = set()
    last_header = -1
    for ri in range(template_idx):
        if row_visible_text(grid[ri]):
            last_header = ri
        else:
            blank_idx.add(ri)
    header_rows = [ri for ri in range(template_idx) if ri not in blank_idx]

    cols = []
    for rec in grid[template_idx]:
        cell_text = cell_full_text(rec)
        for tag in TAG_RE.findall(cell_text):
            if tag == "_stt":
                # biến engine: STT tự sinh 1-based khi render, không phải
                # cột dữ liệu người dùng -> không vào cols/fields
                continue
            cols.append({"col": rec["col"], "tag": tag})

    rows_json = [
        {"cells": [cell_record_to_json(rec) for rec in row]}
        for ri, row in enumerate(grid)
        if ri != template_idx and ri not in blank_idx
    ]

    return {
        "t": "loop",
        "loop": loop_name,
        "headerRows": header_rows,
        "cols": cols,
        "blankRows": len(blank_idx & set(range(last_header + 1, template_idx))),
        "rows": rows_json,
    }


# ---------------------------------------------------------------------------
# Section specs (P2/Lane B2): kiểm định cấu trúc bảng kind mới trong blocks
#
# form_maps.json[fid]["section_specs"] (do tools/build_form_maps.py sinh cho
# kind fixed-matrix / nested-grid / output-table) cho converter biết:
#   1. bảng nào phải tồn tại và ở vị trí nào;
#   2. cấu trúc vật lý phải thế nào (số cột gridCol, rowspan/colspan, vị trí
#      the STT/totals/the dữ liệu theo col_index);
#   3. output-table không được chứa bất kỳ input nào.
# Lệch -> GridSpecMismatch -> exit != 0. Chỉ áp dụng cho form CÓ spec nên các
# mẫu hiện tại giữ hành vi và JSON byte-identical.
#
# CHÚ Ý integration: danh sách kind KHÔNG được ghi vào block JSON bằng khoá
# mới. tools/audit_blocks_check.py (contract QA) khóa chặt ALLOWED_TOP_KEYS =
# {id, ten_mau, blocks, v} và ALLOWED_BLOCK_KEYS[t]; muốn gắn marker vào block
# thì trước hết B3/owner phải mở rộng allowlist đó. Cho đến khi đó, định danh
# của từng bảng lấy bằng JOIN: thứ tự các block t∈{table,loop} trong blocks[]
# trùng với "table_index" ghi trong form_maps[fid]["section_specs"][k].
# ---------------------------------------------------------------------------


class GridSpecMismatch(ValueError):
    """Blocks lệch so với section_spec khai báo — fail loud."""


def _fold_txt(text: str) -> str:
    return layout_roles.fold(text or "")


def _block_row_runs_text(cell: dict) -> str:
    return "".join(r.get("text", "") for r in cell.get("runs", []) if "text" in r)


def _block_cell_inputs(cell: dict) -> list[str]:
    return [r["input"]["key"] for r in cell.get("runs", []) if "input" in r]


def _physical_indexes(cells: list[dict]):
    """[(index_bắt_đầu, colspan, cell)] theo chiều rộng vật lý của một hàng."""
    out = []
    cursor = 0
    for cell in cells:
        span = int(cell.get("colspan") or 1)
        out.append((cursor, span, cell))
        cursor += span
    return out


def _find_table_block(blocks, want_headers, start_at=0):
    """Table block có hàng 0 khớp chuỗi tiêu đề đã gập dấu."""
    want = "|".join(_fold_txt(h) for h in want_headers)
    for bi in range(start_at, len(blocks)):
        blk = blocks[bi]
        if blk.get("t") != "table" or not blk.get("rows"):
            continue
        got = "|".join([_fold_txt(_block_row_runs_text(cell))
                        for _i, _s, cell in _physical_indexes(blk["rows"][0]["cells"])])
        if got == want:
            return bi, blk
    return None, None


def _table_block_signature(blk: dict) -> str:
    """Chuỗi chuẩn hoá của toàn bộ block 'table' (mỗi hàng = các ô gập dấu
    nối bằng '|'). Dùng để định vị lưới output-table mà không phụ thuộc một
    hàng tiêu đề cụ thể."""
    lines = []
    for row in blk.get("rows", []):
        cells = [_fold_txt(_block_row_runs_text(cell))
                 for _pos, _span, cell in _physical_indexes(row.get("cells", []))]
        lines.append("|".join(cells))
    return "#".join(lines)


def _find_output_grid_block(blocks, labels: list, cols: int,
                            rows_hint: int = 0):
    """Tìm block bảng khớp lưới output. labels phẳng ("A", "B", ...) theo
    legacy = N hàng × 1 cột (mỗi hàng nhãn ở ô đầu). rows_hint=1 + cols>1:
    bảng 1 HÀNG × N CỘT (khối ký ngang PLI_M70 sau 5R2 — cặp signer nằm
    ngang) -> nhãn trải theo cột trong cùng một hàng."""
    matrix = bool(labels) and all(isinstance(x, list) for x in labels)
    horizontal = (not matrix and rows_hint == 1 and cols > 1
                  and len(labels) == cols)
    rows_txt = []
    for r in range(len(labels) if not horizontal else 1):
        if matrix:
            cells = [labels[r][c] if c < len(labels[r]) else ""
                     for c in range(cols)]
        elif horizontal:
            cells = [labels[c] if c < len(labels) else "" for c in range(cols)]
        else:
            cells = [labels[r] if (c == 0 and r < len(labels)) else ""
                     for c in range(cols)]
        rows_txt.append("|".join(_fold_txt(x) for x in cells))
    want = "#".join(rows_txt)
    for bi, blk in enumerate(blocks):
        if blk.get("t") != "table" or not blk.get("rows"):
            continue
        if _table_block_signature(blk) == want:
            return bi, blk
    return None, None


def verify_fixed_matrix_spec(blk, spec: dict, sem: str) -> list[str]:
    errs = []
    cols = spec["columns"]
    rows_spec = spec["fixed_rows"]
    width = len(cols)
    body = blk.get("rows") or []
    if len(body) != len(rows_spec) + 1:
        errs.append("số hàng %d != %d (header + %d dòng pháp lý)"
                    % (len(body), len(rows_spec) + 1, len(rows_spec)))
        return errs
    for ri, row in enumerate(body):
        if sum(int(c.get("colspan") or 1) for c in row["cells"]) != width:
            errs.append("hàng %d mất cột vật lý (tổng colspan != %d)"
                        % (ri, width))
    head_cells = {pos: cell for pos, _span, cell
                  in _physical_indexes(body[0]["cells"])}
    for ci, col in enumerate(cols):
        cell = head_cells.get(ci)
        if cell is None or \
                _fold_txt(_block_row_runs_text(cell)) != _fold_txt(col["header"]):
            errs.append("header cột %d != %r" % (ci, col["header"]))
    seen_keys = {}
    for ri, row_spec in enumerate(rows_spec):
        row = body[ri + 1]
        cells_by_pos = {pos: cell for pos, _span, cell
                        in _physical_indexes(row["cells"])}
        for ci, col in enumerate(cols):
            cell = cells_by_pos.get(ci)
            if cell is None:
                errs.append("hàng '%s' thiếu ô tại cột %d" % (row_spec["id"], ci))
                continue
            keys = _block_cell_inputs(cell)
            if col["tag"] is None:
                if keys:
                    errs.append("hàng '%s' cột nhãn %d vẫn còn the %s"
                                % (row_spec["id"], ci, keys))
                continue
            want_key = "%s.%s.%s" % (sem, row_spec["id"], col["tag"])
            seen_keys[want_key] = seen_keys.get(want_key, 0) + 1
            if keys != [want_key]:
                errs.append("ô (%s,%d) có the %s, kỳ vọng [%r]"
                            % (row_spec["id"], ci, keys, want_key))
    duplicated = {k: n for k, n in seen_keys.items() if n > 1}
    if duplicated:
        errs.append("the tổng/dòng bị lặp: %s" % sorted(duplicated))
    for tot in spec.get("totals", []):
        want = "%s.%s.%s" % (sem, tot["target_row"], tot["target_tag"])
        if seen_keys.get(want, 0) < 1:
            errs.append("mất the tổng %r" % want)
    return errs


def verify_nested_grid_spec(blk, spec: dict, sem: str) -> list[str]:
    errs = []
    width = int(spec["width"])
    hdr_spec = spec["header_rows"]
    hdr_n = len(hdr_spec)
    hdr_idx = blk.get("headerRows")
    if list(hdr_idx or []) != list(range(hdr_n)):
        errs.append("headerRows %s != %s" % (hdr_idx, list(range(hdr_n))))
        return errs
    rows = blk.get("rows") or []
    if len(rows) < hdr_n:
        errs.append("blocks thiếu hàng sau header (%d < %d)"
                    % (len(rows), hdr_n))
        return errs
    # Không được có thêm hàng khuôn nào ngoài hàng mẫu chuẩn: {#name} kèm
    # {/name} xuất hiện trong hàng khác là dấu hiệu bảng bị nhân bản khi rebuild.
    def row_text(row: dict) -> str:
        return "".join(r.get("text", "") for c in row.get("cells", [])
                       for r in c.get("runs", []))
    extra_rows = []
    for ri, row in enumerate(rows[hdr_n:], start=hdr_n):
        text = row_text(row)
        for m in LOOP_OPEN_RE.finditer(text):
            if "{/%s}" % m.group(1) in text:
                extra_rows.append((ri, m.group(1)))
    if extra_rows:
        errs.append("bảng chứa %d hàng khuôn lặp nữa ngoài hàng mẫu chuẩn %s "
                    "(phải đúng 1 hàng khuôn)" % (len(extra_rows),
                                                  extra_rows[:4]))
    covered = {}
    expected = {}
    for ri, row_spec in enumerate(hdr_spec):
        cursor = 0
        for src in row_spec:
            while (ri, cursor) in covered:
                cursor += 1
            for r in range(ri, ri + src["rowspan"]):
                for c in range(cursor, cursor + src["colspan"]):
                    covered[(r, c)] = True
            expected[(ri, cursor)] = (src["text"], src["rowspan"],
                                      src["colspan"])
            cursor += src["colspan"]
    missing = [(r, c) for r in range(hdr_n) for c in range(width)
               if not covered.get((r, c))]
    if missing:
        errs.append("header chưa phủ kín %dx%d, thiếu %s"
                    % (hdr_n, width, missing[:5]))
    # Dựng lại VỊ TRÍ vật lý của các ô header trong blocks: ô bị gộp dọc
    # (vMerge continue) bị parse_table lượt khỏi danh sách cell, nên vị trí
    # từng ô phải suy từ occupancy của các ô phía trên mang rowspan.
    actual = {}
    over_wide = []
    filled = set()
    for ri, row in enumerate(rows[:hdr_n]):
        cursor = 0
        for cell in row["cells"]:
            rs = int(cell.get("rowspan") or 1)
            cs = int(cell.get("colspan") or 1)
            while (ri, cursor) in filled:
                cursor += 1
            if cursor + cs > width:
                over_wide.append((ri, cursor, cs))
                break
            actual[(ri, cursor)] = (_fold_txt(_block_row_runs_text(cell)),
                                    rs, cs)
            for r in range(ri, ri + rs):
                for c in range(cursor, cursor + cs):
                    filled.add((r, c))
            cursor += cs
    del filled
    if over_wide:
        errs.append("header ô vượt quá chiều rộng vật lý %d tại %s"
                    % (width, over_wide[:4]))
    for pos, want in sorted(expected.items()):
        text, want_rs, want_cs = want
        got = actual.get(pos)
        if got is None:
            errs.append("header thiếu ô tại %s" % (pos,))
            continue
        if got[0] != _fold_txt(text):
            errs.append("header %s text=%r != %r" % (pos, got[0], text))
        if got[1] != want_rs or got[2] != want_cs:
            errs.append("header %s span=(rs=%d,cs=%d) != kỳ vọng "
                        "(rs=%d,cs=%d)" % (pos, got[1], got[2],
                                           want_rs, want_cs))
    for pos in sorted(actual.keys()):
        if pos not in expected:
            errs.append("header thừa ô tại %s (%r)"
                        % (pos, actual[pos][0]))
    stt_ci = (spec.get("stt") or {}).get("col_index")
    got_cols = {int(c["col"]): c["tag"] for c in blk.get("cols", [])}
    want_cols = {}
    for c in spec["columns"]:
        if c["tag"]:
            want_cols[int(c["col_index"])] = c["tag"]
    if got_cols != want_cols:
        errs.append("cols của hàng khuôn %s != %s (vị trí vật lý sai/mất cột)"
                    % (
                        sorted(got_cols.items()),
                        sorted(want_cols.items())))
    if stt_ci is not None and stt_ci in got_cols:
        errs.append("cột STT %d bị gắn the dữ liệu" % stt_ci)
    return errs


def verify_output_table_spec(blk, spec: dict, sem: str) -> list[str]:
    errs = []
    grid = spec["grid"]
    rows = blk.get("rows") or []
    if len(rows) != grid["rows"]:
        errs.append("số hàng %d != %d" % (len(rows), grid["rows"]))
    inputs = []
    for ri, row in enumerate(rows):
        if sum(int(c.get("colspan") or 1) for c in row["cells"]) != grid["cols"]:
            errs.append("hàng %d lệch số cột vật lý (kỳ vọng %d)"
                        % (ri, grid["cols"]))
        for cell in row["cells"]:
            inputs.extend(_block_cell_inputs(cell))
    if inputs:
        errs.append("output-table KHÔNG được có input, thấy %s" % inputs[:8])
    labels = grid["labels"]
    matrix = bool(labels) and all(isinstance(x, list) for x in labels)
    # Bảng ký NGANG (5R2 PLI_M70): grid rows=1, cols>1, labels phẳng đếm đủ
    # cols -> nhãn trải theo CỘT trong 1 hàng (không phải N hàng × 1 cột).
    horizontal = (not matrix and int(grid.get("rows", 0) or 0) == 1
                  and int(grid.get("cols", 1)) > 1
                  and len(labels) == int(grid["cols"]))
    for ri, row in enumerate(rows):
        for pos, _span, cell in _physical_indexes(row["cells"]):
            if matrix:
                want = (labels[ri][pos] if ri < len(labels) and
                        pos < len(labels[ri]) else "")
            elif horizontal:
                want = labels[pos] if pos < len(labels) else ""
            else:
                want = (labels[ri] if pos == 0 and ri < len(labels) else "")
            if _fold_txt(_block_row_runs_text(cell)) != _fold_txt(want):
                errs.append("nhãn (%d,%d) = %r != %r"
                            % (ri, pos, _block_row_runs_text(cell), want))
    return errs


def annotate_and_verify_section_specs(blocks, section_specs) -> int:
    """Định vị + kiểm định mọi section_spec trên danh sách blocks.

    Trả về số section đã ghép; raise GridSpecMismatch khi có bất kỳ lệch nào
    (thiếu/xa lạ/mất cột/theo vị trí sai). KHÔNG thêm khoá nào vào block —
    xem ghi chú integration phía trên."""
    problems = []
    matched = 0
    for spec_info in section_specs or []:
        kind = spec_info.get("kind")
        sem = spec_info.get("semantic_key")
        sid = spec_info.get("section_id") or spec_info.get("id") or sem
        payload = spec_info
        if kind == "nested-grid":
            loop = spec_info.get("loop") or sem
            cand = [b for b in blocks
                    if b.get("t") == "loop" and b.get("loop") == loop]
            if not cand:
                problems.append("%s (%s): không tìm thấy block loop %r"
                                % (sid, sem, loop))
                continue
            if len(cand) > 1:
                # >1 hàng khuôn cùng tên ({#loop} lặp lại trong bảng) -> engine
                # sẽ nhân bản hàng mơ hồ; bắt ngay ở lúc convert.
                problems.append(
                    "%s (%s): thấy %d block loop %r (phải đúng 1 hàng khuôn)"
                    % (sid, sem, len(cand), loop))
                continue
            blk = cand[0]
            errs = verify_nested_grid_spec(blk, payload, sem)
        elif kind == "fixed-matrix":
            headers = [c["header"] for c in payload["columns"]]
            _bi, blk = _find_table_block(blocks, headers)
            if blk is None:
                problems.append("%s (%s): không tìm thấy bảng khớp header %s"
                                % (sid, sem, headers))
                continue
            errs = verify_fixed_matrix_spec(blk, payload, sem)
        elif kind == "output-table":
            grid = payload.get("grid", {})
            _bi, blk = _find_output_grid_block(
                blocks, grid.get("labels", []),
                int(grid.get("cols", 1)),
                rows_hint=int(grid.get("rows", 0) or 0))
            if blk is None:
                problems.append("%s (%s): không tìm thấy lưới output khớp "
                                "labels %r" % (sid, sem,
                                               grid.get("labels")))
                continue
            errs = verify_output_table_spec(blk, payload, sem)
        else:
            problems.append("%s: kind không hỗ trợ trong converter: %r"
                            % (sid, kind))
            continue
        matched += 1
        if errs:
            problems.append("%s (%s/%s):\n   - %s"
                            % (sid, sem, kind, "\n   - ".join(errs)))
    if problems:
        raise GridSpecMismatch(
            "section_specs lệch cấu trúc (%d):\n - %s"
            % (len(problems), "\n - ".join(problems)))
    return matched


# ---------------------------------------------------------------------------
# Chuyển 1 file
# ---------------------------------------------------------------------------

def load_footnotes(doc) -> dict:
    """Đọc text chú thích từ word/footnotes.xml (python-docx 1.2.0 không có
    API — truy cập part qua package). Trả về {id: text}; bỏ id <= 0
    (separator/continuation). Fail-soft: thiếu part -> {}."""
    out: dict = {}
    try:
        parts = list(doc.part.package.iter_parts())
    except Exception:  # noqa: BLE001
        return out
    for part in parts:
        name = str(part.partname)
        if not name.endswith("/footnotes.xml") and name != "/word/footnotes.xml":
            continue
        try:
            from lxml import etree

            root = etree.fromstring(part.blob)
            for el in root.iter(qn("w:footnote")):
                raw = el.get(qn("w:id"))
                try:
                    fid = int(raw)
                except (TypeError, ValueError):
                    continue
                if fid < 1:
                    continue
                txt = "".join(t.text or "" for t in el.iter(qn("w:t")))
                txt = txt.strip()
                if txt:
                    out.setdefault(fid, txt)
        except Exception:  # noqa: BLE001
            continue  # part hỏng -> coi như không có footnote text
    return out


def convert_document(
    doc, fid: str, ten_mau: str, overrides: dict | None = None,
    footnotes: dict | None = None, section_specs: list | None = None,
) -> dict:
    """Chuyển một Document đã mở thành block JSON (không IO file).

    Tách riêng để test dựng fixture DOCX trong temp và gọi trực tiếp.
    section_specs: metadata cấu trúc của kind mới (fixed-matrix/nested-grid/
    output-table) lấy từ form_maps[fid]["section_specs"]; có -> converter
    kiểm định CHẶT cấu trúc block tương ứng rồi mới ghi JSON.
    """
    blocks: list[dict] = []
    prev_ref = {"block": None, "consumed": False, "emitted": blocks}
    cond_stack: list[str] = []  # flag đang mở theo thứ tự (lồng nhau = AND)
    ctx = {
        "form_id": fid,
        "footnotes": footnotes if footnotes is not None else {},
        "overrides": overrides or {},
        "body_para_index": -1,
    }

    def _attach_cond(b: dict) -> dict:
        if cond_stack:
            b["cond"] = list(cond_stack)
        return b

    for kind, obj in iter_blocks(doc):
        if kind == "p":
            ctx["body_para_index"] += 1
            full = _para_full_text(obj).strip()
            mo = COND_OPEN_FULL_RE.match(full)
            mc = COND_CLOSE_FULL_RE.match(full)
            if mo:
                cond_stack.append(mo.group(1))
                continue
            if mc:
                name = mc.group(1)
                if not cond_stack or cond_stack[-1] != name:
                    raise ValueError(
                        "%s: marker {/%s} sai cặp (đang mở: %s)"
                        % (fid, name, cond_stack))
                cond_stack.pop()
                continue
            new_blocks = para_block(obj, prev_ref, ctx)
            blocks.extend(_attach_cond(b) for b in new_blocks)
            prev_ref["block"] = new_blocks[-1] if new_blocks else prev_ref["block"]
            if not new_blocks:
                # paragraph rỗng: không reset trạng thái consumed
                pass
            else:
                prev_ref["consumed"] = False
        else:
            blk = table_block(obj, fid)
            _attach_cond(blk)
            blocks.append(blk)
            prev_ref["block"] = None
            prev_ref["consumed"] = False

    if cond_stack:
        raise ValueError(
            "%s: marker {?%s} không được đóng trong template" % (fid, cond_stack[-1]))

    # [N-v3] pass mirror TỔNG QUÁT chạy SAU KHI parse toàn tài liệu: run
    # input có key nằm trong registry output_mirrors[fid] -> run mirror
    # (paragraph lẫn table cell; ô (0,0) band đã xử lý riêng nên bỏ qua).
    _output_mirrors_to_mirror(blocks, fid)

    if section_specs:
        # kind mới: kiểm định cấu trúc NGAY TRƯỚC khi ghi JSON (fail loud);
        # chỉ đếm số section đã ghép, không thêm khoá nào vào JSON đầu ra.
        annotate_and_verify_section_specs(blocks, section_specs)

    return {"id": fid, "ten_mau": ten_mau, "blocks": blocks,
            "v": SCHEMA_VERSION}


SCHEMA_VERSION = 3


def convert(fid: str, form_maps: dict, overrides: dict | None = None) -> dict:
    path = TEMPLATES_DIR / ("%s.docx" % fid)
    doc = Document(str(path))
    fmap = form_maps.get(fid, {})
    ten_mau = fmap.get("ten_mau", "")
    return convert_document(
        doc, fid, ten_mau,
        overrides=overrides if overrides is not None else load_role_overrides(),
        footnotes=load_footnotes(doc),
        section_specs=fmap.get("section_specs"),
    )


# ---------------------------------------------------------------------------
# Tự kiểm định
# ---------------------------------------------------------------------------

def collect_docx_tags(path: Path) -> int:
    """Đếm tổng số the {…} (không tính {#..}/{/..}) từ python-docx trực tiếp.

    The {_stt} bị loại vì đây là biến engine (STT tự sinh khi render) không
    xuất hiện dưới dạng input trong blocks — giữ đối xứng với collect_json_tags.
    Duyệt con trực tiếp w:tc của w:tr thay vì row.cells + dedup id(): proxy
    python-docx tạo mới mỗi lần nên id() không ổn định, có thể bỏ nhầm ô."""
    doc = Document(str(path))
    n = 0
    for kind, obj in iter_blocks(doc):
        if kind == "p":
            n += len(TAG_RE.findall(_para_full_text(obj)))
        else:
            for tr in obj._tbl.tr_lst:
                for tc in tc_iter(tr):
                    for pp in _cell_paragraphs(tc, obj.part.document):
                        n += sum(1 for t in TAG_RE.findall(pp.text)
                                 if t != "_stt")
    return n


def collect_json_tags(data: dict) -> int:
    # [integ(F0L) 2026-09-06] đếm thêm run MIRROR. Premise cũ: mỗi the {…}
    # ở template tương ứng đúng 1 run "input" ở blocks. Hiện trạng: wave v7
    # (HEADER-NAME-SYNC) biến the ở ô (0,0) band thành run {"mirror": …} —
    # CỐ ÊN không phải input thứ hai (owner decision), nên side JSON đếm thiếu
    # đúng 1 cho từng mẫu band -> self-check (b) lệch với mọi mẫu có band
    # (evidence: tại e268038, seed 20260824 bắt 4/10 mẫu band: PLI_M13,
    # PLII_M04, PLII_M05, PLI_M27 đều docx=json+1). Đếm mirror làm đối chiếu
    # 1:1 trở lại chính xác (mirror mất -> (b) vẫn bắt được).
    n = 0
    for b in data["blocks"]:
        if b["t"] in ("p", "table"):
            rows = b.get("rows") if b["t"] == "table" else None
            if rows is not None:
                for row in rows:
                    for cell in row["cells"]:
                        n += sum(1 for r in cell["runs"]
                                 if "input" in r or "mirror" in r)
            else:
                n += sum(1 for r in b["runs"]
                         if "input" in r or "mirror" in r)
        elif b["t"] == "chk":
            n += 1
        elif b["t"] == "loop":
            n += len(b["cols"])
            for row in b.get("rows", []):
                for cell in row["cells"]:
                    n += sum(1 for r in cell["runs"] if "input" in r)
    return n


ANY_TAG_RE = re.compile(r"\{[/#?]?[A-Za-z_][A-Za-z0-9_.#/?]*\}")


def collect_docx_text(path: Path) -> str:
    """Toàn bộ text của file, đã bỏ mọi the {…}/{#…}/{/…}.
    Duyệt con trực tiếp w:tc (không dùng row.cells + id() — xem
    collect_docx_tags)."""
    doc = Document(str(path))
    parts = []
    for kind, obj in iter_blocks(doc):
        if kind == "p":
            parts.append(ANY_TAG_RE.sub("", obj.text))
        else:
            for tr in obj._tbl.tr_lst:
                for tc in tc_iter(tr):
                    txt = "\n".join(
                        pp.text for pp in _cell_paragraphs(tc, obj.part.document)
                    )
                    parts.append(ANY_TAG_RE.sub("", txt))
    return "".join(parts)


_WS_RE = re.compile(r"\s+")


def collect_json_text(data: dict) -> str:
    parts = []

    def runs_text(runs):
        for r in runs:
            if "text" in r:
                parts.append(r["text"])

    for b in data["blocks"]:
        if b["t"] == "p":
            runs_text(b["runs"])
        elif b["t"] == "chk":
            parts.append(b.get("label", ""))
        elif b["t"] == "table":
            for row in b["rows"]:
                for cell in row["cells"]:
                    runs_text(cell["runs"])
        elif b["t"] == "loop":
            for row in b.get("rows", []):
                for cell in row["cells"]:
                    runs_text(cell["runs"])
    return "".join(parts)


def summarize_tree(data: dict) -> str:
    lines = []
    for i, b in enumerate(data["blocks"]):
        if b["t"] == "p":
            txt = "".join(r.get("text", "") for r in b["runs"])
            ntag = sum(1 for r in b["runs"] if "input" in r)
            head = _WS_RE.sub(" ", txt).strip()[:60]
            lines.append("[%d] p(%s) tags=%d %r" % (i, b["align"], ntag, head))
        elif b["t"] == "chk":
            lines.append("[%d] chk %s label=%r" % (i, b["key"], b["label"][:40]))
        elif b["t"] == "table":
            nr = len(b["rows"])
            nc = len(b["rows"][0]["cells"]) if b["rows"] else 0
            first = ""
            if b["rows"] and b["rows"][0]["cells"]:
                first = _WS_RE.sub(
                    " ",
                    "".join(
                        r.get("text", "") for r in b["rows"][0]["cells"][0]["runs"]
                    ),
                ).strip()[:30]
            lines.append("[%d] table rows=%d cells(r0)=%d %r" % (i, nr, nc, first))
        elif b["t"] == "loop":
            lines.append(
                "[%d] loop %s headerRows=%s blankRows=%d cols=%s rows=%d"
                % (
                    i,
                    b["loop"],
                    b["headerRows"],
                    b["blankRows"],
                    [(c["col"], c["tag"]) for c in b["cols"]],
                    len(b.get("rows", [])),
                )
            )
    return "\n".join(lines)


def self_check(ids: list[str], form_maps: dict, tree_ids: list[str]) -> int:
    errors = []
    # (a) đủ file, json.load OK
    json_files = sorted(OUT_DIR.glob("*.json"))
    loaded = {}
    for jf in json_files:
        try:
            loaded[jf.stem] = json.loads(jf.read_text(encoding="utf-8"))
        except Exception as exc:  # noqa: BLE001
            errors.append("(a) %s json hỏng: %r" % (jf.name, exc))
    missing = [i for i in ids if i not in loaded]
    if missing:
        errors.append("(a) thiếu %d file: %s" % (len(missing), missing[:5]))

    # (b) 10 mẫu ngẫu nhiên: tổng số the khớp python-docx
    rng = random.Random(20260824)
    sample = rng.sample(sorted(loaded.keys()), min(10, len(loaded)))
    for fid in sample:
        want = collect_docx_tags(TEMPLATES_DIR / ("%s.docx" % fid))
        got = collect_json_tags(loaded[fid])
        status = "OK" if want == got else "LỆCH"
        print("  (b) %-10s the docx=%-4d json=%-4d %s" % (fid, want, got, status))
        if want != got:
            errors.append("(b) %s: docx=%d json=%d" % (fid, want, got))

    # (c) cây block 3 mẫu đối chiếu mắt
    for fid in tree_ids:
        print("----- cây block %s -----" % fid)
        print(summarize_tree(loaded[fid]))

    # (d) không mất text
    rng2 = random.Random(42)
    sample2 = rng2.sample(sorted(loaded.keys()), min(20, len(loaded)))
    for fid in sample2:
        want = _WS_RE.sub("", collect_docx_text(TEMPLATES_DIR / ("%s.docx" % fid)))
        got = _WS_RE.sub("", collect_json_text(loaded[fid]))
        if want != got:
            errors.append(
                "(d) %s mất text: docx=%d ký tự vs json=%d ký tự"
                % (fid, len(want), len(got))
            )
            print("  (d) %-10s MẤT TEXT (%d vs %d)" % (fid, len(want), len(got)))
        else:
            print("  (d) %-10s text khớp (%d ký tự)" % (fid, len(want)))

    print("-" * 60)
    if errors:
        print("TỰ KIỂM ĐỊNH THẤT BẠI:")
        for e in errors:
            print(" -", e)
        return 1
    print("TỰ KIỂM ĐỊNH: %d file JSON đều hợp lệ." % len(json_files))
    return 0


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> int:
    args = sys.argv[1:]
    form_maps = json.loads(FORM_MAPS_PATH.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    if "--one" in args:
        fid = args[args.index("--one") + 1]
        data = convert(fid, form_maps)
        out = OUT_DIR / ("%s.json" % fid)
        out.write_text(
            json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8"
        )
        print("Đã ghi", out)
        return 0

    ids = sorted(form_maps.keys())
    for fid in ids:
        data = convert(fid, form_maps)
        out = OUT_DIR / ("%s.json" % fid)
        out.write_text(
            json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8"
        )
    print("Đã chuyển %d template -> %s" % (len(ids), OUT_DIR))

    stats = {}
    total_blocks = 0
    for jf in sorted(OUT_DIR.glob("*.json")):
        d = json.loads(jf.read_text(encoding="utf-8"))
        cnt = {}
        for b in d["blocks"]:
            cnt[b["t"]] = cnt.get(b["t"], 0) + 1
        stats[jf.stem] = cnt
        total_blocks += len(d["blocks"])
    kinds = ("p", "chk", "table", "loop")
    print("%-10s %6s %5s %7s %6s %7s" % (("file",) + kinds + ("tổng",)))
    for fid in sorted(stats):
        c = stats[fid]
        print(
            "%-10s %6d %5d %7d %6d %7d"
            % ((fid,) + tuple(c.get(k, 0) for k in kinds) + (sum(c.values()),))
        )
    print("Tổng số block mọi mẫu:", total_blocks)

    return self_check(ids, form_maps, ["PLI_M02", "PLI_M06", "PLII_M01"])


if __name__ == "__main__":
    sys.exit(main())
