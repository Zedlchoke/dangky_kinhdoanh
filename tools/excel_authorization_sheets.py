# -*- coding: utf-8 -*-
"""Sheet AUTH_* trong workbook Excel case-level (wave I5).

Đối tác: tools/excel_case_export.py, tools/excel_case_import.py,
tools/authorization_store.py.

Sheet là cấu trúc (không nằm __map) — workbook cũ không có AUTH_* vẫn nạp draft.
"""
from __future__ import annotations

import json
import uuid
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from tools.excel_export import GOLD_HEX, GRAY_HEX, META_FONT_HEX

AUTH_MARKER = "__AUTH_EXCEL_V1__"
AUTH_SCHEMA_VERSION = 1

AUTH_META_SHEET = "AUTH_META"
AUTH_AGENTS_SHEET = "AUTH_AGENTS"
AUTH_SETS_SHEET = "AUTH_SETS"
AUTH_DOCUMENTS_SHEET = "AUTH_DOCUMENTS"
AUTH_ATTACHMENTS_SHEET = "AUTH_ATTACHMENTS"
AUTH_AUDIT_SHEET = "AUTH_AUDIT"

AUTH_SHEETS = (
    AUTH_META_SHEET,
    AUTH_AGENTS_SHEET,
    AUTH_SETS_SHEET,
    AUTH_DOCUMENTS_SHEET,
    AUTH_ATTACHMENTS_SHEET,
    AUTH_AUDIT_SHEET,
)

AGENT_COLUMNS = (
    "agent_id", "ho_ten", "ngay_sinh", "gioi_tinh", "quoc_tich",
    "so_cccd", "ngay_cap", "noi_cap", "dia_chi_lien_he", "dien_thoai",
    "email", "chuc_danh", "don_vi_dich_vu", "trang_thai", "ghi_chu",
    "created_at", "updated_at",
)

SET_COLUMNS = (
    "authorization_id", "case_id", "agent_id", "status", "scope_text",
    "co_quan_tiep_nhan", "thoi_han_den", "template_ids_json",
    "created_at", "updated_at",
)

DOCUMENT_COLUMNS = (
    "document_id", "authorization_id", "template_id", "file_name",
    "sha256", "status", "generated_at",
)

ATTACHMENT_COLUMNS = (
    "attachment_id", "authorization_id", "rel_path", "sha256",
    "mime", "size_bytes",
)

AUDIT_COLUMNS = (
    "event_id", "event_type", "entity_id", "detail", "created_at",
)

REQUIRED_AGENT = ("agent_id", "ho_ten", "so_cccd")
REQUIRED_SET = ("authorization_id", "case_id", "agent_id")


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _set_text(cell, value) -> None:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        cell.value = text
        cell.data_type = "s"
    else:
        cell.value = text


def _write_meta_sheet(ws: Worksheet, case_id: str, meta: dict | None) -> None:
    meta = meta or {}
    ws.cell(row=1, column=1, value=AUTH_MARKER).font = Font(bold=True, size=12)
    line = "case_id=%s | auth_schema_version=%s | exported_at=%s" % (
        case_id, AUTH_SCHEMA_VERSION, meta.get("exported_at") or "")
    cell = ws.cell(row=2, column=1, value=line)
    cell.font = Font(size=10, italic=True, color=META_FONT_HEX)
    note = ws.cell(
        row=3, column=1,
        value="Metadata ủy quyền dịch vụ — app đọc khi nạp workbook case-level.")
    note.font = Font(size=10, color=META_FONT_HEX)
    ws.column_dimensions["A"].width = 90


def _write_data_sheet(
    ws: Worksheet,
    title: str,
    columns: tuple[str, ...],
    rows: list[dict[str, Any]],
) -> None:
    ws.cell(row=1, column=1, value=title).font = Font(
        bold=True, size=13, color="1F3864")
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=max(len(columns), 4))
    note = ws.cell(
        row=2, column=1,
        value="Không đổi tên cột. Ô vàng có thể sửa; hàng dữ liệu bắt đầu từ hàng 4.")
    note.font = Font(size=10, italic=True, color=META_FONT_HEX)
    ws.merge_cells(start_row=2, start_column=1, end_row=2,
                   end_column=max(len(columns), 4))
    header_row = 3
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=name)
        cell.font = Font(bold=True, size=10, color="1F3864")
        cell.fill = _fill(GRAY_HEX)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r_idx, row in enumerate(rows, start=header_row + 1):
        for col_idx, name in enumerate(columns, start=1):
            cell = ws.cell(row=r_idx, column=col_idx)
            _set_text(cell, row.get(name))
            cell.number_format = "@"
            cell.protection = Protection(locked=False)
            cell.fill = _fill(GOLD_HEX)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col_idx, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(
            12, min(36, len(name) + 4))
    ws.protection.sheet = True
    ws.protection.enable()


def _read_meta_case_id(ws: Worksheet) -> str | None:
    if ws.cell(row=1, column=1).value != AUTH_MARKER:
        return None
    line = str(ws.cell(row=2, column=1).value or "")
    for part in line.split("|"):
        part = part.strip()
        if part.startswith("case_id="):
            return part.split("=", 1)[1].strip()
    return None


def _read_data_rows(ws: Worksheet, columns: tuple[str, ...]) -> list[dict[str, str]]:
    header_row = 3
    headers = [
        str(ws.cell(header_row, c).value or "").strip()
        for c in range(1, len(columns) + 1)
    ]
    if headers != list(columns):
        return []
    rows: list[dict[str, str]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        values = [
            str(ws.cell(r, c).value or "").strip()
            for c in range(1, len(columns) + 1)
        ]
        if not any(values):
            continue
        rows.append({columns[i]: values[i] for i in range(len(columns))})
    return rows


def collect_bundle(case_id: str) -> dict[str, Any]:
    from tools.authorization_store import bundle_for_excel  # noqa: WPS433

    return bundle_for_excel(case_id)


def append_auth_sheets(
    wb: Workbook,
    case_id: str,
    meta: dict | None = None,
    bundle: dict[str, Any] | None = None,
) -> None:
    """Gắn sheet AUTH_* vào workbook (trước wb.save)."""
    data = bundle if bundle is not None else collect_bundle(case_id)
    _write_meta_sheet(wb.create_sheet(AUTH_META_SHEET), case_id, meta)
    _write_data_sheet(
        wb.create_sheet(AUTH_AGENTS_SHEET),
        "Profile người làm thủ tục (ủy quyền)",
        AGENT_COLUMNS,
        data.get("agents") or [],
    )
    _write_data_sheet(
        wb.create_sheet(AUTH_SETS_SHEET),
        "Hồ sơ ủy quyền theo case",
        SET_COLUMNS,
        data.get("sets") or [],
    )
    _write_data_sheet(
        wb.create_sheet(AUTH_DOCUMENTS_SHEET),
        "Văn bản đã sinh (metadata)",
        DOCUMENT_COLUMNS,
        data.get("documents") or [],
    )
    _write_data_sheet(
        wb.create_sheet(AUTH_ATTACHMENTS_SHEET),
        "Tệp đính kèm (metadata, không nhúng binary)",
        ATTACHMENT_COLUMNS,
        data.get("attachments") or [],
    )
    _write_data_sheet(
        wb.create_sheet(AUTH_AUDIT_SHEET),
        "Lịch sử sự kiện (không PII)",
        AUDIT_COLUMNS,
        data.get("audit") or [],
    )


def has_auth_sheets(sheetnames: list[str]) -> bool:
    return AUTH_META_SHEET in sheetnames


def read_auth_bundle(wb, case_id: str) -> dict[str, Any]:
    """Đọc và validate sheet AUTH_*; không đụng draft."""
    names = wb.sheetnames
    if not has_auth_sheets(names):
        return {"ok": True, "present": False, "bundle": None, "errors": []}

    errors: list[dict[str, str]] = []
    for sheet in AUTH_SHEETS:
        if sheet not in names:
            errors.append({
                "code": "AUTH_MISSING_SHEET",
                "message": "Thiếu sheet %s trong khối AUTH" % sheet,
                "sheet": sheet,
            })
    if errors:
        return {"ok": False, "present": True, "bundle": None, "errors": errors}

    meta_ws = wb[AUTH_META_SHEET]
    file_case = _read_meta_case_id(meta_ws)
    if file_case and file_case != case_id:
        return {
            "ok": False,
            "present": True,
            "bundle": None,
            "errors": [{
                "code": "AUTH_WRONG_CASE",
                "message": "AUTH_META case_id=%r khác case đang mở %r"
                           % (file_case, case_id),
            }],
        }

    bundle = {
        "schema_version": AUTH_SCHEMA_VERSION,
        "case_id": case_id,
        "agents": _read_data_rows(wb[AUTH_AGENTS_SHEET], AGENT_COLUMNS),
        "sets": _read_data_rows(wb[AUTH_SETS_SHEET], SET_COLUMNS),
        "documents": _read_data_rows(wb[AUTH_DOCUMENTS_SHEET], DOCUMENT_COLUMNS),
        "attachments": _read_data_rows(
            wb[AUTH_ATTACHMENTS_SHEET], ATTACHMENT_COLUMNS),
        "audit": _read_data_rows(wb[AUTH_AUDIT_SHEET], AUDIT_COLUMNS),
    }

    seen_agents: set[str] = set()
    for idx, row in enumerate(bundle["agents"], start=1):
        for field in REQUIRED_AGENT:
            if not str(row.get(field) or "").strip():
                errors.append({
                    "code": "AUTH_MISSING_REQUIRED",
                    "message": "AUTH_AGENTS hàng %d thiếu %s" % (idx + 3, field),
                    "sheet": AUTH_AGENTS_SHEET,
                })
        aid = row.get("agent_id", "")
        if aid in seen_agents:
            errors.append({
                "code": "AUTH_DUPLICATE",
                "message": "agent_id trùng: %s" % aid,
                "sheet": AUTH_AGENTS_SHEET,
            })
        seen_agents.add(aid)

    seen_sets: set[str] = set()
    for idx, row in enumerate(bundle["sets"], start=1):
        for field in REQUIRED_SET:
            if not str(row.get(field) or "").strip():
                errors.append({
                    "code": "AUTH_MISSING_REQUIRED",
                    "message": "AUTH_SETS hàng %d thiếu %s" % (idx + 3, field),
                    "sheet": AUTH_SETS_SHEET,
                })
        if row.get("case_id") and row["case_id"] != case_id:
            errors.append({
                "code": "AUTH_WRONG_CASE",
                "message": "AUTH_SETS hàng %d case_id=%r"
                           % (idx + 3, row["case_id"]),
                "sheet": AUTH_SETS_SHEET,
            })
        auth_id = row.get("authorization_id", "")
        if auth_id in seen_sets:
            errors.append({
                "code": "AUTH_DUPLICATE",
                "message": "authorization_id trùng: %s" % auth_id,
                "sheet": AUTH_SETS_SHEET,
            })
        seen_sets.add(auth_id)
        if row.get("agent_id") and row["agent_id"] not in seen_agents:
            errors.append({
                "code": "AUTH_UNKNOWN_AGENT",
                "message": "AUTH_SETS hàng %d agent_id %r không có trong AUTH_AGENTS"
                           % (idx + 3, row["agent_id"]),
                "sheet": AUTH_SETS_SHEET,
            })

    if errors:
        return {"ok": False, "present": True, "bundle": bundle, "errors": errors}
    return {"ok": True, "present": True, "bundle": bundle, "errors": []}
