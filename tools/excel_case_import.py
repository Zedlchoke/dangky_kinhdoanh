# -*- coding: utf-8 -*-
"""Nạp workbook Excel CASE-LEVEL khách điền -> candidate draft v2 NGUYÊN TỰ
(wave E2 — PLAN_hoan_thien_pfat_hanh §2.2, §5-E2).

Đối tác của `tools/excel_case_export.py` (wave E1): workbook một file cho cả
case với sheet `__map` veryHidden làm hợp đồng vị trí ô + CHECKSUM.

TRÌNH TỰ BẮT BUỘC (verify contract TRƯỚC khi đọc answers):

  0. mở file (raw + values)              -> BAD_FILE
  1. sheet `__map` + marker + metadata   -> BAD_HEADER
  2. contract / map_version              -> BAD_VERSION
  3. case_id == manifest.case_id         -> WRONG_CASE
  4. manifest_hash == manifest           -> BAD_VERSION
  5. CHECKSUM sha256(map) tái tính       -> CHECKSUM
  6. sheet declared tồn tại + không có sheet lạ -> BAD_HEADER
  7. map row khớp manifest: question lạ / key lạ / form không thuộc case /
     when-transform-row_group lệch       -> BAD_HEADER
  8. công thức trong ô điền (`=cmd|…`)   -> FORMULA;
     ô mirror bị gõ đè                   -> MIRROR_CELL
  9. ô lạ ngoài vùng điền                -> UNKNOWN_CELL;
     bảng tràn số hàng                   -> TOO_MANY_ROWS
 10. giá trị ô theo type (select/mark/date/money/integer/text) -> BAD_*;
     ID/text bị Excel đổi số/scientific  -> BAD_ID_TEXT
 11. resolve `when`, dựng store, required trống -> MISSING_REQUIRED
 11b. checkedForms theo lựa chọn (wave E4): nhóm `forms_oneOf` -> form được
     chọn; form tùy chọn ngoài oneOf -> mark `optionalForm.<fid>`=X;
     nhóm không chọn / không còn form nào -> MISSING_REQUIRED; ô của form
     KHÔNG được chọn mà có dữ liệu -> STRAY_CELL (kèm ô + câu hỏi).
     Candidate vẫn nguyên tử; `activeDocId` = form bắt buộc đầu hoặc form
     được chọn.
 12. bảng lặp 0/1/N hàng; hai hàng y hệt -> DUP_ROW

CÓ MỘT lỗi bất kỳ -> `{"ok": False, "errors": [...], "candidate": None}` —
KHÔNG BAO GIỜ trả data một phần; draft đang mở không bị ghi đè. UI chỉ thay
draft bằng candidate khi ok=True VÀ người dùng xác nhận (có nút hủy).

Candidate draft v2 (hình dạng khớp `draftPayload()` UI + `FD.migrateDraft`):
`{v, caseId, checkedForms, activeDocId, autoSum, viewState, savedAt, store:
{values, checks, loops}, source}`. Câu mark scalar key `chk_*` -> `checks`
(bool — đúng cách runtime lưu checkbox); mark trong bảng lặp giữ 'X' trong
hàng; các câu còn lại -> `values`.

Ô mirror (shared key): KHÔNG đọc giá trị — dữ liệu thật ở ô canonical ở
`Thông_tin_chung` (cấm last-write-wins); mirror chỉ được kiểm việc không bị
gõ đè.

Thuần function: chỉ đọc file truyền vào, KHÔNG đụng draft/store/IO khác,
KHÔNG in log dữ liệu khách. Audit metadata (file hash, schema/manifest
version, timestamp, số đếm) trả về ở `audit` — bridge quyết định ghi đâu.

Chuẩn hoá giá trị tái dùng các hàm THUẦN của `tools.excel_import` (X3) để
hai lane không lệch contract parse ngày/tiền/select/guard mã định danh.

Chạy suite đối tác:  python3 -X utf8 -m tools.test_excel_case_negative
"""
from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from tools.excel_case_export import (
    CASE_MAP_VERSION, COMMON_SHEET, CONTRACT, DOCS_SHEET, GUIDE_LINES,
    GUIDE_SHEET, LISTS_SHEET, MAP_COLUMNS, MAP_MARKER, MAP_SHEET,
    ONEOF_KEY_PREFIX, OPTIONAL_FORM_KEY_PREFIX,
    ExcelCaseExportError, is_review_approved, is_selector_question,
)
from tools.excel_authorization_sheets import AUTH_SHEETS
from tools.excel_case_export import (
    _canonical_json,   # cùng contract hash với exporter (một nguồn)
    _validate_manifest,
)
# Chuẩn hoá giá trị: các hàm THUẦN của lane X3 — tái dùng để hai lane không
# lệch nhau về parse (đều pure function, cùng gói `tools`).
from tools.excel_import import (
    _cell_text, _id_text_guard, _json_canon, _match_option, _norm_label,
    _parse_cell_address, _parse_date, _parse_number, _split_loop_key,
    _targets_of, _when_match,
)
from tools.questionnaire.schema import DATE_NOTE, EXAMPLE_PREFIX, SELECT_NOTE

__all__ = [
    "BAD_FILE", "BAD_HEADER", "BAD_VERSION", "WRONG_CASE", "CHECKSUM",
    "UNKNOWN_CELL", "TOO_MANY_ROWS", "DUP_ROW", "FORMULA", "MIRROR_CELL",
    "MISSING_REQUIRED", "BAD_SELECT", "BAD_DATE", "BAD_NUMBER",
    "BAD_ID_TEXT", "STRAY_CELL", "import_case_workbook",
    "manifest_has_select",
]

# Mã lỗi — contract với bridge/UI (toast + panel "còn thiếu")
BAD_FILE = "BAD_FILE"
BAD_HEADER = "BAD_HEADER"
BAD_VERSION = "BAD_VERSION"
WRONG_CASE = "WRONG_CASE"
CHECKSUM = "CHECKSUM"
UNKNOWN_CELL = "UNKNOWN_CELL"
TOO_MANY_ROWS = "TOO_MANY_ROWS"
DUP_ROW = "DUP_ROW"
FORMULA = "FORMULA"
MIRROR_CELL = "MIRROR_CELL"
MISSING_REQUIRED = "MISSING_REQUIRED"
BAD_SELECT = "BAD_SELECT"
BAD_DATE = "BAD_DATE"
BAD_NUMBER = "BAD_NUMBER"
BAD_ID_TEXT = "BAD_ID_TEXT"
# Wave E4 (oneOf/form tùy chọn): ô có dữ liệu nhưng mẫu đó KHÔNG được chọn
STRAY_CELL = "STRAY_CELL"

_INT_COLS = {"row_last"}            # cột map kiểu số nguyên khi hash
_MAYBE_INT_COLS = {"version"}       # int hoặc "" (version câu hỏi)
_MIRROR_NOTE = ("Đã điền ở sheet %s — ô này tự phản chiếu, không "
                "điền lại" % COMMON_SHEET)
_COMMON_EMPTY_NOTE = ("Case này không có trường dùng chung nào giữa các mẫu "
                      "— điền trực tiếp ở sheet các mẫu bên dưới.")
_COMMON_SHARED_NOTE = (" — Các ô dưới đây dùng chung cho nhiều mẫu, chỉ điền "
                       "MỘT lần")

# Các sheet cấu trúc của workbook — không có ô map nhưng LUÔN được phép tồn tại
_STRUCTURAL_SHEETS = (GUIDE_SHEET, COMMON_SHEET, DOCS_SHEET) + tuple(AUTH_SHEETS)


def _err(code: str, message: str, cell: str | None = None,
         question_id: str | None = None, form_id: str | None = None,
         sheet: str | None = None) -> dict:
    """Một phần tử `errors[]` — hình dạng cố định cho bridge/UI."""
    return {
        "code": code,
        "message": message,
        "cell": cell,
        "question_id": question_id,
        "form_id": form_id,
        "sheet": sheet,
    }


def _fail(errors: list[dict], audit: dict | None = None) -> dict:
    """Kết quả lỗi: candidate luôn None (nguyên tắc nguyên tử)."""
    return {"ok": False, "errors": errors, "candidate": None,
            "audit": audit or {}}


@dataclass
class _Entry:
    """Một dòng của sheet __map = MỘT ô điền (đã đối chiếu manifest)."""
    form_id: str
    sheet: str
    question_id: str
    key: str
    address: str
    row: int
    col: int
    row_last: int
    layout: str
    when: dict | None = None
    transform: str | None = None
    row_group: str | None = None
    repeatable: bool = False
    required: bool = False
    mirror: bool = False
    canonical_sheet: str = ""
    canonical_cell: str = ""
    value: object = field(default=None)
    filled: bool = False
    resolved: bool = False
    value_error: bool = False


# ---------------------------------------------------------------------------
# Đọc sheet __map (phòng thủ)
# ---------------------------------------------------------------------------

def _find_header(rows: list[list]) -> tuple[int | None, dict[str, int]]:
    """(chỉ số hàng header, map tên cột -> vị trí) — header có cột question_id."""
    for idx, row in enumerate(rows):
        texts = [_cell_text(v) for v in row]
        if "question_id" in [t.lower() for t in texts if t]:
            col_of: dict[str, int] = {}
            for pos, text in enumerate(texts):
                if text:
                    col_of.setdefault(text.lower(), pos)
            return idx, col_of
    return None, {}


def _read_map(wb) -> tuple[dict[str, str], list[_Entry], list[dict]]:
    """Đọc phòng thủ `__map`. Trả (metadata, entries, lỗi khung)."""
    if MAP_SHEET not in wb.sheetnames:
        return {}, [], [_err(
            BAD_HEADER,
            "File thiếu sheet %r — không phải workbook app xuất" % MAP_SHEET)]
    ws = wb[MAP_SHEET]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows or _cell_text(rows[0][0] if rows[0] else None) != MAP_MARKER:
        return {}, [], [_err(
            BAD_HEADER, "Sheet %r thiếu marker %r — file lạ"
            % (MAP_SHEET, MAP_MARKER))]

    meta: dict[str, str] = {}
    for value in rows[0][1:]:
        text = _cell_text(value)
        key, sep, val = text.partition("=")
        if sep and key.strip():
            meta.setdefault(key.strip(), val.strip())

    header_idx, col_of = _find_header(rows)
    if header_idx is None:
        return meta, [], [_err(
            BAD_HEADER, "Sheet %r không có hàng header (cần cột %s)"
            % (MAP_SHEET, " | ".join(MAP_COLUMNS)))]

    missing = [name for name in ("form_id", "sheet", "question_id", "key",
                                 "cell") if name not in col_of]
    if missing:
        return meta, [], [_err(
            BAD_HEADER, "Sheet %r thiếu cột bắt buộc: %s"
            % (MAP_SHEET, ", ".join(missing)))]

    def opt(col_name: str, texts: list[str]) -> str:
        col = col_of.get(col_name)
        return "" if col is None or col >= len(texts) else texts[col]

    entries: list[_Entry] = []
    errors: list[dict] = []
    seen: set[tuple] = set()
    for line_no, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        texts = [_cell_text(v) for v in row]
        qid = opt("question_id", texts)
        if not qid and not any(texts):
            continue  # hàng trắng — bỏ qua
        if not qid:
            errors.append(_err(BAD_HEADER,
                               "Sheet %r hàng %d thiếu question_id"
                               % (MAP_SHEET, line_no)))
            continue
        form_id = opt("form_id", texts)
        sheet = opt("sheet", texts)
        key = opt("key", texts)
        address = opt("cell", texts)
        if not (form_id and sheet and key and address):
            errors.append(_err(
                BAD_HEADER, "Sheet %r hàng %d (%s): thiếu form_id/sheet/key/cell"
                % (MAP_SHEET, line_no, qid), question_id=qid))
            continue
        coord = _parse_cell_address(address)
        if coord is None:
            errors.append(_err(
                BAD_HEADER, "Sheet %r hàng %d (%s): cell %r không hợp lệ"
                % (MAP_SHEET, line_no, qid, address), cell=address,
                question_id=qid, sheet=sheet))
            continue

        when: dict | None = None
        when_raw = opt("when", texts)
        if when_raw:
            try:
                parsed = json.loads(when_raw)
            except ValueError:
                errors.append(_err(
                    BAD_HEADER, "Sheet %r hàng %d (%s): when không phải JSON"
                    % (MAP_SHEET, line_no, qid), cell=address,
                    question_id=qid, sheet=sheet))
                continue
            if not isinstance(parsed, dict) or not all(
                    isinstance(v, list) for v in parsed.values()):
                errors.append(_err(
                    BAD_HEADER,
                    "Sheet %r hàng %d (%s): when phải là {key: [giá trị]}"
                    % (MAP_SHEET, line_no, qid), cell=address,
                    question_id=qid, sheet=sheet))
                continue
            when = parsed

        transform = opt("transform", texts) or None
        row_last_raw = opt("row_last", texts)
        try:
            row_last = int(row_last_raw) if row_last_raw else coord[0]
        except ValueError:
            errors.append(_err(
                BAD_HEADER, "Sheet %r hàng %d (%s): row_last %r không phải số"
                % (MAP_SHEET, line_no, qid, row_last_raw), cell=address,
                question_id=qid, sheet=sheet))
            continue

        # Khai báo y hệt 2 lần = map hỏng (khác when/transform thì hợp lệ —
        # câu date tách 3 key cùng một ô).
        fingerprint = (form_id, sheet, qid, key, address,
                       transform or "", opt("row_group", texts))
        if fingerprint in seen:
            errors.append(_err(
                BAD_HEADER, "Sheet %r hàng %d (%s): ô %s khai báo trùng lặp"
                % (MAP_SHEET, line_no, qid, address), cell=address,
                question_id=qid, sheet=sheet))
            continue
        seen.add(fingerprint)

        entries.append(_Entry(
            form_id=form_id, sheet=sheet, question_id=qid, key=key,
            address=address, row=coord[0], col=coord[1], row_last=row_last,
            layout=opt("layout", texts) or "single",
            when=when, transform=transform,
            row_group=opt("row_group", texts) or None,
            repeatable=opt("repeatable", texts) == "yes",
            required=opt("required", texts) == "yes",
            mirror=opt("mirror", texts) == "yes",
            canonical_sheet=opt("canonical_sheet", texts),
            canonical_cell=opt("canonical_cell", texts)))
    return meta, entries, errors


def _map_meta(wb) -> dict[str, str]:
    """Metadata `key=value` ở hàng 1 của __map (đọc lại riêng cho checksum)."""
    ws = wb[MAP_SHEET]
    for row in ws.iter_rows(values_only=True, max_row=1):
        meta: dict[str, str] = {}
        for value in row[1:]:
            text = _cell_text(value)
            key, sep, val = text.partition("=")
            if sep and key.strip():
                meta.setdefault(key.strip(), val.strip())
        return meta
    return {}


def _read_map_records(wb) -> list[dict]:
    """Đọc lại __map thành list dict ĐÚNG hình dạng exporter đã hash.

    Chuẩn hoá như lane X3 ghi chú: ô chuỗi rỗng openpyxl trả None -> "";
    `row_last` về int; `version` giữ int nếu là int. Mọi lệch so với lúc ghi
    làm checksum fail.
    """
    ws = wb[MAP_SHEET]
    rows_raw = [list(r) for r in ws.iter_rows(values_only=True)]
    header_idx, col_of = _find_header(rows_raw)
    records: list[dict] = []
    if header_idx is None:
        return records
    for row in rows_raw[header_idx + 1:]:
        if not any(v is not None and str(v) != "" for v in row):
            continue
        record: dict = {}
        for name in MAP_COLUMNS:
            col = col_of.get(name)
            raw = row[col] if col is not None and col < len(row) else None
            if name in _INT_COLS:
                record[name] = int(raw) if raw not in (None, "") else ""
            elif name in _MAYBE_INT_COLS:
                record[name] = raw if isinstance(raw, int) else \
                    ("" if raw is None else str(raw))
            else:
                record[name] = "" if raw is None else str(raw)
        records.append(record)
    return records


def _checksum_of(meta: dict, rows: list[dict]) -> str:
    """sha256 payload canonical {meta (trừ checksum), rows} — như exporter."""
    body = {k: v for k, v in meta.items() if k != "checksum"}
    return hashlib.sha256(_canonical_json({
        "meta": body, "rows": rows,
    }).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Static texts + quét ô lạ (heuristic X3, siết theo CỘT cho case-level)
# ---------------------------------------------------------------------------

def _static_texts(manifest: dict, meta_line: str) -> set[str]:
    """Chuỗi tĩnh app tự render trên các sheet có ô điền — không nhầm là ô lạ."""
    out = {_norm_label(t) for t in (
        [meta_line, meta_line + _COMMON_SHARED_NOTE, SELECT_NOTE, DATE_NOTE,
         _MIRROR_NOTE, _COMMON_EMPTY_NOTE, "Thông tin dùng chung cho cả hồ sơ"]
        + list(GUIDE_LINES))}
    for form in manifest.get("forms", []):
        for name in ("title", "ten_mau"):
            if form.get(name):
                out.add(_norm_label(str(form[name])))
        for q in form.get("questions", []):
            for name in ("group", "text", "example"):
                if q.get(name):
                    out.add(_norm_label(str(q[name])))
                    if name == "example":
                        out.add(_norm_label(EXAMPLE_PREFIX + str(q[name])))
            for opt in q.get("options") or []:
                if isinstance(opt, dict):
                    for name in ("label", "value"):
                        if opt.get(name):
                            out.add(_norm_label(str(opt[name])))
    out.discard("")
    return out


def _find_stray_cells(ws, entries: list[_Entry], static: set[str],
                      sheet: str) -> list[dict]:
    """Ô có dữ liệu ngoài vùng điền trên sheet có ô map.

    Siết hơn X3: chỉ cần NẰM TRONG CỘT có ô map là bị kiểm — bắt được khách
    điền TRÀN bảng (hàng mới dưới row_last). Ô nằm dưới MỘT bảng nào đó trong
    cùng cột -> TOO_MANY_ROWS (bảng gần nhất), ô lạ chỗ khác -> UNKNOWN_CELL.
    """
    mapped = {(e.row, e.col) for e in entries}
    mapped_cols = {c for _, c in mapped}
    # khối bảng theo cột: (hàng đầu, hàng cuối); entry bảng N hàng là N cặp
    # (row_i, row_last) cùng row_last — đầu bảng = min(row) của nhóm đó.
    pairs_by_col: dict[int, set[tuple[int, int]]] = {}
    for e in entries:
        if e.layout == "table":
            pairs_by_col.setdefault(e.col, set()).add((e.row, e.row_last))
    blocks_by_col: dict[int, list[tuple[int, int]]] = {}
    for col, pairs in pairs_by_col.items():
        grouped: dict[int, list[int]] = {}
        for row, last in pairs:
            grouped.setdefault(last, []).append(row)
        blocks_by_col[col] = sorted((min(rows), last)
                                    for last, rows in grouped.items())
    errors: list[dict] = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.column not in mapped_cols:
                continue
            if (cell.row, cell.column) in mapped:
                continue
            text = _cell_text(cell.value)
            if not text or _norm_label(text) in static:
                continue
            # bảng gần nhất phía trên trong cùng cột (nếu có)
            block = None
            for first, last in blocks_by_col.get(cell.column, []):
                if cell.row > last:
                    block = (first, last)
            if block is not None:
                n_rows = block[1] - block[0] + 1
                errors.append(_err(
                    TOO_MANY_ROWS,
                    "Ô %s: bảng chỉ có %d hàng điền — cần thêm dòng hãy yêu "
                    "cầu app xuất lại phiếu, không tự chèn hàng"
                    % (cell.coordinate, n_rows),
                    cell=cell.coordinate, sheet=sheet))
            else:
                errors.append(_err(
                    UNKNOWN_CELL,
                    "Ô %s có dữ liệu %r nhưng không khớp vùng điền nào"
                    % (cell.coordinate, text),
                    cell=cell.coordinate, sheet=sheet))
    return errors


# ---------------------------------------------------------------------------
# Helpers nhỏ
# ---------------------------------------------------------------------------

def manifest_has_select(manifest: dict) -> bool:
    """Case có câu select (-> workbook phải có sheet __lists)."""
    for form in manifest.get("forms", []):
        for q in form.get("questions", []):
            if q.get("type") == "select":
                return True
    return False


def _is_check_key(key: str) -> bool:
    """Key checkbox scalar của runtime (`chk_*` — store.checks lưu bool)."""
    return str(key).startswith("chk_")


def _sheet_of_loop(entries: list[_Entry], loop_name: str) -> str | None:
    """Sheet đầu tiên chứa bảng lặp `loop_name` (cho lỗi DUP_ROW)."""
    for e in entries:
        if e.row_group and (e.row_group == loop_name or
                            e.key.startswith(loop_name + ".")):
            return e.sheet
    return None


def _rebuild_meta_line(manifest: dict, meta: dict) -> str:
    """Dựng lại dòng 'Hồ sơ: …' của exporter để bộ lọc chuỗi tĩnh khớp."""
    hash8 = str(meta.get("manifest_hash") or manifest.get("manifest_hash")
                or "")[:8]
    return "Hồ sơ: %s (%s) — Mã phiếu: %s — Ngày lập: %s" % (
        manifest.get("case_title"), manifest.get("case_id"), hash8,
        meta.get("today") or "")


def _normalize_value(question: dict, entry: "_Entry", raw):
    """Giá trị ô -> (giá trị data, lỗi). Ô trống -> (None, None).

    Cùng contract parse với X3 (tái dùng hàm thuần của excel_import); lỗi có
    kèm form_id/sheet cho panel case-level.
    """
    text = _cell_text(raw)
    if not text:
        return None, None
    qtype = str(question.get("type") or "")
    try:
        if qtype in ("select", "mark"):
            value = _match_option(question, text)
            if value is None:
                raise ValueError(
                    "giá trị %r không nằm trong options (%s)" % (
                        text, " / ".join(
                            str(o.get("label", o.get("value", "")))
                            for o in question.get("options") or [])))
            return value, None
        if qtype == "date":
            day, month, year = _parse_date(text)
            if entry.transform == "date_day":
                return "%02d" % day, None
            if entry.transform == "date_month":
                return "%02d" % month, None
            if entry.transform == "date_year":
                return "%04d" % year, None
            return "%02d/%02d/%04d" % (day, month, year), None
        if qtype in ("money", "integer"):
            return _parse_number(text, qtype), None
        if qtype in ("text", "tel", "email"):
            bad = _id_text_guard(entry, raw, text)
            if bad is not None:
                return None, bad
        return text, None
    except ValueError as exc:
        code = {"select": BAD_SELECT, "mark": BAD_SELECT,
                "date": BAD_DATE, "money": BAD_NUMBER,
                "integer": BAD_NUMBER}.get(qtype, BAD_HEADER)
        return None, _err(code, "Ô %s (%s): %s" % (entry.address,
                                                   entry.question_id, exc),
                          cell=entry.address, question_id=entry.question_id,
                          form_id=entry.form_id, sheet=entry.sheet)


# ---------------------------------------------------------------------------
# API chính
# ---------------------------------------------------------------------------

def import_case_workbook(path: "str | Path", manifest: dict,
                         imported_at: str | None = None) -> dict:
    """Nạp workbook case-level -> candidate draft v2. Thuần function, fail-loud.

    `manifest`: dict case-level wave E0 (app đọc từ
    app/data/case_questionnaires/<case_id>.json). `imported_at`: chuỗi thời
    điểm (bridge truyền) — chỉ vào `audit`/candidate.source, không bắt buộc.

    Trả {"ok", "errors", "candidate", "audit"}; `audit` gồm file hash +
    versions + số đếm (KHÔNG có dữ liệu khách).
    """
    audit: dict = {}
    try:
        _validate_manifest(manifest)
    except ExcelCaseExportError as exc:
        return _fail([_err(BAD_HEADER, "manifest không hợp lệ: %s" % exc)])
    audit["case_id"] = manifest.get("case_id")

    # -- 0. Mở file -----------------------------------------------------------
    p = Path(path)
    try:
        audit["file_sha256"] = hashlib.sha256(p.read_bytes()).hexdigest()
    except OSError as exc:
        return _fail([_err(BAD_FILE,
                           "Không đọc được file: %s" % type(exc).__name__)],
                     audit)
    try:
        wb_raw = load_workbook(filename=str(p), data_only=False)
        wb_vals = load_workbook(filename=str(p), data_only=True)
    except Exception as exc:  # noqa: BLE001 — sai đuôi/file hỏng/không phải xlsx
        return _fail([_err(BAD_FILE,
                           "Không mở được file .xlsx: %s" % type(exc).__name__)],
                     audit)

    # -- 1. Sheet __map ---------------------------------------------------------
    meta, entries, errors = _read_map(wb_raw)
    if errors:
        return _fail(errors, audit)
    if not entries:
        return _fail([_err(BAD_HEADER, "Sheet %r không khai báo ô điền nào"
                           % MAP_SHEET)], audit)
    # [wave-compat 2026-09-10] Sheet_state: `visible` = khả năng cao file đã bị
    # can thiệp cấu trúc -> chặn. `hidden` (LibreOffice Calc hạ veryHidden xuống
    # khi khách re-save — đo được ở gate Calc, test_excel_calc_roundtrip) thì
    # CHẤP NHẬN: mọi rào chống sửa trộm (checksum/presence/mirror) vẫn bắt buộc
    # ở các bước sau, trạng thái ẩn chỉ là trình bày -> ghi audit để UI hiển thị.
    _map_state = wb_raw[MAP_SHEET].sheet_state
    if _map_state == "visible":
        return _fail([_err(BAD_HEADER,
                           "Sheet %r đang hiện (visible) — file không đúng "
                           "định dạng phiếu (map phải ẩn)" % MAP_SHEET)], audit)
    audit["map_sheet_state"] = _map_state

    # -- 2/3/4. Contract + case + hash ------------------------------------------
    if meta.get("contract") != CONTRACT or \
            meta.get("map_version") != str(CASE_MAP_VERSION):
        return _fail([_err(
            BAD_VERSION,
            "File khai báo contract=%r map_version=%r, app cần %r/%r"
            % (meta.get("contract"), meta.get("map_version"), CONTRACT,
               CASE_MAP_VERSION))], audit)
    if meta.get("case_id") != manifest.get("case_id"):
        return _fail([_err(
            WRONG_CASE,
            "Workbook của case %r, đang mở case %r — nạp nhầm hồ sơ"
            % (meta.get("case_id"), manifest.get("case_id")))], audit)
    if meta.get("manifest_hash") != manifest.get("manifest_hash"):
        return _fail([_err(
            BAD_VERSION,
            "Workbook sinh từ manifest khác (hash %s…) khác manifest hiện "
            "tại (%s…) — hãy xuất lại phiếu" % (
                str(meta.get("manifest_hash"))[:8],
                str(manifest.get("manifest_hash"))[:8]))], audit)

    # -- 5. CHECKSUM --------------------------------------------------------------
    records = _read_map_records(wb_raw)
    stored = meta.get("checksum", "")
    computed = _checksum_of(meta, records) if records else ""
    if not stored or stored != computed:
        return _fail([_err(
            CHECKSUM,
            "Checksum khung phiếu lệch (file %s… / tái tính %s…) — file bị "
            "sửa sau khi xuất" % (stored[:8] or "?", computed[:8] or "?"))],
            audit)

    # -- 6. Sheet declared tồn tại + không có sheet lạ ----------------------------
    declared_sheets = sorted({e.sheet for e in entries})
    allowed = set(declared_sheets) | set(_STRUCTURAL_SHEETS) | \
        {MAP_SHEET, LISTS_SHEET}
    for name in declared_sheets:
        if name not in wb_raw.sheetnames:
            return _fail([_err(
                BAD_HEADER, "Thiếu sheet %r đã khai báo trong phiếu" % name)],
                audit)
    extra = [n for n in wb_raw.sheetnames if n not in allowed]
    if extra:
        return _fail([_err(
            BAD_HEADER, "File có sheet lạ không nằm trong phiếu: %s"
            % ", ".join(extra))], audit)

    # -- 7. Map <-> manifest (khung phải khớp từng ô) ------------------------------
    questions: dict[str, tuple[str, dict]] = {}   # qid -> (form_id, question)
    form_ids = {f["form_id"] for f in manifest.get("forms", [])}
    for form in manifest.get("forms", []):
        if not is_review_approved(form.get("content_review")):
            continue
        for q in form.get("questions", []):
            questions[q["question_id"]] = (form["form_id"], q)

    read_entries: list[_Entry] = []       # ô canonical (đọc dữ liệu)
    mirror_entries: list[_Entry] = []     # ô mirror (chỉ kiểm không bị gõ đè)
    by_qid: dict[str, list[_Entry]] = {}
    for entry in entries:
        if entry.question_id not in questions:
            return _fail([_err(
                BAD_HEADER,
                "Ô %s: question_id %r lạ — không thuộc case %r"
                % (entry.address, entry.question_id, manifest.get("case_id")),
                cell=entry.address, question_id=entry.question_id,
                sheet=entry.sheet)], audit)
        if entry.form_id not in form_ids:
            return _fail([_err(
                BAD_HEADER, "Ô %s: form %r không thuộc case"
                % (entry.address, entry.form_id), cell=entry.address,
                question_id=entry.question_id, form_id=entry.form_id)], audit)
        _owner_fid, question = questions[entry.question_id]
        targets = _targets_of(question)
        target = next((t for t in targets if t["key"] == entry.key), None)
        if target is None:
            return _fail([_err(
                BAD_HEADER, "Ô %s (%s): key %r không thuộc câu hỏi"
                % (entry.address, entry.question_id, entry.key),
                cell=entry.address, question_id=entry.question_id,
                sheet=entry.sheet)], audit)
        if entry.when is not None and \
                _json_canon(entry.when) != _json_canon(target["when"] or {}):
            return _fail([_err(
                BAD_HEADER, "Ô %s (%s): when trong map lệch questionnaire"
                % (entry.address, entry.question_id),
                cell=entry.address, question_id=entry.question_id,
                sheet=entry.sheet)], audit)
        entry.when = target["when"]
        if entry.transform is not None and \
                entry.transform != (target["transform"] or None):
            return _fail([_err(
                BAD_HEADER, "Ô %s (%s): transform trong map lệch questionnaire"
                % (entry.address, entry.question_id),
                cell=entry.address, question_id=entry.question_id,
                sheet=entry.sheet)], audit)
        entry.transform = target["transform"] or None
        if (entry.row_group or None) != (question.get("row_group") or None):
            return _fail([_err(
                BAD_HEADER, "Ô %s (%s): row_group lệch questionnaire"
                % (entry.address, entry.question_id),
                cell=entry.address, question_id=entry.question_id,
                sheet=entry.sheet)], audit)
        entry.row_group = question.get("row_group") or None
        if entry.mirror:
            if not entry.canonical_sheet or not entry.canonical_cell:
                return _fail([_err(
                    BAD_HEADER, "Ô %s (%s): mirror thiếu ô canonical"
                    % (entry.address, entry.question_id),
                    cell=entry.address, question_id=entry.question_id,
                    sheet=entry.sheet)], audit)
            mirror_entries.append(entry)
        else:
            read_entries.append(entry)
            by_qid.setdefault(entry.question_id, []).append(entry)

    # Câu KHÔNG lặp: mỗi key đúng một ô, mỗi key đúng một transform (X3)
    for qid, q_entries in by_qid.items():
        if questions[qid][1].get("row_group"):
            continue
        for key in {e.key for e in q_entries}:
            cells = {e.address for e in q_entries if e.key == key}
            transforms = {e.transform for e in q_entries if e.key == key}
            if len(cells) > 1 or len(transforms) > 1:
                return _fail([_err(
                    BAD_HEADER, "Câu %s (không lặp) key %r có %d ô / %d "
                    "transform — khung hỏng" % (qid, key, len(cells),
                                                len(transforms)))], audit)

    # MỘT ô chỉ phục vụ MỘT câu hỏi (tính cả mirror — toàn bộ map)
    cell_owner: dict[tuple[str, str], str] = {}
    for entry in entries:
        owner_key = (entry.sheet, entry.address)
        if cell_owner.setdefault(owner_key, entry.question_id) != \
                entry.question_id:
            return _fail([_err(
                BAD_HEADER,
                "Ô %s!%s phục vụ hai câu hỏi (%s và %s) — khung hỏng"
                % (entry.sheet, entry.address,
                   cell_owner[owner_key], entry.question_id),
                cell=entry.address, sheet=entry.sheet)], audit)

    # -- 8. Công thức trong ô điền / mirror bị gõ đè --------------------------------
    for entry in read_entries:
        raw = wb_raw[entry.sheet].cell(row=entry.row, column=entry.col).value
        if isinstance(raw, str) and raw.startswith("="):
            errors.append(_err(
                FORMULA,
                "Ô %s (%s) chứa CÔNG THỨC %r — chỉ gõ dữ liệu thô, không gõ "
                "công thức" % (entry.address, entry.question_id, raw[:40]),
                cell=entry.address, question_id=entry.question_id,
                form_id=entry.form_id, sheet=entry.sheet))
            entry.value_error = True
    for entry in mirror_entries:
        raw = wb_raw[entry.sheet].cell(row=entry.row, column=entry.col).value
        if raw is not None and not (isinstance(raw, str)
                                    and raw.startswith("=")):
            errors.append(_err(
                MIRROR_CELL,
                "Ô %s (%s) là ô tự phản chiếu từ %s — không được gõ vào "
                "(dữ liệu thật ở ô canonical)" % (
                    entry.address, entry.question_id, entry.canonical_sheet),
                cell=entry.address, question_id=entry.question_id,
                form_id=entry.form_id, sheet=entry.sheet))

    # -- 9. Ô lạ / bảng tràn hàng -----------------------------------------------------
    static = _static_texts(manifest, _rebuild_meta_line(manifest, meta))
    for sheet in declared_sheets:
        sheet_entries = [e for e in entries if e.sheet == sheet]
        stray = _find_stray_cells(wb_vals[sheet], sheet_entries, static, sheet)
        if stray:
            errors.extend(stray)
            return _fail(errors, audit)   # structural: giữ cả lỗi formula ở trên
    if errors:
        return _fail(errors, audit)

    # -- 10. Đọc + chuẩn hoá giá trị (gom đủ lỗi, không dừng sớm) ----------------------
    for qid, q_entries in by_qid.items():
        for entry in q_entries:
            if entry.value_error:
                continue  # ô formula đã lỗi — không đọc nữa
            entry.value, issue = _normalize_value(
                questions[qid][1], entry,
                wb_vals[entry.sheet].cell(row=entry.row,
                                          column=entry.col).value)
            entry.filled = entry.value is not None
            if issue:
                entry.value_error = True
                errors.append(issue)

    # answers cho `when`: câu một data_key đã điền (toàn case — canonical)
    answers: dict[str, object] = {}
    for qid, q_entries in by_qid.items():
        if not isinstance(questions[qid][1].get("data_key"), str):
            continue
        for entry in q_entries:
            if entry.filled:
                answers[entry.key] = entry.value

    # -- 11. Resolve `when` + dựng store -----------------------------------------------
    values: dict[str, object] = {}
    checks: dict[str, bool] = {}
    loops: dict[str, dict[int, dict[str, object]]] = {}
    resolved_by_qid: dict[str, list[_Entry]] = {}
    for qid, q_entries in by_qid.items():
        question = questions[qid][1]
        active = {t["key"] for t in _targets_of(question)
                  if _when_match(t["when"], answers)}
        groups: dict[str, list[_Entry]] = {}
        for entry in q_entries:
            groups.setdefault(entry.address, []).append(entry)
        kept: list[_Entry] = []
        for address, group in groups.items():
            hits = [e for e in group if e.key in active]
            if not hits:
                if any(e.filled for e in group):
                    errors.append(_err(
                        BAD_HEADER,
                        "Ô %s (%s): có dữ liệu nhưng không target nào khớp "
                        "điều kiện when — phiếu lệch questionnaire"
                        % (address, qid), cell=address, question_id=qid,
                        sheet=group[0].sheet))
                continue
            for entry in hits:
                entry.resolved = True
                kept.append(entry)
                if not entry.filled:
                    continue
                if entry.row_group:
                    loop_name, tag = _split_loop_key(entry.key)
                    loops.setdefault(loop_name, {}).setdefault(
                        entry.row, {})[tag] = entry.value
                elif _is_check_key(entry.key):
                    checks[entry.key] = entry.value == "X"
                else:
                    values[entry.key] = entry.value
        resolved_by_qid[qid] = kept

    # -- 11b. oneOf + form tùy chọn -> checkedForms (wave E4) -----------------
    # Design (Lead chốt E4, owner ủy quyền):
    # - Mỗi nhóm `forms_oneOf` có MỘT câu select bắt buộc ở `Thông_tin_chung`
    #   (data_key `oneOf.<n>`, value = form_id được chọn). checkedForms chỉ
    #   chứa form ĐƯỢC CHỌN trong từng nhóm (+ các form bắt buộc ngoài nhóm).
    # - Form bat_buoc=false NGOÀI oneOf vào checkedForms khi câu mark
    #   `optionalForm.<form_id>` được đánh X.
    # - Ô của form KHÔNG thuộc checkedForms mà CÓ dữ liệu -> STRAY_CELL
    #   (kèm ô + câu hỏi); nhóm oneOf không chọn gì / không form nào còn lại
    #   -> MISSING_REQUIRED. Candidate vẫn nguyên tử.
    checked_set: set[str] = set()
    checked: list[str] = []
    approved_forms = [f for f in manifest.get("forms", [])
                      if is_review_approved(f.get("content_review"))]
    group_index_of: dict[str, int] = {}
    for gi, group in enumerate(manifest.get("forms_oneOf") or [], start=1):
        for fid in group.get("ids") or []:
            group_index_of[str(fid)] = gi
    selected_in_group: dict[int, str] = {}
    for key, val in values.items():
        if isinstance(key, str) and key.startswith(ONEOF_KEY_PREFIX):
            try:
                gi = int(key[len(ONEOF_KEY_PREFIX):].split(".")[0])
            except ValueError:
                continue
            if isinstance(val, str):
                selected_in_group.setdefault(gi, val)
    for form in approved_forms:
        fid = form["form_id"]
        gi = group_index_of.get(fid)
        if gi is not None:
            if selected_in_group.get(gi) == fid:
                checked.append(fid)
        elif form.get("bat_buoc"):
            checked.append(fid)
        elif values.get("%s%s" % (OPTIONAL_FORM_KEY_PREFIX, fid)) == "X":
            checked.append(fid)
    checked_set = set(checked)
    for gi in range(1, len(manifest.get("forms_oneOf") or []) + 1):
        if not selected_in_group.get(gi):
            errors.append(_err(
                MISSING_REQUIRED,
                "Chưa chọn trường hợp hồ sơ (câu «Trường hợp hồ sơ của bạn "
                "là trường hợp nào?» ở sheet %s) — chọn MỘT dòng" % COMMON_SHEET,
                sheet=COMMON_SHEET))
    if not checked and not any(e["code"] == MISSING_REQUIRED
                               for e in errors):
        errors.append(_err(
            MISSING_REQUIRED,
            "Không xác định được mẫu nào trong hồ sơ — hãy chọn trường hợp/"
            "đánh dấu ít nhất MỘT mẫu ở sheet %s" % COMMON_SHEET,
            sheet=COMMON_SHEET))

    # ô của form không được chọn mà có dữ liệu -> STRAY_CELL
    for entry in read_entries:
        if entry.form_id in checked_set or not entry.filled:
            continue
        _owner_fid, question = questions[entry.question_id]
        if is_selector_question(question) or entry.sheet == COMMON_SHEET:
            continue
        errors.append(_err(
            STRAY_CELL,
            "Ô %s (%s) có dữ liệu nhưng mẫu %s KHÔNG thuộc trường hợp bạn "
            "chọn — xóa dữ liệu ở ô này hoặc chọn lại trường hợp ở sheet %s"
            % (entry.address, entry.question_id, entry.form_id, COMMON_SHEET),
            cell=entry.address, question_id=entry.question_id,
            form_id=entry.form_id, sheet=entry.sheet))

    # -- 12. Required trống (đã scoping ở header trên) -----------------------------
    # Scoping (E4): bắt buộc áp cho (a) câu của form thuộc checkedForms, và
    # (b) câu canonical của form KHÁC nếu key đó được một form ĐƯỢC CHỌN yêu
    # cầu bắt buộc (shared key: ô canonical nằm ở câu của chủ — ví dụ case
    # chuyển đổi chọn M03 thì tên DN vẫn phải điền dù ô canonical thuộc câu
    # của M01). Câu có `when` không khớp vẫn được bỏ qua như cũ.
    keys_required_checked: set[str] = set()
    for qid in questions:
        form_id_q, question = questions[qid]
        if form_id_q in checked_set and question.get("required"):
            for t in _targets_of(question):
                keys_required_checked.add(t["key"])
    for qid, (form_id, question) in questions.items():
        if form_id not in checked_set and not (
                question.get("required") and
                any(t["key"] in keys_required_checked
                    for t in _targets_of(question))):
            continue  # form không được chọn và không key nào được chọn cần
        if not question.get("required"):
            continue
        kept = resolved_by_qid.get(qid, [])
        if not kept:
            continue  # câu không kích hoạt theo `when`
        if any(e.value_error for e in by_qid.get(qid, [])):
            continue  # đã báo lỗi giá trị — đừng báo "trống" thêm
        if question.get("row_group"):
            rows_map: dict[int, list[_Entry]] = {}
            for entry in kept:
                rows_map.setdefault(entry.row, []).append(entry)
            filled_rows = sorted(r for r, cells in rows_map.items()
                                 if any(c.filled for c in cells))
            if not filled_rows:
                errors.append(_err(
                    MISSING_REQUIRED,
                    "Câu bắt buộc %s (mẫu %s) chưa điền dòng nào" % (qid,
                                                                     form_id),
                    cell=kept[0].address, question_id=qid, form_id=form_id,
                    sheet=kept[0].sheet))
                continue
            for row_no in filled_rows:
                for entry in rows_map[row_no]:
                    if not entry.filled:
                        errors.append(_err(
                            MISSING_REQUIRED,
                            "Ô %s (%s) bắt buộc nhưng trống"
                            % (entry.address, qid),
                            cell=entry.address, question_id=qid,
                            form_id=form_id, sheet=entry.sheet))
            continue
        if not any(e.filled for e in kept):
            errors.append(_err(
                MISSING_REQUIRED, "Ô %s (%s) bắt buộc nhưng trống"
                % (kept[0].address, qid),
                cell=kept[0].address, question_id=qid, form_id=form_id,
                sheet=kept[0].sheet))
    if errors:
        return _fail(errors, audit)

    # -- 13. Gộp bảng lặp + DUP_ROW --------------------------------------------------------
    loop_out: dict[str, list[dict]] = {}
    for loop_name, rows_map in loops.items():
        seen_rows: dict[str, int] = {}
        for row_no in sorted(rows_map):
            row_cells = rows_map[row_no]
            if not any(v not in (None, "") for v in row_cells.values()):
                continue  # hàng trống không tính duplicate
            fingerprint = _json_canon(row_cells)
            if fingerprint in seen_rows:
                errors.append(_err(
                    DUP_ROW,
                    "Bảng %r: hàng sheet %d có nội dung Y HỆT hàng %d — xóa "
                    "bớt một hàng" % (loop_name, row_no, seen_rows[fingerprint]),
                    question_id=loop_name,
                    sheet=_sheet_of_loop(entries, loop_name)))
            else:
                seen_rows[fingerprint] = row_no
        loop_out[loop_name] = [rows_map[r] for r in sorted(rows_map)]
    if errors:
        return _fail(errors, audit)

    # -- 14. Candidate draft v2 nguyên tử + audit -------------------------------------------
    # checkedForms theo lựa chọn oneOf / form tùy chọn (step 11b); các câu
    # SELECTOR pseudo-key không rơi vào store (chỉ dùng lúc nạp).
    mandatory_checked = [f["form_id"] for f in approved_forms
                         if f.get("bat_buoc") and f["form_id"] in checked_set]
    chosen = next((selected_in_group[gi]
                   for gi in sorted(selected_in_group)
                   if selected_in_group[gi] in checked_set), None)
    active = mandatory_checked[0] if mandatory_checked else \
        (chosen or (checked or [None])[0])
    clean_values = {k: v for k, v in values.items()
                    if not (isinstance(k, str) and (
                        k.startswith(ONEOF_KEY_PREFIX) or
                        k.startswith(OPTIONAL_FORM_KEY_PREFIX)))}
    candidate = {
        "v": 2,
        "caseId": manifest["case_id"],
        "checkedForms": checked,
        "activeDocId": active,
        "autoSum": False,
        "viewState": {},
        "savedAt": None,          # UI gắn thời điểm áp dụng
        "store": {"values": clean_values, "checks": checks,
                  "loops": loop_out},
        "source": {
            "kind": CONTRACT,
            "manifest_hash": manifest.get("manifest_hash"),
            "file_sha256": audit.get("file_sha256"),
            "imported_at": imported_at or "",
            "forms": checked,
        },
    }
    audit.update({
        "schema_version": manifest.get("schema_version"),
        "manifest_hash": manifest.get("manifest_hash"),
        "forms": checked,
        "n_values": len(clean_values),
        "n_checks": len(checks),
        "n_loop_rows": sum(len(v) for v in loop_out.values()),
    })
    auth_result: dict = {"ok": True, "present": False, "bundle": None, "errors": []}
    try:
        from tools.excel_authorization_sheets import read_auth_bundle  # noqa: WPS433

        auth_result = read_auth_bundle(wb_raw, str(manifest["case_id"]))
    except Exception as exc:  # noqa: BLE001
        auth_result = {
            "ok": False,
            "present": True,
            "bundle": None,
            "errors": [{
                "code": "AUTH_READ_FAILED",
                "message": "%s: %s" % (type(exc).__name__, exc),
            }],
        }
    return {
        "ok": True,
        "errors": [],
        "candidate": candidate,
        "audit": audit,
        "authorization": auth_result,
    }
