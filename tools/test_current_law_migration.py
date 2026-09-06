# -*- coding: utf-8 -*-
"""Contract tests cho wave F0L current-law migration (Lane A).

Van ban hien hanh doi tuong:
  - TT 121/2026/TT-BTC (hieu luc 21/08/2026): thay Mau I so 1,2,3,4,5,10,12,27,30;
    Bai bo Mau I so 11 (khoan 2 Dieu 1).
  - ND 296/2026/NĐ-CP (hieu luc 23/07/2026): sua ND 168/2025.

Nguon anh: /tmp/tt121_pages (p2, p49-p70), /tmp/nd296_pages (p001-p014) — moi
assertion fingerprint duoi day deu tra ve 1 trang anh cu the, ghi trong
docs/legal_sources_2026/inventory/*.json (truong source_quote).

TRANG THAI TREN NHANH LANE (NOTE QUAN TRONG cho Lead):
  - TestM11RemovedByCurrentLaw::test_a_form_maps_marker la test CO TINH HOP DONG
    DANH (assertion theo HOP DONG DICH): form_maps.json la generated artifact,
    Lead regen tren main va phai them 'current_law_status':
    'REMOVED_BY_CURRENT_LAW' cho entry PLI_M11 (xem
    docs/legal_sources_2026/inventory/PLI_M11.json, muc de_xuat_marker_catalog).
    Test nay DO tren nhanh lane cho toi khi Lead merge marker; hop dong dap ung:
    khong SKIP, khong ha assertion.
  - Cac test con lai phai XANH ngay tren nhanh lane.

Chay: python3 -X utf8 -m tools.test_current_law_migration
"""
import hashlib
import json
import re
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
LEGAL_DIR = ROOT / "docs" / "legal_sources_2026"
SOURCES_MD = LEGAL_DIR / "SOURCES.md"
INVENTORY = LEGAL_DIR / "inventory"
FORM_SPECS = ROOT / "tools" / "form_specs"
USE_CASES = ROOT / "app" / "data" / "use_cases.json"
FORM_MAPS = ROOT / "app" / "data" / "form_maps.json"

SOURCE_M10 = (ROOT / "06_ho_kinh_doanh_va_chuyen_doi" / "02_chuyen_doi_hkd_sang_dn" /
              "DN_PL1_Mau_10_danh_sach_chu_so_huu_huong_loi_cua_doanh_nghiep.docx")
SOURCE_M11 = (ROOT / "06_ho_kinh_doanh_va_chuyen_doi" / "02_chuyen_doi_hkd_sang_dn" /
              "DN_PL1_Mau_11_ke_khai_thong_tin_de_xac_dinh_chu_so_huu_huong_loi_cua_doanh_nghiep.docx")
SOURCE_M12 = ROOT / "02_thay_doi_thong_tin_doanh_nghiep" / "PL1_Mau_12_tach_tu_file_full.docx"
SOURCE_M27 = (ROOT / "04_tam_ngung_cham_dut_giai_the" /
              "PL1_Mau_27_thong_bao_ve_viec_tam_ngung_kinh_doanh_tiep_tuc_kinh_doanh_"
              "truoc_thoi_han_da_thong_bao_cua_doanh.docx")
SOURCE_M30 = (ROOT / "04_tam_ngung_cham_dut_giai_the" /
              "PL1_Mau_30_thong_bao_ve_viec_giai_the_doanh_nghiep.docx")

PDF_FILES = [
    "121-btc.signed.pdf",
    "296_2026_nd-cp_23072026-signed.signed.pdf",
    "68-btc.pdf",
    "168nd.signed.pdf",
]


def docx_text(path):
    """Toan bo text cua file docx (paragraph + bang) de so fingerprint."""
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


# ---------------------------------------------------------------------------
# (a) + (b) + inventory: Mau I so 11 da bi bai bo
# ---------------------------------------------------------------------------
class TestM11RemovedByCurrentLaw(unittest.TestCase):
    """TT 121/2026, khoan 2 Dieu 1 (anh p2): 'Bai bo Mau so 11 tai Phu luc I...'."""

    def test_a_form_maps_marker(self):
        # HOP DONG DICH — Lead them 'current_law_status' khi regen form_maps
        # tren main. DO tren nhanh lane (chua co marker), XANH sau merge.
        # Khong SKIP: contract phai duoc kiem soat boi test.
        fm = load_json(FORM_MAPS)
        self.assertIn("PLI_M11", fm,
                      "form_maps phai con entry PLI_M11 (danh dau trang thai, "
                      "khong xoa entry)")
        self.assertEqual(fm["PLI_M11"].get("current_law_status"),
                         "REMOVED_BY_CURRENT_LAW",
                         "PLI_M11.current_law_status phai la "
                         "'REMOVED_BY_CURRENT_LAW' — Lead merge marker theo "
                         "inventory/PLI_M11.json de_xuat_marker_catalog")

    def test_b_no_usecase_references_mau_phu_luc_i_11(self):
        # use_cases khong duoc con form pair {pl:'I', mau:11} hoac de cap Mau 11 PL I
        uc = load_json(USE_CASES)
        raw = json.dumps(uc, ensure_ascii=False)

        form_pairs = []

        def walk(obj):
            if isinstance(obj, dict):
                if obj.get("pl") == "I" and obj.get("mau") == 11:
                    form_pairs.append(obj)
                for v in obj.values():
                    walk(v)
            elif isinstance(obj, list):
                for v in obj:
                    walk(v)

        walk(uc)
        self.assertEqual(form_pairs, [],
                         "Khong use case nao duoc tham chieu Phu luc I Mau so 11 "
                         "(da bi TT 121/2026 bai bo)")
        self.assertNotIn("Mẫu số 11", raw)
        self.assertNotIn("Mẫu 11", raw)

    def test_c_inventory_verdict_and_legal_quote(self):
        inv = load_json(INVENTORY / "PLI_M11.json")
        self.assertEqual(inv["verdict"], "REMOVED_BY_CURRENT_LAW")
        quote = inv["legal_basis"]["quote_verbatim"]
        self.assertIn("Bãi bỏ Mẫu số 11", quote)
        self.assertIn("Phụ lục I", quote)
        # source M11 giu nguyen theo lenh lane (khong sua, khong xoa)
        self.assertTrue(SOURCE_M11.exists())


# ---------------------------------------------------------------------------
# (c) Hash 4 file PDF nguon khop SOURCES.md
# ---------------------------------------------------------------------------
class TestSourcePdfsMatchSourcesMd(unittest.TestCase):
    def test_sha256_of_four_pdfs(self):
        table = SOURCES_MD.read_text(encoding="utf-8")
        for fname in PDF_FILES:
            with self.subTest(pdf=fname):
                path = LEGAL_DIR / fname
                self.assertTrue(path.exists(),
                                "Thieu file nguon %s trong repo" % fname)
                rows = [ln for ln in table.splitlines()
                        if fname in ln and ln.count("|") >= 4]
                self.assertTrue(rows,
                                "SOURCES.md khong co dong ghi %s" % fname)
                m = re.search(r"[0-9a-f]{64}", rows[0])
                self.assertIsNotNone(m,
                                     "Dong SOURCES.md cua %s khong co sha256"
                                     % fname)
                expected = m.group(0)
                digest = hashlib.sha256(path.read_bytes()).hexdigest()
                self.assertEqual(digest, expected,
                                 "SHA-256 cua %s khong khop SOURCES.md"
                                 % fname)


# ---------------------------------------------------------------------------
# (d) Fingerprint 2026 — chuoi/ cau truc MOI phai CO trong source/spec
# ---------------------------------------------------------------------------
class TestFingerprint2026(unittest.TestCase):
    def test_m10_source_title_and_nd296_footnote(self):
        # p49: tieu de + chan trang chu thich 2 dan Dieu 3 ND 296/2026
        text = docx_text(SOURCE_M10)
        self.assertIn("DANH SÁCH CHỦ SỞ HỮU HƯỞNG LỢI CỦA DOANH NGHIỆP", text)
        self.assertIn("Điều 3 Nghị định số 296/2026/NĐ-CP", text)
        self.assertIn("(Ký và ghi họ tên)9", text)

    def test_m10_spec_18_columns_with_formula_row(self):
        # p49: hang so '1'..'18' voi cong thuc '11=9+10' va '14=12+13'
        spec_file = FORM_SPECS / "formation" / "PLI_M10.json"
        spec = load_json(spec_file)
        self.assertEqual(len(spec["tables"]), 1)
        entry = spec["tables"][0]
        self.assertEqual(entry["spec"]["kind"], "nested-grid")
        cols = entry["spec"]["columns"]
        self.assertEqual(len(cols), 17,
                         "Mau 10 2026 co 18 cot = STT (generated) + 17 cot du lieu")
        self.assertEqual(entry["spec"]["stt"],
                         {"mode": "generated", "col_index": 0})
        number_row = entry["spec"]["header_rows"][3]
        self.assertEqual([c["text"] for c in number_row],
                         ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10",
                          "11=9+10", "12", "13", "14=12+13", "15", "16",
                          "17", "18"])
        tags = {c["tag"] for c in cols}
        expected = {"ho_ten", "ngay_sinh", "gioi_tinh", "so_giay_to",
                    "quoc_tich", "dan_toc", "dia_chi",
                    "ty_le_von_dieule_truc_tiep", "ty_le_von_dieule_gian_tiep",
                    "ty_le_von_dieule_tong", "ty_le_co_phan_truc_tiep",
                    "ty_le_co_phan_gian_tiep", "ty_le_co_phan_tong",
                    "thanh_vien_hop_danh", "nhom_ca_nhan_quan_he",
                    "kiem_soat_doanh_nghiep", "nguoi_quan_ly_doanh_nghiep"}
        self.assertEqual(tags, expected)
        # group header span 9 cot (p49)
        band = entry["spec"]["header_rows"][0][8]
        self.assertEqual(band["text"], "Chủ sở hữu hưởng lợi của doanh nghiệp2")
        self.assertEqual(band["colspan"], 9)
        # loop canonical giu nguyen
        self.assertEqual(entry["spec"]["loop"], "ds_cshhl")

    def test_m27_ndd_block_2026(self):
        # p67: khoi NDD 3 dong + chu thich 3,4 gop
        text = docx_text(SOURCE_M27)
        self.assertIn("Địa chỉ thư điện tử4", text)
        self.assertIn("3,4 Doanh nghiệp kê khai số điện thoại, thư điện tử",
                      text)
        self.assertIn("Trường hợp có nhiều hơn một người đại diện theo pháp luật",
                      text)
        self.assertIn("(Ký và ghi họ tên)6", text)

    def test_m30_so_dang_ky_co_dong(self):
        # p70: cau moi + chu thich 1 gioi han CTCP
        text = docx_text(SOURCE_M30)
        self.assertIn("bản sao sổ đăng ký cổ đông", text)
        self.assertIn("không phải là công ty niêm yết và công ty đăng ký giao "
                      "dịch chứng khoán", text)

    def test_m12_cshhl_items_2026(self):
        # p63: item a) them '(bao gồm người quản lý doanh nghiệp)', item b) moi
        text = docx_text(SOURCE_M12)
        self.assertIn("bao gồm người quản lý doanh nghiệp", text)
        self.assertIn("khoản 1 Điều 3 Luật số 76/2025/QH15", text)

    def test_inventory_files_present_and_parse(self):
        for name in ("PLI_M10.json", "PLI_M11.json", "PLI_M12.json",
                     "PLI_M27.json", "PLI_M30.json", "ND296_impact.json"):
            with self.subTest(inventory=name):
                data = load_json(INVENTORY / name)
                self.assertIsInstance(data, dict)


# ---------------------------------------------------------------------------
# (e) Fingerprint 2025 — chuoi/ cau truc CU phai MAT khoi source/spec
# ---------------------------------------------------------------------------
class TestFingerprint2025Gone(unittest.TestCase):
    def test_m12_no_mau11_references(self):
        # p63: 2 item dan Mau so 11 da xoa; khong con cu phap NĐ168 D52/D18
        text = docx_text(SOURCE_M12)
        self.assertNotIn("Mẫu số 11", text,
                         "M12 2026 khong con dan Mau so 11 (bai bo)")
        self.assertNotIn("Điều 52 Nghị định số 168/2025", text)
        self.assertNotIn("Điều 18 Nghị định số 168/2025", text)

    def test_m27_2025_strings_gone(self):
        text = docx_text(SOURCE_M27)
        # p67: yeu cau SDT trung CSDL dan cu 2025 da bo; khong danh so NDD
        self.assertNotIn("Cơ sở dữ liệu quốc gia về dân cư", text)
        self.assertNotIn("1. Họ, chữ đệm và tên", text)
        self.assertNotIn("2. Họ, chữ đệm và tên", text)
        self.assertNotIn("Ngày, tháng, năm sinh: ……../……./…………", text)
        self.assertNotIn("(Ký và ghi họ tên)5", text)

    def test_m10_spec_2025_columns_gone(self):
        # spec 2026 khong con cot 'Quyền chi phối'/'Ghi chú' hay 2 cot ty le don
        # (chi xet CAU TRUC columns/header_rows, khong xet chuoi trong note —
        # note con giu lich su migration de trace)
        spec = load_json(FORM_SPECS / "formation" / "PLI_M10.json")
        entry = spec["tables"][0]
        cols = entry["spec"]["columns"]
        self.assertNotIn("quyen_chi_phoi", {c["tag"] for c in cols})
        for c in cols:
            self.assertNotIn("Ghi chú", c["header"])
            self.assertNotIn(c["tag"], ("ty_le", "ty_le_co_phan"),
                             "2 cot ty le don 2025 phai tach 3 cot "
                             "trực tiếp/gián tiếp/tổng")
        for row in entry["spec"]["header_rows"]:
            for cell in row:
                self.assertNotIn("Quyền chi phối", cell["text"])
                self.assertNotIn("Ghi chú", cell["text"])

    def test_m10_source_no_2025_convert_garbage(self):
        text = docx_text(SOURCE_M10)
        # paragraph rac convert Mau 9 2025 da thay bang chan mau 2026 (p49)
        self.assertNotIn("15 - Người đại diện theo pháp luật", text)
        self.assertNotIn("Chủ tịch Hội đồng thành viên của công ty sau chuyển "
                         "đổi ký trực tiếp", text)


if __name__ == "__main__":
    unittest.main()
