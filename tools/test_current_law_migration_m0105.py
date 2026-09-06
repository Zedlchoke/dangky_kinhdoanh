#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test_current_law_migration_m0105.py — Lane A2 (F0L current-law migration).

Khoá delta hiến pháp M01-M05 (TT 121/2026/TT-BTC, hiệu lực 21/08/2026, thay
Mẫu 1-5 Phụ lục I TT 68/2025) + N-v4 M07 (dựng lại nguồn Mẫu 7 từ bản in
chính thức TT 68/2025 p49-p50 — M07 KHÔNG bị TT121 thay).

Assertion đọc 3 lớp:
  [SRC]    source DOCX 2025 (06_ho_kinh_doanh_va_chuyen_doi/…) — fingerprint
           2026 phải CÓ, fingerprint 2025 phải MẤT. Danh sách fingerprint TRÍCH
           TỪ docs/legal_sources_2026/inventory/<FORM_ID>.json (test_evidence),
           không hardcode trong test.
  [MAPS]   app/data/form_maps.json — canonical key: key 2025 bị xoá phải MẤT,
           key 2026 phải CÓ. (đỏ-chờ-merge: cần Lead regen trên main)
  [TPL]    app/templates/<FORM_ID>.docx — fingerprint 2026 trong template.
           (đỏ-chờ-merge: generated artifact, Lead regen sau merge)

Lớp [MAPS]/[TPL] tự SKIP (skipUnless with reason rõ ràng) khi template chưa
được regen theo nguồn mới — phát hiện bằng marker 2026 trong template. Lớp
[SRC] luôn chạy: chỉ phụ thuộc source đã commit.

Chạy: python3 -X utf8 -m tools.test_current_law_migration_m0105
"""
from __future__ import annotations

import json
import re
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import docx  # noqa: E402

FORMS = ["PLI_M01", "PLI_M02", "PLI_M03", "PLI_M04", "PLI_M05", "PLI_M07"]
INVENTORY_DIR = ROOT / "docs" / "legal_sources_2026" / "inventory"

# Canonical key delta per form (so form_maps entry hiện hành).
REMOVED_KEYS_2025 = {
    "PLI_M01": ["hkd_cd.so_gcn", "hkd_cd.ngay_cap", "hkd_cd.noi_cap", "hkd_cd.ma_so_thue"],
    "PLI_M02": ["hkd_cd.so_gcn", "hkd_cd.ngay_cap", "hkd_cd.noi_cap", "hkd_cd.ma_so_thue"],
    "PLI_M03": ["hkd_cd.so_gcn", "hkd_cd.ngay_cap", "hkd_cd.noi_cap", "hkd_cd.ma_so_thue"],
    "PLI_M04": ["hkd_cd.so_gcn", "hkd_cd.ngay_cap", "hkd_cd.noi_cap", "hkd_cd.ma_so_thue"],
    "PLI_M05": ["hkd_cd.so_gcn", "hkd_cd.ngay_cap", "hkd_cd.noi_cap", "hkd_cd.ma_so_thue"],
    "PLI_M07": [],
}
ADDED_KEYS_2026 = {
    "PLI_M01": ["hkd_ma_so"],
    "PLI_M02": ["hkd_ma_so"],
    "PLI_M03": ["hkd_ma_so"],
    "PLI_M04": ["hkd_ma_so"],
    "PLI_M05": ["hkd_ma_so", "ap_dung_tu", "so_lao_dong", "chk_bc_hang_thang",
                "chk_bc_03_thang", "chk_bc_06_thang", "chk_ht_doc_lap",
                "chk_ht_hop_nhat", "chk_ht_phu_thuoc", "chk_gtgt_khoan_tru",
                "chk_gtgt_truc_tiep_tren_gtgt", "chk_gtgt_truc_tiep_doanh_so",
                "chk_gtgt_khong_phai_nop", "gd.ho_ten", "kt.ho_ten"],
    "PLI_M07": ["dia_diem_ky", "ngay_ky_vb", "thang_ky_vb", "nam_ky_vb"],
}
# Marker 2026 trong template để phát hiện "đã regen chưa".
TPL_MARKERS = {
    "PLI_M01": "khong khai thac duoc thong tin ve ho kinh doanh",
    "PLI_M02": "cam ket",
    "PLI_M03": "cam ket",
    "PLI_M04": "cam ket",
    "PLI_M05": "cac thanh vien hop danh",
    "PLI_M07": "danh sach co dong sang lap cong ty co phan",
}


def fold(s: str) -> str:
    import unicodedata
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"\s+", " ", s).strip().lower()


def load_inventory(form_id: str) -> dict:
    return json.loads((INVENTORY_DIR / ("%s.json" % form_id)).read_text(
        encoding="utf-8"))


def source_path(form_id: str) -> Path:
    import tools.build_form_maps as bfm
    return Path(bfm.find_source(form_id))


def source_paras(form_id: str) -> list[str]:
    d = docx.Document(str(source_path(form_id)))
    return [p.text for p in d.paragraphs]


def source_folds(form_id: str) -> list[str]:
    return [fold(t) for t in source_paras(form_id)]


def tpl_paras(form_id: str) -> list[str]:
    d = docx.Document(str(ROOT / "app" / "templates" / ("%s.docx" % form_id)))
    return [p.text for p in d.paragraphs]


def tpl_folds(form_id: str) -> list[str]:
    return [fold(t) for t in tpl_paras(form_id)]


def form_maps_keys(form_id: str) -> set[str]:
    fm = json.loads((ROOT / "app" / "data" / "form_maps.json")
                    .read_text(encoding="utf-8"))
    return {f["key"] for f in fm[form_id]["fields"]}


def tpl_regenerated(form_id: str) -> bool:
    """Template đã regen theo nguồn 2026/N-v4 chưa? (marker-based)"""
    try:
        return any(TPL_MARKERS[form_id] in t for t in tpl_folds(form_id))
    except Exception:
        return False


# ---------------------------------------------------------------------------
class TestSourceFingerprints(unittest.TestCase):
    """[SRC] luôn chạy — chỉ đọc source DOCX đã commit."""

    def test_inventory_exists_for_every_form(self):
        for fid in FORMS:
            inv = load_inventory(fid)
            self.assertEqual(inv["form_id"], fid)
            self.assertIn("test_evidence", inv)
            self.assertIn("fingerprint_2026_must_exist", inv["test_evidence"])
            self.assertIn("fingerprint_2025_must_disappear", inv["test_evidence"])

    def test_source_2026_fingerprints_present(self):
        for fid in FORMS:
            inv = load_inventory(fid)
            sf = source_folds(fid)
            for fp in inv["test_evidence"]["fingerprint_2026_must_exist"]:
                # so sánh trên chuỗi đã fold (fingerprint trong inventory là
                # tiếng Việt có dấu); '11.6' là paragraph cell — substring OK
                needle = fold(fp)
                hit = any(needle in t for t in sf)
                self.assertTrue(
                    hit,
                    "%s: fingerprint 2026 KHÔNG có trong nguồn: %r" % (fid, fp))

    def test_source_2025_fingerprints_absent(self):
        for fid in FORMS:
            inv = load_inventory(fid)
            sf = source_folds(fid)
            for fp in inv["test_evidence"]["fingerprint_2025_must_disappear"]:
                fpf = fold(fp)
                # '1.6' (M04) phải so khớp paragraph nguyên vẹn, không substring
                if fid == "PLI_M04" and fp == "1.6":
                    self.assertNotIn(
                        "1.6", sf,
                        "%s: nguồn còn paragraph '1.6' — TT121 in '11.6' (p39)" % fid)
                    continue
                self.assertFalse(
                    any(fpf in t for t in sf),
                    "%s: fingerprint 2025 VẪN còn trong nguồn: %r" % (fid, fp))

    def test_m07_source_is_full_print_not_stub(self):
        """N-v4: nguồn M07 phải là bản in đầy đủ (p49-p50), không còn stub."""
        sp = source_paras("PLI_M07")
        self.assertGreaterEqual(
            len(sp), 20,
            "M07 nguồn vẫn là stub 2 dòng — chưa dựng lại từ TT 68/2025 p49-p50")
        sf = source_folds("PLI_M07")
        for mark in ("i. co dong sang lap la ca nhan",
                     "ii. co dong sang lap la to chuc",
                     "2. thong tin nguoi dai dien theo phap luat/uy quyen cua co dong",
                     "ke khai tuong tu cac co dong sang lap la ca nhan",
                     "13 - nguoi dai dien theo phap luat cua doanh nghiep ky truc tiep vao phan nay"):
            self.assertTrue(any(mark in t for t in sf),
                            "M07 nguồn thiếu marker bản in: %r" % mark)

    def test_uy_ban_raw_glyph(self):
        """Delta 'Ủy ban'->'Uỷ ban' (TT121 p14/p26/p35) — fold-blind, so raw."""
        for fid in ("PLI_M02", "PLI_M03", "PLI_M04"):
            raw = "\n".join(source_paras(fid))
            self.assertIn("Uỷ ban Chứng khoán Nhà nước", raw,
                          "%s: thiếu 'Uỷ ban' theo bản in 2026" % fid)
            self.assertNotIn("Ủy ban Chứng khoán Nhà nước", raw,
                             "%s: còn 'Ủy ban' (chính tả 2025)" % fid)

    def test_m02_m03_m04_cshhl_question_removed_kept_in_bot(self):
        """Câu hỏi CSHHL 'Có/Không' bị TT121 bỏ ở mục 12/11/13 — không còn."""
        for fid in ("PLI_M02", "PLI_M03", "PLI_M04"):
            sf = source_folds(fid)
            self.assertFalse(
                any(t.startswith("doanh nghiep co chu so huu huong loi khong")
                    for t in sf),
                "%s: câu hỏi CSHHL cũ còn trong nguồn" % fid)
            self.assertTrue(
                any("theo mau so 10 ban hanh kem theo thong tu nay" in t
                    and "chu so huu huong loi" in t for t in sf),
                "%s: thiếu dòng khai báo CSHHL theo Mẫu số 10" % fid)


class TestGeneratedArtifacts(unittest.TestCase):
    """[MAPS]+[TPL] — đỏ-chờ-merge: SKIP tới khi Lead regen trên main."""

    def test_form_maps_canonical_keys(self):
        for fid in FORMS:
            if not tpl_regenerated(fid):
                self.skipTest(
                    "%s: template chưa regen theo nguồn current-law trên main "
                    "(Lead chạy build_form_maps sau merge) — assertion [MAPS] "
                    "chờ merge" % fid)
            keys = form_maps_keys(fid)
            for k in REMOVED_KEYS_2025[fid]:
                self.assertNotIn(k, keys,
                                 "%s: key 2025 %r vẫn còn trong form_maps" % (fid, k))
            for k in ADDED_KEYS_2026[fid]:
                self.assertIn(k, keys,
                              "%s: key 2026 %r thiếu trong form_maps" % (fid, k))

    def test_template_2026_fingerprints(self):
        for fid in FORMS:
            if not tpl_regenerated(fid):
                self.skipTest(
                    "%s: template chưa regen — assertion [TPL] đỏ-chờ-merge" % fid)
            inv = load_inventory(fid)
            tf = tpl_folds(fid)
            for fp in inv["test_evidence"]["fingerprint_2026_must_exist"]:
                needle = fold(fp)
                self.assertTrue(any(needle in t for t in tf),
                                "%s: fingerprint 2026 thiếu trong template: %r"
                                % (fid, fp))

    def test_template_2025_fingerprints_absent(self):
        for fid in FORMS:
            if not tpl_regenerated(fid):
                self.skipTest(
                    "%s: template chưa regen — assertion [TPL] đỏ-chờ-merge" % fid)
            inv = load_inventory(fid)
            tf = tpl_folds(fid)
            for fp in inv["test_evidence"]["fingerprint_2025_must_disappear"]:
                fpf = fold(fp)
                if fid == "PLI_M04" and fp == "1.6":
                    self.assertNotIn("1.6", tf,
                                     "M04 template còn paragraph '1.6'")
                    continue
                self.assertFalse(any(fpf in t for t in tf),
                                 "%s: fingerprint 2025 còn trong template: %r"
                                 % (fid, fp))


if __name__ == "__main__":
    unittest.main(verbosity=2)
