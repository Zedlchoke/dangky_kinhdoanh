# -*- coding: utf-8 -*-
"""I5 — sheet AUTH_* trong workbook case-level.

Chạy: python -X utf8 -m tools.test_excel_authorization_sheets
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from tools.excel_authorization_sheets import (  # noqa: E402
    AUTH_AGENTS_SHEET, AUTH_META_SHEET, AUTH_SETS_SHEET,
    read_auth_bundle,
)
from tools.excel_case_export import write_case_workbook  # noqa: E402
from tools.excel_case_import import import_case_workbook  # noqa: E402

FIXTURE_META = {"today": "01/01/2026", "exported_at": "01/01/2026 10:00"}
_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run():
            try:
                fn()
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def _manifest(case_id: str = "thanh_lap_tnhh_1tv") -> dict:
    return json.loads(
        (ROOT / "app" / "data" / "case_questionnaires" / ("%s.json" % case_id))
        .read_text(encoding="utf-8"))


def _with_temp_store():
    tmp = tempfile.mkdtemp(prefix="dkkd_auth_xl_")
    os.environ["DKKD_AUTH_DATA_DIR"] = tmp
    import importlib
    import tools.authorization_store as store
    importlib.reload(store)
    return store, Path(tmp)


@check("AUTH-X01: export workbook có sheet AUTH_*")
def t_export_has_sheets():
    _store, _tmp = _with_temp_store()
    manifest = _manifest()
    out = Path(tempfile.mkdtemp()) / "wb.xlsx"
    write_case_workbook(manifest, out, meta=FIXTURE_META)
    wb = openpyxl.load_workbook(out)
    assert AUTH_META_SHEET in wb.sheetnames
    assert AUTH_AGENTS_SHEET in wb.sheetnames
    assert AUTH_SETS_SHEET in wb.sheetnames


@check("AUTH-X02: round-trip profile + set qua Excel")
def t_roundtrip():
    store, _tmp = _with_temp_store()
    manifest = _manifest()
    agent_id = str(uuid.uuid4())
    store.save_agent_profile({
        "agent_id": agent_id,
        "ho_ten": "Trần Lan",
        "so_cccd": "001234567890",
        "dien_thoai": "0901111222",
    })
    draft = store.create_authorization_draft(
        manifest["case_id"], agent_id, {"co_quan_tiep_nhan": "Phòng ĐKKD"})
    assert draft["ok"] is True
    auth_id = draft["authorization"]["authorization_id"]

    out = Path(tempfile.mkdtemp()) / "wb.xlsx"
    write_case_workbook(manifest, out, meta=FIXTURE_META)
    wb = openpyxl.load_workbook(out)
    read = read_auth_bundle(wb, manifest["case_id"])
    assert read["ok"] is True, read
    assert read["present"] is True
    assert len(read["bundle"]["agents"]) >= 1
    assert any(r["agent_id"] == agent_id for r in read["bundle"]["agents"])
    assert any(r["authorization_id"] == auth_id for r in read["bundle"]["sets"])

    os.environ["DKKD_AUTH_DATA_DIR"] = tempfile.mkdtemp(prefix="dkkd_auth_xl2_")
    import importlib
    import tools.authorization_store as store2
    importlib.reload(store2)
    applied = store2.apply_excel_bundle(read["bundle"], case_id=manifest["case_id"])
    assert applied["ok"] is True, applied
    listed = store2.list_agent_profiles()
    assert any(p["ho_ten"] == "Trần Lan" for p in listed["profiles"])


@check("AUTH-X03: AUTH rỗng — đọc sheet OK, không lỗi khung")
def t_read_empty_auth():
    _store, _tmp = _with_temp_store()
    manifest = _manifest()
    out = Path(tempfile.mkdtemp()) / "wb.xlsx"
    write_case_workbook(manifest, out, meta=FIXTURE_META)
    wb = openpyxl.load_workbook(out)
    read = read_auth_bundle(wb, manifest["case_id"])
    assert read["ok"] is True, read
    assert read["present"] is True
    assert read["bundle"]["agents"] == []
    res = import_case_workbook(out, manifest, imported_at="2026-01-01T00:00:00")
    codes = [e.get("code") for e in (res.get("errors") or [])]
    assert "BAD_HEADER" not in codes


@check("AUTH-X04: workbook cũ không AUTH — tương thích ngược")
def t_legacy_no_auth():
    manifest = _manifest()
    out = Path(tempfile.mkdtemp()) / "legacy.xlsx"
    write_case_workbook(manifest, out, meta=FIXTURE_META, include_authorization=False)
    wb = openpyxl.load_workbook(out)
    assert AUTH_META_SHEET not in wb.sheetnames
    read = read_auth_bundle(wb, manifest["case_id"])
    assert read["present"] is False
    res = import_case_workbook(out, manifest, imported_at="2026-01-01T00:00:00")
    codes = [e.get("code") for e in (res.get("errors") or [])]
    assert "BAD_HEADER" not in codes


def main() -> int:
    t_export_has_sheets()
    t_roundtrip()
    t_read_empty_auth()
    t_legacy_no_auth()
    failed = [n for ok, n in _results if not ok]
    print("\n%d/%d PASS" % (len(_results) - len(failed), len(_results)))
    if failed:
        print("FAILED:", ", ".join(failed))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
