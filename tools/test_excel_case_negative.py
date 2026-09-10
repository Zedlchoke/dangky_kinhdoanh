# -*- coding: utf-8 -*-
"""E3 — negative workbook CASE-LEVEL: khách/lập trình phá khung -> fail-loud,
KHÔNG data một phần, draft cũ không bị ghi đè
(PLAN_hoan_thien_pfat_hanh §2.2 danh sách fail-loud bắt buộc, §5-E3).

[NGC1]  workbook của case khác                              -> WRONG_CASE
[NGC2]  sửa một ô trong __map (không tính lại checksum)     -> CHECKSUM
[NGC3]  xóa sheet __map                                     -> BAD_HEADER
[NGC4]  đổi key trong __map + TÍNH LẠI checksum đúng thuật
        toán (lách lớp 5)                                   -> BAD_HEADER
[NGC5]  xóa sheet form                                      -> BAD_HEADER
[NGC6]  chèn sheet lạ                                       -> BAD_HEADER
[NGC7]  công thức độc hại "=cmd|' /C calc'!A0" vào ô vàng   -> FORMULA
[NGC8]  ô CCCD là ô SỐ (Excel tự đổi, mất số 0 đầu)         -> BAD_ID_TEXT
[NGC9]  ô lạ trong cột có ô điền                            -> UNKNOWN_CELL
[NGC10] bảng tràn: điền hàng thứ row_last+1                 -> TOO_MANY_ROWS
[NGC11] hai hàng bảng y hệt nhau                            -> DUP_ROW
[NGC12] select gõ tay ngoài danh sách                       -> BAD_SELECT
[NGC13] ngày 32/13/2020 + tiền "năm tỷ"                     -> BAD_DATE + BAD_NUMBER
[NGC14] manifest_hash/map_version lệch                      -> BAD_VERSION
[NGC15] gõ đè ô mirror (shared key)                         -> MIRROR_CELL
[NGC16] file lớn: điền ĐỦ mọi ô điền của workbook (valid)   -> ok=True
[NGC17] file không phải xlsx                                -> BAD_FILE

Mọi case lỗi: `candidate is None` (nguyên tắc nguyên tử) và `errors[]` có
{cell, question_id} để UI trỏ đúng ô.

Chạy:  python3 -X utf8 -m tools.test_excel_case_negative   (exit 0 = PASS)
"""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from tools.excel_case_export import (  # noqa: E402
    MAP_COLUMNS, MAP_SHEET, write_case_workbook,
)
from tools.excel_case_import import (  # noqa: E402
    BAD_DATE, BAD_FILE, BAD_HEADER, BAD_ID_TEXT, BAD_NUMBER, BAD_SELECT,
    BAD_VERSION, CHECKSUM, DUP_ROW, FORMULA, MISSING_REQUIRED,
    MIRROR_CELL, TOO_MANY_ROWS, UNKNOWN_CELL, WRONG_CASE,
    _checksum_of, _read_map_records, import_case_workbook,
)
from tools.test_excel_case_roundtrip import _shared_manifest  # noqa: E402

FIXTURE_META = {"today": "01/01/2026", "exported_at": "01/01/2026 10:00"}

CASE = "dang_ky_ho_kinh_doanh_lan_dau"

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run():
            try:
                fn()
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def _manifest(case_id: str = CASE) -> dict:
    return json.loads(
        (ROOT / "app" / "data" / "case_questionnaires" / ("%s.json" % case_id))
        .read_text(encoding="utf-8"))


def _export(tmp: Path, name: str = "wb.xlsx", manifest: dict | None = None):
    manifest = manifest or _manifest()
    out = write_case_workbook(manifest, tmp / name, meta=FIXTURE_META)
    return out, openpyxl.load_workbook(out)


def _map_col(mp) -> dict:
    return {mp.cell(3, i + 1).value: i + 1 for i in range(len(MAP_COLUMNS))}


def _save(tmp: Path, wb, name: str = "mut.xlsx") -> Path:
    path = tmp / name
    wb.save(path)
    return path


def _codes(errors: list[dict]) -> list[str]:
    return [e["code"] for e in errors]


def _cell_of(mp, col, qid, layout="single", occurrence=0):
    """Địa chỉ + sheet của ô điền thứ `occurrence` của câu (theo map)."""
    seen = 0
    for r in range(4, mp.max_row + 1):
        if mp.cell(r, col["question_id"]).value != qid or \
                mp.cell(r, col["layout"]).value != layout:
            continue
        if seen == occurrence:
            return (mp.cell(r, col["sheet"]).value,
                    mp.cell(r, col["cell"]).value,
                    mp.cell(r, col["row_last"]).value)
        seen += 1
    raise AssertionError("không thấy ô của %s" % qid)


@check("[NGC1] workbook case khác -> WRONG_CASE, candidate None")
def ngc1():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, _ = _export(tmp, manifest=_manifest("cham_dut_hoat_dong_hkd"))
        res = import_case_workbook(out, _manifest())  # nạp vào case hộ KD
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [WRONG_CASE], res["errors"]


@check("[NGC2] sửa ô trong __map (không rehash) -> CHECKSUM")
def ngc2():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        sheet, cell, _ = _cell_of(mp, col, "q_ten_hkd")
        # sửa trực tiếp dòng map tương ứng (cột key) — KHÔNG tính lại checksum
        for r in range(4, mp.max_row + 1):
            if mp.cell(r, col["cell"]).value == cell and \
                    mp.cell(r, col["question_id"]).value == "q_ten_hkd":
                mp.cell(r, col["key"]).value = "hkd_ten_bi_sua"
                break
        else:
            raise AssertionError("không tìm thấy dòng map của q_ten_hkd")
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [CHECKSUM], res["errors"]


@check("[NGC3] xóa sheet __map -> BAD_HEADER")
def ngc3():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        del wb[MAP_SHEET]
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [BAD_HEADER], res["errors"]


@check("[NGC4] đổi key trong map + TÍNH LẠI checksum đúng -> BAD_HEADER (key lạ)")
def ngc4():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        for r in range(4, mp.max_row + 1):
            if mp.cell(r, col["question_id"]).value == "q_so_cccd":
                mp.cell(r, col["key"]).value = "key_khong_ton_tai"
                break
        # tính lại checksum bằng ĐÚNG helper của importer (lách lớp 5)
        meta = {}
        for v in mp.cell(1, 1).value and [mp.cell(1, c).value
                                          for c in range(2, mp.max_column + 1)]:
            text = str(v or "")
            k, sep, val = text.partition("=")
            if sep and k.strip():
                meta.setdefault(k.strip(), val.strip())
        records = _read_map_records(wb)
        new_checksum = _checksum_of(meta, records)
        for c in range(2, mp.max_column + 1):
            if str(mp.cell(1, c).value or "").startswith("checksum="):
                mp.cell(1, c).value = "checksum=%s" % new_checksum
                break
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert BAD_HEADER in _codes(res["errors"]), res["errors"]
        bad = [e for e in res["errors"]
               if e["code"] == BAD_HEADER and "key_khong_ton_tai"
               in e["message"]]
        assert bad, res["errors"]


@check("[NGC5] xóa sheet form -> BAD_HEADER (thiếu sheet declared)")
def ngc5():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        sheet, _, _ = _cell_of(mp, col, "q_ten_hkd")
        del wb[sheet]
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [BAD_HEADER], res["errors"]
        assert sheet in res["errors"][0]["message"], res["errors"][0]


@check("[NGC6] chèn sheet lạ -> BAD_HEADER")
def ngc6():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        wb.create_sheet("Trang_Khach_Them")
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [BAD_HEADER], res["errors"]


@check("[NGC7] công thức '=cmd|…' vào ô vàng -> FORMULA, không data")
def ngc7():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        sheet, cell, _ = _cell_of(mp, col, "q_ten_hkd")
        wb[sheet][cell] = "=cmd|' /C calc'!A0"
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [FORMULA], res["errors"]
        assert res["errors"][0]["cell"] == cell


@check("[NGC8] ô CCCD là ô SỐ (mất số 0 đầu) -> BAD_ID_TEXT")
def ngc8():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        sheet, cell, _ = _cell_of(mp, col, "q_so_cccd")
        wb[sheet][cell] = 79085001234          # Excel đã tự đổi thành số
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert BAD_ID_TEXT in _codes(res["errors"]), res["errors"]


@check("[NGC9] ô lạ trong cột có ô điền -> UNKNOWN_CELL")
def ngc9():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        sheet, cell, _ = _cell_of(mp, col, "q_ten_hkd")
        row = int("".join(ch for ch in cell if ch.isdigit()))
        wb[sheet]["C%d" % row] = "phần thừa khách gõ"   # cột C có ô map khác
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [UNKNOWN_CELL], res["errors"]


@check("[NGC10] điền tràn bảng (hàng row_last+1) -> TOO_MANY_ROWS")
def ngc10():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        sheet, cell, row_last = _cell_of(mp, col, "q_nganh_ten", "table")
        col_letter = "".join(ch for ch in cell if ch.isalpha())
        wb[sheet]["%s%d" % (col_letter, row_last + 1)] = "hàng tràn"
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [TOO_MANY_ROWS], res["errors"]


@check("[NGC11] hai hàng bảng y hệt -> DUP_ROW (đã điền đủ bắt buộc)")
def ngc11():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        # điền đủ bắt buộc trước — required trống được báo TRƯỚC DUP_ROW
        required = {
            "q_ho_ten": "nguyễn văn a", "q_so_cccd": "079085001234",
            "q_dia_chi_chi_tiet": "25 Nguyễn Trãi",
            "q_xa_phuong": "Phường Bến Thành",
            "q_tinh_thanh": "TP. Hồ Chí Minh", "q_ten_hkd": "Hộ KD E",
        }
        for r in range(4, mp.max_row + 1):
            qid = mp.cell(r, col["question_id"]).value
            if mp.cell(r, col["layout"]).value == "mirror":
                continue
            sheet = mp.cell(r, col["sheet"]).value
            cell = mp.cell(r, col["cell"]).value
            if qid in required:
                wb[sheet][cell] = required[qid]
            elif qid == "q_nganh_ten":
                # hai dòng ĐẦU bảng y hệt nhau, các dòng khác phân biệt
                row_no = int("".join(ch for ch in cell if ch.isdigit()))
                wb[sheet][cell] = "Ngành trùng nhau" if row_no <= 51 \
                    else "Ngành %d" % row_no
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [DUP_ROW], res["errors"]


@check("[NGC12] select gõ tay ngoài danh sách -> BAD_SELECT")
def ngc12():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        sheet, cell, _ = _cell_of(mp, col, "q_gioi_tinh")
        wb[sheet][cell] = "Nam Nữ 123"
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert BAD_SELECT in _codes(res["errors"]), res["errors"]


@check("[NGC13] ngày 32/13/2020 + tiền 'năm tỷ' -> BAD_DATE + BAD_NUMBER")
def ngc13():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        sheet, cell, _ = _cell_of(mp, col, "q_ngay_sinh")
        wb[sheet][cell] = "32/13/2020"
        sheet, cell, _ = _cell_of(mp, col, "q_von")
        wb[sheet][cell] = "năm tỷ"
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        codes = _codes(res["errors"])
        assert BAD_DATE in codes and BAD_NUMBER in codes, res["errors"]


@check("[NGC14] manifest_hash/map_version lệch -> BAD_VERSION")
def ngc14():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        # (a) nạp workbook vào manifest CÙNG case nhưng hash khác (giả lập
        # nội dung case đổi sau khi khách đã tải phiếu cũ)
        manifest = _manifest()
        manifest = json.loads(json.dumps(manifest, ensure_ascii=False))
        manifest["manifest_hash"] = "0" * 64
        out, _ = _export(tmp)
        res = import_case_workbook(out, manifest)
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [BAD_VERSION], res["errors"]
        # (b) sửa map_version trong file -> contract lệch
        out2, wb = _export(tmp, name="wb2.xlsx")
        mp = wb[MAP_SHEET]
        for c in range(2, mp.max_column + 1):
            if str(mp.cell(1, c).value or "").startswith("map_version="):
                mp.cell(1, c).value = "map_version=999"
                break
        res2 = import_case_workbook(_save(tmp, wb, "mut14.xlsx"), _manifest())
        assert res2["ok"] is False and res2["candidate"] is None, res2
        assert _codes(res2["errors"]) == [BAD_VERSION], res2["errors"]


@check("[NGC15] gõ đè ô mirror (shared key) -> MIRROR_CELL")
def ngc15():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        manifest = _shared_manifest()
        out, wb = _export(tmp, manifest=manifest)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        for r in range(4, mp.max_row + 1):
            if mp.cell(r, col["layout"]).value == "mirror" and \
                    mp.cell(r, col["question_id"]).value == "q_hkd_ma_so":
                sheet = mp.cell(r, col["sheet"]).value
                cell = mp.cell(r, col["cell"]).value
                wb[sheet][cell] = "0399999999"   # khách tự ý gõ vào ô mirror
                break
        res = import_case_workbook(_save(tmp, wb), manifest)
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [MIRROR_CELL], res["errors"]


@check("[NGC16] file lớn: điền ĐỦ mọi ô điền (giá trị hợp lệ) -> ok=True")
def ngc16():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        manifest = _manifest()
        out, wb = _export(tmp)
        mp = wb[MAP_SHEET]
        col = _map_col(mp)
        questions = {}
        for form in manifest["forms"]:
            for q in form["questions"]:
                questions[q["question_id"]] = q

        def valid_value(q, seq=0):
            t = q["type"]
            if t == "select":
                return q["options"][0]["label"]
            if t == "mark":
                return "X"
            if t == "date":
                return "01/02/2026"
            if t == "money":
                return "1.000"
            if t == "integer":
                return "5"
            return "Nội dung %d" % seq   # từng hàng bảng KHÁC nhau

        n_filled = 0
        seq_by_qid: dict[str, int] = {}
        for r in range(4, mp.max_row + 1):
            if mp.cell(r, col["layout"]).value == "mirror":
                continue
            qid = mp.cell(r, col["question_id"]).value
            sheet = mp.cell(r, col["sheet"]).value
            cell = mp.cell(r, col["cell"]).value
            seq = seq_by_qid.get(qid, 0)
            wb[sheet][cell] = valid_value(questions[qid], seq)
            seq_by_qid[qid] = seq + 1
            n_filled += 1
        assert n_filled >= 100, n_filled   # workbook đầy đủ, nhiều ô
        res = import_case_workbook(_save(tmp, wb), _manifest())
        assert res["ok"], res["errors"]
        st = res["candidate"]["store"]
        assert st["values"] and st["loops"]["ds_nganh"], st
        # bảng 10 hàng đầy đủ theo thứ tự
        assert len(st["loops"]["ds_nganh"]) == 10, \
            len(st["loops"]["ds_nganh"])


@check("[NGC17] file không phải xlsx -> BAD_FILE")
def ngc17():
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        junk = tmp / "junk.xlsx"
        junk.write_text("không phải zip", encoding="utf-8")
        res = import_case_workbook(junk, _manifest())
        assert res["ok"] is False and res["candidate"] is None, res
        assert _codes(res["errors"]) == [BAD_FILE], res["errors"]


def main() -> int:
    fns = [ngc1, ngc2, ngc3, ngc4, ngc5, ngc6, ngc7, ngc8, ngc9, ngc10,
           ngc11, ngc12, ngc13, ngc14, ngc15, ngc16, ngc17]
    for fn in fns:
        fn()
    passed = sum(1 for ok, _ in _results if ok)
    failed = len(_results) - passed
    print("-" * 64)
    print("test_excel_case_negative: %d PASS / %d FAIL" % (passed, failed))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
