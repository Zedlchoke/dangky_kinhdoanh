# -*- coding: utf-8 -*-
"""Wave X2 — test xuất phiếu .xlsx (PLAN.md mục 4.4, file `tools/excel_export.py`).

Contract chốt cứng (lane x3-import + bridge X5 code đối xứng, KHÔNG được lệch):

[XC] Hàm `write_questionnaire(questionnaire, out_path, meta=None) -> Path`,
     hằng `MAP_SHEET = "___map"` import được bằng đúng dòng
     `from tools.excel_export import MAP_SHEET, write_questionnaire`;
     openpyxl thuần; KHÔNG mutate dict truyền vào; không print trong library.

[XL] Bố cục sheet "Phieu": hàng 1 tiêu đề = title (+ dòng hồ sơ từ `meta` ở
     hàng 2), hàng 3..7 = 5 dòng hướng dẫn (vàng/xám, không sửa ô khác,
     dropdown, ngày dd/mm/yyyy, ví dụ mờ), hàng 9+ nội dung: tiêu đề nhóm có
     màu riêng; 1 câu = 1 hàng (A nhãn, B ô điền, C ví dụ mờ/ghi chú, D ghi
     chú date); câu `row_group` = một bảng nhỏ (header + 10 hàng trống).

[XP] Quy ước màu + khóa: Ô VÀNG = khách điền (locked=False), Ô XÁM = app tự
     làm (`autofill`, locked=True). Gate "0 ô data ngoài locked-region" kiểm
     MÁY: mọi ô unlocked trong vùng quét phải là ô điền vàng đã khai báo, và
     ngược lại.

[XD] Dropdown chọn-là-chọn: `DataValidation(type="list")` từ `options[].label`
     cho `select`, luôn "X" cho `mark`; `showDropDown` KHÔNG bật (attribute
     OOXML nghịch dấu — bật là mất mũi tên); bảng lặp áp cho cả cột.

[XM] Sheet ẩn `___map`: A1 marker, B1 JSON thông tin file, bảng
     `MAP_COLUMNS` với MỘT dòng cho MỖI target `data_key` (question_id, key,
     when_json, transform, row_group, layout, ô điền, hàng đầu/cuối) — X3
     dịch ngược ô → key mà không phải dò lại layout.

[XR] Round-trip đọc lại bằng openpyxl: text tiếng Việt, protection, fill,
     dropdown, sheet ẩn giữ nguyên; ghi 2 lần cho cùng cấu trúc (deterministic).

[XN] Negative fail-loud: `ExcelExportError` có `code` đúng cho khung sai.

Chạy:  python3 -X utf8 -m tools.test_excel_export      (exit 0 = PASS)
"""
from __future__ import annotations

import inspect
import json
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from tools.excel_export import (  # noqa: E402
    COL_DATE_NOTE, COL_FILL, COL_LABEL, COL_NOTE, DATE_NOTE, FIRST_BODY_ROW,
    GOLD_HEX, GRAY_HEX, GROUP_PALETTE, GUIDE_FIRST_ROW, GUIDE_ROWS,
    MAP_COLUMNS, MAP_FIRST_ROW, MAP_HEADER_ROW, MAP_MARKER, MAP_SHEET,
    MAP_VERSION, MARK_VALUE, NOTE_FONT_HEX, REPEAT_ROWS, SELECT_NOTE,
    SHEET_NAME, ExcelExportError, write_questionnaire,
)
from tools.questionnaire.schema import validate  # noqa: E402

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run(*args):
            try:
                fn(*args)
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


# ---------------------------------------------------------------------------
# Fixture: questionnaire schema v1 (cùng shape artifact PLI_M72, tự dựng)
# ---------------------------------------------------------------------------

FIXTURE: dict = {
    "doc_id": "X2_DEMO",
    "title": "Phiếu thu thập thông tin — demo X2",
    "questionnaire_version": 1,
    "cluster": "demo_x2",
    "questions": [
        {"question_id": "q_loai", "group": "Về doanh nghiệp",
         "text": "Bạn muốn thành lập loại hình nào?",
         "type": "select", "required": True, "version": 1,
         "data_key": "dn.loai_hinh",
         "options": [
             {"value": "DNTN", "label": "Doanh nghiệp tư nhân"},
             {"value": "TNHH", "label": "Công ty TNHH"},
             {"value": "CTCP", "label": "Công ty cổ phần"},
         ]},
        {"question_id": "q_ten", "group": "Về doanh nghiệp",
         "text": "Bạn muốn đặt tên công ty là gì?",
         "type": "text", "example": "Công ty TNHH Minh Long",
         "required": True, "version": 1, "data_key": "dn.ten_tieng_viet"},
        {"question_id": "q_nganh_ten", "group": "Về ngành nghề",
         "text": "Bạn kinh doanh ngành nào? Ghi tên ngành",
         "type": "text", "example": "Bán lẻ hàng điện tử",
         "repeatable": True, "row_group": "nganh", "version": 1,
         "data_key": [
             {"key": "ds_nganh_dntn.ten_nganh",
              "when": {"dn.loai_hinh": ["DNTN"]}},
             {"key": "ds_nganh_ctcp.ten_nganh",
              "when": {"dn.loai_hinh": ["CTCP"]}},
         ]},
        {"question_id": "q_nganh_ma", "group": "Về ngành nghề",
         "text": "Mã ngành (app tra VSIC giúp — không cần điền)",
         "type": "text", "autofill": "vsic", "repeatable": True,
         "row_group": "nganh", "version": 1,
         "data_key": [
             {"key": "ds_nganh_dntn.ma_nganh",
              "when": {"dn.loai_hinh": ["DNTN", "CTCP"]}},
         ]},
        {"question_id": "q_nganh_chinh", "group": "Về ngành nghề",
         "text": "Đây có phải ngành nghề chính không?",
         "type": "mark",
         "options": [{"value": "X", "label": "Ngành chính"}],
         "repeatable": True, "row_group": "nganh", "version": 1,
         "data_key": "ds_nganh_dntn.la_nganh_chinh"},
        {"question_id": "q_vsic", "group": "Về ngành nghề",
         "text": "Mã ngành doanh nghiệp đã có (app tự điền)",
         "type": "text", "autofill": "vsic", "version": 1,
         "data_key": "dn.ma_so_dn"},
        {"question_id": "q_ngay_sinh", "group": "Về người đại diện",
         "text": "Ngày, tháng, năm sinh?",
         "type": "date", "example": "05/12/1985", "version": 1,
         "data_key": [
             {"key": "ngay_sinh", "transform": "date_day"},
             {"key": "nam_sinh", "transform": "date_year"},
         ]},
        {"question_id": "q_ngay_cap", "group": "Về người đại diện",
         "text": "CCCD ngày cấp? (dd/mm/yyyy)",
         "type": "date", "version": 1, "data_key": "ngay_cap"},
    ],
}

# Bố cục kỳ vọng của fixture (tính tay, ghim để bắt lệch layout contract)
ROW_SECTION_1 = FIRST_BODY_ROW                     # 9  "Về doanh nghiệp"
ROW_Q_LOAI = FIRST_BODY_ROW + 1                    # 10
ROW_Q_TEN = FIRST_BODY_ROW + 3                     # 12
ROW_SECTION_2 = FIRST_BODY_ROW + 5                 # 14 "Về ngành nghề"
ROW_TABLE_HEAD = FIRST_BODY_ROW + 6                # 15
ROW_TABLE_FIRST = ROW_TABLE_HEAD + 1               # 16
ROW_TABLE_LAST = ROW_TABLE_HEAD + REPEAT_ROWS      # 25
ROW_Q_VSIC = ROW_TABLE_LAST + 2                    # 27 (cùng nhóm -> không lặp header)
ROW_SECTION_3 = ROW_TABLE_LAST + 4                 # 29 "Về người đại diện"
ROW_NGAY_SINH = ROW_TABLE_LAST + 5                 # 30
ROW_NGAY_CAP = ROW_TABLE_LAST + 7                  # 32

# Ô điền VÀNG (khách điền) của fixture — cột B, D của bảng lặp là vàng,
# cột C (q_nganh_ma, autofill) là XÁM KHÓA.
EXPECTED_GOLD = (
    ["B%d" % ROW_Q_LOAI, "B%d" % ROW_Q_TEN]
    + ["B%d" % r for r in range(ROW_TABLE_FIRST, ROW_TABLE_LAST + 1)]
    + ["D%d" % r for r in range(ROW_TABLE_FIRST, ROW_TABLE_LAST + 1)]
    + ["B%d" % ROW_NGAY_SINH, "B%d" % ROW_NGAY_CAP]
)


def must_raise(code: str, fn, *args) -> None:
    try:
        fn(*args)
    except ExcelExportError as exc:
        assert exc.code == code, "đòi mã %s, nhận %s (%s)" % (
            code, exc.code, exc)
        return
    raise AssertionError("không raise (đòi mã %s)" % code)


# ---------------------------------------------------------------------------
# Helpers đọc lại file
# ---------------------------------------------------------------------------

def write_tmp(tmp: dict, meta: dict | None = None,
              questionnaire: dict | None = None) -> Path:
    return write_questionnaire(
        questionnaire or FIXTURE, tmp["out"] / "phieu.xlsx", meta)


def open_book(path: Path):
    wb = load_workbook(str(path))
    return wb, wb[SHEET_NAME], wb[MAP_SHEET]


def unlocked_cells(ws, extra_rows: int = 20, max_col: int = 10) -> set:
    """Tất cả ô `locked=False` trong vùng quét (thân + đệm ngoài)."""
    found = set()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row + extra_rows,
                            max_col=max_col):
        for cell in row:
            if cell.protection is not None and cell.protection.locked is False:
                found.add(cell.coordinate)
    return found


def rgb6(cell) -> str:
    """Màu nền 6 ký tự (openpyxl trả '00FFF2CC' hoặc 'FF...' tùy nguồn)."""
    return str(cell.fill.fgColor.rgb)[-6:]


def validations(ws) -> dict:
    """DataValidation theo vùng áp dụng (sqref string -> object)."""
    return {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}


def map_rows(mp) -> list[dict]:
    """Đọc bảng map như X3 sẽ đọc: header theo tên cột, JSON parse when_json.

    openpyxl BỎ chuỗi rỗng khi ghi (ô không có giá trị) -> đọc lại là None;
    chuẩn hóa None -> "" cho cột chữ để điều kiện so sánh không phụ thuộc giá
    trị None của ô trống (X3 làm y như vậy).
    """
    header = [mp.cell(row=MAP_HEADER_ROW, column=c + 1).value
              for c in range(len(MAP_COLUMNS))]
    assert header == list(MAP_COLUMNS), header
    rows = []
    r = MAP_FIRST_ROW
    while mp.cell(row=r, column=1).value not in (None, ""):
        row_dict = {}
        for c, name in enumerate(MAP_COLUMNS):
            value = mp.cell(row=r, column=c + 1).value
            row_dict[name] = "" if value is None else value
        row_dict["when"] = json.loads(row_dict["when_json"])
        rows.append(row_dict)
        r += 1
    return rows


def filled_cells(ws) -> set:
    """Tất cả ô có nền quy ước (VÀNG hoặc XÁM) — tức 'ô data' của phiếu."""
    found = set()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=8):
        for cell in row:
            if rgb6(cell) in (GOLD_HEX, GRAY_HEX):
                found.add(cell.coordinate)
    return found


def gold_cells(ws) -> set:
    found = set()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=8):
        for cell in row:
            if rgb6(cell) == GOLD_HEX:
                found.add(cell.coordinate)
    return found


# ---------------------------------------------------------------------------
# [XC] Contract hàm + hằng (x3-import / X5 code đối xứng)
# ---------------------------------------------------------------------------

@check("[XC1] signature write_questionnaire(q, out_path, meta=None) -> Path")
def xc_signature(tmp):
    sig = inspect.signature(write_questionnaire)
    params = list(sig.parameters)
    assert params == ["questionnaire", "out_path", "meta"], params
    assert sig.parameters["meta"].default is None, sig.parameters["meta"]
    path = write_tmp(tmp)
    assert isinstance(path, Path), type(path)
    assert path.exists() and path.stat().st_size > 0
    assert path.suffix == ".xlsx"


@check("[XC2] import đúng dòng contract của X3/X5 + MAP_SHEET = '___map'")
def xc_constants():
    assert MAP_SHEET == "___map", MAP_SHEET
    assert SHEET_NAME == "Phieu", SHEET_NAME
    assert MARK_VALUE == "X" and REPEAT_ROWS == 10
    from tools.excel_export import MAP_SHEET as m2       # noqa: F401
    from tools.excel_export import write_questionnaire as w2  # noqa: F401
    assert m2 == "___map" and callable(w2)


@check("[XC3] không mutate questionnaire; library code không print")
def xc_pure(tmp):
    before = json.dumps(FIXTURE, ensure_ascii=False, sort_keys=True)
    write_tmp(tmp)
    after = json.dumps(FIXTURE, ensure_ascii=False, sort_keys=True)
    assert before == after, "write_questionnaire sửa dict truyền vào"
    src = (ROOT / "tools" / "excel_export.py").read_text(encoding="utf-8")
    assert "print(" not in src, "library code không được print"
    assert src.count('"""') >= 2, "thiếu module docstring"


# ---------------------------------------------------------------------------
# [XL] Bố cục sheet "Phieu"
# ---------------------------------------------------------------------------

@check("[XL1] hàng 1 tiêu đề, hàng 2 dòng hồ sơ khi có meta")
def xl_header(tmp):
    wb, ws, _ = open_book(write_tmp(tmp, meta={
        "case_id": "PLI-2026-0001", "today": "28/08/2026"}))
    assert ws["A1"].value == FIXTURE["title"], ws["A1"].value
    meta_line = ws["A2"].value
    assert "PLI-2026-0001" in meta_line and "28/08/2026" in meta_line, meta_line
    assert ws["A1"].font.bold is True

    # không có meta: hàng 2 giữ chỗ trống, bố cục các hàng sau KHÔNG trôi
    wb2, ws2, _ = open_book(write_tmp(tmp, meta=None))
    assert ws2["A1"].value == FIXTURE["title"]
    assert ws2["A2"].value in (None, "")
    assert ws2["A%d" % ROW_Q_LOAI].value == \
        ws["A%d" % ROW_Q_LOAI].value, "layout phải không phụ thuộc meta"


@check("[XL2] 5 dòng hướng dẫn: vàng/xám, không sửa ô khác, dropdown, ngày, ví dụ")
def xl_guide(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    lines = [ws.cell(row=GUIDE_FIRST_ROW + i, column=COL_LABEL).value
             for i in range(GUIDE_ROWS)]
    assert all(isinstance(x, str) and x.strip() for x in lines), lines
    joined = "\n".join(lines)
    for needle in ("VÀNG", "XÁM", "dd/mm/yyyy", "mũi tên"):
        assert needle in joined, "hướng dẫn thiếu %r: %s" % (needle, joined)
    assert "Không sửa" in joined or "không sửa" in joined, joined
    assert "VÍ DỤ" in joined or "ví dụ" in joined, joined
    # hàng ngay sau hướng dẫn là đệm, không dính nội dung
    assert ws.cell(row=GUIDE_FIRST_ROW + GUIDE_ROWS, column=COL_LABEL).value \
        in (None, "")


@check("[XL3] câu đơn: A nhãn, B ô điền, C ví dụ mờ/ghi chú, D ghi chú date")
def xl_question_rows(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    assert ws["A%d" % ROW_SECTION_1].value == "Về doanh nghiệp"
    assert ws["A%d" % ROW_Q_LOAI].value == \
        FIXTURE["questions"][0]["text"]
    assert ws["B%d" % ROW_Q_LOAI].value in (None, ""), "ô điền phải trống"
    assert ws["C%d" % ROW_Q_LOAI].value == SELECT_NOTE, \
        "select phải có ghi chú cách chọn, không gắn tiền tố 'ví dụ:'"
    assert ws["A%d" % ROW_Q_TEN].value == FIXTURE["questions"][1]["text"]
    assert ws["C%d" % ROW_Q_TEN].value == \
        "ví dụ: %s" % FIXTURE["questions"][1]["example"]
    # date có example -> C là ví dụ mờ, D là ghi chú định dạng
    assert ws["C%d" % ROW_NGAY_SINH].value == "ví dụ: 05/12/1985"
    assert ws["D%d" % ROW_NGAY_SINH].value == DATE_NOTE == "dd/mm/yyyy"
    # date không có example -> D vẫn ghi chú định dạng
    assert ws["C%d" % ROW_NGAY_CAP].value in (None, "")
    assert ws["D%d" % ROW_NGAY_CAP].value == "dd/mm/yyyy"
    # câu không phải date thì không có ghi chú định dạng
    assert ws["D%d" % ROW_Q_TEN].value in (None, "")


@check("[XL4] bảng row_group: header = text các câu, 10 hàng trống điền")
def xl_table(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    items = [q for q in FIXTURE["questions"] if q.get("row_group") == "nganh"]
    assert [x["question_id"] for x in items] == \
        ["q_nganh_ten", "q_nganh_ma", "q_nganh_chinh"]
    for offset, item in enumerate(items):
        assert ws.cell(row=ROW_TABLE_HEAD, column=COL_FILL + offset).value == \
            item["text"], "header cột %d lệch" % offset
    for row in range(ROW_TABLE_FIRST, ROW_TABLE_LAST + 1):
        for offset in range(len(items)):
            cell = ws.cell(row=row, column=COL_FILL + offset)
            assert cell.value in (None, ""), \
                "hàng điền phải trống: %s" % cell.coordinate
    # nhóm chỉ in tiêu đề MỘT lần: câu đứng sau bảng trong cùng nhóm không
    # lặp lại hàng tiêu đề nhóm
    assert ws["A%d" % ROW_Q_VSIC].value == \
        FIXTURE["questions"][5]["text"]
    assert ws["A%d" % (ROW_Q_VSIC - 1)].value in (None, ""), \
        "không được chèn thêm header nhóm"


# ---------------------------------------------------------------------------
# [XP] Quy ước màu + khóa (gate "0 ô data ngoài locked-region")
# ---------------------------------------------------------------------------

@check("[XP1] Ô VÀNG unlocked = đúng tập ô điền khai báo; Ô XÁM khóa cho autofill")
def xp_colors(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    got_gold, got_gray = [], []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=8):
        for cell in row:
            color = rgb6(cell)
            if color == GOLD_HEX:
                got_gold.append(cell.coordinate)
            elif color == GRAY_HEX:
                got_gray.append(cell.coordinate)
    assert sorted(got_gold) == sorted(EXPECTED_GOLD), \
        "ô vàng lệch: %s" % sorted(set(got_gold) ^ set(EXPECTED_GOLD))
    # cột autofill của bảng lặp + câu autofill đơn -> XÁM
    expected_gray = (
        ["C%d" % r for r in range(ROW_TABLE_FIRST, ROW_TABLE_LAST + 1)]
        + ["B%d" % ROW_Q_VSIC])
    assert sorted(got_gray) == sorted(expected_gray), \
        "ô xám lệch: %s" % sorted(set(got_gray) ^ set(expected_gray))


@check("[XP2] sheet protection bật; MỌI ô unlocked là ô điền vàng đã khai báo")
def xp_protection(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    assert ws.protection.sheet is True, "sheet phải được khóa"
    unlocked = unlocked_cells(ws)
    assert unlocked == set(EXPECTED_GOLD), \
        "gate 0-ô-lạ: unlocked thừa/thiếu %s" % sorted(unlocked ^ set(EXPECTED_GOLD))
    for coordinate in EXPECTED_GOLD:
        cell = ws[coordinate]
        assert rgb6(cell) == GOLD_HEX, coordinate
        assert cell.number_format == "@", \
            "%s phải là ô Text để round-trip giữ nguyên chuỗi" % coordinate
    # ô nhãn/ghi chú/tiêu đề/kể cả ô ngoài vùng nội dung đều khóa
    for coordinate in ("A1", "A%d" % ROW_Q_TEN, "C%d" % ROW_Q_TEN,
                       "F%d" % ROW_TABLE_HEAD, "A%d" % (ROW_NGAY_CAP + 20),
                       "H40"):
        assert ws[coordinate].protection.locked is True, coordinate


@check("[XP3] câu autofill: ô XÁM KHÓA đứng cạnh ô VÀNG mở (quy ước màu)")
def xp_autofill(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    assert ws["B%d" % ROW_Q_VSIC].protection.locked is True
    assert rgb6(ws["B%d" % ROW_Q_VSIC]) == GRAY_HEX
    for row in range(ROW_TABLE_FIRST, ROW_TABLE_LAST + 1):
        assert ws["C%d" % row].protection.locked is True, row
        assert rgb6(ws["C%d" % row]) == GRAY_HEX, row
        assert ws["B%d" % row].protection.locked is False, row


# ---------------------------------------------------------------------------
# [XS] Trình bày: màu nhóm, viền, cột rộng vừa, freeze, gridlines
# ---------------------------------------------------------------------------

@check("[XS1] mỗi nhóm một màu riêng, không trùng vàng/xám, có viền")
def xs_group_colors(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    colors = [rgb6(ws["A%d" % r]) for r in
              (ROW_SECTION_1, ROW_SECTION_2, ROW_SECTION_3)]
    assert len(set(colors)) == 3, "hai nhóm liền nhau trùng màu: %s" % colors
    for color, row in zip(colors, (ROW_SECTION_1, ROW_SECTION_2, ROW_SECTION_3)):
        assert color not in (GOLD_HEX, GRAY_HEX), row
        assert color in GROUP_PALETTE, color
        assert ws["A%d" % row].font.bold is True, row
        assert ws["A%d" % row].border.left.style is not None, "thiếu viền"
    assert ws["B%d" % ROW_Q_TEN].border.left.style is not None
    assert ws["B%d" % ROW_Q_TEN].border.right.style is not None
    # header bảng lặp mang màu nhóm của nó (nhóm 2)
    assert rgb6(ws.cell(row=ROW_TABLE_HEAD, column=COL_FILL)) == colors[1]


@check("[XS2] cột rộng vừa + wrap text + freeze panes + ẩn gridlines")
def xs_readability(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    widths = {}
    for letter in ("A", "B", "C", "D", "E", "F", "G"):
        dim = ws.column_dimensions.get(letter)
        if dim is not None and dim.width:
            widths[letter] = dim.width
    assert set(("A", "B", "C", "D")) <= set(widths), widths
    # vùng đọc thường xuyên A..D phải vừa màn hình thường (không cuộn ngang);
    # tổng cả bảng (E..G chỉ dùng khi row_group > 3 cột) cũng bị chặn trên.
    main_width = sum(widths[l] for l in ("A", "B", "C", "D"))
    assert main_width <= 120, \
        "A..D quá rộng (%.0f) -> cuộn ngang: %s" % (main_width, widths)
    assert sum(widths.values()) <= 170, widths
    assert ws["A%d" % ROW_Q_TEN].alignment.wrap_text is True
    assert ws["C%d" % ROW_Q_TEN].alignment.wrap_text is True
    assert ws.cell(row=ROW_TABLE_HEAD, column=COL_FILL).alignment.wrap_text \
        is True
    assert ws.freeze_panes == "A%d" % FIRST_BODY_ROW, ws.freeze_panes
    assert ws.sheet_view.showGridLines is False
    # font mờ cho ví dụ
    example = ws["C%d" % ROW_Q_TEN]
    assert example.font.italic is True, "ví dụ phải là chữ mờ"
    assert example.font.color.rgb[-6:] == NOTE_FONT_HEX, \
        "ví dụ phải màu xám %s, nhận %s" % (NOTE_FONT_HEX, example.font.color.rgb)


# ---------------------------------------------------------------------------
# [XD] Dropdown chọn-là-chọn
# ---------------------------------------------------------------------------

@check("[XD1] select -> DataValidation list từ options[].label, mũi tên còn")
def xd_select(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    dv_by_sqref = validations(ws)
    coordinate = "B%d" % ROW_Q_LOAI
    matches = [dv for sq, dv in dv_by_sqref.items() if coordinate in sq]
    assert len(matches) == 1, "ô select phải có đúng 1 validation: %s" % dv_by_sqref
    dv = matches[0]
    assert dv.type == "list"
    labels = [o["label"] for o in FIXTURE["questions"][0]["options"]]
    assert dv.formula1 == '"%s"' % ",".join(labels), dv.formula1
    assert dv.allow_blank is True, "phải cho phép bỏ trống"
    assert dv.showDropDown is not True, \
        "showDropDown=True là mất mũi tên trong Excel (OOXML nghịch dấu)"
    assert dv.showErrorMessage is True


@check("[XD2] mark -> danh sách chỉ 'X'; bảng lặp áp validation cả cột")
def xd_mark_and_range(tmp):
    wb, ws, _ = open_book(write_tmp(tmp))
    dv_by_sqref = validations(ws)
    mark_range = "D%d:D%d" % (ROW_TABLE_FIRST, ROW_TABLE_LAST)
    assert mark_range in dv_by_sqref, sorted(dv_by_sqref)
    dv = dv_by_sqref[mark_range]
    assert dv.formula1 == '"X"', dv.formula1
    assert dv.allow_blank is True
    # cột select trong bảng (nếu có) cũng phải phủ đủ 10 hàng — dùng cột D
    gold_range = "B%d:B%d" % (ROW_TABLE_FIRST, ROW_TABLE_LAST)
    assert gold_range not in dv_by_sqref, \
        "cột text gõ tay không được bị khóa bằng dropdown: %s" % sorted(dv_by_sqref)


# ---------------------------------------------------------------------------
# [XM] Sheet map ẩn cho X3
# ---------------------------------------------------------------------------

@check("[XM1] sheet ẩn '___map': marker + JSON head (doc_id, version, sheet)")
def xm_header(tmp):
    wb, ws, mp = open_book(write_tmp(tmp, meta={"case_id": "PLI-1"}))
    assert wb.sheetnames == [SHEET_NAME, MAP_SHEET], wb.sheetnames
    assert mp.sheet_state == "hidden", mp.sheet_state
    assert mp["A1"].value == MAP_MARKER, mp["A1"].value
    head = json.loads(mp["B1"].value)
    assert head["doc_id"] == "X2_DEMO"
    assert head["questionnaire_version"] == 1
    assert head["map_version"] == MAP_VERSION == 1
    assert head["sheet"] == SHEET_NAME
    assert head["case_id"] == "PLI-1"
    assert head["first_body_row"] == FIRST_BODY_ROW
    assert mp.protection.sheet is True, "map cũng phải khóa"


@check("[XM2] một dòng map cho MỖI target key; khi/transform/row_group đúng")
def xm_rows(tmp):
    wb, ws, mp = open_book(write_tmp(tmp))
    rows = map_rows(mp)
    expected = 0
    for q in FIXTURE["questions"]:
        dk = q["data_key"]
        expected += len(dk) if isinstance(dk, list) else 1
    assert len(rows) == expected == 10, "%d != %d" % (len(rows), expected)

    by_id = {r["question_id"]: r for r in rows}
    loai = by_id["q_loai"]
    assert loai["data_key"] == "dn.loai_hinh" and loai["type"] == "select"
    assert loai["cell"] == "B%d" % ROW_Q_LOAI
    assert loai["cell_row"] == ROW_Q_LOAI and loai["row_last"] == ROW_Q_LOAI
    assert loai["layout"] == "single" and loai["cell_col"] == COL_FILL
    assert loai["when"] == {} and loai["transform"] == ""
    assert loai["required"] == "yes" and loai["repeatable"] == ""

    # câu 2 target: 2 dòng, when_json parse lại đúng, thứ tự giữ theo file
    ngay = [r for r in rows if r["question_id"] == "q_ngay_sinh"]
    assert [r["data_key"] for r in ngay] == ["ngay_sinh", "nam_sinh"]
    assert [r["transform"] for r in ngay] == ["date_day", "date_year"]
    assert all(r["when"] == {} for r in ngay)
    assert all(r["cell"] == "B%d" % ROW_NGAY_SINH for r in ngay)

    # target có when: JSON string ổn định (sort_keys) và đọc lại thành dict
    ten = [r for r in rows if r["question_id"] == "q_nganh_ten"]
    assert len(ten) == 2
    assert [r["when"] for r in ten] == [
        {"dn.loai_hinh": ["DNTN"]}, {"dn.loai_hinh": ["CTCP"]}]
    assert all(r["layout"] == "table" for r in ten)
    assert all(r["row_group"] == "nganh" for r in ten)
    assert all(r["repeatable"] == "yes" for r in ten)
    assert all(r["multi_key"] == "yes" for r in ten)
    # ô điền của bảng = hàng dữ liệu ĐẦU TIÊN, row_last = hàng thứ REPEAT_ROWS
    assert ten[0]["cell"] == "B%d" % ROW_TABLE_FIRST
    assert ten[0]["cell_row"] == ROW_TABLE_FIRST
    assert ten[0]["row_last"] == ROW_TABLE_LAST
    # cột xám autofill cũng được map (X3 biết cột nào app tự điền)
    ma = [r for r in rows if r["question_id"] == "q_nganh_ma"][0]
    assert ma["cell_col"] == COL_FILL + 1 and ma["autofill"] == "vsic"
    assert ma["cell"] == "C%d" % ROW_TABLE_FIRST


@check("[XM3] map KHỚP 2 CHIỀU với ô data (gate 0 ô data ngoài locked-region)")
def xm_two_way(tmp):
    wb, ws, mp = open_book(write_tmp(tmp))
    rows = map_rows(mp)
    mapped = set()
    for r in rows:
        col = r["cell_col"]
        for row in range(r["cell_row"], r["row_last"] + 1):
            mapped.add("%s%d" % (chr(ord("A") + col - 1), row))
    # (1) mọi ô có nền quy ước (vàng + xám) phải được khai báo trong map
    assert filled_cells(ws) == mapped, \
        "ô data ngoài map: %s | ô map không có nền: %s" % (
            sorted(filled_cells(ws) - mapped), sorted(mapped - filled_cells(ws)))
    # (2) ô unlocked phải là TẬP CON vàng của map (ô xám = app tự làm, khóa)
    unlocked = unlocked_cells(ws)
    assert unlocked <= mapped, "ô unlocked ngoài map: %s" % sorted(unlocked - mapped)
    assert unlocked == gold_cells(ws), "ô unlocked != ô vàng: %s" % sorted(
        unlocked ^ gold_cells(ws))


# ---------------------------------------------------------------------------
# [XR] Round-trip + deterministic
# ---------------------------------------------------------------------------

@check("[XR1] đọc lại file: text tiếng Việt, protection, fill, dropdown, ẩn")
def xr_roundtrip(tmp):
    path = write_tmp(tmp, meta={"case_id": "PLI-2", "today": "01/01/2027"})
    wb, ws, mp = open_book(path)
    assert ws["A1"].value == FIXTURE["title"]
    assert ws["C%d" % ROW_Q_TEN].value.endswith("Minh Long")
    assert ws["D%d" % ROW_NGAY_SINH].value == "dd/mm/yyyy"
    assert ws.protection.sheet is True
    assert ws["B%d" % ROW_Q_TEN].protection.locked is False
    assert rgb6(ws["B%d" % ROW_Q_TEN]) == GOLD_HEX
    assert mp.sheet_state == "hidden"
    # fixture có đúng 2 ô chọn-là-chọn: q_loai (select) và cột mark của bảng
    assert len(ws.data_validations.dataValidation) == 2, \
        [str(dv.sqref) for dv in ws.data_validations.dataValidation]
    # ô number_format='@' giúp Excel giữ nguyên chuỗi khách gõ — mô phỏng:
    ws["B%d" % ROW_Q_TEN] = "Công ty TNHH Minh Long"
    ws["B%d" % ROW_NGAY_SINH] = "05/12/1985"
    again = tmp["out"] / "phieu_filled.xlsx"
    wb.save(str(again))
    wb2, ws2, _ = open_book(again)
    assert ws2["B%d" % ROW_Q_TEN].value == "Công ty TNHH Minh Long"
    assert ws2["B%d" % ROW_NGAY_SINH].value == "05/12/1985", \
        "giá trị chuỗi phải giữ nguyên khi đọc lại"


@check("[XR2] ghi 2 lần -> cùng cấu trúc (không random, không trôi hàng)")
def xr_deterministic(tmp):
    p1 = write_questionnaire(FIXTURE, tmp["out"] / "a.xlsx")
    p2 = write_questionnaire(FIXTURE, tmp["out"] / "b.xlsx")
    wb1, ws1, mp1 = open_book(p1)
    wb2, ws2, mp2 = open_book(p2)
    assert ws1.max_row == ws2.max_row and ws1.max_column == ws2.max_column
    for row in ws1.iter_rows(min_row=1, max_row=ws1.max_row, max_col=8):
        for cell in row:
            other = ws2[cell.coordinate]
            assert cell.value == other.value, cell.coordinate
            assert rgb6(cell) == rgb6(other), cell.coordinate
            assert cell.protection.locked == other.protection.locked, \
                cell.coordinate
    assert map_rows(mp1) == map_rows(mp2)


# ---------------------------------------------------------------------------
# [XN] Negative fail-loud
# ---------------------------------------------------------------------------

def _mutated(**changes) -> dict:
    q = json.loads(json.dumps(FIXTURE))
    q.update(changes)
    return q


@check("[XN1] khung sai: NOT_OBJECT / MISSING_FIELD / NO_QUESTIONS / DUP_ID")
def xn_structure(tmp):
    out = tmp["out"] / "neg.xlsx"
    must_raise("NOT_OBJECT", write_questionnaire, [], out)          # type: ignore
    must_raise("MISSING_FIELD", write_questionnaire, {"questions": []}, out)
    must_raise("MISSING_FIELD", write_questionnaire, _mutated(title="   "), out)
    must_raise("NO_QUESTIONS", write_questionnaire, _mutated(questions=[]), out)
    dup = json.loads(json.dumps(FIXTURE))
    dup["questions"][1]["question_id"] = dup["questions"][0]["question_id"]
    must_raise("DUP_ID", write_questionnaire, dup, out)
    assert not out.exists(), "fail-loud phải xảy ra TRƯỚC khi ghi file"


@check("[XN2] câu sai: BAD_TYPE / MISSING_FIELD / BAD_QUESTION / BAD_ROW_GROUP")
def xn_question(tmp):
    out = tmp["out"] / "neg2.xlsx"
    bad = json.loads(json.dumps(FIXTURE))
    bad["questions"][1]["type"] = "sobolan"
    must_raise("BAD_TYPE", write_questionnaire, bad, out)
    no_text = json.loads(json.dumps(FIXTURE))
    del no_text["questions"][1]["text"]
    must_raise("MISSING_FIELD", write_questionnaire, no_text, out)
    no_group = json.loads(json.dumps(FIXTURE))
    del no_group["questions"][1]["group"]
    must_raise("MISSING_FIELD", write_questionnaire, no_group, out)
    bad_rg = json.loads(json.dumps(FIXTURE))
    bad_rg["questions"][2]["row_group"] = ""
    must_raise("BAD_ROW_GROUP", write_questionnaire, bad_rg, out)


@check("[XN3] dropdown sai: NO_OPTIONS / BAD_OPTION_LABEL / LIST_TOO_LONG")
def xn_dropdown(tmp):
    out = tmp["out"] / "neg3.xlsx"
    no_opt = json.loads(json.dumps(FIXTURE))
    del no_opt["questions"][0]["options"]
    must_raise("NO_OPTIONS", write_questionnaire, no_opt, out)
    comma = json.loads(json.dumps(FIXTURE))
    comma["questions"][0]["options"][0]["label"] = "Doanh nghiệp, tư nhân"
    must_raise("BAD_OPTION_LABEL", write_questionnaire, comma, out)
    dup_label = json.loads(json.dumps(FIXTURE))
    dup_label["questions"][0]["options"][1]["label"] = \
        dup_label["questions"][0]["options"][0]["label"]
    must_raise("BAD_OPTION_LABEL", write_questionnaire, dup_label, out)
    long_q = json.loads(json.dumps(FIXTURE))
    long_q["questions"][0]["options"] = [
        {"value": "V%d" % i, "label": "Nhãn rất dài %d %s" % (i, "x" * 40)}
        for i in range(10)]
    must_raise("LIST_TOO_LONG", write_questionnaire, long_q, out)
    # mark KHÔNG cần options (mặc định 'X')
    no_mark_opt = json.loads(json.dumps(FIXTURE))
    del no_mark_opt["questions"][4]["options"]
    path = write_questionnaire(no_mark_opt, out)
    wb, ws, _ = open_book(path)
    dv_by_sqref = validations(ws)
    mark = [dv for sq, dv in dv_by_sqref.items()
            if "D%d" % ROW_TABLE_FIRST in sq]
    assert mark and mark[0].formula1 == '"X"', mark


@check("[XN4] bảng quá rộng -> TOO_MANY_COLUMNS; meta sai -> BAD_META")
def xn_limits(tmp):
    out = tmp["out"] / "neg4.xlsx"
    wide = json.loads(json.dumps(FIXTURE))
    base = wide["questions"][2]
    extra = []
    for i in range(6):                       # 1 + 6 = 7 câu trong 1 row_group
        item = json.loads(json.dumps(base))
        item["question_id"] = "q_extra_%d" % i
        item["data_key"] = [{"key": "k.%d" % i}]
        extra.append(item)
    wide["questions"][2:3] = [base] + extra
    must_raise("TOO_MANY_COLUMNS", write_questionnaire, wide, out)

    must_raise("BAD_META", write_questionnaire, FIXTURE, out, {"case_id": 123})
    must_raise("BAD_META", write_questionnaire, FIXTURE, out, ["x"])
    # meta dư key lạ: bỏ qua, không lỗi
    path = write_questionnaire(FIXTURE, out, {"case_id": "C1", "today": "02/03/2027",
                                              "extra": "bỏ qua"})
    wb, ws, _ = open_book(path)
    assert "C1" in ws["A2"].value and "02/03/2027" in ws["A2"].value


# ---------------------------------------------------------------------------
# [X2-M72] Tích hợp với artifact questionnaire thật (read-only)
# ---------------------------------------------------------------------------

@check("[X2-M72] PLI_M72 thật: 4 nhóm màu, dropdown 5 lựa chọn, map đủ target")
def x2_pli_m72(tmp):
    artifact = ROOT / "app" / "data" / "questionnaire" / "PLI_M72.json"
    questionnaire = json.loads(artifact.read_text(encoding="utf-8"))
    assert validate(questionnaire, json.loads(
        (ROOT / "app" / "data" / "form_maps.json").read_text(encoding="utf-8")
    )) == [], "artifact X1 phải sạch trước khi xuất"

    path = write_questionnaire(questionnaire, tmp["out"] / "PLI_M72.xlsx",
                               meta={"case_id": "PLI-2026-0007",
                                     "today": "28/08/2026"})
    wb, ws, mp = open_book(path)
    assert ws["A1"].value == questionnaire["title"]
    assert "PLI-2026-0007" in ws["A2"].value

    # tổng target key của phiếu thật = số dòng map
    rows = map_rows(mp)
    expected = sum(
        len(q["data_key"]) if isinstance(q["data_key"], list) else 1
        for q in questionnaire["questions"])
    assert len(rows) == expected, "%d != %d" % (len(rows), expected)

    # dropdown loại hình: đủ 5 label theo thứ tự artifact
    loai_row = next(r for r in rows if r["data_key"] == "dn.loai_hinh")
    dv_by_sqref = validations(ws)
    dv = dv_by_sqref["B%d" % loai_row["cell_row"]]
    labels = [o["label"] for o in questionnaire["questions"][0]["options"]]
    assert dv.formula1 == '"%s"' % ",".join(labels), dv.formula1

    # bảng ngành: 3 cột (ten/ma/chinh), 10 hàng, cột ma là XÁM KHÓA
    nganh = [r for r in rows if r["row_group"] == "nganh"]
    cols = sorted({r["cell_col"] for r in nganh})
    assert cols == [2, 3, 4], cols
    first, last = nganh[0]["cell_row"], nganh[0]["row_last"]
    assert last - first + 1 == REPEAT_ROWS
    ma_cols = {r["cell_col"] for r in nganh if r["autofill"] == "vsic"}
    assert len(ma_cols) == 1, ma_cols
    for row in range(first, last + 1):
        for col in cols:
            cell = ws.cell(row=row, column=col)
            if col in ma_cols:
                assert cell.protection.locked is True and \
                    rgb6(cell) == GRAY_HEX, cell.coordinate
            else:
                assert cell.protection.locked is False and \
                    rgb6(cell) == GOLD_HEX, cell.coordinate

    # gate 2 chiều trên toàn phiếu thật: map phủ đủ ô data, unlocked = ô vàng
    mapped = set()
    for r in rows:
        for row in range(r["cell_row"], r["row_last"] + 1):
            mapped.add("%s%d" % (chr(ord("A") + r["cell_col"] - 1), row))
    assert filled_cells(ws) == mapped, "ô data và map lệch nhau: %s" % sorted(
        filled_cells(ws) ^ mapped)
    assert unlocked_cells(ws) <= mapped, "ô unlocked ngoài map"

    # câu date của phiếu thật có ghi chú dd/mm/yyyy
    for r in rows:
        if r["type"] == "date":
            assert ws.cell(row=r["cell_row"], column=COL_DATE_NOTE).value == \
                "dd/mm/yyyy", r["question_id"]


# ---------------------------------------------------------------------------
# [XPIN] Pin openpyxl trong requirements-dev.txt
# ---------------------------------------------------------------------------

@check("[XPIN] requirements-dev.txt pin openpyxl==3.1.5 kèm ghi chú wave X2")
def xpin():
    text = (ROOT / "requirements-dev.txt").read_text(encoding="utf-8")
    assert "openpyxl==3.1.5" in text, "thiếu pin openpyxl"
    assert "X2" in text, "phải ghi rõ Excel wave X2 trong mục mới"
    # version pin khớp version đang chạy (test thật sự chạy trên bản đó)
    import openpyxl
    assert openpyxl.__version__ == "3.1.5", openpyxl.__version__
    # pin phải đứng ở mục riêng, không lẫn vào phần python-docx
    lines = text.splitlines()
    pin_line = next(i for i, ln in enumerate(lines)
                    if ln.strip() == "openpyxl==3.1.5")
    head = "\n".join(lines[:pin_line])
    assert "X2" in head.split("python-docx==1.2.0")[-1], \
        "mục pin openpyxl phải có ghi chú wave X2 ngay phía trên"


# ---------------------------------------------------------------------------

def main() -> int:
    tmp_root = Path(tempfile.mkdtemp(prefix="x2_export_"))
    tmp = {"out": tmp_root / "generated"}
    tmp["out"].mkdir(parents=True, exist_ok=True)
    tests = [
        lambda: xc_signature(tmp),
        xc_constants,
        lambda: xc_pure(tmp),
        lambda: xl_header(tmp),
        lambda: xl_guide(tmp),
        lambda: xl_question_rows(tmp),
        lambda: xl_table(tmp),
        lambda: xp_colors(tmp),
        lambda: xp_protection(tmp),
        lambda: xp_autofill(tmp),
        lambda: xs_group_colors(tmp),
        lambda: xs_readability(tmp),
        lambda: xd_select(tmp),
        lambda: xd_mark_and_range(tmp),
        lambda: xm_header(tmp),
        lambda: xm_rows(tmp),
        lambda: xm_two_way(tmp),
        lambda: xr_roundtrip(tmp),
        lambda: xr_deterministic(tmp),
        lambda: xn_structure(tmp),
        lambda: xn_question(tmp),
        lambda: xn_dropdown(tmp),
        lambda: xn_limits(tmp),
        lambda: x2_pli_m72(tmp),
        xpin,
    ]
    import shutil as _shutil
    try:
        for test in tests:
            test()
    finally:
        _shutil.rmtree(tmp_root, ignore_errors=True)
    passed = sum(1 for ok, _ in _results if ok)
    failed = len(_results) - passed
    print("-" * 64)
    print("test_excel_export: %d PASS / %d FAIL" % (passed, failed))
    for ok, name in _results:
        if not ok:
            print(" - FAIL:", name)
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
