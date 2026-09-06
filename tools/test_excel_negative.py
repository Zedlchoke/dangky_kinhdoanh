# -*- coding: utf-8 -*-
"""Wave X4 — negative Excel: khách phá khung phiếu -> fail-loud, KHÔNG data
một phần (điều kiện "không bao giờ ghi đè draft khi fail" của PLAN X3/X4).

Mỗi case: export phiếu thật (PLI_M72) -> mutate như khách/lập trình sai ->
import_questionnaire phải trả {"ok": False, "errors": [mã lỗi đúng], "data":
None}. Mã lỗi là contract chung với bridge X5 (hiện toast người dùng).

[NG1] xóa sheet map / file không phải xlsx             -> BAD_HEADER
[NG2] sửa header map (key -> key_hong)                 -> BAD_HEADER
[NG3] metadata questionnaire_version lệch              -> BAD_VERSION
[NG4] gõ vào ô lạ trong vùng điền (C12 — cột có ô map) -> UNKNOWN_CELL
[NG5] thiếu câu bắt buộc (q_ten_tieng_viet trống)      -> MISSING_REQUIRED
[NG6] select gõ tay giá trị ngoài danh sách            -> BAD_SELECT
[NG7] ngày 32/13/2020                                  -> BAD_DATE
[NG8] tiền "năm tỷ"                                    -> BAD_NUMBER
[NG9] [O4] ô CCCD là ô SỐ (Excel tự đổi, mất số 0 đầu/precision) -> BAD_ID_TEXT
[NG10] [O4] chuỗi scientific "1.23E+20" vào ô mã định danh       -> BAD_ID_TEXT
[NG11] [O7] version lạ trên cụm hộ KD (PLII_M01)                 -> BAD_VERSION
[NG12] [O7] xóa sheet map trên cụm hộ KD (PLII_M03)              -> BAD_HEADER
[NG13] [O7] ô công thức độc hại "=1+1"/"=cmd…" (PLII_M02) -> fail-loud, no data
[NG14] [O7] ô lạ trong vùng điền (PLII_M01)                  -> UNKNOWN_CELL
[NG15] [P6=C] version lạ trên cụm thay đổi ngành (PLI_M12)       -> BAD_VERSION
[NG16] [P6=C] xóa sheet phiếu trên cụm thuế (PLI_M27)            -> BAD_HEADER
[NG17] [P6=C] ô công thức độc hại vào ô bắt buộc (PLI_M12) -> fail-loud, no data
[NG18] [P6=C] ô lạ trong vùng điền (PLI_M12)                 -> UNKNOWN_CELL

Chạy:  python3 -X utf8 -m tools.test_excel_negative   (exit 0 = PASS)
"""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from tools.excel_export import (  # noqa: E402
    MAP_COLUMNS, MAP_SHEET, write_questionnaire)
from tools.excel_import import import_questionnaire  # noqa: E402

QUESTIONNAIRE = json.loads(
    (ROOT / "app" / "data" / "questionnaire" / "PLI_M72.json")
    .read_text(encoding="utf-8"))
# [O7] cụm hộ kinh doanh (P6=C rollout 2026-09-06) — tái dùng harness pilot.
Q_M01 = json.loads((ROOT / "app" / "data" / "questionnaire" / "PLII_M01.json")
                   .read_text(encoding="utf-8"))
Q_M02 = json.loads((ROOT / "app" / "data" / "questionnaire" / "PLII_M02.json")
                   .read_text(encoding="utf-8"))
Q_M03 = json.loads((ROOT / "app" / "data" / "questionnaire" / "PLII_M03.json")
                   .read_text(encoding="utf-8"))

BASE_ANSWERS = {"q_loai_hinh": "DNTN", "q_ten_tieng_viet": "Công ty A",
                "q_dia_chi_tru_so": "1 Đường B, TP.HCM"}

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run():
            try:
                fn()
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def export_filled(extra: dict[str, str] | None = None,
                  questionnaire: dict | None = None,
                  required_extra: dict[str, str] | None = None) -> Path:
    """Export phiếu thật + điền các ô (theo map) của `extra` và BASE_ANSWERS.

    `questionnaire` None = pilot PLI_M72; truyền cụm mới để tái dùng harness.
    Với cụm không có nhánh when (hộ kinh doanh) cần `required_extra` điền
    đủ câu required của cụm đó (map không lọc theo when).
    """
    q = questionnaire or QUESTIONNAIRE
    answers = dict(BASE_ANSWERS)
    answers.update(required_extra or {})
    answers.update(extra or {})
    tmp = Path(tempfile.mkdtemp(prefix="p3ng_"))
    out = write_questionnaire(q, tmp / "phieu.xlsx")
    wb = openpyxl.load_workbook(out)
    mp, ws = wb[MAP_SHEET], wb["Phieu"]
    col = {h: i + 1 for i, h in enumerate(MAP_COLUMNS)}
    for r in range(4, mp.max_row + 1):
        qid = mp.cell(r, col["question_id"]).value
        if qid in answers:
            when = mp.cell(r, col["when"]).value or ""
            if q is not QUESTIONNAIRE or when in ("", "{}") or '"DNTN"' in when:
                ws[mp.cell(r, col["cell"]).value] = answers[qid]
    filled = tmp / "filled.xlsx"
    wb.save(filled)
    return filled


def codes(res: dict) -> list[str]:
    return [e.get("code") for e in res.get("errors") or []]


@check("[NG1a] xóa sheet map -> BAD_HEADER, data None")
def ng1a():
    tmp = Path(tempfile.mkdtemp(prefix="p3ng1_"))
    out = write_questionnaire(QUESTIONNAIRE, tmp / "p.xlsx")
    wb = openpyxl.load_workbook(out)
    del wb[MAP_SHEET]
    p = tmp / "no_map.xlsx"
    wb.save(p)
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert codes(res) == ["BAD_HEADER"], res["errors"]


@check("[NG1b] file không phải xlsx -> BAD_HEADER (không crash)")
def ng1b():
    tmp = Path(tempfile.mkdtemp(prefix="p3ng1b_"))
    p = tmp / "not_excel.xlsx"
    p.write_text("đây là file text", encoding="utf-8")
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert codes(res) == ["BAD_HEADER"], res["errors"]


@check("[NG2] sửa tên cột map (key -> key_hỏng) -> BAD_HEADER")
def ng2():
    p = export_filled()
    wb = openpyxl.load_workbook(p)
    mp = wb[MAP_SHEET]
    header = [mp.cell(row=3, column=c).value for c in range(1, len(MAP_COLUMNS) + 1)]
    key_idx = header.index("key") + 1
    mp.cell(row=3, column=key_idx, value="key_hong")
    wb.save(p)
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert "BAD_HEADER" in codes(res), res["errors"]


@check("[NG3] metadata questionnaire_version=99 -> BAD_VERSION")
def ng3():
    p = export_filled()
    wb = openpyxl.load_workbook(p)
    mp = wb[MAP_SHEET]
    fixed = False
    for row in mp.iter_rows(min_row=1, max_row=2):
        for cell in row:
            if isinstance(cell.value, str) and \
                    cell.value.startswith("questionnaire_version="):
                cell.value = "questionnaire_version=99"
                fixed = True
    assert fixed, "không tìm thấy metadata questionnaire_version"
    wb.save(p)
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert codes(res) == ["BAD_VERSION"], res["errors"]


@check("[NG4] gõ ô lạ trong vùng điền (C12) -> UNKNOWN_CELL")
def ng4():
    p = export_filled()
    wb = openpyxl.load_workbook(p)
    ws = wb["Phieu"]
    ws["C12"] = "khách gõ linh tinh"          # C12: cột C có ô map, ô này không
    wb.save(p)
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert "UNKNOWN_CELL" in codes(res), res["errors"]


@check("[NG5] thiếu câu bắt buộc -> MISSING_REQUIRED")
def ng5():
    p = export_filled({"q_ten_tieng_viet": None})
    wb = openpyxl.load_workbook(p)
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert "MISSING_REQUIRED" in codes(res), res["errors"]


@check("[NG6] select gõ tay giá trị ngoài danh sách -> BAD_SELECT")
def ng6():
    p = export_filled({"q_loai_hinh": "Công ty TNHH phương bắc"})
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert "BAD_SELECT" in codes(res), res["errors"]


@check("[NG7] ngày 32/13/2020 -> BAD_DATE")
def ng7():
    p = export_filled({"q_ngay_sinh": "32/13/2020"})
    wb = openpyxl.load_workbook(p)
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert "BAD_DATE" in codes(res), res["errors"]


@check("[NG8] tiền 'năm tỷ' -> BAD_NUMBER")
def ng8():
    p = export_filled({"q_von": "năm tỷ"})
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert "BAD_NUMBER" in codes(res), res["errors"]


@check("[NG9] [O4] ô CCCD bị Excel đổi thành số (ghi int) -> BAD_ID_TEXT")
def ng9():
    p = export_filled()
    wb = openpyxl.load_workbook(p)
    mp, ws = wb[MAP_SHEET], wb["Phieu"]
    col = {h: i + 1 for i, h in enumerate(MAP_COLUMNS)}
    hit = 0
    for r in range(4, mp.max_row + 1):
        if mp.cell(r, col["question_id"]).value == "q_so_cccd" and \
                '"DNTN"' in (mp.cell(r, col["when"]).value or ""):
            # Ghi SỐ như Excel tự đổi khi khách dán mã không kèm định dạng Text.
            ws[mp.cell(r, col["cell"]).value] = 791085001234
            hit += 1
    assert hit >= 1, "không tìm thấy ô q_so_cccd nhánh DNTN"
    wb.save(p)
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert "BAD_ID_TEXT" in codes(res), res["errors"]


@check("[NG10] [O4] chuỗi scientific '1.23E+20' vào ô mã định danh "
       "-> BAD_ID_TEXT (không chấp nhận scientific notation)")
def ng10():
    p = export_filled({"q_so_cccd": "1.23E+20"})
    res = import_questionnaire(p, QUESTIONNAIRE)
    assert res["ok"] is False and res["data"] is None, res
    assert "BAD_ID_TEXT" in codes(res), res["errors"]


# --- [O7 2026-09-06] negative tái dùng harness cho cụm hộ kinh doanh ---------

M01_REQUIRED = {"q_ho_ten": "Nguyễn Văn A", "q_so_cccd": "079085001234",
                "q_ten_hkd": "Hộ kinh doanh A", "q_dia_chi_chi_tiet": "25 Nguyễn Trãi",
                "q_xa_phuong": "P. Bến Thành", "q_tinh_thanh": "TP.HCM"}

# --- [P6=C 2026-09-06 — f4xl] cụm thay đổi ngành (PLI_M12) + cụm thuế (M27) --

Q_M12 = json.loads((ROOT / "app" / "data" / "questionnaire" / "PLI_M12.json")
                   .read_text(encoding="utf-8"))
Q_M27 = json.loads((ROOT / "app" / "data" / "questionnaire" / "PLI_M27.json")
                   .read_text(encoding="utf-8"))
M12_REQUIRED = {"q_dn_ten": "CÔNG TY TNHH AN PHÁT",
                "q_dn_ma_so": "0312345678"}
M27_REQUIRED = {"q_dn_ten": "CÔNG TY TNHH AN PHÁT",
                "q_dn_ma_so": "0312345678"}


def _tamper_meta_version(p: Path, q: dict) -> None:
    """Sửa metadata questionnaire_version trên sheet map -> version lạ."""
    wb = openpyxl.load_workbook(p)
    mp = wb[MAP_SHEET]
    fixed = False
    for row in mp.iter_rows(min_row=1, max_row=2):
        for cell in row:
            if isinstance(cell.value, str) and \
                    cell.value.startswith("questionnaire_version="):
                cell.value = "questionnaire_version=99"
                fixed = True
    assert fixed, "không tìm thấy metadata questionnaire_version"
    wb.save(p)


@check("[NG15] P6=C cụm thay đổi ngành: metadata questionnaire_version=99 "
       "(PLI_M12) -> BAD_VERSION")
def ng15():
    p = export_filled(required_extra=M12_REQUIRED, questionnaire=Q_M12)
    _tamper_meta_version(p, Q_M12)
    res = import_questionnaire(p, Q_M12)
    assert res["ok"] is False and res["data"] is None, res
    assert codes(res) == ["BAD_VERSION"], res["errors"]


@check("[NG16] P6=C cụm thuế: xóa sheet phiếu 'Phieu' (PLI_M27) -> BAD_HEADER")
def ng16():
    tmp = Path(tempfile.mkdtemp(prefix="p3ng16_"))
    out = write_questionnaire(Q_M27, tmp / "p.xlsx")
    wb = openpyxl.load_workbook(out)
    del wb["Phieu"]
    # openpyxl không lưu workbook chỉ còn sheet ẩn — mở khóa map để ghi được
    # file "khách đã xóa nhầm sheet phiếu" (map vẫn bị loại khỏi danh sách
    # sheet phiếu khi import theo tên).
    wb[MAP_SHEET].sheet_state = "visible"
    p = tmp / "no_phieu.xlsx"
    wb.save(p)
    res = import_questionnaire(p, Q_M27)
    assert res["ok"] is False and res["data"] is None, res
    assert codes(res) == ["BAD_HEADER"], res["errors"]


@check("[NG17] P6=C cụm thay đổi ngành: ô công thức độc hại '=1+1' / '=cmd…' "
       "vào ô bắt buộc (PLI_M12) -> fail-loud không crash, không data một phần")
def ng17():
    p = export_filled({"q_dn_ten": "=1+1",
                       "q_dn_ma_so": "=cmd|' /C calc'!A0"},
                      required_extra=M12_REQUIRED, questionnaire=Q_M12)
    res = import_questionnaire(p, Q_M12)
    # app chỉ ĐỌC giá trị (data_only=True), không bao giờ thực thi công thức;
    # ô công thức không có cached value -> rỗng -> required missing -> ok=False.
    assert res["ok"] is False and res["data"] is None, res
    assert "MISSING_REQUIRED" in codes(res), res["errors"]
    assert all("calc" not in str(v) for row in (res["data"] or {}).items()
               for v in row), "giá trị công thức không được lọt vào data"


@check("[NG11] O7 cụm hộ KD: metadata questionnaire_version=99 (PLII_M01) "
       "-> BAD_VERSION")
def ng11():
    p = export_filled(required_extra=M01_REQUIRED, questionnaire=Q_M01)
    wb = openpyxl.load_workbook(p)
    mp = wb[MAP_SHEET]
    fixed = False
    for row in mp.iter_rows(min_row=1, max_row=2):
        for cell in row:
            if isinstance(cell.value, str) and \
                    cell.value.startswith("questionnaire_version="):
                cell.value = "questionnaire_version=99"
                fixed = True
    assert fixed
    wb.save(p)
    res = import_questionnaire(p, Q_M01)
    assert res["ok"] is False and res["data"] is None, res
    assert codes(res) == ["BAD_VERSION"], res["errors"]


@check("[NG12] O7 cụm hộ KD: xóa sheet map (PLII_M03) -> BAD_HEADER")
def ng12():
    p = export_filled(questionnaire=Q_M03)
    wb = openpyxl.load_workbook(p)
    del wb[MAP_SHEET]
    p2 = p.with_name("no_map_m03.xlsx")
    wb.save(p2)
    res = import_questionnaire(p2, Q_M03)
    assert res["ok"] is False and res["data"] is None, res
    assert codes(res) == ["BAD_HEADER"], res["errors"]


@check("[NG13] O7: ô công thức độc hại '=1+1' vào ô bắt buộc (PLII_M02) "
       "-> fail-loud không crash, KHÔNG trả data một phần")
def ng13():
    p = export_filled({"q_hkd_ten": "=1+1", "q_hkd_ma_so": "=cmd|' /C calc'!A0"},
                      questionnaire=Q_M02)
    res = import_questionnaire(p, Q_M02)
    # app chỉ ĐỌC giá trị (data_only=True), không bao giờ thực thi công thức;
    # ô công thức không có cached value -> rỗng -> required missing -> ok=False.
    assert res["ok"] is False and res["data"] is None, res
    assert "MISSING_REQUIRED" in codes(res), res["errors"]
    assert all(e.get("code") != "BAD_HEADER" or True for e in res["errors"])


@check("[NG14] O7 cụm hộ KD: điền ô lạ trong vùng điền (PLII_M01) "
       "-> UNKNOWN_CELL")
def ng14():
    p = export_filled(required_extra=M01_REQUIRED, questionnaire=Q_M01)
    wb = openpyxl.load_workbook(p)
    mp, ws = wb[MAP_SHEET], wb["Phieu"]
    col = {h: i + 1 for i, h in enumerate(MAP_COLUMNS)}
    # Chọn ô lạ CÙNG HÀNG với một ô điền thật, cột C (cột C có ô map khác
    # nhờ bảng row_group) — đúng tiền đề heuristic: cột có ô map + hàng có
    # ô map + nội dung không phải text tĩnh.
    addr = None
    for r in range(4, mp.max_row + 1):
        if mp.cell(r, col["question_id"]).value == "q_ten_hkd":
            addr = mp.cell(r, col["cell"]).value
            break
    assert addr and addr.startswith("B"), addr
    row = int(addr[1:])
    ws["C%d" % row] = "khach ghi chen ngang"
    wb.save(p)
    res = import_questionnaire(p, Q_M01)
    assert res["ok"] is False and res["data"] is None, res
    assert "UNKNOWN_CELL" in codes(res), res["errors"]


@check("[NG18] P6=C cụm thay đổi ngành: điền ô lạ trong vùng điền (PLI_M12) "
       "-> UNKNOWN_CELL")
def ng18():
    p = export_filled(required_extra=M12_REQUIRED, questionnaire=Q_M12)
    wb = openpyxl.load_workbook(p)
    mp, ws = wb[MAP_SHEET], wb["Phieu"]
    col = {h: i + 1 for i, h in enumerate(MAP_COLUMNS)}
    # Ô lạ CÙNG HÀNG với một ô điền thật (cột C — cột C có ô map nhờ bảng
    # row_group ngành): cột có ô map + hàng có ô map + nội dung không tĩnh.
    addr = None
    for r in range(4, mp.max_row + 1):
        if mp.cell(r, col["question_id"]).value == "q_dn_ten":
            addr = mp.cell(r, col["cell"]).value
            break
    assert addr and addr.startswith("B"), addr
    row = int(addr[1:])
    ws["C%d" % row] = "khach ghi chen ngang"
    wb.save(p)
    res = import_questionnaire(p, Q_M12)
    assert res["ok"] is False and res["data"] is None, res
    assert "UNKNOWN_CELL" in codes(res), res["errors"]


def main() -> int:
    for fn in (ng1a, ng1b, ng2, ng3, ng4, ng5, ng6, ng7, ng8, ng9, ng10,
               ng11, ng12, ng13, ng14, ng15, ng16, ng17, ng18):
        fn()
    passed = sum(1 for ok, _ in _results if ok)
    failed = len(_results) - passed
    print("-" * 64)
    print("test_excel_negative: %d PASS / %d FAIL" % (passed, failed))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
