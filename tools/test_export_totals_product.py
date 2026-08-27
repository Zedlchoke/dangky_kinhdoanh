# -*- coding: utf-8 -*-
"""Formula `product(<tag>[, <tag>...])` cho fixed-matrix (Lane W2).

Quyết định owner D1.4=A: "giá trị loại cổ phần = số lượng × mệnh giá" là
projection — ô tổng outputOnly PHẢI được tính lại lúc export từ chính các ô dữ
liệu, không tin draft. Grammar mới, khai báo lại ở
tools/form_specs/schema/table_kinds.schema.json (fixed_matrix.totals):

    product(<col_tag>[, <col_tag>...])

Mệnh đề chứng minh ở đây (lớp Python tools/docx_render.py):

[PRT] product trong recompute_totals()/render()
     - ô tổng = cộng dồn tích từng dòng (SL × MG) của các dòng input;
     - dòng KHÔNG kê khai gì bị bỏ qua (một loại cổ phần không dùng != 0 đồng);
     - dòng có kê khai mà THIẾU thừa số -> tổng KHÔNG XÁC ĐỊNH -> in TRỐNG,
       tuyệt đối không in '0' (trống = "chưa đủ dữ liệu nguồn", không phải 0);
     - draft trắng tinh: sum vẫn '0', product vẫn trống;
     - stored != computed -> ExportTotalsError, CHỈ dùng mã lỗi có sẵn
       (EXPORT_TOTALS_RECOMPUTE_CONFLICT / FORMULA_UNSUPPORTED);
     - formula lạ (`sqrt(x)`, `sum(a,b)`, `PRODUCT(...)`) và tag lạ bị chặn;
     - render() không mutate payload caller và chạy hai lần ra cùng kết quả.

[SCHEMA] example mới `examples["fixed-matrix-product"]` (F3 thu gọn) phải qua
     tools/build_form_maps._json_spec() như mọi instance khai báo khác.

[PARITY] CÙNG input -> JS production module (app_web/ui/fixed-matrix.js chạy qua
     node subprocess) và Python phải ra CÙNG số / cùng mã lỗi blocking.

Template DOCX dùng ở đây là sandbox tự dựng bằng python-docx trong tmp (mô phỏng
đúng the mà generator gắn vào bảng fixed-matrix), KHÔNG đụng app/templates.
Exit 0 = PASS / 1 = FAIL.

Chạy:  python -X utf8 tools/test_export_totals_product.py
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document                     # noqa: E402

from tools import build_form_maps as bfm      # noqa: E402
from tools.docx_render import (               # noqa: E402
    LEFTOVER_RE, ExportTotalsError, render, recompute_totals,
)

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run(*args):
            try:
                fn(*args)
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


# ---------------------------------------------------------------------------
# Fixture F3 thu gọn: co_cau_loai_co_phan
# ---------------------------------------------------------------------------

SEM = "co_cau_loai_co_phan"
ROWS = ("cp_pho_thong", "cp_uu_dai_khac", "cp_quy_dac_biet", "tong_so")

PRODUCT_META = {
    "kind": "fixed-matrix",
    "id": "co_cau_loai_co_phan_demo",
    "semantic_key": SEM,
    "columns": [
        {"header": "Loại cổ phần", "tag": None, "input_type": "readonly"},
        {"header": "Số lượng", "tag": "so_luong", "input_type": "money"},
        {"header": "Mệnh giá cổ phần (VNĐ)", "tag": "menh_gia_co_phan",
         "input_type": "money"},
        {"header": "Giá trị (VNĐ)", "tag": "gia_tri_vnd",
         "input_type": "money"},
    ],
    "fixed_rows": [
        {"id": "cp_pho_thong", "label": "Cổ phần phổ thông"},
        {"id": "cp_uu_dai_khac", "label": "Các cổ phần ưu đãi khác"},
        {"id": "cp_quy_dac_biet", "label": "Cổ phần quỹ đặc biệt"},
        {"id": "tong_so", "label": "Tổng số", "output_only": True},
    ],
    # Giá trị = SL × MG (D1.4=A); Tổng số lượng vẫn là sum thường để chứng minh
    # hai formula sống chung mà không đổi hành vi của nhau.
    "totals": [
        {"formula": "product(so_luong, menh_gia_co_phan)",
         "target_row": "tong_so", "target_tag": "gia_tri_vnd"},
        {"formula": "sum(so_luong)", "target_row": "tong_so",
         "target_tag": "so_luong"},
    ],
}

THREE_FACTOR_META = {
    "kind": "fixed-matrix",
    "id": PRODUCT_META["id"],
    "semantic_key": SEM,
    "columns": PRODUCT_META["columns"] + [
        {"header": "Tỷ lệ (%)", "tag": "ty_le", "input_type": "percent"}],
    "fixed_rows": PRODUCT_META["fixed_rows"],
    "totals": [{"formula": "product(so_luong, menh_gia_co_phan, ty_le)",
                "target_row": "tong_so", "target_tag": "gia_tri_vnd"}],
}

EXPECTED_GIA_TRI = "200000000"      # 10.000 x 10.000 + 5.000 x 20.000


def nest(flat_values: dict) -> dict:
    """Key phẳng '<sem>.<row>.<tag>' -> cây dict như app_web/main.py đựng."""
    out: dict = {}
    for key, value in flat_values.items():
        cur = out
        parts = key.split(".")
        for part in parts[:-1]:
            cur = cur.setdefault(part, {})
        cur[parts[-1]] = value
    return out


def draft(extra: dict | None = None) -> dict:
    """Draft chuẩn: hai dòng phổ thông/ưu đãi đã kê khai đủ, dòng quỹ để trống
    hoàn toàn; ghi đè bằng key phẳng tùy mệnh đề."""
    values = {
        "%s.cp_pho_thong.so_luong" % SEM: "10.000",
        "%s.cp_pho_thong.menh_gia_co_phan" % SEM: "10.000",
        "%s.cp_uu_dai_khac.so_luong" % SEM: "5.000",
        "%s.cp_uu_dai_khac.menh_gia_co_phan" % SEM: "20.000",
    }
    values.update(extra or {})
    return nest(values)


def total_key(tag: str) -> str:
    return "%s.tong_so.%s" % (SEM, tag)


def report_by_target(report: list) -> dict:
    return {r["target"]: r for r in report}


def computed_of(data: dict, meta=PRODUCT_META) -> tuple[str | None, list]:
    report = recompute_totals(data, [meta], strict=False)
    row = report_by_target(report)[total_key("gia_tri_vnd")]
    return row["computed"], report


def build_template(directory: Path) -> Path:
    """Sandbox template đúng cấu trúc generator xuất cho fixed-matrix."""
    path = directory / "W2_PRODUCT_DEMO.docx"
    doc = Document()
    doc.add_paragraph("DEMO MAU W2 - CO CAU LOAI CO PHAN")
    tbl = doc.add_table(rows=len(ROWS) + 1, cols=4)
    for ci, text in enumerate(["Loại cổ phần", "Số lượng",
                               "Mệnh giá cổ phần (VNĐ)", "Giá trị (VNĐ)"]):
        tbl.rows[0].cells[ci].text = text
    labels = {
        "cp_pho_thong": "Cổ phần phổ thông",
        "cp_uu_dai_khac": "Các cổ phần ưu đãi khác",
        "cp_quy_dac_biet": "Cổ phần quỹ đặc biệt",
        "tong_so": "Tổng số",
    }
    tags = ["so_luong", "menh_gia_co_phan", "gia_tri_vnd"]
    for ri, rid in enumerate(ROWS):
        row = tbl.rows[ri + 1]
        row.cells[0].text = labels[rid]
        for ci, tag in enumerate(tags, start=1):
            row.cells[ci].text = "{%s.%s.%s}" % (SEM, rid, tag)
    doc.save(str(path))
    return path


def document_texts(docx_path: Path) -> tuple[list[str], list[str]]:
    with zipfile.ZipFile(docx_path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    import re
    return re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml), LEFTOVER_RE.findall(xml)


# ---------------------------------------------------------------------------
# [PRT] product() trong recompute_totals / render
# ---------------------------------------------------------------------------

@check("[PRT1] product TINH LAI tong gia tri tu so luong x menh gia")
def prt_product_computes_total():
    data = draft()
    report = recompute_totals(data, [PRODUCT_META], strict=False)
    by_target = report_by_target(report)
    assert by_target[total_key("gia_tri_vnd")]["computed"] == EXPECTED_GIA_TRI, \
        by_target[total_key("gia_tri_vnd")]
    assert by_target[total_key("gia_tri_vnd")]["had_input"] is True
    assert by_target[total_key("gia_tri_vnd")]["stored"] == ""
    assert data[SEM]["tong_so"]["gia_tri_vnd"] == EXPECTED_GIA_TRI, \
        "ô tổng phải được điền trên data truyền vào"
    assert data[SEM]["cp_pho_thong"]["menh_gia_co_phan"] == "10.000", \
        "ô input không được đụng tới"
    assert by_target[total_key("so_luong")]["computed"] == "15000", \
        "sum(so_luong) giữ hành vi cũ (15000): %r" % \
        by_target[total_key("so_luong")]


@check("[PRT2] dong khong ke khai gi -> bo qua, khong tinh thanh 0")
def prt_undeclared_row_skipped():
    data = draft()
    # thêm một dòng nữa chỉ điền số lượng mà KHÔNG điền... -> đó là thiếu thừa
    # số (PRT4). Ở đây ta thêm dòng rỗng hoàn toàn bằng cặp giá trị rỗng.
    data_with_blank_row = draft({
        "%s.cp_quy_dac_biet.so_luong" % SEM: "",
        "%s.cp_quy_dac_biet.menh_gia_co_phan" % SEM: "",
    })
    computed, _ = computed_of(data_with_blank_row)
    assert computed == EXPECTED_GIA_TRI, \
        "dòng rỗng hoàn toàn phải bị bỏ qua, thấy %r" % computed


@check("[PRT3] chua ke khai gi: sum van '0' nhung product van TRONG")
def prt_nothing_declared():
    meta_sum_only = json.loads(json.dumps(PRODUCT_META))
    empty = nest({})

    product_report = recompute_totals(empty, [PRODUCT_META], strict=False)
    assert report_by_target(product_report)[
        total_key("gia_tri_vnd")]["computed"] == "", \
        "product trên draft trắng phải trống"

    sum_report = recompute_totals(
        nest({}), [dict(meta_sum_only, totals=[meta_sum_only["totals"][1]])],
        strict=False)
    assert report_by_target(sum_report)[
        total_key("so_luong")]["computed"] == "0", \
        "sum cũ vẫn phải '0' khi không kê khai gì"


@check("[PRT4] dong khai thieu mot thua so -> tong TRONG (''), khong in '0'")
def prt_partial_factor_blank():
    data = draft({
        "%s.cp_uu_dai_khac.menh_gia_co_phan" % SEM: "",
        # stored cũ còn số trong draft cũng phải bị nêu ra ở strict mode
    })
    computed, report = computed_of(data)
    assert computed == "", "thiếu mệnh giá phải cho tổng trống, thấy %r" % computed
    assert data[SEM]["tong_so"]["gia_tri_vnd"] == ""
    assert report_by_target(report)[
        total_key("so_luong")]["computed"] == "15000", \
        "cột sum không đổi hành vi"


@check("[PRT5] render end-to-end: o tong in dung so tinh lai, het the")
def prt_render_prints_recomputed(tmp):
    payload = draft()
    out = tmp["out"] / "prt_render_full.docx"
    out.write_bytes(render(str(tmp["tpl"]), payload,
                           tables_meta=[PRODUCT_META]))
    texts, left = document_texts(out)
    joined = "".join(texts)
    assert EXPECTED_GIA_TRI in joined, texts[:40]
    assert left == [], "còn tag sót: %s" % left[:6]
    assert payload.get(SEM, {}).get("tong_so") is None, \
        "render không được mutate payload caller"
    assert payload[SEM]["cp_pho_thong"]["so_luong"] == "10.000"

    # stored cũ trong draft thì không thắng: strict phải raise, chế độ mềm
    # (totals_strict=False) vẫn ghi số tính lại vào DOCX
    stale = draft({total_key("gia_tri_vnd"): "777"})
    raised = None
    try:
        render(str(tmp["tpl"]), stale, tables_meta=[PRODUCT_META])
    except ExportTotalsError as exc:
        raised = exc
    assert isinstance(raised, ExportTotalsError), raised

    soft = draft({total_key("gia_tri_vnd"): "777"})
    soft_out = tmp["out"] / "prt_render_soft.docx"
    soft_out.write_bytes(render(str(tmp["tpl"]), soft,
                                tables_meta=[PRODUCT_META],
                                totals_strict=False))
    texts_soft, _ = document_texts(soft_out)
    assert EXPECTED_GIA_TRI in "".join(texts_soft) and \
        "777" not in "".join(texts_soft), texts_soft[:40]


@check("[PRT6] render: tong khong xac dinh -> o trong that su, khong in 0")
def prt_render_blank_cell(tmp):
    payload = draft({"%s.cp_uu_dai_khac.menh_gia_co_phan" % SEM: ""})
    out = tmp["out"] / "prt_render_blank.docx"
    out.write_bytes(render(str(tmp["tpl"]), payload,
                           tables_meta=[PRODUCT_META]))
    texts, left = document_texts(out)
    assert not [t for t in texts if t.strip() == "0"], \
        "không được in 0 thay cho ô trống: %s" % texts
    assert left == [], "còn tag sót: %s" % left[:6]
    assert EXPECTED_GIA_TRI not in "".join(texts), texts


@check("[PRT7] stored != computed -> ExportTotalsError, chi ma loi co san")
def prt_conflict_codes():
    data = draft({total_key("gia_tri_vnd"): "12345"})
    raised = None
    try:
        recompute_totals(data, [PRODUCT_META], strict=True)
    except ExportTotalsError as exc:
        raised = exc
    assert isinstance(raised, ExportTotalsError), \
        "lệch tổng phải fail loud, nhận %r" % (raised,)
    codes = {p["code"] for p in raised.problems}
    assert codes == {"EXPORT_TOTALS_RECOMPUTE_CONFLICT"}, codes
    detail = str(raised)
    assert "tong_so" in detail and "12345" in detail, detail

    data_soft = draft({total_key("gia_tri_vnd"): "12345"})
    soft = recompute_totals(data_soft, [PRODUCT_META], strict=False)
    assert len(soft) == 2, soft
    assert data_soft[SEM]["tong_so"]["gia_tri_vnd"] == EXPECTED_GIA_TRI, \
        "chế độ mềm vẫn phải dùng số tính lại"


@check("[PRT8] tong phai TRONG ma draft con giu so -> cung la xung dot")
def prt_conflict_when_blank_overrides_number():
    data = draft({"%s.cp_uu_dai_khac.menh_gia_co_phan" % SEM: "",
                  total_key("gia_tri_vnd"): "999"})
    raised = None
    try:
        recompute_totals(data, [PRODUCT_META], strict=True)
    except ExportTotalsError as exc:
        raised = exc
    assert isinstance(raised, ExportTotalsError), \
        "draft giữ số khi tổng phải trống phải bị nêu ra, nhận %r" % (raised,)
    conflict = [p for p in raised.problems
                if p["code"] == "EXPORT_TOTALS_RECOMPUTE_CONFLICT"]
    assert conflict and any(p["row"] == "tong_so" and p["stored"] == "999"
                            for p in conflict), raised.problems


@check("[PRT9] formula/tag la van FORMULA_UNSUPPORTED (sqrt/sum 2 tham so)")
def prt_unsupported_still_blocked():
    cases = [
        "sqrt(so_luong)",
        "product()",
        "sum(so_luong, menh_gia_co_phan)",
        "PRODUCT(so_luong, menh_gia_co_phan)",
        "product(so_luong, tag_khong_ton_tai)",
    ]
    for bad_formula in cases:
        meta = json.loads(json.dumps(PRODUCT_META))
        meta["totals"] = [{"formula": bad_formula, "target_row": "tong_so",
                           "target_tag": "menh_gia_co_phan"}]
        raised = None
        try:
            recompute_totals({}, [meta])
        except ExportTotalsError as exc:
            raised = exc
        assert isinstance(raised, ExportTotalsError), \
            "%r phải fail loud, nhận %r" % (bad_formula, raised)
        assert raised.problems[0]["code"] == "FORMULA_UNSUPPORTED", \
            raised.problems


@check("[PRT10] khong doi hanh vi sum cu: bo o rong, khong du lieu -> '0'")
def prt_sum_unchanged():
    meta = json.loads(json.dumps(PRODUCT_META))
    meta["totals"] = [{"formula": "sum(menh_gia_co_phan)",
                       "target_row": "tong_so",
                       "target_tag": "menh_gia_co_phan"}]
    empty = nest({})
    rep = recompute_totals(empty, [meta], strict=False)
    assert report_by_target(rep)["%s.tong_so.menh_gia_co_phan"
                                 % SEM]["computed"] == "0"

    partial = draft({
        "%s.cp_uu_dai_khac.menh_gia_co_phan" % SEM: "12.500",
        "%s.cp_uu_dai_khac.so_luong" % SEM: "",
    })
    rep = recompute_totals(partial, [meta], strict=False)
    got = report_by_target(rep)["%s.tong_so.menh_gia_co_phan" % SEM]["computed"]
    assert got == "22500", "ô rỗng phải bị bỏ qua như cũ, thấy %r" % got


@check("[PRT11] thap phan dau phay va dau am: khong lam tron, giu dau")
def prt_decimal_and_negative():
    data = draft({
        "%s.cp_pho_thong.menh_gia_co_phan" % SEM: "10.000,5",
        "%s.cp_uu_dai_khac.menh_gia_co_phan" % SEM: "-1000",
        "%s.cp_uu_dai_khac.so_luong" % SEM: "2",
    })
    computed, _ = computed_of(data)
    # 10000 x 10000.5 + 2 x (-1000) = 100005000 - 2000 = 100003000
    assert computed == "100003000", computed

    decimal_only = draft({
        "%s.cp_pho_thong.so_luong" % SEM: "3",
        "%s.cp_pho_thong.menh_gia_co_phan" % SEM: "0,5",
        "%s.cp_uu_dai_khac.so_luong" % SEM: "",
        "%s.cp_uu_dai_khac.menh_gia_co_phan" % SEM: "",
    })
    computed, _ = computed_of(decimal_only)
    assert computed == "1.5", computed


@check("[PRT12] ba thua so: so luong x menh gia x ty le")
def prt_three_factors():
    data = draft({
        "%s.cp_pho_thong.ty_le" % SEM: "50",
        "%s.cp_uu_dai_khac.ty_le" % SEM: "25",
    })
    computed, _ = computed_of(data, THREE_FACTOR_META)
    # 10000 x 10000 x 50 + 5000 x 20000 x 25 = 5e9 + 2.5e9
    assert computed == "7500000000", computed

    lacking = draft({"%s.cp_pho_thong.ty_le" % SEM: "50"})
    computed, _ = computed_of(lacking, THREE_FACTOR_META)
    assert computed == "", "thiếu ty_le ở dòng ưu đãi -> trống"


@check("[PRT13] ownership guard + docstring/ký hiệu engine hai phía")
def prt_ownership_guard():
    src_py = (ROOT / "tools" / "docx_render.py").read_text(encoding="utf-8")
    src_js = (ROOT / "app_web" / "ui" /
              "fixed-matrix.js").read_text(encoding="utf-8")
    schema_src = (ROOT / "tools" / "form_specs" / "schema" /
                  "table_kinds.schema.json").read_text(encoding="utf-8")
    # mã chốt đường export thật thuộc về exportLoopsStrict (JS): lớp Py không đặt tên nó
    assert "EXPORT_LOOP_CONTRACT_CONFLICT" not in src_py, \
        "Python layer không được chép/dùng reserved code của exportLoopsStrict"
    for own in ("EXPORT_TOTALS_RECOMPUTE_CONFLICT", "FORMULA_UNSUPPORTED",
                "_parse_totals_formula"):
        assert own in src_py, "thiếu %s ở tools/docx_render.py" % own
    assert "function productFormula" in src_js, \
        "phía JS thiếu productFormula()"
    assert "product(" in schema_src and "D1.4" in schema_src, \
        "schema chưa ghi grammar product() / quyết định D1.4=A"


@check("[PRT14] schema examples deu qua validator, co example F3 thu gon")
def prt_schema_examples_and_note():
    schema_path = ROOT / "tools" / "form_specs" / "schema" / \
        "table_kinds.schema.json"
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    examples = schema.get("examples") or {}
    kinds = ("fixed-matrix", "fixed-matrix-product", "nested-grid",
             "output-table")
    for kind in kinds:
        entry = examples.get(kind)
        assert isinstance(entry, dict), "schema thiếu example %s" % kind
        spec = bfm._json_spec(entry["spec"], "PLI_SCHEMA_DEMO")
        assert spec["kind"] == entry["spec"]["kind"], kind
    product_example = examples["fixed-matrix-product"]["spec"]
    assert product_example["totals"][0]["formula"].startswith("product("), \
        product_example["totals"]
    totals_doc = " ".join(schema["fixed_matrix"]["spec_fields"]["totals"])
    for needle in ("product(", "TRỐNG"):
        assert needle in totals_doc, "schema chưa mô tả: %s" % needle


@check("[PRT15] idempotent: goi lan 2 cung ket qua, khong dong mau")
def prt_idempotent():
    first = draft()
    second = json.loads(json.dumps(first))
    recompute_totals(first, [PRODUCT_META], strict=False)
    one = recompute_totals(first, [PRODUCT_META], strict=False)
    two = recompute_totals(second, [PRODUCT_META], strict=False)
    def fingerprints(report):
        return sorted((r["target"], r["computed"]) for r in report)
    assert fingerprints(one) == fingerprints(two), \
        "%s != %s" % (fingerprints(one), fingerprints(two))
    assert fingerprints(one) == [
        (total_key("gia_tri_vnd"), EXPECTED_GIA_TRI),
        (total_key("so_luong"), "15000")], fingerprints(one)
    # gọi lại lần nữa vẫn ra cùng số, không bị ảnh hưởng bởi giá trị đã ghi
    three = json.loads(json.dumps(first))
    recompute_totals(three, [PRODUCT_META], strict=False)
    assert json.dumps(
        {r["target"]: r["computed"] for r in recompute_totals(
            three, [PRODUCT_META], strict=False)}, sort_keys=True) == \
        json.dumps({r["target"]: r["computed"] for r in two},
                   sort_keys=True)


# ---------------------------------------------------------------------------
# [PARITY] CÙNG input -> node (production module) và Python ra CÙNG kết quả
# ---------------------------------------------------------------------------

# probe chạy module JS PRODUCTION, đọc stdin JSON {runs:[{spec, cases}]} rồi
# trả về giá trị ô tổng + các lỗi blocking của validate().
PROBE_JS = """
const Matrix = require(process.argv[2]);
let raw = '';
process.stdin.on('data', d => raw += d);
process.stdin.on('end', () => {
  const payload = JSON.parse(raw);
  function codes(errs) {
    return errs.filter(e => e.severity === 'blocking').map(e => e.code);
  }
  const runs = payload.runs.map(function (run) {
    return {
      blank_errors: codes(Matrix.validate(run.spec, {values: {}})),
      cases: run.cases.map(function (c) {
        var draft = {values: c.values};
        var res = Matrix.recompute(run.spec, draft);
        var row = res.rows.filter(function (r) {
          return String(r.id) === String(c.target_row); })[0];
        return {value: row ? row.cells[c.target_tag] : null,
                errors: codes(Matrix.validate(run.spec, draft))};
      })
    };
  });
  process.stdout.write(JSON.stringify({runs: runs}));
});
"""


def js_spec(meta: dict, formulas: list) -> dict:
    """Map metadata Python sang shape camelCase mà fixed-matrix.js đọc; dict
    kết quả vẫn hợp lệ cho recompute_totals() nên CÙNG fixture chạy cả hai
    phía (giữ cả `target_row` cho Python lẫn `id` cho JS)."""
    columns = []
    for column in meta["columns"]:
        converted = {}
        if column.get("input_type") is not None:
            converted["inputType"] = column["input_type"]
        if column.get("tag") is not None:
            converted["tag"] = column["tag"]
        columns.append(converted)
    rows = [{"id": r["id"], "label": r["label"],
             **({"outputOnly": True} if r.get("output_only") else {})}
            for r in meta["fixed_rows"]]
    totals = [{"id": t["target_row"], "target_row": t["target_row"],
               "target_tag": t["target_tag"], "formula": formula}
              for formula, t in zip(formulas, meta["totals"])]
    return {"form_id": "W2_PARITY_DEMO", "kind": "fixed-matrix",
            "semantic_key": SEM, "columns": columns, "fixed_rows": rows,
            "totals": totals}


def run_probe(runs: list) -> list:
    """Chạy module JS production trên cùng fixture; key value trong case được
    tự thêm prefix semantic_key cho đúng shape draft của fixed-matrix.js."""
    payload_runs = []
    for run in runs:
        cases = []
        for case in run["cases"]:
            cases.append(dict(case, values={
                "%s.%s" % (SEM, key): value
                for key, value in case["values"].items()
            }))
        payload_runs.append(dict(run, spec=run["spec"], cases=cases))
    script_dir = Path(tempfile.mkdtemp(prefix="w2_probe_"))
    script = script_dir / "probe.js"
    script.write_text(PROBE_JS, encoding="utf-8")
    try:
        proc = subprocess.run(
            ["node", str(script),
             str(ROOT / "app_web" / "ui" / "fixed-matrix.js")],
            input=json.dumps({"runs": payload_runs}).encode("utf-8"),
            capture_output=True, timeout=180)
    finally:
        shutil.rmtree(script_dir, ignore_errors=True)
    if proc.returncode != 0:
        raise AssertionError("node probe thất bại: %s"
                             % proc.stderr.decode("utf-8", "replace")[:800])
    return json.loads(proc.stdout.decode("utf-8"))["runs"]


def assert_parity(case_name: str, expect: str, py_value, js_value) -> None:
    """JS recompute trả NUMBER (projection thô), Python trả CHUỖI theo
    format_total(); parity bắt buộc là CÙNG SỐ và cùng rỗng/không rỗng."""
    assert (expect == "") == (py_value == ""), \
        "[%s] Python=%r kỳ vọng %r" % (case_name, py_value, expect)
    assert (expect == "") == (js_value in ("", None)), \
        "[%s] JS=%r kỳ vọng %r" % (case_name, js_value, expect)
    if expect != "":
        assert float(js_value) == float(py_value), \
            "[%s] JS=%r vs Python=%r lệch nhau" % (
                case_name, js_value, py_value)


PARITY_RUNS = [
    {
        "name": "du du lieu 2 dong",
        "spec": js_spec(PRODUCT_META, ["product(so_luong, menh_gia_co_phan)"]),
        "cases": [
            {"values": {
                "cp_pho_thong.so_luong": "10.000",
                "cp_pho_thong.menh_gia_co_phan": "10.000",
                "cp_uu_dai_khac.so_luong": "5.000",
                "cp_uu_dai_khac.menh_gia_co_phan": "20.000",
             }, "target_row": "tong_so", "target_tag": "gia_tri_vnd"},
            {"values": {
                "cp_pho_thong.so_luong": "10.000",
                "cp_pho_thong.menh_gia_co_phan": "10.000",
                "cp_uu_dai_khac.so_luong": "5.000",
             }, "target_row": "tong_so", "target_tag": "gia_tri_vnd"},
            {"values": {}, "target_row": "tong_so",
             "target_tag": "gia_tri_vnd"},
            {"values": {
                "cp_pho_thong.so_luong": "3",
                "cp_pho_thong.menh_gia_co_phan": "0,5",
             }, "target_row": "tong_so", "target_tag": "gia_tri_vnd"},
        ],
    },
    {
        "name": "ba thua so",
        "spec": js_spec(THREE_FACTOR_META,
                        ["product(so_luong, menh_gia_co_phan, ty_le)"]),
        "cases": [
            {"values": {
                "cp_pho_thong.so_luong": "2",
                "cp_pho_thong.menh_gia_co_phan": "1.000",
                "cp_pho_thong.ty_le": "50",
                "cp_uu_dai_khac.so_luong": "4",
                "cp_uu_dai_khac.menh_gia_co_phan": "2.500",
                "cp_uu_dai_khac.ty_le": "20",
             }, "target_row": "tong_so", "target_tag": "gia_tri_vnd"},
            {"values": {
                "cp_pho_thong.so_luong": "2",
                "cp_pho_thong.menh_gia_co_phan": "1.000",
                "cp_pho_thong.ty_le": "50",
                "cp_uu_dai_khac.so_luong": "4",
             }, "target_row": "tong_so", "target_tag": "gia_tri_vnd"},
        ],
    },
]

PY_EXPECTS = [
    ["200000000", "", "", "1.5"],
    ["300000", ""],
]


@check("[PARITY1] JS <-> Py: cung input ra cung gia tri")
def parity_same_input_same_value():
    runs = run_probe(PARITY_RUNS)
    for expected_list, run_spec, js_run in zip(PY_EXPECTS, PARITY_RUNS, runs):
        for expect, case, js_case in zip(expected_list, run_spec["cases"],
                                         js_run["cases"]):
            values = nest({
                "%s.%s" % (SEM, key): value
                for key, value in case["values"].items()
            })
            py_value, _ = computed_of(values, run_spec["spec"])
            assert_parity(run_spec["name"], expect, py_value,
                          js_case["value"])


@check("[PARITY2] formula la ('sqrt') fail loud ca HAI phia, blocking")
def parity_unsupported_blocking_both_sides():
    bad_formula = "sqrt(gia_tri_vnd)"
    meta = json.loads(json.dumps(PRODUCT_META))
    meta["totals"] = [{"formula": bad_formula, "target_row": "tong_so",
                       "target_tag": "gia_tri_vnd"}]
    raised = None
    try:
        recompute_totals(draft(), [meta])
    except ExportTotalsError as exc:
        raised = exc
    assert isinstance(raised, ExportTotalsError), raised
    assert raised.problems[0]["code"] == "FORMULA_UNSUPPORTED"

    js_run = run_probe([{"name": "bad", "spec": js_spec(meta, [bad_formula]),
                         "cases": []}])[0]
    assert js_run["blank_errors"] == ["FORMULA_UNSUPPORTED"], js_run


@check("[PARITY3] tag la / chu HOA: hai phia deu tu choi, khong lech nhau")
def parity_unknown_tag_and_case_both_sides():
    for bad_formula in ("product(so_luong, tag_khong_ton_tai)",
                        "PRODUCT(so_luong, menh_gia_co_phan)"):
        meta = json.loads(json.dumps(PRODUCT_META))
        meta["totals"] = [{"formula": bad_formula, "target_row": "tong_so",
                           "target_tag": "gia_tri_vnd"}]
        raised = None
        try:
            recompute_totals(draft(), [meta])
        except ExportTotalsError as exc:
            raised = exc
        assert isinstance(raised, ExportTotalsError) and \
            raised.problems[0]["code"] == "FORMULA_UNSUPPORTED", \
            "Python chấp nhận sai %r" % bad_formula
        js_run = run_probe([{"name": bad_formula,
                             "spec": js_spec(meta, [bad_formula]),
                             "cases": []}])[0]
        assert js_run["blank_errors"] == ["FORMULA_UNSUPPORTED"], \
            "JS không từ chối %r: %s" % (bad_formula, js_run)


@check("[PARITY4] du lieu hop le: JS validate() khong con blocking error")
def parity_clean_data_no_blocking_errors():
    js_run = run_probe([PARITY_RUNS[0]])[0]
    filled = js_run["cases"][0]
    assert filled["errors"] == [], \
        "dữ liệu hợp lệ mà JS vẫn blocking: %s" % filled["errors"]


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main() -> int:
    tmp_root = Path(tempfile.mkdtemp(prefix="w2_product_"))
    tmp = {"root": tmp_root, "out": tmp_root / "rendered"}
    tmp["tpl"] = build_template(tmp_root)
    tmp["out"].mkdir(parents=True, exist_ok=True)

    tests = [
        prt_product_computes_total,
        prt_undeclared_row_skipped,
        prt_nothing_declared,
        prt_partial_factor_blank,
        lambda: prt_render_prints_recomputed(tmp),
        lambda: prt_render_blank_cell(tmp),
        prt_conflict_codes,
        prt_conflict_when_blank_overrides_number,
        prt_unsupported_still_blocked,
        prt_sum_unchanged,
        prt_decimal_and_negative,
        prt_three_factors,
        prt_ownership_guard,
        prt_schema_examples_and_note,
        prt_idempotent,
        parity_same_input_same_value,
        parity_unsupported_blocking_both_sides,
        parity_unknown_tag_and_case_both_sides,
        parity_clean_data_no_blocking_errors,
    ]
    try:
        for test in tests:
            test()
    finally:
        shutil.rmtree(tmp_root, ignore_errors=True)
    passed = sum(1 for ok, _ in _results if ok)
    failed = len(_results) - passed
    print("-" * 60)
    print("test_export_totals_product: %d PASS / %d FAIL" % (passed, failed))
    for ok, name in _results:
        if not ok:
            print(" - FAIL:", name)
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
