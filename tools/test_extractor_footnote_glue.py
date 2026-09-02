# -*- coding: utf-8 -*-
"""Test lane fix-extractor (Wave D batch 1) — glue chú thích/superscript,
STT-glue + digit-run, và đối chiếu nguồn cho chain gap đánh số.

Phủ 4 lớp bằng chứng của audit Wave D trên blocks SANDBOX (regen từ
app/templates vào thư mục temp BÊN NGOÀI repo — KHÔNG BAO GIỜ ghi vào
app/data/blocks):

  1. Class 3 — footnote GLUE trên 23 form (numbering_report.json):
     mỗi hit của heuristic GLUE phẳng (numbering_audit.GLUE,
     "[^\\d\\s:]\\d{1,2}\\s*:") phải có TOÀN BỘ chữ số nằm trong run mang
     vertAlign (superscript/subscript) — tức "đã tách semantic" đúng contract
     L1 (tools/census/check_layout_l1.py rule (d)). Toàn bộ hit của 23 form
     đều đã tách từ wave trước; heuristic phẳng của census không nhìn thấy
     vertAlign nên vẫn flag — đề nghị Lead cập nhật census (ngoài ownership
     lane này), KHÔNG chèn khoảng trắng/di chuyển chữ số ra khỏi runs vì làm
     lệch render (app_web/ui/render.js vẽ <sup> theo run.vertAlign).

  2. Negative fixture — chữ số thuộc NỘI DUNG THẬT ("Mã số thuế 10 chữ số",
     chuỗi dữ liệu "12345678910") KHÔNG được bị tách/marked superscript.

  3. Class 12 — STT-glue + digit-run (PLII_M02 #69/#122, PLI_M72 #109/#112):
     phân loại root cause. Text dính tồn tại NGUYÊN VĂN trong template như
     w:p thân (không phải w:tbl) — extractor trung thực, lỗi thuộc template
     (BLOCKED_OWNER). Test pin tính "traceable về template": nếu template
     được lane khác sửa lại thành bảng thật, số block flag giảm và assertion
     traceability tự hết đo (không dùng allowlist che).

  4. Class 2 — chain gap PLI_M03->M06 (thiếu mục 7), PLI_M04->M08 (thiếu
     mục 8-9): blocks phải phản chiếu 1:1 các heading đánh số của TEMPLATE
     (fidelity — extractor không mất gì so với nguồn đầu vào của nó). Nguồn
     chính thức (_Bieu_mau_PL_I_II_full.docx: Mẫu 3 có mục 7 "7. Thành viên
     công ty: (kê khai theo Mẫu số 6 Phụ lục I ...) Gửi kèm."; Mẫu 4 có mục
     8-9 "8. Cổ đông sáng lập (kê khai theo Mẫu số 7 ...) : Gửi kèm (nếu
     có)." / "9. Cổ đông là nhà đầu tư nước ngoài (kê khai theo Mẫu số 8
     ...): Gửi kèm (nếu có).") bị mất tại BƯỚC TÁCH nguồn, không phải ở
     extractor.

Chạy:
    python3 -X utf8 tools/test_extractor_footnote_glue.py
    python3 -X utf8 -m tools.test_extractor_footnote_glue   # exit 0 = PASS
"""
from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import unittest

TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(TOOLS_DIR)
sys.path.insert(0, TOOLS_DIR)

import docx_to_blocks as d2b          # noqa: E402
from docx import Document             # noqa: E402
from docx.oxml.ns import qn           # noqa: E402

# Heuristic GLUE của tools/census/numbering_audit.py (đọc-only, sao chép để
# test đối sánh — đổi bên census phải đổi cả bản này).
GLUE_RE = re.compile(r"[^\d\s:]\d{1,2}\s*:")
# Rule (e)/(f) của tools/census/check_layout_l1.py (đọc-only).
STT_GLUE_RUN_RE = re.compile(r"(?i)^stt(?=[^\W\d_])")
LONG_DIGIT_RUN_RE = re.compile(r"\d{5,}")
# Heading đánh số "N. <chữ>" — cùng quy ước census/numbering_audit.TOPNUM.
TOPNUM_RE = re.compile(r"^(\d{1,2})\s*[\.\)]\s*\S")
# Bản sao GLUE_RE (chữ+dính số) của check_layout_l1.py rule (d).
GLUE_LETTER_DIGIT_RE = re.compile(
    r"([A-Za-zÀ-ỹăâđêôơưÁ-ỸĂÂĐÊÔƠƯ]{2,})([0-9]+)")

# 23 form có GLUE trong generated_ho_so/_visual_layout/census/numbering_report.json
GLUE_23_FORMS = [
    "PLII_M01", "PLII_M02", "PLI_M01", "PLI_M02", "PLI_M03", "PLI_M04",
    "PLI_M05", "PLI_M06", "PLI_M08", "PLI_M09", "PLI_M12", "PLI_M13",
    "PLI_M15", "PLI_M17", "PLI_M19", "PLI_M21", "PLI_M27", "PLI_M32",
    "PLI_M66", "PLI_M67", "PLI_M68", "PLI_M69", "PLI_M72",
]

# Class 12: form chứa STT-glue/digit-run tại HEAD 65c7346 (số block flag).
STT_GLUE_FORMS = {
    "PLII_M02": 2,   # block #69, #122
    "PLI_M72": 2,    # block #109, #112
}

# Class 2: chuỗi companion có GAP (numbering_report.continuation_chains).
CHAIN_FORMS = {
    "PLI_M03": {"main_last": 6, "comp": "PLI_M06", "comp_first": 8,
                "missing": "mục 7 (pointer 'Gửi kèm' Mẫu số 6)"},
    "PLI_M04": {"main_last": 7, "comp": "PLI_M08", "comp_first": 10,
                "missing": "mục 8-9 (pointer 'Gửi kèm' Mẫu số 7/8)"},
}


def para_text(block):
    return "".join(
        r.get("text", "") for r in block.get("runs", [])
        if isinstance(r, dict) and "text" in r
    )


def glue_hits_unresolved(block):
    """Các hit GLUE_RE trên text phẳng mà chữ số KHÔNG nằm trong run vertAlign.

    Từng ký tự text được đánh dấu bằng vertAlign của run chứa nó — đối xứng
    với cách check_layout_l1.unresolved_glue_count coi vertAlign = đã phân
    loại xong (semantic)."""
    runs = [r for r in block.get("runs", [])
            if isinstance(r, dict) and "text" in r]
    marks = []
    for r in runs:
        va = bool(r.get("vertAlign"))
        marks.extend([va] * len(r["text"]))
    full = "".join(r["text"] for r in runs)
    out = []
    for m in GLUE_RE.finditer(full):
        if any(m.group(0)[k].isdigit() and not marks[m.start() + k]
               for k in range(len(m.group(0)))):
            out.append(m.group(0))
    return out


def cl1_unresolved_glue_count(block, eff_role=None):
    """Sao chép 1:1 check_layout_l1.unresolved_glue_count (census đọc-only)
    để test tự chứa, không phụ thuộc import từ tools/census."""
    if eff_role in ("footnote", "form-number", "source-indicator"):
        return 0
    count = 0
    prev_letter_tail = False
    for r in block.get("runs", []):
        if not isinstance(r, dict) or "input" in r:
            prev_letter_tail = False
            continue
        txt = r.get("text", "")
        va = r.get("vertAlign")
        hits = []
        for m in GLUE_LETTER_DIGIT_RE.finditer(txt):
            letters = m.group(1)
            before = txt[m.start() - 1] if m.start() > 0 else ""
            if before == "/" and letters.isupper():
                continue  # mã văn bản pháp lý "76/2025/QH15"
            hits.append(m.group(0))
        if va and re.match(r"^[0-9]+(?![0-9])", txt) and prev_letter_tail:
            hits = [h for h in hits if not txt.startswith(h)]
        count += len(hits)
        prev_letter_tail = bool(txt) and re.search(r"[A-Za-zÀ-ỹà-ỹ]$", txt)
    return count


def template_body_paragraph_texts(fid):
    """Text các w:p là CON TRỰC TIẾP của body trong template (Class 12 phải
    nằm ở đây — paragraph scrape dính, không phải w:tbl)."""
    doc = Document(os.path.join(ROOT, "app", "templates", "%s.docx" % fid))
    out = []
    for child in doc.element.body.iterchildren():
        if child.tag != qn("w:p"):
            continue
        out.append("".join(t.text or "" for t in child.iter(qn("w:t"))))
    return out


def stt_flagged_blocks(data):
    """Các block "p" bị flag rule (e)/(f) — quét như structural_findings của
    check_layout_l1 (run mang vertAlign được bỏ qua)."""
    out = []
    for bi, blk in enumerate(data.get("blocks", [])):
        if blk.get("t") != "p":
            continue
        for r in blk.get("runs", []):
            if not isinstance(r, dict) or r.get("vertAlign"):
                continue
            txt = r.get("text", "")
            if not isinstance(txt, str):
                continue
            if STT_GLUE_RUN_RE.match(txt) or LONG_DIGIT_RUN_RE.search(txt):
                out.append((bi, txt))
                break
    return out


class BaseSandbox(unittest.TestCase):
    """Regen sandbox dùng chung: 23 form GLUE -> temp dir ngoài repo."""

    @classmethod
    def setUpClass(cls):
        cls.sandbox = tempfile.mkdtemp(prefix="extractor_sandbox_glue_")
        form_maps = json.loads(
            d2b.FORM_MAPS_PATH.read_text(encoding="utf-8"))
        cls.blocks = {}
        for fid in GLUE_23_FORMS:
            data = d2b.convert(fid, form_maps)
            with open(os.path.join(
                    cls.sandbox, "%s.json" % fid), "w",
                    encoding="utf-8") as fh:
                json.dump(data, fh, ensure_ascii=False, indent=1)
            cls.blocks[fid] = data


class TestFootnoteGlueSandbox23(BaseSandbox):
    """Class 3: GLUE 23 form -> 0 UNRESOLVED trên sandbox (vertAlign-aware)."""

    def test_khong_con_glue_chua_tach_semantic(self):
        total_unresolved = 0
        total_raw = 0
        for fid, data in self.blocks.items():
            for bi, blk in enumerate(data["blocks"]):
                if blk.get("t") != "p":
                    continue
                total_raw += len(GLUE_RE.findall(para_text(blk)))
                unr = glue_hits_unresolved(blk)
                total_unresolved += len(unr)
                self.assertFalse(
                    unr,
                    "%s block#%d còn GLUE chưa tách (chữ số thuộc run "
                    "thường): %s" % (fid, bi, unr))
        # heuristic phẳng phải vẫn thấy hit (23 form đang được census flag);
        # nếu =0 thì form/census đã đổi — cần rà lại bài test này
        self.assertGreater(
            total_raw, 0,
            "không còn hit GLUE phẳng nào: 23 form có thể đã hết flag ở "
            "census — kiểm tra lại numbering_report và cập nhật test")
        self.assertEqual(total_unresolved, 0)

    def test_l1_unresolved_glue_count_0_tren_23_form(self):
        """Đối xứng với rule (d) của census/check_layout_l1 (đọc-only)."""
        for fid, data in self.blocks.items():
            for bi, blk in enumerate(data["blocks"]):
                if blk.get("t") != "p":
                    continue
                n = cl1_unresolved_glue_count(blk, blk.get("role"))
                self.assertEqual(
                    n, 0, "%s block#%d rule (d) glue còn %d" % (fid, bi, n))

    def test_noi_dung_chu_so_van_nguyen_ven_trong_text(self):
        """KHÔNG xóa chữ số: tách = đánh dấu run, text phẳng vẫn còn marker."""
        texts = [para_text(b) for b in self.blocks["PLII_M01"]["blocks"]
                 if b.get("t") == "p"]
        self.assertTrue(
            any("kinh doanh1:" in t for t in texts),
            "mất marker 'kinh doanh1:' — tách sai cách (xóa chữ số)")


class TestNegativeRealDigits(unittest.TestCase):
    """Chữ số thuộc NỘI DUNG THẬT không được bị tách/marked superscript."""

    def _build_runs(self, *parts):
        """parts: (text, is_superscript) -> paragraph -> build_runs."""
        doc = Document()
        p = doc.add_paragraph()
        for text, sup in parts:
            r = p.add_run(text)
            if sup:
                r.font.superscript = True
        return d2b.build_runs(p)

    def test_ma_so_thue_10_chu_so_giu_nguyen_run_thuong(self):
        runs = self._build_runs(("Mã số thuế: ", False), ("0312345678", False))
        self.assertFalse(
            any(r.get("vertAlign") for r in runs),
            "MST 10 chữ số là dữ liệu, không được đánh dấu superscript")
        self.assertEqual("".join(r["text"] for r in runs),
                         "Mã số thuế: 0312345678")

    def test_chuoi_du_lieu_12345678910_khong_bi_danh_dau(self):
        # build_runs gộp 2 run kề nhau cùng metadata — contract là: text
        # nguyên vẹn (chuỗi digit liền) và KHÔNG run nào mang vertAlign.
        runs = self._build_runs(("STT cột đếm scrape: ", False),
                                ("12345678910", False))
        self.assertFalse(any(r.get("vertAlign") for r in runs))
        self.assertIn("12345678910", "".join(
            r["text"] for r in runs if "text" in r))

    def test_so_trong_noi_dung_that_khong_bi_danh_dau(self):
        runs = self._build_runs(
            ("Số định danh cá nhân gồm ", False), ("12", False),
            (" chữ số", False))
        self.assertFalse(any(r.get("vertAlign") for r in runs))
        self.assertEqual("".join(r["text"] for r in runs),
                         "Số định danh cá nhân gồm 12 chữ số")

    def test_superscript_that_thi_tach_rieng_run(self):
        """Positive: marker chú thích thật (superscript) phải là run riêng
        có vertAlign, kề sát chữ body — đúng schema blocks-v3."""
        runs = self._build_runs(
            ("3. Ngành, nghề kinh doanh", False), ("1", True), (":", False))
        self.assertEqual(
            [r.get("vertAlign") for r in runs],
            [None, "superscript", None])
        self.assertEqual(runs[1]["text"], "1")
        self.assertEqual(runs[0]["text"], "3. Ngành, nghề kinh doanh")


class TestSttGlueClassification(BaseSandbox):
    """Class 12: STT-glue + digit-run là dữ liệu template, KHÔNG phải lỗi
    extract. Mỗi run của block bị flag phải traceable NGUYÊN VĂN về một w:p
    thân của template (scrape dính chung một w:t)."""

    def test_so_block_flag_khop_head_va_thuoc_dung_form(self):
        for fid, expected in STT_GLUE_FORMS.items():
            flagged = stt_flagged_blocks(self.blocks[fid])
            self.assertEqual(
                len(flagged), expected,
                "%s: số block flag (e)/(f) = %d, kỳ vọng %d (đổi khi template "
                "được sửa lại thành bảng thật -> cập nhật STT_GLUE_FORMS)"
                % (fid, len(flagged), expected))

    def test_run_flag_co_mat_nguyen_van_trong_template(self):
        for fid in STT_GLUE_FORMS:
            doc_texts = template_body_paragraph_texts(fid)
            flagged = stt_flagged_blocks(self.blocks[fid])
            self.assertTrue(flagged)
            for bi, _txt in flagged:
                for r in self.blocks[fid]["blocks"][bi].get("runs", []):
                    txt = r.get("text", "")
                    if not txt.strip():
                        continue
                    self.assertTrue(
                        any(txt in tp for tp in doc_texts),
                        "%s block#%d run %r KHÔNG tìm thấy trong w:p thân "
                        "template -> extractor tự ghép? Phải điều tra lại"
                        % (fid, bi, txt[:40]))

    def test_stt_glue_khong_nam_trong_bang_that(self):
        """Nếu text dính nằm trong w:tbl của template -> extractor gộp sai ô;
        hiện trạng (blocker): nó nằm ở w:p thân, extractor chỉ phản chiếu."""
        for fid in STT_GLUE_FORMS:
            doc = Document(os.path.join(
                ROOT, "app", "templates", "%s.docx" % fid))
            found = False
            for p in doc.element.body.iter(qn("w:p")):
                txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
                if STT_GLUE_RUN_RE.search(txt.strip()):
                    found = True
                    anc = [a.tag for a in p.iterancestors()]
                    self.assertNotIn(
                        qn("w:tbl"), anc,
                        "%s: STT-glue nằm TRONG bảng — extractor gộp sai ô?"
                        % fid)
            self.assertTrue(
                found, "%s: template không còn paragraph STT-glue -> đã được "
                "lane template sửa; cập nhật STT_GLUE_FORMS" % fid)


class TestChainGapSourceFidelity(BaseSandbox):
    """Class 2: blocks phản chiếu 1:1 heading đánh số của TEMPLATE cho 4
    form trong 2 chuỗi gap -> extractor không mất mục nào so với nguồn đầu
    vào; mục 7 / 8-9 bị mất từ BƯỚC TÁCH nguồn (BLOCKED_OWNER template)."""

    def _numbered(self, fid, texts):
        return [t.strip() for t in texts
                if TOPNUM_RE.match(t.strip()) and len(t.strip()) < 140]

    def _template_numbered(self, fid):
        return self._numbered(fid, template_body_paragraph_texts(fid))

    def _blocks_numbered(self, fid):
        return self._numbered(
            fid, [para_text(b) for b in self.blocks[fid]["blocks"]
                  if b.get("t") == "p"])

    def test_heading_danh_so_blocks_khop_template(self):
        for fid in ("PLI_M03", "PLI_M06", "PLI_M04", "PLI_M08"):
            self.assertEqual(
                self._blocks_numbered(fid), self._template_numbered(fid),
                "%s: blocks lệch heading đánh số so với template" % fid)

    def test_gap_con_nguyen_o_template_khong_phai_do_extractor(self):
        # subitem thập phân ("9.1", scrape lệch "1.6") không phải mục top-level
        subnum = re.compile(r"^\d{1,2}\s*\.\s*\d{1,2}")

        def top_nums(heads):
            out = []
            for t in heads:
                m = TOPNUM_RE.match(t)
                if m and not subnum.match(t):
                    out.append(int(m.group(1)))
            return out

        for fid, info in CHAIN_FORMS.items():
            nums = top_nums(self._template_numbered(fid))
            for n in range(info["main_last"] + 1, info["comp_first"]):
                self.assertNotIn(
                    n, nums,
                    "%s ĐÃ có mục %d trong template nhưng blocks thiếu -> "
                    "đây là lỗi extractor, phải sửa converter" % (fid, n))
            self.assertIn(
                info["main_last"], nums,
                "%s: mất cả mục cuối %d ở template?" % (fid, info["main_last"]))
            comp_nums = top_nums(self._template_numbered(info["comp"]))
            self.assertEqual(
                min(comp_nums), info["comp_first"],
                "%s: đầu chuỗi bên companion đã đổi -> cập nhật CHAIN_FORMS"
                % info["comp"])


if __name__ == "__main__":
    unittest.main()
