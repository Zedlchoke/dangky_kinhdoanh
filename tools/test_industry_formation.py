# -*- coding: utf-8 -*-
"""Wave 2 Agent 2A: bảng ngành nghề ds_nganh cho nhóm thành lập (production).

Mẫu kiểm: PLI_M01, PLI_M03, PLI_M04, PLI_M05 (đăng ký mới/thành lập DN).

Với TỪNG mẫu:
1. blocks app/data/blocks/<ID>.json có đúng 1 block loop "ds_nganh":
   - hàng tiêu đề 4 cột khớp số ô body/template;
   - cols zero-based {ten_nganh, ma_nganh, la_nganh_chinh};
   - cột 0 (STT) KHÔNG tag (STT tự sinh qua biến engine {_stt});
   - rows + blankRows + 1 khuôn = tổng hàng bảng trong template.
   Cách assert giống tools/test_pli_m02_industry.py (mẫu tham chiếu).
2. form_maps.json khai báo dynamic_tables tương ứng (col/tag khớp).
3. Render template bằng tools/docx_render.write_rendered với 3 dòng mẫu
   (mã VSIC thật lấy từ app/data/nganh_vsic.json, dòng đầu là ngành chính):
   - mỗi giá trị nằm đúng ô cột (so text từng cell);
   - STT hiển thị 1/2/3 (tự sinh);
   - ĐÚNG MỘT ô "X" ở cột ngành chính, các ô còn lại rỗng;
   - không còn sót '{#ds_nganh}', '{/ds_nganh}', tag cột hay {_stt}.

File test thuần Python, exit code 0 = PASS / 1 = FAIL.

Chạy: python -X utf8 -m tools.test_industry_formation
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document                      # noqa: E402
from docx.oxml.ns import qn                    # noqa: E402
from docx.table import Table                   # noqa: E402
from docx.text.paragraph import Paragraph      # noqa: E402

from tools.docx_render import write_rendered   # noqa: E402

FORM_IDS = ["PLI_M01", "PLI_M03", "PLI_M04", "PLI_M05"]

FORM_MAPS_PATH = ROOT / "app" / "data" / "form_maps.json"
VSIC_PATH = ROOT / "app" / "data" / "nganh_vsic.json"
OUT_DIR = ROOT / "generated_ho_so" / "_test_industry_formation"

LOOP_NAME = "ds_nganh"
EXPECTED_HEADERS = ["STT", "Tên ngành", "Mã ngành",
                    "Ngành, nghề kinh doanh chính"]
EXPECTED_COLS = [(1, "ten_nganh"), (2, "ma_nganh"), (3, "la_nganh_chinh")]

# Mã VSIC cấp 4 thật; tên lấy từ nganh_vsic.json khi chạy.
SAMPLE_VSIC_CODES = ["0111", "4711", "6201"]

_results: list[tuple[bool, str]] = []


def blocks_path(form_id: str) -> Path:
    return ROOT / "app" / "data" / "blocks" / ("%s.json" % form_id)


def template_path(form_id: str) -> Path:
    return ROOT / "app" / "templates" / ("%s.docx" % form_id)


def check(name: str):
    def deco(fn):
        def run():
            try:
                fn()
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def fold(text: str) -> str:
    nfkd = unicodedata.normalize("NFD", text.lower())
    out = "".join(c for c in nfkd if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", out.replace("đ", "d")).strip()


def cell_runs_text(cell: dict) -> str:
    return "".join(r.get("text", "") for r in cell["runs"])


def tc_iter(tr):
    for child in tr.iterchildren():
        if child.tag == qn("w:tc"):
            yield child


def tr_cell_texts(tbl, ri: int) -> list[str]:
    """Text từng ô của hàng ri — duyệt trực tiếp w:tc (mỗi ô đúng một lần)."""
    out = []
    for tc in tc_iter(tbl._tbl.tr_lst[ri]):
        txt = "\n".join(
            Paragraph(p_el, tbl.part.document).text
            for p_el in tc.iterchildren()
            if p_el.tag == qn("w:p")
        )
        out.append(txt.strip())
    return out


def body_tables(doc):
    """Các bảng theo thứ tự thân tài liệu (duyệt con trực tiếp body)."""
    return [Table(child, doc) for child in doc.element.body.iterchildren()
            if child.tag == qn("w:tbl")]


def find_industry_table(doc):
    """Tìm bảng có hàng tiêu đề khớp EXPECTED_HEADERS."""
    want = fold(" ".join(EXPECTED_HEADERS))
    for ti, tbl in enumerate(doc.tables):
        if len(tbl.rows) and fold(" ".join(tr_cell_texts(tbl, 0))) == want:
            return ti, tbl
    return None, None


def load_sample_rows() -> list[dict]:
    """3 dòng dữ liệu mẫu với mã VSIC thật từ app/data/nganh_vsic.json."""
    vsic = json.loads(VSIC_PATH.read_text(encoding="utf-8"))
    by_ma = {e["ma"]: e["ten"] for e in vsic.get("nganh", []) if e.get("cap") == 4}
    rows = []
    for i, ma in enumerate(SAMPLE_VSIC_CODES):
        assert ma in by_ma, \
            "mã VSIC mẫu %r không có trong nganh_vsic.json" % ma
        rows.append({
            "ma_nganh": ma,
            "ten_nganh": by_ma[ma],
            "la_nganh_chinh": "X" if i == 0 else "",
        })
    return rows


@check("nganh_vsic.json: 3 ma VSIC mau deu la nganh cap 4 that")
def t_sample_data_real():
    rows = load_sample_rows()
    assert len(rows) == 3
    assert len({r["ma_nganh"] for r in rows}) == 3, "mã mẫu bị trùng nhau"


def make_blocks_check(form_id: str):
    @check("%s blocks: dung 1 loop ds_nganh, grid 4 cot, cot 0 STT khong tag" % form_id)
    def t_blocks():
        data = json.loads(blocks_path(form_id).read_text(encoding="utf-8"))
        loops = [b for b in data["blocks"] if b.get("t") == "loop"]
        assert len(loops) == 1, "kỳ vọng đúng 1 block loop, thấy %d" % len(loops)
        blk = loops[0]
        assert blk["loop"] == LOOP_NAME, \
            "loop=%r != %r" % (blk["loop"], LOOP_NAME)
        assert blk["headerRows"] == [0], \
            "headerRows=%r != [0]" % (blk["headerRows"],)

        # header 4 cột, text khớp thiết kế
        hdr_cells = blk["rows"][blk["headerRows"][0]]["cells"]
        assert len(hdr_cells) == len(EXPECTED_HEADERS), \
            "header có %d ô != %d" % (len(hdr_cells), len(EXPECTED_HEADERS))
        got_headers = [cell_runs_text(c).strip() for c in hdr_cells]
        assert got_headers == EXPECTED_HEADERS, \
            "header %r != %r" % (got_headers, EXPECTED_HEADERS)

        # mọi hàng trong rows đều khớp số ô header ("header khớp body")
        ncols = len(EXPECTED_HEADERS)
        for ri, row in enumerate(blk["rows"]):
            assert len(row["cells"]) == ncols, \
                "rows[%d] có %d ô != %d" % (ri, len(row["cells"]), ncols)

        # cols zero-based đúng thứ tự; cột 0 không tag
        got_cols = sorted((c["col"], c["tag"]) for c in blk["cols"])
        assert got_cols == EXPECTED_COLS, \
            "cols %r != %r" % (got_cols, EXPECTED_COLS)
        assert all(c["col"] != 0 for c in blk["cols"]), \
            "cột 0 (STT tự sinh) không được có tag trong cols"

        # form_maps khai báo tương ứng; blankRows hợp lý; khớp bảng template
        fm = json.loads(FORM_MAPS_PATH.read_text(encoding="utf-8"))[form_id]
        dts = [t for t in fm.get("dynamic_tables") or []
               if t["loop"] == LOOP_NAME]
        assert len(dts) == 1, \
            "kỳ vọng 1 dynamic_tables ds_nganh, thấy %d" % len(dts)
        dt = dts[0]
        got_fm_cols = sorted((c["col"], c["tag"]) for c in dt["columns"])
        assert got_fm_cols == EXPECTED_COLS, \
            "form_maps columns %r != %r" % (got_fm_cols, EXPECTED_COLS)
        fld_keys = {f["key"] for f in fm.get("fields", [])}
        for _, tag in EXPECTED_COLS:
            assert ("%s.%s" % (LOOP_NAME, tag)) in fld_keys, \
                "thiếu field %s.%s trong form_maps" % (LOOP_NAME, tag)

        doc = Document(str(template_path(form_id)))
        tables = body_tables(doc)
        tbl = tables[dt["table_index"]]
        n_tr = len(tbl._tbl.tr_lst)
        total = len(blk["rows"]) + blk["blankRows"] + 1
        assert total == n_tr, \
            "rows(%d)+blank(%d)+khuôn(1)=%d != %d hàng template" % (
                len(blk["rows"]), blk["blankRows"], total, n_tr)
        assert blk["blankRows"] >= 1, "cần ít nhất 1 hàng trắng để in"

    return t_blocks


def make_render_check(form_id: str, sample_rows: list[dict]):
    @check("%s render DOCX: 3 dong vao dung cot, STT 1/2/3, dung mot X nganh chinh" % form_id)
    def t_render():
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        tpl = template_path(form_id)
        out_path = OUT_DIR / ("%s_ds_nganh.docx" % form_id)
        write_rendered(tpl, {LOOP_NAME: sample_rows}, out_path)

        with zipfile.ZipFile(out_path) as z:
            xml = z.read("word/document.xml").decode("utf-8")

        doc = Document(str(out_path))
        ti, tbl = find_industry_table(doc)
        assert tbl is not None, "không tìm thấy bảng ngành nghề trong file render"
        texts = [tr_cell_texts(tbl, ri) for ri in range(len(tbl.rows))]

        # hàng tiêu đề giữ nguyên
        assert [t.strip() for t in texts[0]] == EXPECTED_HEADERS, \
            "header render %r != %r" % (texts[0], EXPECTED_HEADERS)

        # [F3T.2 max(K,N)] dữ liệu NGAY SAU header; hàng trắng residual
        # nằm SAU dữ liệu (pin cũ "cuối bảng" là vị trí khiếm khuyết)
        ten_col = next(c for c, tag in EXPECTED_COLS if tag == "ten_nganh")
        ma_col = next(c for c, tag in EXPECTED_COLS if tag == "ma_nganh")
        x_col = next(c for c, tag in EXPECTED_COLS if tag == "la_nganh_chinh")
        n_data = len(sample_rows)
        data_texts = texts[1:1 + n_data]
        x_count = 0
        for i, row in enumerate(data_texts):
            want = sample_rows[i]
            assert len(row) == len(EXPECTED_HEADERS), \
                "hàng dữ liệu %d có %d ô != %d" % (
                    i, len(row), len(EXPECTED_HEADERS))
            assert row[0] == str(i + 1), \
                "STT dòng %d = %r, kỳ vọng %r" % (i, row[0], str(i + 1))
            assert row[ma_col] == want["ma_nganh"], \
                "Mã ngành dòng %d = %r != %r" % (i, row[ma_col], want["ma_nganh"])
            assert row[ten_col] == want["ten_nganh"], \
                "Tên ngành dòng %d = %r != %r" % (i, row[ten_col], want["ten_nganh"])
            assert row[x_col] == want["la_nganh_chinh"], \
                "Ngành chính dòng %d = %r != %r" % (
                    i, row[x_col], want["la_nganh_chinh"])
            x_count += 1 if row[x_col] == "X" else 0
        assert x_count == 1, \
            "cột ngành chính phải có ĐÚNG MỘT ô 'X', thấy %d" % x_count

        # không còn bất kỳ the loop/cột/biến engine nào sót trong XML
        tokens = ["{#%s}" % LOOP_NAME, "{/%s}" % LOOP_NAME, "{_stt}"]
        tokens += ["{%s}" % tag for _, tag in EXPECTED_COLS]
        for tok in tokens:
            assert tok not in xml, \
                "còn sót token %r trong document.xml" % tok

    return t_render


def main() -> int:
    t_sample_data_real()
    sample_ok = all(ok for ok, _ in _results)
    sample_rows = load_sample_rows() if sample_ok else [
        {"ma_nganh": "", "ten_nganh": "", "la_nganh_chinh": ""} for _ in range(3)]
    for form_id in FORM_IDS:
        make_blocks_check(form_id)()
        make_render_check(form_id, sample_rows)()
    passed = sum(1 for ok, _ in _results if ok)
    failed = len(_results) - passed
    print("-" * 60)
    print("test_industry_formation: %d PASS / %d FAIL" % (passed, failed))
    for ok, name in _results:
        if not ok:
            print(" - FAIL:", name)
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
