# -*- coding: utf-8 -*-
"""Kiểm tra Wave 2 Agent 2C: lưới ngành ds_nganh trong 2 mẫu HỘ KINH DOANH
PLII_M02 (giấy đề nghị thay đổi nội dung ĐKKD hộ KD) và PLII_M07 (giấy chứng
nhận đăng ký hộ kinh doanh) - production.

P1-C (2026-08-26) mở rộng cho PLII_M02 — 4 semantic loop ngành theo PLAN mục A:
  nganh_bo_sung / nganh_xoa_bo / nganh_sua_doi_chi_tiet / nganh_hien_tai.
Mỗi nghiệp vụ được test ĐỘC LẬP: vùng tồn tại, đúng cột, dữ liệu không trộn
giữa các loop (draft tách key). Các test post-integration chạy capability
probe TRƯỚC khi assert: artifact chưa có loop tương ứng thì ghi nhận
EXPECTED-FAIL-PRE-INTEGRATION (không nuốt lỗi assert, không pass giả). Khi
generator đã tích hợp occurrences, probe thấy loop -> assert thật chạy hết.
Biến môi trường P1C_STRICT=1 biến mọi EXPECTED-FAIL-PRE-INTEGRATION thành
FAIL (dùng để gate sau integration).

Phạm vi TỪNG mẫu:
1. blocks app/data/blocks/<ID>.json:
   - PLII_M07: đúng 1 loop "ds_nganh";
   - PLII_M02: loop "ds_nganh" phải tồn tại DUY NHẤT (các loop ngành khác
     của P1 không được trùng tên), cấu trúc như cũ:
     * hàng tiêu đề 4 cột khớp số ô của body/template;
     * cols zero-based {ten_nganh, ma_nganh, la_nganh_chinh};
     * cột 0 (STT) KHÔNG tag (STT tự sinh qua biến engine {_stt}).
2. form_maps.json khai báo dynamic_tables tương ứng (col/tag khớp).
3. Render template bằng tools/docx_render.write_rendered với 3 dòng mẫu
   (dòng đầu đánh dấu X ngành chính):
   - đủ 3 hàng dữ liệu nhân bản, mỗi giá trị nằm đúng ô cột;
   - STT hiển thị 1/2/3 (tự sinh);
   - đúng MỘT dấu X ở cột "Ngành, nghề kinh doanh chính";
   - không còn sót '{#ds_nganh}', '{/ds_nganh}', tag cột hay {_stt}.

Bằng chứng cột ngành chính trên template gốc (python-docx, đã ghi vào
tools/form_specs/household/*.json):
- PLII_M02: ô tiêu đề mục C.3 'Ngành, nghề kinh doanh chính9' + chú thích
  '9 Đánh dấu X để chọn một trong các ngành, nghề đã kê khai làm ngành,
  nghề kinh doanh chính.'
- PLII_M07: ô tiêu đề mục 3 'Ngành, nghề kinh doanh chính' (mẫu này không
  kèm số chú thích).

File test thuần Python (chỉ DOCX/blocks), exit code 0 = PASS / 1 = FAIL
(FAIL thật; EXPECTED-FAIL-PRE-INTEGRATION không làm fail trừ khi P1C_STRICT=1).

Chạy:  python -X utf8 -m tools.test_industry_household
"""
from __future__ import annotations

import json
import os
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document                      # noqa: E402
from docx.oxml.ns import qn                   # noqa: E402
from docx.table import Table                  # noqa: E402
from docx.text.paragraph import Paragraph     # noqa: E402

from tools.docx_render import write_rendered  # noqa: E402

FORM_MAPS_PATH = ROOT / "app" / "data" / "form_maps.json"
OUT_DIR = ROOT / "generated_ho_so" / "_test_industry_household"

FORMS = ["PLII_M02", "PLII_M07"]
LOOP_NAME = "ds_nganh"
EXPECTED_HEADERS = ["STT", "Tên ngành", "Mã ngành",
                    "Ngành, nghề kinh doanh chính"]
EXPECTED_COLS = [(1, "ten_nganh"), (2, "ma_nganh"), (3, "la_nganh_chinh")]

# P1-C: 4 vùng ngành của PLII_M02 (template anchor -> block hiện trạng).
# region 4 (nganh_hien_tai) = bảng C.3 cấp đổi/đăng ký lại — SAU integration
# giữ tên loop canonical ds_nganh (dùng chung PLI_M02/PLII_M07) kèm metadata
# occurrence_id=nganh_hien_tai; match_heading là literal prefix sau fold.
SPEC_PATH = ROOT / "tools" / "form_specs" / "household" / "PLII_M02.json"
M02_OCCURRENCES = [
    ("nganh_bo_sung", "1. bo sung nganh, nghe kinh doanh",
     ["^2\\. bo nganh, nghe kinh doanh"],
     ["STT", "Tên ngành, nghề kinh doanh được bổ sung", "Mã ngành",
      "Ngành, nghề kinh doanh chính"]),
    ("nganh_xoa_bo", "2. bo nganh, nghe kinh doanh",
     ["^3\\. sua doi nganh, nghe kinh doanh"],
     ["STT", "Tên ngành, nghề kinh doanh được bỏ khỏi danh sách đã đăng ký",
      "Mã ngành", "Ngành, nghề kinh doanh chính"]),
    ("nganh_sua_doi_chi_tiet", "3. sua doi nganh, nghe kinh doanh",
     ["^luu y: truong hop ho kinh doanh thay doi nganh"],
     ["STT", "Tên ngành, nghề kinh doanh", "Mã ngành"]),
    ("nganh_hien_tai", "3. nganh, nghe kinh doanh",
     ["^4\\. von kinh doanh"], EXPECTED_HEADERS),
]


def blocks_path(fid: str) -> Path:
    return ROOT / "app" / "data" / "blocks" / ("%s.json" % fid)


def template_path(fid: str) -> Path:
    return ROOT / "app" / "templates" / ("%s.docx" % fid)


SAMPLE_ROWS = [
    {"ten_nganh": "Hoạt động lập trình máy vi tính",
     "ma_nganh": "6201", "la_nganh_chinh": "X"},
    {"ten_nganh": "Trồng lúa",
     "ma_nganh": "0111", "la_nganh_chinh": ""},
    {"ten_nganh": "Bán lẻ lương thực trong cửa hàng",
     "ma_nganh": "4711", "la_nganh_chinh": ""},
]

_results: list[tuple[str, str]] = []  # (status, label): PASS/FAIL/EFPI
STRICT = os.environ.get("P1C_STRICT") == "1"


def check(name):
    """name: chuỗi tĩnh hoặc hàm nhận cùng tham số với fn -> nhãn test."""
    def deco(fn):
        def run(*args):
            label = name(*args) if callable(name) else name
            try:
                fn(*args)
                _results.append(("PASS", label))
                print("PASS %s" % label)
            except EFPI as exc:
                _results.append(("EXPECTED-FAIL-PRE-INTEGRATION", label))
                print("EXPECTED-FAIL-PRE-INTEGRATION %s: %s" % (label, exc))
            except AssertionError as exc:
                _results.append(("FAIL", "%s: %s" % (label, exc)))
                print("FAIL %s: %s" % (label, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append(("FAIL", "%s: exception %r" % (label, exc)))
                print("FAIL %s: exception %r" % (label, exc))
        return run
    return deco


class EFPI(Exception):
    """Capability probe thất bại: artifact CHƯA có tính năng post-integration.

    Được ném CHỈ từ probe (trước khi vào phần assert thật); mọi assert bên
    dưới vẫn chạy đầy đủ khi probe pass -> không thể pass giả."""
    pass


def form_maps() -> dict:
    return json.loads(FORM_MAPS_PATH.read_text(encoding="utf-8"))


def fold(text: str) -> str:
    nfkd = unicodedata.normalize("NFD", text.lower())
    out = "".join(c for c in nfkd if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", out.replace("đ", "d")).strip()


def cell_runs_text(cell: dict) -> str:
    return "".join(r.get("text", "") for r in cell["runs"])


def tc_iter(tr):
    for child in tr.iterchildren():
        if child.tag == qn("w:tc"):
            yield child


def tr_cell_texts(tbl, ri: int) -> list[str]:
    """Text từng ô của hàng ri — duyệt trực tiếp w:tc (mỗi ô đúng một lần)."""
    out = []
    for tc in tc_iter(tbl._tbl.tr_lst[ri]):
        txt = "\n".join(
            Paragraph(p_el, tbl.part.document).text
            for p_el in tc.iterchildren()
            if p_el.tag == qn("w:p")
        )
        out.append(txt.strip())
    return out


def body_tables(doc):
    """Các bảng theo thứ tự thân tài liệu (duyệt con trực tiếp body)."""
    return [Table(child, doc) for child in doc.element.body.iterchildren()
            if child.tag == qn("w:tbl")]


def find_industry_table(doc):
    """Tìm bảng có hàng tiêu đề khớp EXPECTED_HEADERS."""
    want = fold(" ".join(EXPECTED_HEADERS))
    for ti, tbl in enumerate(doc.tables):
        if len(tbl.rows) and fold(" ".join(tr_cell_texts(tbl, 0))) == want:
            return ti, tbl
    return None, None


@check(lambda fid: "blocks %s: loop ds_nganh dung 1 lan, 4 cot, cols zero-based, cot 0 khong tag" % fid)
def t_blocks_loop_structure(fid: str):
    data = json.loads(blocks_path(fid).read_text(encoding="utf-8"))
    loops = [b for b in data["blocks"] if b.get("t") == "loop"]
    cands = [b for b in loops if b.get("loop") == LOOP_NAME]
    if fid == "PLII_M02":
        # P1: M02 co nhieu semantic loop nganh (bo_sung/xoa_bo/sua_doi/hien_tai)
        # nhung loop ds_nganh van phải DUY NHẤT và không được trùng tên.
        assert len(cands) == 1, \
            "kỳ vọng đúng 1 block loop %r, thấy %d (tổng %d loop: %r)" % (
                LOOP_NAME, len(cands), len(loops),
                [b.get("loop") for b in loops])
        names = [b["loop"] for b in loops]
        dupes = sorted({n for n in names if names.count(n) > 1})
        assert not dupes, "loop bị trùng tên: %r" % (dupes,)
    else:
        assert len(loops) == 1 and len(cands) == 1, \
            "kỳ vọng đúng 1 block loop, thấy %d" % len(loops)
    blk = cands[0]
    assert blk["loop"] == LOOP_NAME, "loop=%r != %r" % (blk["loop"], LOOP_NAME)
    assert blk["headerRows"] == [0], "headerRows=%r != [0]" % (blk["headerRows"],)

    hdr_cells = blk["rows"][blk["headerRows"][0]]["cells"]
    assert len(hdr_cells) == 4, "header có %d ô != 4" % len(hdr_cells)
    got_headers = [cell_runs_text(c).strip() for c in hdr_cells]
    assert got_headers == EXPECTED_HEADERS, \
        "header %r != %r" % (got_headers, EXPECTED_HEADERS)

    for ri, row in enumerate(blk["rows"]):
        assert len(row["cells"]) == 4, \
            "rows[%d] có %d ô != 4" % (ri, len(row["cells"]))

    got_cols = sorted((c["col"], c["tag"]) for c in blk["cols"])
    assert got_cols == EXPECTED_COLS, "cols %r != %r" % (got_cols, EXPECTED_COLS)
    assert all(c["col"] != 0 for c in blk["cols"]), \
        "cột 0 (STT tự sinh) không được có tag trong cols"

    # grid khớp template: rows + blank + 1 khuôn = tổng hàng bảng trong template
    fm = json.loads(FORM_MAPS_PATH.read_text(encoding="utf-8"))[fid]
    dt = next(t for t in fm["dynamic_tables"] if t["loop"] == LOOP_NAME)
    doc = Document(str(template_path(fid)))
    tbl = body_tables(doc)[dt["table_index"]]
    n_tr = len(tbl._tbl.tr_lst)
    total = len(blk["rows"]) + blk["blankRows"] + 1
    assert total == n_tr, \
        "rows(%d)+blank(%d)+khuôn(1)=%d != %d hàng template" % (
            len(blk["rows"]), blk["blankRows"], total, n_tr)
    assert blk["blankRows"] >= 1, "cần ít nhất 1 hàng trắng để in"


@check(lambda fid: "form_maps %s: dynamic_tables khai bao dung col/tag cho ds_nganh" % fid)
def t_form_maps_metadata(fid: str):
    fm = json.loads(FORM_MAPS_PATH.read_text(encoding="utf-8"))[fid]
    dts = [t for t in fm.get("dynamic_tables") or [] if t["loop"] == LOOP_NAME]
    assert len(dts) == 1, "kỳ vọng 1 dynamic_tables ds_nganh, thấy %d" % len(dts)
    dt = dts[0]
    got_cols = sorted((c["col"], c["tag"]) for c in dt["columns"])
    assert got_cols == EXPECTED_COLS, \
        "columns %r != %r" % (got_cols, EXPECTED_COLS)
    fld_keys = {f["key"] for f in fm.get("fields", [])}
    for _, tag in EXPECTED_COLS:
        assert ("%s.%s" % (LOOP_NAME, tag)) in fld_keys, \
            "thiếu field %s.%s trong form_maps" % (LOOP_NAME, tag)


@check(lambda fid: "render DOCX %s: 3 dong vao dung cot, STT 1/2/3, dung mot X nganh chinh, khong con token" % fid)
def t_render_three_rows(fid: str):
    out_dir = OUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / ("%s_ds_nganh.docx" % fid)
    write_rendered(str(template_path(fid)),
                   {LOOP_NAME: SAMPLE_ROWS}, out_path)

    with zipfile.ZipFile(out_path) as z:
        xml = z.read("word/document.xml").decode("utf-8")

    doc = Document(str(out_path))
    _, tbl = find_industry_table(doc)
    assert tbl is not None, "không tìm thấy bảng ngành nghề trong file render"
    texts = [tr_cell_texts(tbl, ri) for ri in range(len(tbl.rows))]

    assert [t.strip() for t in texts[0]] == EXPECTED_HEADERS, \
        "header render %r != %r" % (texts[0], EXPECTED_HEADERS)

    n_data = len(SAMPLE_ROWS)
    data_texts = texts[-n_data:]
    n_x = 0
    for i, row in enumerate(data_texts):
        want = SAMPLE_ROWS[i]
        assert len(row) == 4, "hàng dữ liệu %d có %d ô != 4" % (i, len(row))
        assert row[0] == str(i + 1), \
            "STT dòng %d = %r, kỳ vọng %r" % (i, row[0], str(i + 1))
        assert row[1] == want["ten_nganh"], \
            "Tên ngành dòng %d = %r != %r" % (i, row[1], want["ten_nganh"])
        assert row[2] == want["ma_nganh"], \
            "Mã ngành dòng %d = %r != %r" % (i, row[2], want["ma_nganh"])
        assert row[3] == want["la_nganh_chinh"], \
            "Ngành chính dòng %d = %r != %r" % (i, row[3], want["la_nganh_chinh"])
        n_x += 1 if row[3] == "X" else 0
    assert n_x == 1, "kỳ vọng đúng 1 dòng có X ngành chính, thấy %d" % n_x
    assert data_texts[0][3] == "X", "dòng đầu phải là ngành chính (X)"

    # không còn bất kỳ the loop/cột/biến engine nào sót trong XML
    for tok in ("{#%s}" % LOOP_NAME, "{/%s}" % LOOP_NAME,
                "{ten_nganh}", "{ma_nganh}", "{la_nganh_chinh}", "{_stt}"):
        assert tok not in xml, "còn sót token %r trong document.xml" % tok


def body_paras(doc):
    """[(chi_so, Paragraph)] theo thứ tự thân tài liệu (bỏ qua w:tbl)."""
    out = []
    pi = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            out.append((pi, Paragraph(child, doc)))
            pi += 1
    return out


def body_seq(doc):
    """[("p", text gập dấu) / ("tbl", chỉ số bảng)] theo thứ tự thân tài liệu."""
    seq = []
    ti = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            txt = fold(re.sub(r"\s+", " ", Paragraph(child, doc).text)).strip()
            seq.append(("p", txt))
        elif child.tag == qn("w:tbl"):
            seq.append(("tbl", ti))
            ti += 1
    return seq


def region_spans(seq):
    """[(occ_id, start_idx, end_idx)] của 4 vùng ngành theo M02_OCCURRENCES.

    Assert khi heading không duy nhất / không theo thứ tự / thiếu end_marker.
    """
    spans = []
    prev_start = -1
    for occ_id, heading, end_markers, _headers in M02_OCCURRENCES:
        hits = [i for i, (k, t) in enumerate(seq)
                if k == "p" and re.search(heading, t)]
        assert len(hits) == 1, \
            "%s: match_heading %r khớp %d đoạn != 1" % (occ_id, heading,
                                                        len(hits))
        h = hits[0]
        assert h > prev_start, \
            "%s: heading (%d) không đứng sau vùng trước (%d)" % (
                occ_id, h, prev_start)
        end = None
        for i in range(h + 1, len(seq)):
            if seq[i][0] == "p" and any(re.search(m, seq[i][1])
                                        for m in end_markers):
                end = i
                break
        assert end is not None, \
            "%s: không tìm thấy end_marker %r sau heading" % (occ_id,
                                                              end_markers)
        spans.append((occ_id, h, end))
        prev_start = h
    return spans


def spec_json() -> dict:
    return json.loads(SPEC_PATH.read_text(encoding="utf-8-sig"))


@check("template NGUON PLII_M02: dung 4 vung nganh co header cot roi o Muc A, 0 w:tbl (moi luoi mat khi convert)")
def t_m02_source_regions():
    # Đối chiếu template NGUỒN gốc (chưa qua generator): nơi duy nhất phản
    # ánh trung thực cấu trúc mẫu; template làm việc app/templates/ được
    # generator chèn bảng nên số lượng w:tbl thay đổi theo từng lần build.
    from tools.build_form_maps import find_source
    doc = Document(str(find_source("PLII_M02")))
    seq = body_seq(doc)
    n_tbl = sum(1 for k, _ in seq if k == "tbl")
    assert n_tbl == 0, \
        "template nguồn phải có 0 w:tbl (lưới bị converter mất), thấy %d" % n_tbl

    spans = region_spans(seq)  # đã assert duy nhất/thứ tự/end_marker
    for (occ_id, h, end), (_oid, _hd, _ed, headers) in zip(spans,
                                                           M02_OCCURRENCES):
        if occ_id == "nganh_hien_tai":
            continue  # nguồn mất cả lưới C.3 -> vùng chỉ còn heading
        hdr_folds = [fold(re.sub(r"\s+", " ", hh)).strip() for hh in headers]
        texts = [t for k, t in seq[h + 1:end] if k == "p" and t]
        matched = set()
        extra = []
        for t in texts:
            hit_at = next((fi for fi, hf in enumerate(hdr_folds)
                           if t.startswith(hf)), None)
            if hit_at is None:
                extra.append(t)
            else:
                matched.add(hit_at)
        assert not extra, "%s: vùng chứa paragraph lạ %r" % (occ_id, extra)
        assert matched == set(range(len(headers))), \
            "%s: thiếu tiêu đề cột ở template nguồn, chỉ thấy %d/%d" % (
                occ_id, len(matched), len(headers))


@check("template lam viec PLII_M02: moi w:tbl nam trong 1 trong 4 vung nganh va dong bo 1-1 voi dynamic_tables")
def t_m02_working_template_regions():
    doc = Document(str(template_path("PLII_M02")))
    fm_loops = sorted(t.get("loop") for t in
                      form_maps()["PLII_M02"].get("dynamic_tables") or []
                      if t.get("loop"))
    seq = body_seq(doc)
    spans = region_spans(seq)

    loop_to_occ = {"nganh_bo_sung": "nganh_bo_sung",
                   "nganh_xoa_bo": "nganh_xoa_bo",
                   "nganh_sua_doi_chi_tiet": "nganh_sua_doi_chi_tiet",
                   LOOP_NAME: "nganh_hien_tai"}
    unknown = [l for l in fm_loops if l not in loop_to_occ]
    assert not unknown, \
        "form_maps khai báo loop ngành lạ: %r (chỉ được 4 semantic loop)" % unknown

    tbl_region = []
    for i, (k, _v) in enumerate(seq):
        if k != "tbl":
            continue
        owner = next((oid for oid, h, e in spans if h < i < e), None)
        assert owner is not None, \
            "w:tbl thứ %d nằm ngoài 4 vùng ngành (phần phụ lục không có bảng)" % i
        tbl_region.append(owner)
    declared_regions = sorted(loop_to_occ[l] for l in fm_loops)
    assert sorted(tbl_region) == declared_regions, \
        "mất đồng bộ template<->form_maps: bảng nằm ở %r nhưng dynamic_tables " \
        "khai báo %r" % (sorted(tbl_region), declared_regions)


@check("spec PLII_M02: 4 entry nganh, occurrence_id ngu ngu nghia, match_heading/nth/end_markers hop le")
def t_m02_spec_occurrence_identity():
    raw = spec_json()
    entries = raw.get("tables") or []
    occ_ids = []
    for e in entries:
        occs = e.get("occurrences")
        assert occs, "entry %r thiếu occurrences" % e.get("spec", {}).get("id")
        assert len(occs) == 1, "M02: mỗi entry đúng 1 occurrence, thấy %d" % (
            len(occs),)
        occ = occs[0]
        for key in ("occurrence_id", "match_heading", "nth", "sync_group",
                    "end_markers"):
            assert key in occ, "occurrence thiếu %r" % key
        occ_ids.append(occ["occurrence_id"])
        assert isinstance(occ["end_markers"], list) and occ["end_markers"]
    assert sorted(occ_ids) == sorted(oid for oid, *_ in M02_OCCURRENCES), \
        "occurrence_id %r != %r" % (occ_ids, [o[0] for o in M02_OCCURRENCES])

    doc = Document(str(template_path("PLII_M02")))
    paras = [fold(re.sub(r"\s+", " ", p.text)).strip()
             for _, p in body_paras(doc)]
    by_id = {e["occurrences"][0]["occurrence_id"]: e for e in entries}
    for occ_id, heading, end_markers, _ in M02_OCCURRENCES:
        e = by_id[occ_id]
        occ = e["occurrences"][0]
        assert occ["match_heading"] == heading, \
            "%s: match_heading %r != %r (đối chiếu template)" % (
                occ_id, occ["match_heading"], heading)
        assert occ["nth"] == 1, "%s: nth phải là 1 (heading duy nhất)" % occ_id
        assert occ["sync_group"] is None, \
            "%s: sync_group phải null (PLAN A: không auto-sync loop khác nghĩa)" % occ_id
        hits = sum(1 for t in paras if re.search(occ["match_heading"], t))
        assert hits == occ["nth"], \
            "%s: heading khớp %d lần != nth=%d" % (occ_id, hits, occ["nth"])
        for m in occ["end_markers"]:
            assert any(re.search(m, t) for t in paras), \
                "%s: end_marker %r không tồn tại trong template" % (occ_id, m)


@check("blocks PLII_M02 hien trang: 3 vung Muc A chua co loop (MISSING), hoac da du 4 loop theo occurrence")
def t_m02_blocks_region_state():
    data = json.loads(blocks_path("PLII_M02").read_text(encoding="utf-8"))
    names = [b.get("loop") for b in data["blocks"] if b.get("t") == "loop"]
    dupes = sorted({n for n in names if n and names.count(n) > 1})
    assert not dupes, "loop bị trùng tên: %r" % (dupes,)
    loops = set(names)
    # Vùng 4 (nganh_hien_tai) dùng đúng loop production ds_nganh (bảng C.3);
    # trạng thái tích hợp được xác định qua 3 loop MỚI của Mục A.
    new_required = {oid for oid, *_ in M02_OCCURRENCES
                    if oid != "nganh_hien_tai"}
    new_present = new_required & loops
    if new_present == new_required:
        # Post-integration: đủ 4 vùng; hành vi sâu hơn do test render/isolation
        # và field-key isolation đảm nhiệm — ở đây chỉ cần cấu trúc hợp lệ.
        assert LOOP_NAME in loops, "thiếu loop %r cho vùng nganh_hien_tai" % (
            LOOP_NAME,)
        return
    if new_present:
        raise AssertionError(
            "trạng thái lai: mới có %r, thiếu %r — integration chưa xong "
            "hoặc blocks lệch spec" % (sorted(new_present),
                                       sorted(new_required - loops)))
    # Pre-integration (hiện trạng production): 3 vùng Mục A chỉ có paragraph
    # rơi, KHÔNG được tự sinh loop nào ngoài ds_nganh.
    assert loops == {LOOP_NAME}, \
        "pre-integration: chỉ được có loop %r, thấy %r" % (LOOP_NAME,
                                                           sorted(loops))


def _m02_dt(loop_name):
    fm = form_maps()["PLII_M02"]
    dts = [t for t in fm.get("dynamic_tables") or []
           if t.get("loop") == loop_name]
    if not dts:
        raise EFPI("form_maps PLII_M02 chưa có dynamic_tables cho %r "
                   "(generator chưa hỗ trợ occurrences)" % loop_name)
    assert len(dts) == 1, "nhiều hơn 1 dynamic_tables %r" % loop_name
    return fm, dts[0]


@check("form_maps PLII_M02: moi loop nganh co field key rieng duoi prefix <loop>., khong dung chung key giua cac nghiep vu")
def t_m02_field_key_isolation():
    fm = form_maps()["PLII_M02"]
    dts = [t for t in fm.get("dynamic_tables") or [] if t.get("loop")]
    names = [t["loop"] for t in dts]
    dupes = sorted({n for n in names if names.count(n) > 1})
    assert not dupes, "dynamic_tables bị trùng loop name: %r" % (dupes,)
    fld_keys = {f["key"] for f in fm.get("fields", [])}
    seen_keys = set()
    for dt in dts:
        ln = dt["loop"]
        ks = set()
        for c in dt.get("columns", []):
            tg = c.get("tag")
            if tg:
                k = "%s.%s" % (ln, tg)
                ks.add(k)
                if ln != LOOP_NAME:
                    assert k in fld_keys, \
                        "thiếu field %s (draft phải tách theo loop)" % k
        inter = ks & seen_keys
        assert not inter, \
            "các loop dùng chung field key (dữ liệu sẽ trộn): %r" % sorted(inter)
        seen_keys |= ks


@check("render PLII_M02 post-integration: du 4 vung, moi loop nhan dung du lieu cua minh, khong tron du lieu giua cac nghiep vu")
def t_m02_render_isolation():
    fm = form_maps()["PLII_M02"]
    loops = {t.get("loop") for t in fm.get("dynamic_tables") or []}
    # vùng 4 dùng đúng loop production ds_nganh
    required = ({oid for oid, *_ in M02_OCCURRENCES
                 if oid != "nganh_hien_tai"} | {LOOP_NAME})
    missing = sorted(required - loops)
    if missing:
        raise EFPI("chưa đủ 4 loop trong form_maps (thiếu %r)" % missing)

    data = {
        "nganh_bo_sung": [
            {"ten_nganh_bo_sung": "BS-DỮ-LIỆU-1", "ma_nganh": "111111",
             "la_nganh_chinh": "X"},
            {"ten_nganh_bo_sung": "BS-DỮ-LIỆU-2", "ma_nganh": "222222",
             "la_nganh_chinh": ""},
        ],
        "nganh_xoa_bo": [
            {"ten_nganh_bo": "XB-DỮ-LIỆU-1", "ma_nganh": "333333",
             "la_nganh_chinh": ""},
        ],
        "nganh_sua_doi_chi_tiet": [
            {"ten_nganh": "SD-DỮ-LIỆU-1", "ma_nganh": "444444"},
        ],
        "ds_nganh": SAMPLE_ROWS,
    }
    out_dir = OUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "PLII_M02_occurrences.docx"
    write_rendered(str(template_path("PLII_M02")), data, out_path)

    doc = Document(str(out_path))
    tables = body_tables(doc)
    by_header = {}
    for tbl in tables:
        if not tbl.rows:
            continue
        hdr = fold(" ".join(tr_cell_texts(tbl, 0)))
        by_header[hdr] = tbl
    assert len(by_header) == 4, \
        "kỳ vọng 4 bảng ngành với 4 header phân biệt, thấy %d" % len(by_header)

    def rows_of(tbl):
        return [tr_cell_texts(tbl, ri) for ri in range(len(tbl.rows))]

    # mỗi bảng: STT tự sinh bắt đầu từ 1; dữ liệu KHÔNG trộn giữa các bảng;
    # số cột đúng contract từng region (sửa đổi 3 cột, còn lại 4)
    checks = [
        ("ten nganh, nghe kinh doanh duoc bo sung", 4,
         ["BS-DỮ-LIỆU-1", "BS-DỮ-LIỆU-2"], ["XB-DỮ-LIỆU-1", "SD-DỮ-LIỆU-1"]),
        ("ten nganh, nghe kinh doanh duoc bo khoi danh sach da dang ky", 4,
         ["XB-DỮ-LIỆU-1"], ["BS-DỮ-LIỆU-1", "SD-DỮ-LIỆU-1"]),
        ("ten nganh, nghe kinh doanh ma nganh", 3,
         ["SD-DỮ-LIỆU-1"], ["BS-DỮ-LIỆU-1", "XB-DỮ-LIỆU-1"]),
    ]
    for hdr_frag, ncols, own_vals, foreign_vals in checks:
        matches = [h for h in by_header if hdr_frag in h]
        assert len(matches) == 1, "header %r khớp %d bảng" % (hdr_frag,
                                                              len(matches))
        rows = rows_of(by_header[matches[0]])
        assert all(len(row) == ncols for row in rows), \
            "bảng %r lệch số ô: %r" % (hdr_frag, [len(r) for r in rows])
        body_txt = " ".join(c for row in rows[1:] for c in row)
        for v in own_vals:
            assert v in body_txt, "bảng %r thiếu dữ liệu của chính nó %r" % (
                hdr_frag, v)
        for v in foreign_vals:
            assert v not in body_txt, \
                "NHIỄM DỮ LIỆU: bảng %r chứa dữ liệu của vùng khác %r" % (
                    hdr_frag, v)
        n_data = sum(1 for row in rows[1:] if any(c.strip() for c in row))
        stts = [row[0].strip() for row in rows[1:]
                if any(c.strip() for c in row)]
        assert stts == [str(k + 1) for k in range(n_data)], \
            "STT tự sinh sai: %r" % (stts,)


@check("semantic identity PLII_M02: loop ds_nganh (C.3) mang occurrence_id/semantic_key 'nganh_hien_tai' sau integration")
def t_m02_semantic_identity_hien_tai():
    fm, dt = _m02_dt(LOOP_NAME)
    got = dt.get("occurrence_id") or dt.get("semantic_key")
    if not got:
        raise EFPI("dynamic_tables ds_nganh chưa mang occurrence_id/"
                   "semantic_key (generator chưa hỗ trợ occurrences)")
    assert got == "nganh_hien_tai", \
        "semantic của bảng C.3 phải là 'nganh_hien_tai', thấy %r" % got


def main() -> int:
    for fid in FORMS:
        t_blocks_loop_structure(fid)
        t_form_maps_metadata(fid)
        t_render_three_rows(fid)
    t_m02_source_regions()
    t_m02_working_template_regions()
    t_m02_spec_occurrence_identity()
    t_m02_blocks_region_state()
    t_m02_field_key_isolation()
    t_m02_semantic_identity_hien_tai()
    t_m02_render_isolation()
    counts = {"PASS": 0, "FAIL": 0, "EXPECTED-FAIL-PRE-INTEGRATION": 0}
    for status, _ in _results:
        counts[status] += 1
    print("-" * 60)
    print("test_industry_household: %d PASS / %d FAIL / %d "
          "EXPECTED-FAIL-PRE-INTEGRATION%s" % (
              counts["PASS"], counts["FAIL"],
              counts["EXPECTED-FAIL-PRE-INTEGRATION"],
              "" if not STRICT else " [P1C_STRICT=1]"))
    for status, name in _results:
        if status != "PASS":
            print(" - %s:" % status, name)
    fail = counts["FAIL"]
    efpi = counts["EXPECTED-FAIL-PRE-INTEGRATION"]
    if STRICT and efpi:
        print("P1C_STRICT=1: %d EXPECTED-FAIL-PRE-INTEGRATION tính là FAIL"
              % efpi)
        return 1
    return 0 if fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
