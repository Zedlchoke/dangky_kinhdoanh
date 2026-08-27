# -*- coding: utf-8 -*-
"""Test Wave L1 extractor/schema (tools/docx_to_blocks.py + tools/layout_roles.py).

Fixture tự sinh DOCX bằng python-docx trong THƯ MỤC TEMP BÊN NGOÀI repo
(tempfile.mkdtemp) — KHÔNG BAO GIỜ ghi vào app/data/blocks. Chạy:
    python3 -X utf8 -m tools.test_layout_l1_extractor     # exit 0 = PASS
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import unittest

TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(TOOLS_DIR)
sys.path.insert(0, TOOLS_DIR)

sys.path.insert(0, os.path.join(TOOLS_DIR, "census"))

import docx_to_blocks as d2b          # noqa: E402
import check_layout_l1 as cl1         # noqa: E402
import layout_roles                   # noqa: E402
from layout_roles import classify_role, fold  # noqa: E402

from docx import Document             # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.opc.constants import CONTENT_TYPE as CT    # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.opc.packuri import PackURI  # noqa: E402
from docx.opc.part import Part        # noqa: E402
from docx.oxml.ns import qn           # noqa: E402
from docx.shared import Pt            # noqa: E402

FIXTURE_ID = "FIX_L1_T01"

FOOTNOTES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="1"><w:p><w:r><w:t> Chu ho kinh doanh ky truc tiep vao phan nay.</w:t></w:r></w:p></w:footnote>'
    '<w:footnote w:id="7"><w:p><w:r><w:t>Ghi chú chân trang số 7.</w:t></w:r></w:p></w:footnote>'
    '</w:footnotes>'
).encode("utf-8")


def _add_raw_child(run, tag, attrs=None):
    el = run._element.makeelement(qn(tag), attrs or {})
    run._element.append(el)
    return el


def build_fixture_doc():
    """DOCX tổng hợp phủ mọi tính năng L1."""
    doc = Document()

    # 1) hard break br + cr + tab trong cùng paragraph
    p = doc.add_paragraph()
    r = p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    _add_raw_child(r, "w:br")
    t = _add_raw_child(r, "w:t")
    t.text = "Độc lập - Tự do - Hạnh phúc"
    _add_raw_child(r, "w:cr")
    _add_raw_child(r, "w:tab")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 2) vertAlign superscript + subscript
    p2 = doc.add_paragraph()
    p2.add_run("Tôi là")
    sup = p2.add_run("1")
    sup.font.superscript = True
    p2.add_run(" và H")
    sub = p2.add_run("2")
    sub.font.subscript = True
    p2.add_run("O bình thường")

    # 3) footnoteReference chèn tay + text chưa biết id 99
    p3 = doc.add_paragraph()
    r31 = p3.add_run("Ký tên")
    _add_raw_child(r31, "w:footnoteReference", {qn("w:id"): "7"})
    p3.add_run(" cuối dòng")
    p4 = doc.add_paragraph()
    r41 = p4.add_run("Mất chú thích")
    _add_raw_child(r41, "w:footnoteReference", {qn("w:id"): "99"})

    # 4) indent/spacing/line-spacing multiple + keep-with-next
    p5 = doc.add_paragraph("Văn bản có format đầy đủ.")
    pf = p5.paragraph_format
    pf.left_indent = Pt(12)
    pf.right_indent = Pt(6)
    pf.first_line_indent = Pt(-9)  # hanging indent
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    pf.keep_with_next = True

    # 5) line spacing exact (Pt) -> lineSpacingPt
    p6 = doc.add_paragraph("Giãn dòng exact.")
    p6.paragraph_format.line_spacing = Pt(14)

    # 6) marker kỹ thuật vs dòng pháp lý trong CÙNG một biểu mẫu
    doc.add_paragraph("PL1 - Mau so 1: Giấy đề nghị đăng ký doanh nghiệp")
    doc.add_paragraph("Mẫu số 1")

    # 7) the biến để kiểm tra w giữ nguyên công thức cũ
    p8 = doc.add_paragraph("Họ tên: {dn.ten_tieng_viet} và {chk_dong_y}")

    return doc


def attach_footnotes_part(doc):
    package = doc.part.package
    part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES,
                FOOTNOTES_XML, package)
    doc.part.relate_to(part, RT.FOOTNOTES)


class TestHardBreakNewline(unittest.TestCase):
    def setUp(self):
        self.doc = build_fixture_doc()
        attach_footnotes_part(self.doc)

    def test_br_cr_tab_bao_toan_ky_tu(self):
        blocks = d2b.convert_document(self.doc, FIXTURE_ID, "fixture")["blocks"]
        txt = "".join(
            r["text"] for r in blocks[0]["runs"] if "text" in r
        )
        self.assertEqual(
            txt,
            "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n\t",
        )
        self.assertIn("\n", txt)
        self.assertIn("\t", txt)

    def test_char_stream_khop_run_text_python_docx(self):
        for p in self.doc.paragraphs:
            chars, _fns = d2b._char_stream(p)
            self.assertEqual("".join(c for c, _m in chars), "".join(r.text for r in p.runs))


class TestVertAlignFootnote(unittest.TestCase):
    def setUp(self):
        self.doc = build_fixture_doc()
        attach_footnotes_part(self.doc)
        self.fns = d2b.load_footnotes(self.doc)
        self.blocks = d2b.convert_document(
            self.doc, FIXTURE_ID, "fixture", footnotes=self.fns
        )["blocks"]

    def _p_runs(self, idx):
        return [r for r in self.blocks[idx]["runs"] if "text" in r]

    def test_superscript_run_co_vertalign(self):
        runs = self._p_runs(1)
        self.assertEqual(runs[0]["text"], "Tôi là")
        self.assertNotIn("vertAlign", runs[0])
        self.assertEqual(runs[1]["text"], "1")
        self.assertEqual(runs[1]["vertAlign"], "superscript")

    def test_subscript_run_co_vertalign(self):
        runs = self._p_runs(1)
        self.assertEqual(runs[-2]["text"], "2")
        self.assertEqual(runs[-2]["vertAlign"], "subscript")
        self.assertNotIn("vertAlign", runs[-1])

    def test_footnote_ref_run_dict(self):
        blk = self.blocks[2]
        refs = [r for r in blk["runs"] if r.get("role") == "footnote-ref"]
        self.assertEqual(len(refs), 1)
        self.assertEqual(refs[0]["footnoteId"], 7)
        self.assertEqual(blk.get("footnotes"), {"7": "Ghi chú chân trang số 7."})

    def test_footnote_id_mat_text_van_giu_id_fail_soft(self):
        blk = self.blocks[3]
        refs = [r for r in blk["runs"] if r.get("role") == "footnote-ref"]
        self.assertEqual(len(refs), 1)
        self.assertEqual(refs[0]["footnoteId"], 99)
        self.assertNotIn("footnotes", blk)  # không có text -> không emit map

    def test_footnotes_separator_id_am_bo_qua(self):
        fns = d2b.load_footnotes(self.doc)
        self.assertNotIn(-1, fns)
        self.assertEqual(fns.get(1), "Chu ho kinh doanh ky truc tiep vao phan nay.")


class TestParagraphFormat(unittest.TestCase):
    def setUp(self):
        self.doc = build_fixture_doc()
        self.blocks = d2b.convert_document(self.doc, FIXTURE_ID, "fixture")["blocks"]

    def test_format_day_du_pt_va_multiple(self):
        fmt = self.blocks[4].get("format")
        self.assertEqual(
            fmt,
            {
                "leftIndentPt": 12.0,
                "rightIndentPt": 6.0,
                "firstLineIndentPt": -9.0,
                "spaceBeforePt": 6.0,
                "spaceAfterPt": 3.0,
                "lineSpacing": 1.15,
                "keepWithNext": True,
            },
        )

    def test_line_spacing_exact_thanh_linespacingpt(self):
        fmt = self.blocks[5].get("format")
        self.assertEqual(fmt, {"lineSpacingPt": 14.0})

    def test_default_khong_emit_format(self):
        # paragraph marker kỹ thuật không set format -> không có key "format"
        self.assertNotIn("format", self.blocks[6])

    def test_emu_chinh_xac_120_twips_6pt(self):
        from docx.shared import Twips

        doc = Document()
        p = doc.add_paragraph("x")
        p.paragraph_format.space_before = Twips(120)  # 120 twentieths = 6pt
        self.assertEqual(d2b._para_format_dict(p), {"spaceBeforePt": 6.0})


class TestSemanticRoles(unittest.TestCase):
    def test_source_indicator_pl_marker(self):
        role = classify_role(fold("PL1 - Mau so 1: Giấy đề nghị đăng ký doanh nghiệp"))
        self.assertEqual(role, "source-indicator")

    def test_form_number_hop_le_khong_bi_an(self):
        self.assertEqual(classify_role(fold("Mẫu số 1")), "form-number")
        self.assertEqual(classify_role(fold("Mẫu số 72")), "form-number")
        self.assertNotEqual(classify_role(fold("Mẫu số 72")), "source-indicator")

    def test_cung_mau_ca_hai_dang_marker_va_mauso(self):
        doc = Document()
        doc.add_paragraph("PL2 - Mau so 5: Giấy chứng nhận đăng ký hộ kinh doanh")
        doc.add_paragraph("Mẫu số 5")
        blocks = d2b.convert_document(doc, FIXTURE_ID, "f")["blocks"]
        self.assertEqual(blocks[0].get("role"), "source-indicator")
        self.assertEqual(blocks[1].get("role"), "form-number")
        self.assertNotEqual(blocks[0]["role"], blocks[1]["role"])

    def test_ten_file_docx_la_source_indicator(self):
        self.assertEqual(
            classify_role(fold("Nguồn: PLI_M01_tach.docx")), "source-indicator"
        )

    def test_national_header_motto_separator(self):
        multi = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n-----"
        runs = [{"text": multi}]
        self.assertEqual(classify_role(fold(multi), None, runs), "national-header")
        self.assertEqual(
            classify_role(fold("Độc lập - Tự do - Hạnh phúc")),
            "national-motto",
        )
        sep = "-----------------"
        self.assertEqual(
            classify_role(fold(sep), None, [{"text": sep}]),
            "national-separator",
        )

    def test_date_salutation_note_signature(self):
        date_runs = [
            {"input": {"key": "dia_diem_ky", "w": 12}},
            {"text": ", ngày ", "i": True},
            {"input": {"key": "ngay_ky_vb", "w": 12}},
            {"input": {"key": "thang_ky_vb", "w": 12}},
            {"input": {"key": "nam_ky_vb", "w": 12}},
        ]
        self.assertEqual(classify_role(", ngay thang nam", None, date_runs), "date-line")
        self.assertEqual(classify_role(fold("Kính gửi: Ủy ban nhân dân cấp xã")), "salutation")
        self.assertEqual(classify_role(fold("Lưu ý: kê khai đầy đủ")), "note")
        sig = classify_role(fold("CHỦ HỘ KINH DOANH\n(Ký và ghi họ tên)"))
        self.assertEqual(sig, "signature")

    def test_title_subtitle_section_body(self):
        title = "GIẤY ĐỀ NGHỊ ĐĂNG KÝ HỘ KINH DOANH"
        self.assertEqual(
            classify_role(fold(title), None, [{"text": title}]), "form-title"
        )
        sub = "DOANH NGHIỆP TƯ NHÂN"
        self.assertEqual(
            classify_role(fold(sub), None, [{"text": sub}]), "form-subtitle"
        )
        self.assertEqual(classify_role(fold("1. Tên hộ kinh doanh: ")), "section-heading")
        self.assertEqual(
            classify_role(fold("Đăng ký hộ kinh doanh do tôi là chủ hộ với các nội dung sau:")),
            "body",
        )
        body = "Số định danh cá nhân: "
        self.assertEqual(
            classify_role(
                fold(body), None,
                [{"text": body}, {"input": {"key": "chu_hkd.so_cccd", "w": 20}}],
            ),
            "body",
        )

    def test_footnote_body_digit_prefix(self):
        self.assertEqual(
            classify_role(fold("6 Chủ hộ kinh doanh ký trực tiếp vào phần này.")),
            "footnote",
        )
        # mục lục "1." không được nhầm footnote
        self.assertEqual(classify_role(fold("1. Chủ hộ kinh doanh trước khi thay đổi:")), "section-heading")

    def test_fold_giong_p1_c(self):
        self.assertEqual(fold("Mẫu số 72"), "mau so 72")
        self.assertEqual(fold("Đắk Lắk"), "dak lak")
        self.assertEqual(fold("  A\n B  "), "a b")


# ---------------------------------------------------------------------------
# V2/V3 — heading đánh số bị gộp vào đuôi paragraph trước (audit Wave D)
#
# Bằng chứng thật (PLI_M72 blk110/112, PLI_M09 #88/#132/#163): template scrape
# lại gộp cả đuôi bảng lưới và heading mục kế tiếp vào CÙNG một w:p — heading
# chỉ là <w:r> riêng ở cuối (M72) hoặc dính chung cột đếm "123456789" (M09).
# Converter phải tách ranh giới này thành block "p" riêng; detector Wave L1
# giữ bản predicate đối xứng để báo đỏ khi blocks còn dính.
# ---------------------------------------------------------------------------

GLUE_TAIL = ("(VNĐ và giá trị tương đương theo đơn vị tiền nước ngoài, nếu có)"
             "Tỷ lệ (%)Ghi chú (nếu có)")
GLUE_HEAD_8 = "8. Cổ đông là nhà đầu tư nước ngoài: "
GLUE_HEAD_9 = "9. Người đại diện theo pháp luật của công ty"


def _blocks_text(block):
    return "".join(r.get("text", "") for r in block["runs"] if "text" in r)


class TestGluedNumberedHeadingSplit(unittest.TestCase):
    """(V2) extractor tách ranh giới heading đánh số dính đuôi paragraph."""

    def _convert_runs(self, runs):
        doc = Document()
        p = doc.add_paragraph()
        for r in runs:
            run = p.add_run(r.get("text", ""))
            if r.get("b"):
                run.bold = True
            if r.get("vertAlign"):
                run.font.superscript = r["vertAlign"] == "superscript"
        blocks = d2b.convert_document(doc, FIXTURE_ID, "fx")["blocks"]
        return [b for b in blocks if b.get("t") == "p"]

    def test_m72_heading_run_rieng_duoc_tach(self):
        """Dạng (1) của audit: heading là run riêng ở CUỐI w:p (evidence M72)."""
        groups = d2b.split_glued_numbered_headings([
            {"text": "("}, {"text": GLUE_TAIL[:60]}, {"text": ")Tỷ lệ ("},
            {"text": "%"}, {"text": GLUE_TAIL[60:]}, {"text": GLUE_HEAD_8},
        ])
        self.assertEqual(len(groups), 2)
        self.assertEqual(
            "".join(r["text"] for r in groups[1]), GLUE_HEAD_8)

    def test_docx_thuc_tach_thanh_block_section_heading(self):
        blocks = self._convert_runs(
            [{"text": GLUE_TAIL}, {"text": GLUE_HEAD_8, "b": True}])
        self.assertEqual(len(blocks), 2)
        self.assertEqual(blocks[0].get("role"), "body")
        self.assertEqual(blocks[1].get("role"), "section-heading")
        self.assertEqual(_blocks_text(blocks[1]), GLUE_HEAD_8)

    def test_heading_long_trong_mot_run_cung_tach(self):
        """Dạng glue khác: ')Ghi chú (nếu có)8. Cổ đông…' nằm chung MỘT run."""
        groups = d2b.split_glued_numbered_headings(
            [{"text": GLUE_TAIL + GLUE_HEAD_8}])
        self.assertEqual(len(groups), 2)
        self.assertEqual("".join(r["text"] for r in groups[0]), GLUE_TAIL)
        self.assertEqual("".join(r["text"] for r in groups[1]), GLUE_HEAD_8)

    def test_metadata_align_giu_nguyen_cho_moi_block_con(self):
        doc = Document()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(GLUE_TAIL)
        run = p.add_run(GLUE_HEAD_9)
        run.bold = True
        blocks = [b for b in d2b.convert_document(doc, FIXTURE_ID, "fx")["blocks"]
                  if b.get("t") == "p"]
        self.assertEqual(len(blocks), 2)
        self.assertTrue(all(b.get("align") == "center" for b in blocks))

    def test_superscript_chu_so_khong_mo_heading(self):
        groups = d2b.split_glued_numbered_headings([
            {"text": GLUE_TAIL}, {"text": GLUE_HEAD_9},
            {"text": "3", "vertAlign": "superscript"}, {"text": ":"},
        ])
        self.assertEqual(len(groups), 2)
        self.assertEqual(
            "".join(r["text"] for r in groups[1]), GLUE_HEAD_9 + "3:")

    def test_so_nhieu_chu_so_khong_bi_cat_roi_giua_hai_run(self):
        """Số nhiều chữ số phải giữ nguyên về một nhóm ("1"+"0." = mục "10."),
        không bao giờ tách thành đuôi mồ côi bắt đầu bằng "0. …"."""
        for runs in (
            # số ở ĐẦU đoạn: heading tự nó là block riêng, không phải glue
            [{"text": "1"}, {"text": "0. Thông tin đăng ký thuế:"}],
            # số có chữ đứng trước trong cùng run: chỉ được cắt TRƯỚC "10."
            [{"text": "Điều kiện"}, {"text": "10. Thông tin đăng ký thuế:"}],
            # chú thích số mũ/giá trị số ngay trước: không suy ra ranh giới
            [{"text": "Giấy tờ hợp lệ 76"}, {"text": GLUE_HEAD_8}],
        ):
            groups = d2b.split_glued_numbered_headings(runs)
            joined = ["|".join(r["text"] for r in g) for g in groups]
            orphans = [t for t in joined[1:] if not d2b.NUMBERED_HEADING_RUN_RE.match(t)]
            self.assertEqual(orphans, [], "%r -> %r" % (runs, joined))

    def test_dau_cham_nhap_cap_khong_tach(self):
        """'…mục 5.' + '1. Địa chỉ…' là đánh số nhiều cấp con (PLII_M01)."""
        groups = d2b.split_glued_numbered_headings(
            [{"text": "Thông tin đăng ký thuế5."},
             {"text": "1. Địa chỉ nhận thông báo thuế:"}])
        self.assertEqual(len(groups), 1)

    def test_hard_break_da_ngat_dong_khong_tach(self):
        groups = d2b.split_glued_numbered_headings(
            [{"text": "5.\n"}, {"text": "1. Địa chỉ nhận thông báo thuế:"}])
        self.assertEqual(len(groups), 1)

    def test_heading_dau_paragraph_khong_tach(self):
        self.assertEqual(
            len(d2b.split_glued_numbered_headings([{"text": GLUE_HEAD_8}])), 1)

    def test_doan_van_dai_khong_phai_heading_khong_tach(self):
        long_prose = "1. " + "nội dung pháp lý kéo dài nhiều câu bình thường " * 40
        self.assertGreater(len(long_prose), d2b.MAX_HEADING_TAIL_CHARS)
        groups = d2b.split_glued_numbered_headings(
            [{"text": "Căn cứ văn bản hợp lệ "}, {"text": long_prose}])
        self.assertEqual(len(groups), 1)

    def test_dong_ke_cham_trong_chu_khong_tach(self):
        groups = d2b.split_glued_numbered_headings(
            [{"text": "Họ tên chủ hộ "}, {"text": "1. ...................."}])
        self.assertEqual(len(groups), 1)

    def test_counter_scrape_chua_the_xac_dinh_thi_khong_doan(self):
        """PLI_M09: '123456789' + '2. Thông tin…' chung cột đếm scrape.

        Ranh giới thật KHÔNG suy ra được từ text nên converter chủ động không
        đoán; vùng lỗi do lane sửa template xử lý và vẫn bị detector L1 bắt
        qua rule (e)/(f). Test khoá hành vi để mọi thay đổi là quyết định có
        ý thức."""
        groups = d2b.split_glued_numbered_headings(
            [{"text": "Giấy chứng nhận đăng ký doanh nghiệp; "},
             {"text": "123456789"},
             {"text": "2. Thông tin người đại diện theo pháp luật"}])
        self.assertEqual(len(groups), 1)


class TestDetectorStructuralRules(unittest.TestCase):
    """(V3) rule (e)(f)(g) của tools/census/check_layout_l1.py."""

    @staticmethod
    def _block(runs, role="body"):
        return {"t": "p", "align": "left", "role": role, "runs": runs}

    def _findings(self, block):
        return cl1.structural_findings(block, block.get("role"))

    def test_e_stt_glue_bat_dau_ten_cot_lien(self):
        blk = self._block([{"text": "STTTên cổ đôngQuốc tịch"}])
        self.assertIn(("e",), tuple(f[:1] for f in self._findings(blk)))

    def test_e_stt_hop_le_dung_mot_minh_khong_bao(self):
        runs = [
            {"text": "STT"}, {"input": {"key": "ds.ho_ten", "w": 20}},
            {"text": "\tHọ và tên"},
        ]
        self.assertEqual(self._findings(self._block(runs)), [])

    def test_e_stt_theo_sau_dau_cham_va_ngoac_khong_bao(self):
        for txt in ("xếp hạng theo STT.", "(mã STT)", "STT:"):
            self.assertFalse(cl1.is_stt_glue_run(txt), txt)
        self.assertTrue(cl1.is_stt_glue_run("SttTÊN người"))
        # chú thích số mũ được miễn
        blk = self._block([{"text": "STTHọ tên", "vertAlign": "superscript"}])
        self.assertEqual(self._findings(blk), [])

    def test_f_digit_run_dai_bi_bat(self):
        blk = self._block([{"text": "12345678910"}])
        rules = [f[0] for f in self._findings(blk)]
        self.assertIn("f", rules)
        self.assertNotIn("e", rules)

    def test_f_mien_role_metadata_va_violation_trong_cung_block(self):
        meta = self._block([{"text": "PL1_M01.docx 12345"}], role="source-indicator")
        self.assertEqual(self._findings(meta), [])
        fm = self._block([{"text": "Mẫu số 72 1234567"}], role="form-number")
        self.assertEqual(self._findings(fm), [])
        still = self._block([{"text": "Số điện thoại cơ quan 1900561258 trực" }])
        self.assertEqual([f[0] for f in self._findings(still)], ["f"])

    def test_g1_numbered_text_voi_role_body(self):
        blk = self._block([{"text": "1. Đối với trường hợp bị thu hồi Giấy "
                                     "chứng nhận đăng ký hoạt động"}])
        self.assertEqual([f[0] for f in self._findings(blk)], ["g1"])

    def test_g1_hop_le_khi_role_da_la_heading(self):
        blk = self._block([{"text": "1. Tên doanh nghiệp: "}],
                          role="section-heading")
        self.assertEqual(self._findings(blk), [])

    def test_g1_mien_dong_ke_cham_diEN_tay(self):
        for txt in ("1. ....................", "2. ....................",
                    "1.\t............"):
            self.assertFalse(cl1.numbered_head_misrole(txt, "body"), txt)

    def test_g2_heading_dinh_duoiblock_truoc_pli_m72(self):
        runs = [{"text": GLUE_TAIL}, {"text": GLUE_HEAD_8}]
        out = self._findings(self._block(runs))
        self.assertEqual([f[0] for f in out], ["g2"])
        self.assertIn(GLUE_HEAD_8.strip(), out[0][1])

    def test_block_rieng_role_heading_thi_khong_bao_g1_lan_g2(self):
        self.assertEqual(
            self._findings(
                self._block([{"text": GLUE_HEAD_8}], role="section-heading")), [])
        # cùng text nhưng role=body thì đúng luật là bị (g1) nhắc — role là
        # phần dữ liệu khác ngoài cấu trúc chạy (run).
        self.assertEqual(
            [f[0] for f in self._findings(self._block([{"text": GLUE_HEAD_8}]))],
            ["g1"])

    def test_predicate_detector_khop_converter_chong_drift(self):
        cases = [
            GLUE_TAIL + GLUE_HEAD_8,
            "Điều kiện1" + "0. Thông tin đăng ký thuế:",
            "Thông tin đăng ký thuế5." + "1. Địa chỉ nhận thông báo thuế:",
            "5.\n" + "1. Địa chỉ nhận thông báo thuế:",
            "Giấy chứng nhận; 1234567892. Thông tin người đại diện",
            "Họ tên chủ hộ 1. ....................",
            GLUE_TAIL + GLUE_HEAD_9,
            "Văn bản dài " + "1. " + "nội dung thường " * 90,
        ]
        for full in cases:
            want = []
            got = []
            for m in d2b.NUMBERED_HEADING_FIND_RE.finditer(full):
                pos = m.start()
                if pos and full[:pos].strip() and \
                        d2b._is_glued_heading_cut(full, pos):
                    want.append(pos)
            for m in cl1.NUMBERED_HEAD_OFFSET_RE.finditer(full):
                pos = m.start()
                if pos and full[:pos].strip() and \
                        cl1._is_glued_heading_cut(full, pos):
                    got.append(pos)
            self.assertEqual(want, got, full)
        # guard độ dài phải đồng bộ giữa hai module
        self.assertEqual(d2b.MAX_HEADING_TAIL_CHARS,
                         cl1.GLUED_HEADING_MAX_TAIL_CHARS)
        self.assertEqual(sorted(d2b.HEADING_ROLES),
                         sorted(cl1.HEADING_ROLES_FOR_NUMBERING))


REPO_OVERRIDES_PATH = os.path.join(TOOLS_DIR, "layout_role_overrides.json")

# 12 vi tri (g1) cua audit round 2: heading danh so THAT bi guard len(folded)
# <=120 cua layout_roles tier-11 xep body — xu ly bang override theo PREFIX GAP
# DAU (khong by_index vi truot sau regen). Chi muc bang chuoi goc de doc duoc.
G1_HEADING_CASES = [
    ("PLII_M02",
     "1. Bổ sung ngành, nghề kinh doanh sau (kê khai trong trường hợp hộ kinh "
     "doanh đăng ký bổ sung ngành, nghề kinh doanh vào danh sách đã đăng ký):"),
    ("PLII_M02",
     "2. Bỏ ngành, nghề kinh doanh sau (kê khai trong trường hợp hộ kinh doanh "
     "đăng ký bỏ ngành, nghề kinh doanh khỏi danh sách đã đăng ký):"),
    ("PLII_M02",
     "3. Sửa đổi ngành, nghề kinh doanh chính6 (kê khai trong trường hợp hộ "
     "kinh doanh chọn một trong các ngành nghề đã đăng ký làm nghề chính):"),
    ("PLI_M12",
     "1. Bổ sung ngành, nghề kinh doanh sau (kê khai trong trường hợp doanh "
     "nghiệp thông báo bổ sung ngành, nghề kinh doanh vào danh sách):"),
    ("PLI_M12",
     "2. Bỏ ngành, nghề kinh doanh sau (kê khai trong trường hợp doanh nghiệp "
     "thông báo bỏ ngành, nghề kinh doanh khỏi danh sách):"),
    ("PLI_M12",
     "3. Sửa đổi chi tiết ngành, nghề kinh doanh sau (kê khai trong trường hợp "
     "doanh nghiệp thông báo thay đổi nội dung chi tiết ngành nghề):"),
    ("PLII_M07", "4. Vốn kinh doanh:" + "." * 96),
    ("PLII_M07",
     "7. Danh sách thành viên hộ gia đình đăng ký thành lập hộ kinh doanh "
     "(chỉ ghi trong trường hợp các thành viên hộ gia đình cùng đăng ký):"),
    ("PLII_M14",
     "1. Đối với trường hợp hộ kinh doanh bị thu hồi Giấy chứng nhận đăng ký "
     "hộ kinh doanh để thực hiện biện pháp cưỡng chế nợ thuế theo đề nghị của "
     "Cơ quan quản lý thuế:"),
    ("PLI_M23",
     "3. Thông tin về Giấy phép đầu tư/Giấy chứng nhận đầu tư/Giấy chứng nhận "
     "đăng ký hoạt động của chi nhánh/văn phòng đại diện/giấy tờ có giá trị "
     "tương đương khác do Cơ quan đăng ký đầu tư cấp:"),
    ("PLI_M28",
     "3. Chi nhánh chủ quản (chỉ kê khai đối với trường hợp đăng ký chấm dứt "
     "hoạt động của địa điểm kinh doanh trực thuộc chi nhánh):"),
    ("PLI_M63",
     "1. Đối với trường hợp bị thu hồi Giấy chứng nhận đăng ký hoạt động chi "
     "nhánh, văn phòng đại diện, địa điểm kinh doanh thuộc chi nhánh để thực "
     "hiện biện pháp cưỡng chế nợ thuế:"),
]

# Cùng dạng số nhưng KHÔNG thuộc nhãn nào được override -> phải giữ nguyên role
NEGATIVE_OVERRIDE_CASES = [
    ("PLII_M14", "2. Đối với trường hợp còn lại:"),
    ("PLI_M23", "4. Địa chỉ chi nhánh/văn phòng đại diện/địa điểm kinh doanh:"),
    ("PLI_M28", "4. Lý do chấm dứt hoạt động:\t"),
    ("PLII_M07",
     "5. Chủ thể thành lập hộ kinh doanh (ghi một trong các chủ thể sau):"),
    ("PLI_M28", "1. Tình trạng thành lập "),
    ("PLI_M02", "10. Thông tin đăng ký thuế"),
]


def _repo_overrides():
    return d2b.load_role_overrides(REPO_OVERRIDES_PATH)


def _effective_role(fid, raw_text):
    folded = fold(raw_text)
    role = classify_role(folded, None, [{"text": raw_text}])
    return d2b.apply_role_override(_repo_overrides(), fid, -1, folded, role)


def _blocks_corpus():
    bd = os.path.join(ROOT, "app", "data", "blocks")
    out = []
    for name in sorted(os.listdir(bd)):
        if name.endswith(".json"):
            with io.open(os.path.join(bd, name), encoding="utf-8-sig") as fh:
                out.append((name[:-5], json.load(fh)))
    return out


class TestOverridesAndSchema(unittest.TestCase):
    def test_overrides_file_temp_trong_hop_le(self):
        # File override rỗng/thiếu vẫn hợp lệ (fail-soft của converter).
        tmpdir = tempfile.mkdtemp(prefix="l1_override_empty_")
        path = os.path.join(tmpdir, "empty.json")
        with io.open(path, "w", encoding="utf-8") as fh:
            fh.write("{}")
        self.assertEqual(d2b.load_role_overrides(path), {})

    def test_repo_overrides_shape_hop_le(self):
        data = _repo_overrides()
        self.assertIsInstance(data, dict)
        entries = data.get("by_folded_prefix") or []
        self.assertTrue(entries, "cần entry heading đánh số dài (round 2)")
        for e in entries:
            self.assertIsInstance(e.get("prefix"), str)
            self.assertTrue(e["prefix"].strip())
            self.assertIn(e.get("role"), layout_roles.ROLES)
            self.assertIn("reason", e)   # mọi entry phải ghi rõ lý do
            forms = e.get("forms")
            self.assertIsInstance(forms, list)
            self.assertTrue(forms, "entry override phải khoá theo form")
            for f in forms:
                self.assertRegex(f, r"^PLI+I?_M\d+$")

    def test_repo_overrides_prefix_dinh_vi_doc_nhat(self):
        # Mỗi prefix phải khớp ĐÚNG 1 block "p" trong từng form khai báo —
        # chống prefix trôi/khớp lan khi Lead regen.
        ov = _repo_overrides()
        corpus = dict(_blocks_corpus())
        for e in ov["by_folded_prefix"]:
            for fid in e["forms"]:
                blks = corpus[fid]["blocks"]
                hits = []
                for b in blks:
                    if b.get("t") != "p":
                        continue
                    txt = "".join(r.get("text", "") for r in b["runs"]
                                  if "text" in r)
                    if fold(txt).startswith(e["prefix"]):
                        hits.append(b.get("role"))
                self.assertEqual(
                    len(hits), 1,
                    "%s %r khớp %d block" % (fid, e["prefix"], len(hits)))

    def test_g1_heading_dai_day_role_section_heading(self):
        for fid, raw in G1_HEADING_CASES:
            folded = fold(raw)
            self.assertGreater(len(folded), 0)
            self.assertEqual(
                _effective_role(fid, raw), "section-heading",
                "%s %r" % (fid, raw[:50]))

    def test_g1_so_khac_giu_nguyen_role_hoac_heading_ngan(self):
        for fid, raw in NEGATIVE_OVERRIDE_CASES:
            role = _effective_role(fid, raw)
            self.assertIn(role, {"body", "section-heading"}, "%s %r" % (fid, raw))

    def test_g1_all_toan_bang_blocks_chi_doi_dung_12_vitri(self):
        # Áp toàn bộ bảng override lên app/data/blocks hiện tại: đúng 12 block
        # trên 7 form được chuyển sang section-heading, không block nào khác.
        ov = _repo_overrides()
        flips = {}
        ALLOWED_FORMS = {f for e in ov["by_folded_prefix"] for f in e["forms"]}
        for fid, data in _blocks_corpus():
            for b in data.get("blocks", []):
                if b.get("t") != "p":
                    continue
                raw = "".join(r.get("text", "") for r in b["runs"]
                              if "text" in r)
                if not raw.strip():
                    continue
                base = classify_role(fold(raw), None, b["runs"])
                after = d2b.apply_role_override(ov, fid, -1, fold(raw), base)
                if after != base:
                    # [Lead sync sau regen P0] corpus blocks giờ ĐÃ chứa 12 flip
                    # nên proxy cũ "stored phải là body" hết nghĩa. Thay bằng
                    # mệnh đề MẠNH HƠN: nơi override bắn, blocks phải lưu ĐÚNG
                    # role hiệu lực (base+override); override đè lên heading sẵn
                    # có hoặc lệch giữa corpus/override vẫn đỏ như cũ.
                    self.assertEqual(
                        b.get("role"), after,
                        "%s: override bắn (%r -> %r) nhưng blocks lưu %r — "
                        "corpus không khớp role hiệu lực" % (
                            fid, base, after, b.get("role")))
                    flips.setdefault(fid, []).append((base, after))
                    self.assertIn(fid, ALLOWED_FORMS)
        total = sum(len(v) for v in flips.values())
        self.assertEqual(total, 12, "số vị trí đổi role: %s" %
                         {k: len(v) for k, v in flips.items()})
        self.assertEqual(sorted(flips), sorted([
            "PLII_M02", "PLII_M07", "PLII_M14",
            "PLI_M12", "PLI_M23", "PLI_M28", "PLI_M63"]))
        for fid, pairs in flips.items():
            self.assertTrue(all(a == "body" and b == "section-heading"
                                for a, b in pairs), fid)


    def test_override_by_index_va_prefix(self):
        ov = {
            "by_index": {FIXTURE_ID: {"3": "note"}},
            "by_folded_prefix": [
                {"prefix": "kinh gui", "role": "salutation", "forms": ["KHAC"]}
            ],
        }
        self.assertEqual(d2b.apply_role_override(ov, FIXTURE_ID, 3, "bat ky", "body"), "note")
        # form khác + prefix khớp nhưng bị lọc theo forms -> giữ nguyên
        self.assertEqual(d2b.apply_role_override(ov, FIXTURE_ID, 9, "kinh gui xy", "body"), "body")
        # role lạ trong override -> bỏ qua, trả rule role
        bad = {"by_index": {FIXTURE_ID: {"3": "khong-phai-role"}}}
        self.assertEqual(d2b.apply_role_override(bad, FIXTURE_ID, 3, "x", "body"), "body")

    def test_schema_version_v3(self):
        doc = build_fixture_doc()
        data = d2b.convert_document(doc, FIXTURE_ID, "fixture")
        self.assertEqual(data.get("v"), 3)
        self.assertIn("blocks", data)
        self.assertIn("id", data)
        self.assertIn("ten_mau", data)

    def test_block_cu_thieu_field_moi_van_doc_duoc(self):
        # block cũ schema v2: thiếu role/format/vertAlign — hợp lệ, mặc định body
        old_block = {"t": "p", "align": "left",
                     "runs": [{"text": "Nội dung cũ"}]}
        folded = fold("Nội dung cũ")
        self.assertEqual(layout_roles.classify_role(folded, None, old_block["runs"]), "body")
        self.assertNotIn("role", old_block)


class TestIdempotentAndCompat(unittest.TestCase):
    def test_convert_2_lan_byte_giong_nhau(self):
        fm = json.loads(io.open(
            os.path.join(ROOT, "app", "data", "form_maps.json"),
            encoding="utf-8").read())
        d1 = json.dumps(d2b.convert_document(build_fixture_doc(), FIXTURE_ID, "fx"),
                        ensure_ascii=False, indent=1)
        d2 = json.dumps(d2b.convert_document(build_fixture_doc(), FIXTURE_ID, "fx"),
                        ensure_ascii=False, indent=1)
        self.assertEqual(d1, d2)

        # template thật: convert 2 lần (Document mở riêng) phải byte-giống nhau
        a = json.dumps(d2b.convert("PLI_M01", fm), ensure_ascii=False, indent=1)
        b = json.dumps(d2b.convert("PLI_M01", fm), ensure_ascii=False, indent=1)
        self.assertEqual(a, b)

    def test_random_khong_vao_output_chi_dung_trong_self_check_seed(self):
        # `import random` trong docx_to_blocks chỉ dùng ở self_check() với seed
        # cố định (20260824 / 42) — không nằm trên đường sinh output.
        src = io.open(os.path.join(TOOLS_DIR, "docx_to_blocks.py"), encoding="utf-8").read()
        uses = [ln.strip() for ln in src.splitlines() if "random" in ln and "import random" not in ln]
        self.assertTrue(all(("rng" in u or "Random" in u) for u in uses),
                        "random phải giới hạn trong rng seed của self_check: %r" % uses)

    def test_fixture_ghi_temp_ngoai_repo_khong_dung_blocks_dir(self):
        tmpdir = tempfile.mkdtemp(prefix="l1_extract_test_")
        self.assertFalse(os.path.abspath(tmpdir).startswith(os.path.abspath(ROOT)))
        path = os.path.join(tmpdir, "%s.docx" % FIXTURE_ID)
        doc = build_fixture_doc()
        attach_footnotes_part(doc)
        doc.save(path)
        doc2 = __import__("docx").Document(path)
        data = d2b.convert_document(doc2, FIXTURE_ID, "fx")
        out = os.path.join(tmpdir, "%s.json" % FIXTURE_ID)
        with io.open(out, "w", encoding="utf-8") as fh:
            fh.write(json.dumps(data, ensure_ascii=False))
        self.assertTrue(os.path.exists(out))
        self.assertEqual(data["v"], 3)


if __name__ == "__main__":
    unittest.main(verbosity=2)
