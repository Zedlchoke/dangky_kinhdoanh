#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test_m09_decontaminate.py — Lane C (laneC(M09 decontamination, O1 verdict
REMOVED_BY_SOURCE_CLEANUP cho section vung_nhiem_mau_6789_10, 2026-09-06).

Khoanh vùng nhiễm đã đo bằng python-docx trên file nguồn
06_ho_kinh_doanh_va_chuyen_doi/02_chuyen_doi_hkd_sang_dn/
DN_PL1_Mau_09_danh_sach_thanh_vien_cong_ty_hop_danh.docx (trước khi cắt):

  - Mẫu 9 THẬT (PL1): body child 0..102 — tiêu đề, 'Mẫu số 9', mục 8-10,
    cam kết, khối ký 'CÁC THÀNH VIÊN HỢP DANH' (child 89-90) và chân trang
    chú thích 1-9 (child 91..102). Child 102 GLUE: nửa đầu là chân trang thật
    của Mẫu 9, đuôi run 'Mẫu số 6' là điểm bắt đầu nhiễm.
  - VÙNG NHIỄM (đã xóa khỏi nguồn): run 'Mẫu số 6' đuôi child 102 + child
    103..231 — toàn văn Mẫu 6 (tiêu đề child 103, không có nhãn 'Mẫu số 6'
    tách dòng vì nó dính chân trang Mẫu 9), Mẫu 7 (child 131), Mẫu 8 (child
    156), Mẫu 9 lần 2 header 14 cột (child 182) và Mẫu 10 (child 214) nằm
    sau chân trang Mẫu 9 thật. sectPr cuối body được giữ nguyên.

Mệnh đề kiểm chứng (docs/legal_review_backlog.md mục 12, quyết định 1A):

  [SRC]  nguồn M09 KHÔNG còn >=5 fingerprint đặc trưng của từng mẫu nhiễm
         (4 tiêu đề danh sách, marker section Mẫu 9 lần 2, header glue 14 cột),
         nhãn 'Mẫu số N' đứng riêng chỉ còn đúng 1 ('Mẫu số 9' thật), không
         còn đường '_' tàn dư.
  [TPL]  template production app/templates/PLI_M09.docx cũng sạch fingerprint
         + tàn dư residue đã bị xóa khỏi nguồn không hồi sinh. Assertion
         residue này ĐỎ trên nhánh lane cho tới khi Lead merge xóa nguồn rồi
         regen template (lệnh ở README báo cáo lane).
  [KEY]  key thật của Mẫu 9 vẫn còn trong form_maps: dynamic table
         ds_thanh_vien_hop_danh đủ 14 cột (kể cả 3 cột mới loai_tai_san_
         gop_von / thoi_han_gop_von / ghi_chu), fields đăng ký theo loop,
         template chứa thẻ mở/đóng loop và bảng 14 cột vật lý.

KHÔNG phải sandbox build — test này soi ARTIFACT thật (nguồn + template +
form_maps) để chốt decontamination "xóa thật khỏi nguồn", không ngụy trang
output. Chạy: python3 -X utf8 -m tools.test_m09_decontaminate (từ root repo).
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC = (ROOT / "06_ho_kinh_doanh_va_chuyen_doi" / "02_chuyen_doi_hkd_sang_dn"
       / "DN_PL1_Mau_09_danh_sach_thanh_vien_cong_ty_hop_danh.docx")
TPL = ROOT / "app" / "templates" / "PLI_M09.docx"
FORM_MAPS = ROOT / "app" / "data" / "form_maps.json"
LOOP = "ds_thanh_vien_hop_danh"

# Fingerprint vùng nhiễm — mỗi mục là chuỗi đặc trưng của MỘT mẫu nhiễm, đã
# đối chiếu KHÔNG xuất hiện trong vùng Mẫu 9 thật (child 0..102 cũ). So sánh
# trên NFC để chịu glyph tổ hợp (vd 'NGHIỆP' với dấu trên glyph rời).
CONTAMINATION_FINGERPRINTS = (
    "DANH SÁCH THÀNH VIÊN CÔNG TY TRÁCH NHIỆM HỮU HẠN HAI THÀNH VIÊN TRỞ LÊN",  # Mẫu 6
    "DANH SÁCH CỔ ĐÔNG SÁNG LẬP CÔNG TY CỔ PHẦN",                              # Mẫu 7
    "DANH SÁCH CỔ ĐÔNG LÀ NHÀ ĐẦU TƯ NƯỚC NGOÀI",                              # Mẫu 8
    "DANH SÁCH CHỦ SỞ HỮU HƯỞNG LỢI CỦA DOANH NGHI",                           # Mẫu 10
    "B. Thành viên góp vốn (nếu có)",        # section marker Mẫu 9 lần 2
    "STTTên thành viênNgày, tháng, năm sinh đối với thành viên là cá nhân",     # header glue lưới đã mất
)
MAU_LABEL_RE = re.compile(r"Mẫu số\s*\d+$")

# Nội dung thật Mẫu 9 phải còn — chống cắt lố sang vùng hợp pháp.
KEEP_IN_SOURCE = (
    "PL1 - Mau so 9",
    "Mã số dự án",
    "8. Thông tin đăng ký thuế",
    "Các thành viên hợp danh cam kết",
    "CÁC THÀNH VIÊN HỢP DANH",
    "Trường hợp Tòa án hoặc Trọng tài chỉ định người thực hiện thủ tục",
)

# 14 cột vật lý của bảng thật (khớp tools/form_specs/formation/PLI_M09.json).
EXPECTED_COLS = [
    ("ho_ten", "Tên thành viên"),
    ("ngay_sinh", "ngày, tháng, năm sinh"),
    ("gioi_tinh", "giới tính"),
    ("so_cccd", "giấy tờ pháp lý"),
    ("quoc_tich", "quốc tịch"),
    ("dan_toc", "dân tộc"),
    ("dia_chi", "địa chỉ liên lạc"),
    ("von_dong_gop", "phần vốn góp"),
    ("ty_le_von", "tỷ lệ"),
    ("loai_tai_san_gop_von", "loại tài sản"),
    ("thoi_han_gop_von", "thời hạn góp vốn"),
    ("chu_ky", "chữ ký"),
    ("ghi_chu", "ghi chú"),
]


def _nfc(s: str) -> str:
    # NFC cho glyph tổ hợp + thống nhất non-breaking space (\xa0) của nguồn
    return unicodedata.normalize("NFC", s or "").replace("\u00a0", " ")


def _doc_texts(path: Path) -> list[str]:
    """Text mọi paragraph w:p của body, BAO GỂM cả ô bảng (iter đi xuống)."""
    doc = Document(str(path))
    return ["".join(t.text or "" for t in p.findall(".//" + qn("w:t")))
            for p in doc.element.body.iter(qn("w:p"))]


def _table_shape(path: Path) -> list[tuple[int, int]]:
    doc = Document(str(path))
    return [(len(t.rows), len(t.columns)) for t in doc.tables]


class M09Decontaminate(unittest.TestCase):
    maxDiff = None

    @classmethod
    def setUpClass(cls):
        assert SRC.exists(), "mất nguồn M09: %s" % SRC
        assert TPL.exists(), "mất template PLI_M09: %s" % TPL
        cls.src_texts = _doc_texts(SRC)
        cls.tpl_texts = _doc_texts(TPL)
        cls.src_joined = _nfc("\n".join(cls.src_texts))
        cls.tpl_joined = _nfc("\n".join(cls.tpl_texts))

    # ---------------------------------------------------------- [SRC] nguồn
    def test_src_no_contamination_fingerprints(self):
        """Nguồn đã bị cắt thật: 6 fingerprint vùng nhiễm phải vắng mặt."""
        for fp in CONTAMINATION_FINGERPRINTS:
            self.assertNotIn(_nfc(fp), self.src_joined,
                             "nguồn M09 còn nhiễm: %r — vùng nhiễm chưa bị xóa "
                             "khỏi SOURCE (cấm ngụy trang output)" % fp)

    def test_src_single_mau_label_and_no_residuals(self):
        """Nhãn 'Mẫu số N' đứng riêng chỉ còn đúng 1 ('Mẫu số 9' thật)."""
        labels = [t.strip() for t in self.src_texts
                  if MAU_LABEL_RE.fullmatch(_nfc(t.strip()) or "")]
        self.assertEqual(labels, ["Mẫu số 9"],
                         "nhãn 'Mẫu số N' lẻo: %r (Mẫu 9 lần 2/Mẫu 10 phải đã "
                         "bị xóa)" % labels)
        underscores = [t for t in self.src_texts
                       if re.fullmatch(r"_+", t.strip() or "x")]
        self.assertEqual(len(underscores), 1,
                         "số đường '_' trong nguồn = %d, kỳ vọng 1 (đường ký "
                         "CÁC THÀNH VIÊN HỢP DANH; 4 tàn dư phải đã bị xóa)"
                         % len(underscores))

    def test_src_keeps_real_mau9_content(self):
        """Cắt không lố: nội dung thật Mẫu 9 vẫn còn nguyên sau vùng tiêu đề."""
        for marker in KEEP_IN_SOURCE:
            self.assertIn(_nfc(marker), self.src_joined,
                          "mất nội dung thật Mẫu 9 do cắt lố: %r" % marker)
        # khối ký phải còn và vùng sau nó chỉ còn chân trang chú thích —
        # paragraph không rỗng CUỐI tài liệu phải là chân trang 9 của Mẫu 9
        # thật (toàn bộ vùng nhiễm nằm sau nó đã bị xóa khỏi nguồn)
        nonempty = [t for t in self.src_texts if t.strip()]
        self.assertIn(_nfc("CÁC THÀNH VIÊN HỢP DANH"), self.src_joined)
        last = _nfc(nonempty[-1])
        self.assertTrue(last.startswith("- Trường hợp Tòa án hoặc Trọng tài"),
                        "paragraph cuối nguồn = %r — phải là chân trang 9 của "
                        "Mẫu 9 thật (vùng nhiễm nằm sau nó đã bị xóa)" % last[:60])

    # ------------------------------------------------------ [TPL] template
    def test_tpl_no_contamination_fingerprints(self):
        """Template production cũng không chứa fingerprint vùng nhiễm."""
        for fp in CONTAMINATION_FINGERPRINTS:
            self.assertNotIn(_nfc(fp), self.tpl_joined,
                             "template PLI_M09 còn nhiễm: %r" % fp)

    def test_tpl_single_mau_label(self):
        labels = [t.strip() for t in self.tpl_texts
                  if MAU_LABEL_RE.fullmatch(_nfc(t.strip()) or "")]
        self.assertEqual(labels, ["Mẫu số 9"],
                         "template còn nhãn mẫu lạ: %r" % labels)

    def test_tpl_no_residual_from_removed_source_region(self):
        """Tàn dư vùng nhiễm đã xóa khỏi nguồn KHÔNG hồi sinh trong template.

        ĐỎ trên nhánh lane cho tới khi Lead merge commit xóa nguồn rồi regen
        template (python3 -X utf8 -m tools.build_form_maps): template trong
        repo vẫn là bản build từ nguồn NHIỄM nên còn 4 đường '_' tàn dư
        (nguồn cũ p112/p140/p166/p219) cạnh 1 đường ký hợp pháp.
        """
        underscores = [t for t in self.tpl_texts
                       if re.fullmatch(r"_+", t.strip() or "x")]
        self.assertEqual(len(underscores), 1,
                         "template còn %d đường '_' (kỳ vọng 1 đường ký hợp "
                         "pháp) — template chưa được regen từ nguồn đã xóa "
                         "vùng nhiễm (chờ Lead merge + regen)"
                         % len(underscores))

    # ------------------------------------------------------ [KEY] form_maps
    def test_form_maps_keeps_real_m09_keys(self):
        fm = json.loads(FORM_MAPS.read_text(encoding="utf-8"))
        self.assertIn("PLI_M09", fm, "form_maps mất PLI_M09")
        entry = fm["PLI_M09"]
        loops = [d["loop"] for d in entry.get("dynamic_tables", [])]
        self.assertIn(LOOP, loops, "form_maps mất dynamic table %r" % LOOP)
        dt = next(d for d in entry["dynamic_tables"] if d["loop"] == LOOP)
        got_tags = [c["tag"] for c in dt["columns"]]
        for tag, _label in EXPECTED_COLS:
            self.assertIn(tag, got_tags, "mất cột %r của bảng thành viên hợp danh" % tag)
        self.assertEqual(len(got_tags), 13,  # 14 cột vật lý, STT tự sinh {_stt}
                         "số tag bảng = %d, kỳ vọng 13 (+STT tự sinh = 14 cột)"
                         % len(got_tags))
        keys = {f["key"] for f in entry.get("fields", [])}
        for tag, _label in EXPECTED_COLS:
            if tag == "chu_ky":
                continue  # outputOnly, không đăng ký field nhập
            self.assertIn("%s.%s" % (LOOP, tag), keys,
                          "form_maps mất field %s.%s" % (LOOP, tag))

    def test_template_still_has_loop_and_14_col_table(self):
        joined = "\n".join(self.tpl_texts)
        self.assertIn("{#%s}" % LOOP, joined, "mất thẻ mở loop %r" % LOOP)
        self.assertIn("{/%s}" % LOOP, joined, "mất thẻ đóng loop %r" % LOOP)
        shapes = _table_shape(TPL)
        self.assertIn((6, 14), shapes,
                      "không còn bảng 14 cột vật lý trong template: %r" % shapes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
