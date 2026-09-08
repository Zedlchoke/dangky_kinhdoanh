# -*- coding: utf-8 -*-
"""Kiểm chứng P1.2 Option D trên PLI_M72: export CHỈ chứa biến thể được chọn.

Với từng loại hình v (dntn / tnhh_1tv / tnhh_2tv / ctcp / dhdn):
  - data: dn.ten_tieng_viet='CONG TY TEST', flag của v='X', 4 flag còn lại='',
    loops[loop_v]=3 dòng VSIC thật (dòng 2 la_nganh_chinh='X');
  - render bằng tools.docx_render.write_rendered ra file tạm rồi assert:
    1. heading '(TÊN <biến thể>)' của v CÓ mặt; heading 4 biến thể khác VẮNG;
    2. nội dung đặc trưng của biến thể khác VẮNG ('Chủ doanh nghiệp' chỉ ở
       DNTN, 'Thông tin về chủ sở hữu' chỉ ở TNHH 1TV, 'Danh sách thành viên:'
       chỉ ở TNHH 2TV+, 'Cổ đông sáng lập' chỉ ở CTCP, 'Danh sách thành viên
       hợp danh' chỉ ở DHDN);
    3. luật footnote: tnhh_1tv -> dòng 'Ghi tên...' thứ nhất có mặt;
       tnhh_2tv -> '2 Ghi tên...'; ctcp -> '3 Ghi tên...'; dntn/dhdn ->
       KHÔNG dòng ghi chú nào;
    4. body còn ĐÚNG số w:tbl theo biến thể (P2 member grids: TNHH_2TV=2,
       CTCP=3, DHDN=2 — [P2 pin sync 2026-09-03]); bảng ĐẦU là bảng ngành
       4 cột chứa 3 dòng dữ liệu STT 1/2/3 và đúng một X;
    5. LEFTOVER_RE (từ tools.docx_render) tìm 0 tag sót trong document.xml;
    6. file mở được bằng python-docx; CondTagError KHÔNG xảy ra.
Thêm case KHÔNG chọn flag nào -> chỉ còn 'Mẫu số 72' + tiêu đề chung,
0 w:tbl, 0 leftover, CondTagError KHÔNG xảy ra.

File test thuần Python, style @check giống các test khác.
Exit code 0 = PASS / 1 = FAIL.

Chạy:  python -X utf8 tools/test_m72_conditional_export.py
"""
from __future__ import annotations

import re
import sys
import zipfile
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document                    # noqa: E402
from docx.oxml.ns import qn                  # noqa: E402
from docx.table import Table                 # noqa: E402
from docx.text.paragraph import Paragraph    # noqa: E402

from tools.docx_render import CondTagError, LEFTOVER_RE, write_rendered  # noqa: E402

TEMPLATE_PATH = ROOT / "app" / "templates" / "PLI_M72.docx"
OUT_DIR = ROOT / "generated_ho_so" / "_test_m72_cond_export"

HDR_NGANH = ["STT", "Tên ngành", "Mã ngành",
             "Ngành, nghề kinh doanh chính"]
FOOTNOTE_TEXT = "Ghi tên của tất cả người đại diện"

# Thứ tự biến thể theo thứ tự tài liệu; anchor + nội dung đặc trưng riêng
# (chỉ tồn tại trong đúng một biến thể trên template).
VARIANTS = [
    {
        "label": "DNTN", "flag": "loai_hinh_dntn", "loop": "ds_nganh_dntn",
        "anchor": "(TÊN DOANH NGHIỆP TƯ NHÂN)",
        "unique": "Chủ doanh nghiệp",
        "footnote": None,  # DNTN không có dòng ghi chú chân trang
    },
    {
        "label": "TNHH_1TV", "flag": "loai_hinh_tnhh_1tv",
        "loop": "ds_nganh_tnhh_1tv",
        "anchor": "(TÊN CÔNG TY TNHH MỘT THÀNH VIÊN)",
        "unique": "Thông tin về chủ sở hữu",
        # dòng 'Ghi tên...' thứ nhất đã mất số ở nguồn (không có prefix)
        "footnote": FOOTNOTE_TEXT,
    },
    {
        "label": "TNHH_2TV", "flag": "loai_hinh_tnhh_2tv",
        "loop": "ds_nganh_tnhh_2tv",
        "anchor": "(TÊN CÔNG TY TNHH HAI THÀNH VIÊN)",
        "unique": "Danh sách thành viên:",
        "footnote": "2 " + FOOTNOTE_TEXT,
    },
    {
        "label": "CTCP", "flag": "loai_hinh_ctcp", "loop": "ds_nganh_ctcp",
        "anchor": "(TÊN CÔNG TY CỔ PHẦN)",
        "unique": "Cổ đông sáng lập",
        "footnote": "3 " + FOOTNOTE_TEXT,
    },
    {
        "label": "DHDN", "flag": "loai_hinh_dhdn", "loop": "ds_nganh_dhdn",
        "anchor": "(TÊN CÔNG TY HỢP DANH)",
        "unique": "Danh sách thành viên hợp danh",
        "footnote": None,  # DHDN không có dòng ghi chú chân trang
    },
]

LOOP_ROWS = [
    {"ten_nganh": "Trồng lúa", "ma_nganh": "0111", "la_nganh_chinh": ""},
    {"ten_nganh": "Lập trình máy vi tính", "ma_nganh": "6201",
     "la_nganh_chinh": "X"},
    {"ten_nganh": "Bán lẻ vải, len, sợi, chỉ khâu và hàng dệt khác "
                  "trong các cửa hàng chuyên doanh", "ma_nganh": "4751",
     "la_nganh_chinh": ""},
]

# [P2 pin sync 2026-09-03] P2 rebuild 4 lưới thành viên/cổ đông (nested-grid)
# trong template M72 — mỗi vùng biến thể giờ có bảng ngành + lưới thành viên:
# DNTN/TNHH_1TV=1 bảng, TNHH_2TV=2, CTCP=3, DHDN=2 khi biến thể được chọn.
EXPECTED_TABLES = {
    "loai_hinh_dntn": 1,
    "loai_hinh_tnhh_1tv": 1,
    "loai_hinh_tnhh_2tv": 2,
    "loai_hinh_ctcp": 3,
    "loai_hinh_dhdn": 2,
}

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run(*args):
            try:
                fn(*args)
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def tc_iter(tr):
    for child in tr.iterchildren():
        if child.tag == qn("w:tc"):
            yield child


def tr_cell_texts(tbl: Table, ri: int) -> list[str]:
    """Text từng ô của hàng ri — duyệt trực tiếp w:tc (mỗi ô đúng một lần)."""
    out = []
    for tc in tc_iter(tbl._tbl.tr_lst[ri]):
        txt = "\n".join(
            Paragraph(p_el, tbl.part.document).text
            for p_el in tc.iterchildren()
            if p_el.tag == qn("w:p")
        )
        out.append(txt.strip())
    return out


def body_tables(doc) -> list[Table]:
    return [Table(child, doc) for child in doc.element.body.iterchildren()
            if child.tag == qn("w:tbl")]


def doc_all_text(path: Path) -> str:
    """Toàn bộ text paragraph + ô bảng của file render."""
    doc = Document(str(path))  # python-docx phải mở được
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def render_payload(variant: dict | None) -> dict:
    """Payload export: đúng 1 flag truthy (hoặc không flag nào), loop tương ứng."""
    data: dict = {"dn": {"ten_tieng_viet": "CONG TY TEST"}}
    for v in VARIANTS:
        data[v["flag"]] = "X" if (variant is not None
                                  and v is variant) else ""
        data[v["loop"]] = (list(LOOP_ROWS)
                           if variant is not None and v is variant else [])
    return data


def safe_render(data: dict, out_path: Path) -> None:
    """Render bắt buộc KHÔNG được raise CondTagError (template sai cặp marker)."""
    try:
        write_rendered(TEMPLATE_PATH, data, out_path)
    except CondTagError as exc:
        raise AssertionError(
            "CondTagError KHÔNG được xảy ra với payload hợp lệ: %r" % (exc,))


def leftover_tokens(out_path: Path) -> list[str]:
    with zipfile.ZipFile(out_path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    return LEFTOVER_RE.findall(xml)


def run_variant_case(variant: dict, out_path: Path) -> None:
    safe_render(render_payload(variant), out_path)

    all_text = doc_all_text(out_path)  # cũng chứng minh python-docx mở được
    label = variant["label"]

    # 1. heading của biến thể được chọn CÓ mặt, 4 biến thể khác VẮNG
    assert variant["anchor"] in all_text, \
        "%s: thiếu heading %r trong output" % (label, variant["anchor"])
    for other in VARIANTS:
        if other is variant:
            continue
        assert other["anchor"] not in all_text, \
            "%s: heading biến thể khác bị lọt: %r" % (label, other["anchor"])
        # 2. nội dung đặc trưng của biến thể khác VẮNG
        assert other["unique"] not in all_text, \
            ("%s: nội dung đặc trưng của %s bị lọt: %r"
             % (label, other["label"], other["unique"]))
    assert variant["unique"] in all_text, \
        "%s: mất nội dung đặc trưng riêng %r" % (label, variant["unique"])

    # 3. luật footnote chân trang: đúng 0/1 dòng ghi chú, đúng số mũ 1/2/3
    note_lines = [ln.strip() for ln in all_text.splitlines()
                  if FOOTNOTE_TEXT in ln]
    if variant["footnote"] is not None:
        assert variant["footnote"] in all_text, \
            "%s: thiếu dòng ghi chú %r" % (label, variant["footnote"])
        assert len(note_lines) == 1, \
            "%s: kỳ vọng đúng 1 dòng ghi chú, thấy %d: %r" % (
                label, len(note_lines), note_lines)
        got_num = re.match(r"^(\d*)", note_lines[0]).group(1)
        want_num = {"loai_hinh_tnhh_1tv": "", "loai_hinh_tnhh_2tv": "2",
                    "loai_hinh_ctcp": "3"}[variant["flag"]]
        assert got_num == want_num, \
            ("%s: số mũ dòng ghi chú = %r != %r"
             % (label, got_num, want_num))
    else:
        assert not note_lines, \
            "%s: KHÔNG được có dòng ghi chú chân trang nào, thấy %r" % (
                label, note_lines)

    # 4. body còn đúng số w:tbl theo biến thể [P2 pin sync 2026-09-03]:
    #    bảng ĐẦU là bảng ngành (header 4 cột + 3 dòng dữ liệu STT 1/2/3, 1 X),
    #    các bảng sau là lưới thành viên/cổ đông P2 (rỗng vì payload không có)
    doc = Document(str(out_path))
    tables = body_tables(doc)
    want_tbls = EXPECTED_TABLES[variant["flag"]]
    assert len(tables) == want_tbls, \
        "%s: kỳ vọng %d w:tbl, thấy %d" % (label, want_tbls, len(tables))
    tbl = tables[0]
    texts = [tr_cell_texts(tbl, ri) for ri in range(len(tbl.rows))]
    assert fold_join(texts[0]) == fold_join(HDR_NGANH), \
        "%s: header bảng %r != %r" % (label, texts[0], HDR_NGANH)
    # [Pin sync — owner N-only 2026-09-08] bảng ngành chỉ gồm header + N dòng
    # dữ liệu (STT 1..N, đứng ngay sau header) — blankRows không còn render,
    # không còn K−N hàng trắng đuôi (bỏ công thức max(K,N))
    n_data = len(LOOP_ROWS)
    assert len(texts) == 1 + n_data, \
        ("%s: tổng hàng bảng ngành %d != 1 + N=%d (không hàng trắng)"
         % (label, len(texts), n_data))
    data_texts = texts[1:1 + n_data]
    for i, row in enumerate(data_texts):
        want = LOOP_ROWS[i]
        assert len(row) == 4, \
            "%s: hàng dữ liệu %d có %d ô != 4" % (label, i, len(row))
        assert row[0] == str(i + 1), \
            "%s: STT dòng %d = %r, kỳ vọng %r" % (label, i, row[0], i + 1)
        assert row[1] == want["ten_nganh"], \
            "%s: Tên ngành dòng %d = %r != %r" % (
                label, i, row[1], want["ten_nganh"])
        assert row[2] == want["ma_nganh"], \
            "%s: Mã ngành dòng %d = %r != %r" % (
                label, i, row[2], want["ma_nganh"])
        assert row[3] == want["la_nganh_chinh"], \
            "%s: Ngành chính dòng %d = %r != %r" % (
                label, i, row[3], want["la_nganh_chinh"])
    n_x = sum(1 for row in texts for cell in row if cell == "X")
    assert n_x == 1, \
        "%s: kỳ vọng đúng một dấu X trong bảng, thấy %d" % (label, n_x)

    # 5. không còn tag nào sót trong document.xml
    leftovers = leftover_tokens(out_path)
    assert leftovers == [], \
        "%s: còn %d tag sót: %s" % (label, len(leftovers),
                                    sorted(set(leftovers))[:10])


def fold_join(cells: list[str]) -> str:
    def fold(text: str) -> str:
        nfkd = unicodedata.normalize("NFD", text.lower())
        out = "".join(c for c in nfkd if unicodedata.category(c) != "Mn")
        return re.sub(r"\s+", " ", out.replace("đ", "d")).strip()

    return "|".join(fold(c) for c in cells)


def make_variant_check(variant: dict):
    @check("export %s: CHI xuat bien te duoc chon (heading/noi dung/bang/"
           "footnote/leftover)" % variant["label"])
    def _t():
        out_path = OUT_DIR / ("PLI_M72_%s.docx" % variant["label"].lower())
        run_variant_case(variant, out_path)
    return _t


@check("export khong chon loai hinh nao: chi con 'Mau so 72' + tieu de chung, "
       "0 w:tbl, 0 leftover, khong CondTagError")
def t_none_selected():
    out_path = OUT_DIR / "PLI_M72_khong_chon.docx"
    safe_render(render_payload(None), out_path)

    all_text = doc_all_text(out_path)
    assert "Mẫu số 72" in all_text, "mất dòng 'Mẫu số 72'"
    assert "CÔNG BỐ NỘI DUNG ĐĂNG KÝ THÀNH LẬP DOANH NGHIỆP" in all_text, \
        "mất tiêu đề chung"
    for v in VARIANTS:
        assert v["anchor"] not in all_text, \
            "không chọn gì nhưng vẫn lọt heading %r" % v["anchor"]
        assert v["unique"] not in all_text, \
            "không chọn gì nhưng vẫn lọt nội dung %r" % v["unique"]
    assert FOOTNOTE_TEXT not in all_text, "lọt dòng ghi chú chân trang"

    doc = Document(str(out_path))
    tables = body_tables(doc)
    assert len(tables) == 0, \
        "kỳ vọng 0 w:tbl khi không chọn loại hình, thấy %d" % len(tables)

    leftovers = leftover_tokens(out_path)
    assert leftovers == [], \
        "còn %d tag sót: %s" % (len(leftovers), sorted(set(leftovers))[:10])


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    tests = [make_variant_check(v) for v in VARIANTS] + [t_none_selected]
    for t in tests:
        t()
    passed = sum(1 for ok, _ in _results if ok)
    failed = len(_results) - passed
    print("-" * 60)
    print("test_m72_conditional_export: %d PASS / %d FAIL" % (passed, failed))
    for ok, name in _results:
        if not ok:
            print(" - FAIL:", name)
    print("Thư mục kết quả:", OUT_DIR)
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
