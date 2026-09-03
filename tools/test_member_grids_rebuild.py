# -*- coding: utf-8 -*-
"""Lane fix-member-grids (Owner P2=A 2026-09-03, D2.9=A, D7.4=A) — kiểm tra
rebuild 2 lưới thành viên PLII_M02 + 4 lưới cổ đông/thành viên PLI_M72.

Phạm vi (đúng 3 tiêu chí gate của lane + contract sâu hơn):
  (a) PLII_M02 có ĐÚNG 2 w:tbl lưới thành viên (a4 `tv_hgd_sau_thay_doi`,
      c7 `tv_hgd_cap_doi`) — 10 cột vật lý đếm từ digit-run + chân trang mẫu
      ('cột số 1..10'), header dán chữ `STTHọ tên…Chữ ký` biến mất khỏi
      template/blocks, câu cam kết pháp lý 'Hộ kinh doanh cam kết…' được GIỮ
      nguyên văn (trim keep_from);
  (b) digit-run `12345678910` hết trên template + blocks của cả 2 mẫu;
  (c) PLI_M72: 4 lưới N2–N5 (ds_thanh_vien_tnhh_2tv / ds_co_dong_sang_lap_ctcp
      / ds_co_dong_nuoc_ngoai_ctcp / ds_thanh_vien_hop_danh_dhdn) mỗi lưới mang
      ĐÚNG cond biến thể của nó (P1.2 Option D — marker template-level),
      heading 8/9 bị dán chữ trong nguồn được tách lại đúng vị trí, render
      từng biến thể KHÔNG trộn dữ liệu giữa các lưới.

Sandbox: render vào tempfile BÊN NGOÀI repo (không ghi app/**, không regen).

Chạy:  python3 -X utf8 -m tools.test_member_grids_rebuild   # exit 0 = PASS
"""
from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import unittest
import zipfile

TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(TOOLS_DIR)
sys.path.insert(0, TOOLS_DIR)

from docx import Document                    # noqa: E402
from docx.oxml.ns import qn                 # noqa: E402
from docx.table import Table                # noqa: E402

from tools.build_form_maps import load_form_spec   # noqa: E402
from tools.docx_render import LEFTOVER_RE, write_rendered  # noqa: E402

BLOCKS_DIR = os.path.join(ROOT, "app", "data", "blocks")
FORM_MAPS_PATH = os.path.join(ROOT, "app", "data", "form_maps.json")
TEMPLATES_DIR = os.path.join(ROOT, "app", "templates")

# 10 cột của lưới thành viên hộ gia đình M02 — đếm từ digit-run '12345678910'
# của header nguồn và chân trang chính thức của mẫu (P179: 'chỉ kê khai các
# cột số 1, 2, 3, 4, 5, 10, không phải kê khai các cột số 6, 7, 8, 9').
M02_TV_TAGS = ["ho_ten", "ngay_sinh", "so_cccd", "gioi_tinh", "quoc_tich",
               "dan_toc", "noi_thuong_tru", "noi_o_hien_tai", "chu_ky"]
M02_TV_HEADERS = ["STT", "Họ tên", "Ngày, tháng, năm sinh",
                  "Số định danh cá nhân", "Giới tính", "Quốc tịch", "Dân tộc",
                  "Nơi thường trú", "Nơi ở hiện tại", "Chữ ký"]
M02_OCC = {
    "tv_hgd_sau_thay_doi": "Danh sách các thành viên hộ gia đình đăng ký hộ "
                           "kinh doanh sau khi thay đổi",
    "tv_hgd_cap_doi": "7. Thông tin về các thành viên hộ gia đình đăng ký hộ "
                      "kinh doanh",
}
# 4 lưới N2–N5 M72: (occurrence_id, cond biến thể, tags không STT)
M72_GRIDS = [
    ("ds_thanh_vien_tnhh_2tv", "loai_hinh_tnhh_2tv",
     ["ten_thanh_vien", "quoc_tich", "so_gtpl_ca_nhan_to_chuc",
      "dia_chi_lien_lac", "phan_von_gop", "ty_le", "ghi_chu"]),
    ("ds_co_dong_sang_lap_ctcp", "loai_hinh_ctcp",
     ["ten_co_dong", "quoc_tich", "so_gtpl", "dia_chi_lien_lac",
      "loai_co_phan", "so_co_phan", "gia_tri_co_phan", "ty_le", "ghi_chu"]),
    ("ds_co_dong_nuoc_ngoai_ctcp", "loai_hinh_ctcp",
     ["ten_co_dong", "quoc_tich", "so_gtpl", "dia_chi_lien_lac",
      "loai_co_phan", "so_co_phan", "gia_tri_co_phan", "ty_le", "ghi_chu"]),
    ("ds_thanh_vien_hop_danh_dhdn", "loai_hinh_dhdn",
     ["ten_thanh_vien_hop_danh", "quoc_tich", "so_dinh_danh_ho_chieu",
      "dia_chi_lien_lac", "phan_von_gop", "ty_le"]),
]
M72_FLAGS = ["loai_hinh_dntn", "loai_hinh_tnhh_1tv", "loai_hinh_tnhh_2tv",
             "loai_hinh_ctcp", "loai_hinh_dhdn"]
GLUE_RUN_RE = re.compile(r"(?i)^stt(?=[^\W\d_])")
DIGIT_RUN = "12345678910"


def fold(text: str) -> str:
    import unicodedata
    nfkd = unicodedata.normalize("NFD", text.lower())
    out = "".join(c for c in nfkd if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", out.replace("đ", "d")).strip()


def body_paragraph_texts(fid: str) -> list[str]:
    """Text các w:p CON TRỰC TIẾP của body trong template làm việc."""
    doc = Document(os.path.join(TEMPLATES_DIR, "%s.docx" % fid))
    return ["".join(t.text or "" for t in child.iter(qn("w:t")))
            for child in doc.element.body.iterchildren()
            if child.tag == qn("w:p")]


def body_tables(doc) -> list[Table]:
    return [Table(child, doc) for child in doc.element.body.iterchildren()
            if child.tag == qn("w:tbl")]


def tr_cell_texts(tbl: Table, ri: int) -> list[str]:
    """Text từng ô của hàng ri — duyệt trực tiếp w:tc (mỗi ô đúng một lần)."""
    out = []
    for tc in (c for c in tbl._tbl.tr_lst[ri].iterchildren()
               if c.tag == qn("w:tc")):
        out.append("".join(t.text or "" for t in tc.iter(qn("w:t"))).strip())
    return out


def blocks(fid: str) -> dict:
    with open(os.path.join(BLOCKS_DIR, "%s.json" % fid), encoding="utf-8") as f:
        return json.load(f)


def form_maps() -> dict:
    with open(FORM_MAPS_PATH, encoding="utf-8") as f:
        return json.load(f)


class BaseSandbox(unittest.TestCase):
    """Render sandbox — tempfile ngoài repo, không đụng app/**."""

    @classmethod
    def setUpClass(cls):
        cls.sandbox = tempfile.mkdtemp(prefix="member_grids_sandbox_")


class TestSpecContract(BaseSandbox):
    """Spec 2 form khai báo đủ entry + occurrence hợp lệ qua loader."""

    def test_m02_co_2_entry_thanh_vien(self):
        ext = load_form_spec("PLII_M02")
        ids = sorted(ext["specs_by_id"])
        for sid in ("tv_hgd_sau_thay_doi", "tv_hgd_cap_doi"):
            self.assertIn(sid, ids)
            spec = ext["specs_by_id"][sid]
            self.assertEqual(spec["kind"], "nested-grid")
            self.assertEqual(spec["section_id"],
                             ("a4_tv_hgd_sau_thay_doi" if sid ==
                              "tv_hgd_sau_thay_doi" else "c7_tv_hgd_cap_doi"))
            # 10 cột: STT generated col 0 + 9 the col 1..9
            self.assertEqual(spec["stt"], {"mode": "generated",
                                           "col_index": 0})
            self.assertEqual(spec["width"], 10)
            self.assertEqual([c["tag"] for c in spec["columns"]], M02_TV_TAGS)
            self.assertEqual([c["col_index"] for c in spec["columns"]],
                             list(range(1, 10)))
            # chu_ky outputOnly (D2.4=A)
            chu = spec["columns"][-1]
            self.assertTrue(chu["output_only"])

    def test_m02_sync_group_null(self):
        """PLAN A: 2 lưới ngữ cảnh khác nhau (a4 vs c7) KHÔNG auto-sync."""
        ext = load_form_spec("PLII_M02")
        for ent in ext["entries"]:
            for occ in ent.get("_occurrences") or []:
                if occ["occurrence_id"] in M02_OCC:
                    self.assertIsNone(occ["sync_group"])

    def test_m72_co_4_entry_n2_n5(self):
        ext = load_form_spec("PLI_M72")
        ids = sorted(ext["specs_by_id"])
        want = sorted(oid for oid, _c, _t in M72_GRIDS)
        self.assertEqual([i for i in ids if i in set(want)], want)
        for ent in ext["entries"]:
            oid_list = [o["occurrence_id"]
                        for o in ent.get("_occurrences") or []]
            for oid in oid_list:
                if oid in dict((o, c) for o, c, _t in M72_GRIDS):
                    spec = ext["specs_by_id"][oid]
                    self.assertEqual(spec["kind"], "nested-grid")
                    tags = [c["tag"] for c in spec["columns"]]
                    self.assertEqual(tags, dict(
                        (o, t) for o, _c, t in M72_GRIDS)[oid])

    def test_m72_khong_cond_flag_spec_level(self):
        """Marker biến thể là template-level (P1.2 Option D, tagger là nguồn
        duy nhất); cond_flag spec-level cùng flag làm existing_markers() của
        tools/tag_m72_sections.py lệch sequence -> tagger từ chối chạy."""
        for fid in ("PLI_M72", "PLII_M02"):
            raw_path = os.path.join(ROOT, "tools", "form_specs",
                                    "household" if fid == "PLII_M02"
                                    else "change", "%s.json" % fid)
            with open(raw_path, encoding="utf-8-sig") as f:
                raw = json.load(f)
            for e in raw["tables"]:
                self.assertNotIn(
                    "cond_flag", e,
                    "%s/%s: không được bọc cond_flag trùng marker variant"
                    % (fid, e.get("spec", {}).get("id")))

    def test_fields_dang_ky_loop_tag(self):
        fm = form_maps()
        cases = ([(fid, oid, None) for oid, fid in
                  [(o, "PLII_M02") for o in M02_OCC]] +
                 [("PLI_M72", oid, tags) for oid, _c, tags in M72_GRIDS])
        for fid, oid, explicit_tags in cases:
            keys = {f["key"] for f in fm[fid]["fields"]}
            tags = explicit_tags or (M02_TV_TAGS if fid == "PLII_M02" else [])
            for tag in tags:
                if tag == "chu_ky":  # outputOnly: không đăng ký input
                    continue
                self.assertIn("%s.%s" % (oid, tag), keys,
                              "thiếu field %s.%s" % (oid, tag))


class TestM02MemberGrids(BaseSandbox):
    """Tiêu chí (a): 2 w:tbl thành viên M02, glue hết, cam kết giữ nguyên."""

    def test_blocks_co_dung_2_loop_10_cot(self):
        data = blocks("PLII_M02")
        loops = {b["loop"]: b for b in data["blocks"] if b.get("t") == "loop"}
        for oid in M02_OCC:
            self.assertIn(oid, loops)
            blk = loops[oid]
            self.assertEqual(blk.get("headerRows"), [0])
            cols = [(c["col"], c["tag"]) for c in blk["cols"]]
            self.assertEqual(cols, list(enumerate(M02_TV_TAGS, start=1)))
            hdr = ["".join(r.get("text", "") for r in cell["runs"]).strip()
                   for cell in blk["rows"][0]["cells"]]
            self.assertEqual(hdr, M02_TV_HEADERS, oid)
            self.assertGreaterEqual(blk.get("blankRows", 0), 1)

    def test_glue_va_digit_run_mat_khoi_blocks(self):
        data = blocks("PLII_M02")
        joined = json.dumps(data, ensure_ascii=False)
        self.assertNotIn(DIGIT_RUN, joined)
        for b in data["blocks"]:
            if b.get("t") != "p":
                continue
            txt = "".join(r.get("text", "") for r in b.get("runs", []))
            self.assertFalse(GLUE_RUN_RE.match(txt.strip()),
                             "còn paragraph STT-glue: %r" % txt[:60])

    def test_glue_va_digit_run_mat_khoi_template(self):
        texts = body_paragraph_texts("PLII_M02")
        self.assertTrue(
            all(DIGIT_RUN not in t for t in texts),
            "template còn digit-run %s" % [t[:60] for t in texts
                                           if DIGIT_RUN in t])
        self.assertFalse(
            [t for t in texts if GLUE_RUN_RE.match(t.strip())],
            "template còn paragraph STT-glue")

    def test_cau_cam_ket_duoc_giu_nguyen_van(self):
        """Trim keep_from của tv_hgd_cap_doi phải giữ câu cam kết pháp lý —
        chống 'xóa sạch paragraph dán chữ' làm mất văn bản của mẫu."""
        texts = body_paragraph_texts("PLII_M02")
        self.assertTrue(
            any("Hộ kinh doanh cam kết về tính hợp pháp, chính xác, trung "
                "thực" in t for t in texts),
            "mất câu cam kết 'Hộ kinh doanh cam kết…' sau rebuild")

    def test_template_co_dung_2_w_tbl_10_cot(self):
        doc = Document(os.path.join(TEMPLATES_DIR, "PLII_M02.docx"))
        want = fold(" ".join(M02_TV_HEADERS))
        member_tbls = [t for t in body_tables(doc)
                       if t.rows and fold(" ".join(tr_cell_texts(t, 0)))
                       == want]
        self.assertEqual(len(member_tbls), 2,
                         "M02 phải có đúng 2 w:tbl thành viên, thấy %d"
                         % len(member_tbls))
        for tbl in member_tbls:
            grid = tbl._tbl.find(qn("w:tblGrid"))
            ncols = len([e for e in grid.iterchildren()
                         if e.tag == qn("w:gridCol")])
            self.assertEqual(ncols, 10)
            # hàng mẫu: {#loop}…{/loop} + {_stt}
            tpl = " ".join(tr_cell_texts(tbl, len(tbl.rows) - 1))
            self.assertIn("{_stt}", tpl)
            self.assertTrue(re.search(r"\{#\w+\}", tpl) and
                            re.search(r"\{/\w+\}", tpl))

    def test_render_m02_du_2_bang_nhan_du_lieu_rieng(self):
        row = {"ho_ten": "NGUYEN VAN A", "ngay_sinh": "01/01/1980",
               "so_cccd": "012345678901", "gioi_tinh": "Nam",
               "quoc_tich": "Việt Nam", "dan_toc": "Kinh",
               "noi_thuong_tru": "Hà Nội", "noi_o_hien_tai": "TP.HCM"}
        out = os.path.join(self.sandbox, "m02.docx")
        write_rendered(
            os.path.join(TEMPLATES_DIR, "PLII_M02.docx"),
            {"tv_hgd_sau_thay_doi": [row],
             "tv_hgd_cap_doi": [dict(row, ho_ten="TRAN THI B")]},
            out)
        with zipfile.ZipFile(out) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        self.assertEqual(LEFTOVER_RE.findall(xml), [])
        self.assertIn("NGUYEN VAN A", xml)
        self.assertIn("TRAN THI B", xml)
        doc = Document(out)
        want = fold(" ".join(M02_TV_HEADERS))
        member_tbls = [t for t in body_tables(doc)
                       if t.rows and fold(" ".join(tr_cell_texts(t, 0)))
                       == want]
        self.assertEqual(len(member_tbls), 2)
        # dữ liệu KHÔNG trộn giữa 2 bảng (hàng dữ liệu nằm sau các hàng trắng)
        body0 = " ".join(c for ri in range(1, len(member_tbls[0].rows))
                         for c in tr_cell_texts(member_tbls[0], ri))
        body1 = " ".join(c for ri in range(1, len(member_tbls[1].rows))
                         for c in tr_cell_texts(member_tbls[1], ri))
        self.assertIn("NGUYEN VAN A", body0)
        self.assertNotIn("TRAN THI B", body0)
        self.assertIn("TRAN THI B", body1)
        self.assertNotIn("NGUYEN VAN A", body1)
        # STT tự sinh 1-based trên hàng có dữ liệu
        for tbl in member_tbls:
            stts = [tr_cell_texts(tbl, ri)[0]
                    for ri in range(1, len(tbl.rows))
                    if any(c.strip() for c in tr_cell_texts(tbl, ri))]
            self.assertEqual(stts, [str(k + 1) for k in range(len(stts))])


class TestM72MemberGrids(BaseSandbox):
    """Tiêu chí (c): 4 lưới N2–N5 đúng biến thể, không trộn."""

    def test_blocks_co_4_loop_dung_cond(self):
        data = blocks("PLI_M72")
        loops = {b["loop"]: b for b in data["blocks"] if b.get("t") == "loop"}
        for oid, cond, tags in M72_GRIDS:
            self.assertIn(oid, loops)
            blk = loops[oid]
            self.assertEqual(blk.get("cond"), [cond],
                             "%s: cond %r != [%r]" % (oid, blk.get("cond"),
                                                      cond))
            cols = [(c["col"], c["tag"]) for c in blk["cols"]]
            self.assertEqual(cols, list(enumerate(tags, start=1)))
            self.assertTrue(all(c["col"] != 0 for c in blk["cols"]),
                            "%s: cột 0 STT không được gắn the" % oid)

    def test_glue_va_digit_run_mat_khoi_m72(self):
        joined = json.dumps(blocks("PLI_M72"), ensure_ascii=False)
        self.assertNotIn(DIGIT_RUN, joined)
        self.assertNotIn("STTTên cổ đông", joined)
        texts = body_paragraph_texts("PLI_M72")
        self.assertTrue(all(DIGIT_RUN not in t for t in texts))
        self.assertFalse([t for t in texts if GLUE_RUN_RE.match(t.strip())])

    def test_heading_8_ctcp_duoc_tach_va_heading_9_giu_superscript(self):
        """P146 nguồn dán '(VNĐ…)…8. Cổ đông…' — trim tách heading 8 về
        paragraph riêng, đúng 1 lần. P148 (heading 9 kèm SUPERSCRIPT '3' của
        chân trang '3 Ghi tên…') phải được GIỮ NGUYÊN run structure: trim làm
        phẳng run sẽ tạo GLUE MỚI 'y3:' (rule (d) check_layout_l1) — đây là
        pin chống hồi quy."""
        texts = [t.strip() for t in body_paragraph_texts("PLI_M72") if t.strip()]
        h8 = [i for i, t in enumerate(texts)
              if fold(t).startswith("8. co dong la nha dau tu nuoc ngoai")]
        self.assertEqual(len(h8), 1, "heading 8 xuất hiện %d lần" % len(h8))
        # heading 9 còn tồn tại (trong paragraph dán của nguồn) và số '3'
        # vẫn nằm trong run superscript
        doc = Document(os.path.join(TEMPLATES_DIR, "PLI_M72.docx"))
        h9_paras = []
        for child in doc.element.body.iterchildren():
            if child.tag != qn("w:p"):
                continue
            txt = "".join(t.text or "" for t in child.iter(qn("w:t")))
            if "9. Người đại diện theo pháp luật của công ty" in txt:
                h9_paras.append(child)
        self.assertEqual(len(h9_paras), 1,
                         "heading 9 xuất hiện %d lần" % len(h9_paras))
        sup_digits = []
        for r in h9_paras[0].iter(qn("w:r")):
            txt = "".join(t.text or "" for t in r.iter(qn("w:t")))
            rpr = r.find(qn("w:rPr"))
            va = rpr is not None and rpr.find(qn("w:vertAlign")) is not None
            if va and txt.strip():
                sup_digits.append(txt.strip())
        self.assertIn("3", sup_digits,
                      "chú thích '3' của heading 9 mất superscript — trim đã "
                      "làm phẳng run (regression rule (d))")

    def test_junk_row2_ctcp_con_dung_1_dong_da_ghi_gap(self):
        """Gap đã ghi (BLOCKER_OWNERSHIP build_form_maps: trim không giữ
        vertAlign): dòng dán '(VNĐ…)…9. Người đại diện…' của P148 còn đúng 1
        — KHÔNG được tăng thêm (P146 đã trim sạch)."""
        texts = [t.strip() for t in body_paragraph_texts("PLI_M72") if t.strip()]
        junk = [t for t in texts if t.startswith("(VNĐ và giá trị")]
        self.assertEqual(len(junk), 1,
                         "còn %d dòng dán '(VNĐ…)…' (kỳ vọng đúng 1 dòng "
                         "heading 9 đã ghi gap): %r" % (len(junk), junk))

    def test_header_paragraph_tan_du_da_xoa(self):
        """9+8 header paragraph rời của N2/N5 và 2 row-glue của N3/N4 phải
        biến mất khỏi template (không còn 'Tên thành viên hợp danh' đứng
        trần ngoài bảng)."""
        texts = [t.strip() for t in body_paragraph_texts("PLI_M72") if t.strip()]
        for leak in ("Tên thành viên hợp danh", "Tên cổ đông",
                     "Phần vốn góp", "Số định danh cá nhân/Số Hộ chiếu"):
            stray = [t for t in texts if t == leak]
            self.assertFalse(stray, "còn header rời %r ngoài bảng" % leak)

    def test_template_co_4_w_tbl_n2_n5(self):
        doc = Document(os.path.join(TEMPLATES_DIR, "PLI_M72.docx"))
        count = 0
        for tbl in body_tables(doc):
            hdr0 = fold(" ".join(tr_cell_texts(tbl, 0)))
            if hdr0.startswith("stt ten thanh vien") or \
                    hdr0.startswith("stt ten co dong"):
                count += 1
        self.assertEqual(count, 4, "M72 phải có đúng 4 w:tbl N2–N5, thấy %d"
                         % count)

    def test_render_tung_variant_khong_tron_du_lieu(self):
        payload_by_cond = {
            "loai_hinh_tnhh_2tv": ("ds_thanh_vien_tnhh_2tv",
                                   {"ten_thanh_vien": "LE VAN C",
                                    "quoc_tich": "Việt Nam",
                                    "so_gtpl_ca_nhan_to_chuc": "0123456789",
                                    "dia_chi_lien_lac": "Hà Nội",
                                    "phan_von_gop": "500000000",
                                    "ty_le": "40", "ghi_chu": ""}),
            "loai_hinh_ctcp": ("ds_co_dong_sang_lap_ctcp",
                               {"ten_co_dong": "NGUYEN VAN B",
                                "quoc_tich": "Việt Nam", "so_gtpl": "0123",
                                "dia_chi_lien_lac": "Hà Nội",
                                "loai_co_phan": "Phổ thông",
                                "so_co_phan": "1000",
                                "gia_tri_co_phan": "100000000",
                                "ty_le": "50", "ghi_chu": ""}),
            "loai_hinh_dhdn": ("ds_thanh_vien_hop_danh_dhdn",
                               {"ten_thanh_vien_hop_danh": "PHAM VAN E",
                                "quoc_tich": "Việt Nam",
                                "so_dinh_danh_ho_chieu": "012345678",
                                "dia_chi_lien_lac": "Hải Phòng",
                                "phan_von_gop": "300", "ty_le": "10"}),
        }
        for cond, (loop, row) in payload_by_cond.items():
            data = {f: ("X" if f == cond else "") for f in M72_FLAGS}
            data[loop] = [row]
            out = os.path.join(self.sandbox, "m72_%s.docx" % cond)
            write_rendered(os.path.join(TEMPLATES_DIR, "PLI_M72.docx"),
                           data, out)
            with zipfile.ZipFile(out) as z:
                xml = z.read("word/document.xml").decode("utf-8")
            self.assertEqual(LEFTOVER_RE.findall(xml), [])
            self.assertIn(list(row.values())[0], xml,
                          "%s: dữ liệu lưới không render" % cond)
            # dữ liệu KHÔNG lọt sang biến thể khác
            for other_loop, other_row in \
                    ((l, r) for l, r in payload_by_cond.values()
                     if l != loop):
                self.assertNotIn(list(other_row.values())[0], xml,
                                 "%s: NHIỄM dữ liệu của %s" % (cond,
                                                               other_loop))


if __name__ == "__main__":
    unittest.main()
