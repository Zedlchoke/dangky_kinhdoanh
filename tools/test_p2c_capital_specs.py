#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test_p2c_capital_specs.py — Lane P2C (capital/assets: vốn + tài sản góp vốn).

Mệnh đề kiểm chứng cho đợt spec fixed-matrix vốn/tài sản (docs/
decision_registry.md packet D1: D1.1=A, D1.2=B CHO M01, D1.4=A), theo
docs/p2_component_contracts.md mục 2 + census manifest:

  [EVD]      census manifest có evidence cho 3 section làm spec
             (PLI_M01 tai_san_gop_von_von_dau_tu, PLI_M02 nguon_von_dieule +
             tai_san_gop_von_von_dieu_le): kind/status/semantic_key + evidence
             heading/anchor trỏ đúng vùng template; row id + tag lấy ĐÚNG
             manifest (không tự chế key).
  [SPEC]     2 file spec (formation/PLI_M01.json, core/PLI_M02.json) qua
             validator v1 của tools/build_form_maps.py; D1.2=B cho M01:
             "Các tài sản khác" là Ô MÔ TẢ TEXT (cột chi tiết), KHÔNG sub-grid;
             D1.1=A: dòng tổng output_only + totals sum() đúng target.
  [ANCHOR]   anchor khớp ĐÚNG 1 paragraph đã fold trong blocks HIỆN TẠI và
             trong template, trúng đúng heading evidence của manifest.
  [DROP]     drop chỉ áp trong khoảng [anchor, drop_until) — không nuốt
             paragraph của mục kế (mục 6 Thuế / mục 8 / mục 9 NĐDPL).
  [SANDBOX]  rebuild THẬT trên sandbox tmp (OUT_TPL/OUT_DATA ghi vào temp,
             KHÔNG đụng app/templates / app/data): dựng w:tbl đủ dòng pháp lý,
             convert về blocks kèm annotate_and_verify_section_specs, glue
             biến mất, chuỗi pháp lý dính kèm GIỮ, idempotent 2 lần chạy.
  [EXPORT]   D1.1=A: recompute_totals()/render() tính lại tổng từ chính các ô
             dữ liệu, ghi đè stored; stored != computed -> raise
             EXPORT_TOTALS_RECOMPUTE_CONFLICT (blocking).
  [BLOCKED]  những section CHƯA có quyết định/nguồn thì PHẢI KHÔNG có spec:
             2 fixed-matrix already_complete của PLI_M01 (W1 gap guard),
             mo_hinh_to_chuc_cong_ty PLI_M02 (chưa packet), D1.2 policy ngoài
             M01 (không sub-grid/tài sản khác cho M02), M03/M04/M05/M12.
  [GAP]      những gì contract v1 KHÔNG khai được phải được GHI NHẬN là gap
             (không im lặng): cột STT in sẵn, % tự tính theo giá trị,
             cross_field đối chiếu vốn điều lệ blocking.

Chạy: python3 -X utf8 tools/test_p2c_capital_specs.py   (từ root repo)
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
import tempfile
import unicodedata
from pathlib import Path

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import tools.build_form_maps as bfm          # noqa: E402
import tools.docx_to_blocks as dxtb          # noqa: E402
from tools.docx_render import ExportTotalsError, recompute_totals  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
MANIFEST_PATH = ROOT / "app" / "data" / "table_specs" / "manifest.json"
BLOCKS_DIR = ROOT / "app" / "data" / "blocks"

RESULTS: list[tuple[bool, str]] = []

# ---------------------------------------------------------------------------
# Nguồn chân lý của lane: section làm spec + row id/tag theo census manifest
# ---------------------------------------------------------------------------
SECTIONS = {
    ("PLI_M01", "tai_san_gop_von_von_dau_tu"): {
        "file": "tools/form_specs/formation/PLI_M01.json",
        "semantic_key": "co_cau_tai_san_gop_von_dau_tu",
        "anchor_fold": "tai san gop von:",
        "next_heading_fold": "6. thong tin dang ky thue:",
        "value_header_fold":
            "gia tri von cua tung tai san trong von dau tu (bang so, vnd)",
        "rows": ["dong_viet_nam", "ngoai_te_tu_do_chuyen_doi", "vang",
                 "quyen_su_dung_dat", "quyen_so_huu_tri_tue", "tai_san_khac",
                 "tong_so"],
        "total_row": "tong_so",
        "tags": ["chi_tiet_ngoai_te_tai_san_khac", "gia_tri_vnd", "ty_le"],
        "detail_tag": "chi_tiet_ngoai_te_tai_san_khac",
        "table_shape": (8, 4),
        "tail_tokens": ["{gd.ho_ten}", "{ap_dung_tu}", "{von_dau_tu_so}"],
    },
    ("PLI_M02", "nguon_von_dieule"): {
        "file": "tools/form_specs/core/PLI_M02.json",
        "semantic_key": "nguon_von_dieu_le",
        "anchor_fold": "7. nguon von dieu le:",
        "next_heading_fold": "8. tai san gop von:",
        "value_header_fold":
            "so tien (bang so; vnd va gia tri tuong duong theo don vi tien "
            "nuoc ngoai, neu co)",
        "rows": ["von_ngan_sach_nha_nuoc", "von_tu_nhan", "von_nuoc_ngoai",
                 "von_khac", "tong_cong"],
        "total_row": "tong_cong",
        "tags": ["so_tien_vnd_va_tuong_duong", "ty_le"],
        "detail_tag": None,          # mẫu KHÔNG có ô mô tả cho 'Vốn khác'
        "table_shape": (6, 3),
        "tail_tokens": [],
    },
    ("PLI_M02", "tai_san_gop_von_von_dieu_le"): {
        "file": "tools/form_specs/core/PLI_M02.json",
        "semantic_key": "co_cau_tai_san_gop_von_dieu_le",
        "anchor_fold": "8. tai san gop von:",
        "next_heading_fold": "9. nguoi dai dien theo phap luat",
        "value_header_fold":
            "gia tri von cua tung tai san trong von dieu le (bang so, vnd)",
        "rows": ["dong_viet_nam", "ngoai_te_tu_do_chuyen_doi", "vang",
                 "quyen_su_dung_dat", "quyen_so_huu_tri_tue", "tai_san_khac",
                 "tong_so"],
        "total_row": "tong_so",
        "tags": ["chi_tiet_ngoai_te_tai_san_khac", "gia_tri_vnd", "ty_le"],
        "detail_tag": "chi_tiet_ngoai_te_tai_san_khac",
        "table_shape": (8, 4),
        "tail_tokens": ["{ndd.ho_ten}"],
    },
}
FORM_IDS = ("PLI_M01", "PLI_M02")


def check(name: str):
    def deco(fn):
        def run(*args, **kw):
            try:
                fn(*args, **kw)
                RESULTS.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                RESULTS.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:                        # noqa: BLE001
                RESULTS.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def fold(text: str) -> str:
    """Chung quy tắc với build_form_maps.fold (thường, bỏ dấu, đ->d)."""
    s = re.sub(r"\s+", " ", text or "").strip()
    nfkd = unicodedata.normalize("NFD", s.lower())
    return "".join(c for c in nfkd if unicodedata.category(c) != "Mn").replace(
        "đ", "d")


def blocks_json(form_id: str) -> dict:
    return json.loads((BLOCKS_DIR / ("%s.json" % form_id)).read_text(
        encoding="utf-8"))


def block_texts(blocks: list[dict]) -> list[str]:
    out = []
    for b in blocks:
        runs = b.get("runs") or []
        txt = "".join(str(r.get("text", "")) for r in runs)
        out.append(txt if b.get("t") != "table" else "[TBL]")
    return out


def block_folds(form_id: str) -> list[str]:
    return [fold(t) for t in block_texts(blocks_json(form_id)["blocks"])]


def body_folds(path: Path) -> list[str]:
    doc = docx.Document(str(path))
    out = []
    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}p"):
            out.append(fold(Paragraph(el, doc).text))
        elif el.tag.endswith("}tbl"):
            out.append("[TBL]")
    return out


def sha256_file(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def nest(flat_values: dict) -> dict:
    """Key phẳng '<sem>.<row>.<tag>' -> cây dict như app_web/main.py đựng."""
    out: dict = {}
    for key, value in flat_values.items():
        cur = out
        parts = key.split(".")
        for part in parts[:-1]:
            cur = cur.setdefault(part, {})
        cur[parts[-1]] = value
    return out


def entries_of(form_id: str):
    ext = bfm.load_form_spec(form_id)
    assert ext, "%s: không đọc được file spec" % form_id
    return ext


def spec_entry(form_id: str, section_id: str) -> dict:
    ext = entries_of(form_id)
    hits = [e for e in ext["entries"]
            if (ext["specs_by_id"][e["spec"]].get("section_id")
                or ext["specs_by_id"][e["spec"]]["id"]) == section_id]
    assert len(hits) == 1, "%s: kỳ vọng 1 entry cho %s, thấy %d" % (
        form_id, section_id, len(hits))
    return hits[0]


def manifest_sections() -> dict:
    data = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    return {f["form_id"]: {s["section_id"]: s for s in f["sections"]}
            for f in data["forms"]}


# ---------------------------------------------------------------------------
# Evidence
# ---------------------------------------------------------------------------
@check("[EVD] census manifest co evidence cho 3 section capital/assets")
def t_evidence_manifest():
    forms = manifest_sections()
    for (fid, sid), want in SECTIONS.items():
        sec = forms[fid].get(sid)
        assert sec is not None, "%s: thiếu section %r trong manifest" % (fid, sid)
        assert sec["kind"] == "fixed-matrix", "%s/%s kind=%r" % (fid, sid,
                                                                sec["kind"])
        assert sec["status"] == "missing_grid", \
            "%s/%s status=%r (section đã map rồi thì không rebuild)" % (
                fid, sid, sec["status"])
        assert sec["semantic_key"] == want["semantic_key"], \
            "%s/%s semantic_key=%r != %r" % (fid, sid, sec["semantic_key"],
                                             want["semantic_key"])
        ev = sec.get("source_evidence") or {}
        assert ev.get("heading") and ev.get("anchor"), \
            "%s/%s thiếu evidence nguồn" % (fid, sid)
        assert fold(ev["heading"]).startswith(want["anchor_fold"]), \
            "%s/%s heading evidence %r không khớp anchor %r" % (
                fid, sid, ev["heading"], want["anchor_fold"])


@check("[EVD] row id + tag lay DUNG nguon: manifest fixed_rows va template")
def t_rows_and_tags_from_source():
    forms = manifest_sections()
    listed = forms["PLI_M02"]["nguon_von_dieule"].get("fixed_rows") or []
    ent = spec_entry("PLI_M02", "nguon_von_dieule")
    spec = entries_of("PLI_M02")["specs_by_id"][ent["spec"]]
    assert [r["id"] for r in spec["fixed_rows"]] == list(listed), \
        "PLI_M02/nguon_von_dieule: fixed_rows lệch manifest (%r)" % listed
    # heading value từng mẫu phải là header cột giá trị thật của template
    for (_fid, sid), want in SECTIONS.items():
        ent = spec_entry(_fid, sid)
        spec = entries_of(_fid)["specs_by_id"][ent["spec"]]
        heads = [c["header"] for c in spec["columns"]]
        assert any(fold(h) == want["value_header_fold"] for h in heads), \
            "%s/%s: thiếu header cột giá trị theo mẫu %r" % (_fid, sid, heads)
        # nhãn dòng pháp lý phải là text thật của template (không bị chế)
        for r in spec["fixed_rows"]:
            assert fold(r["label"]), "%s/%s: label rỗng" % (_fid, sid)


# ---------------------------------------------------------------------------
# Spec
# ---------------------------------------------------------------------------
@check("[SPEC] 2 file spec qua validator v1 (load_form_spec) dung kind/section")
def t_spec_validate():
    for fid in FORM_IDS:
        ext = entries_of(fid)
        for e in ext["entries"]:
            spec = ext["specs_by_id"][e["spec"]]
            if spec.get("kind") != "fixed-matrix":
                continue
            assert e.get("no_default_drops") is True, \
                "%s/%s: cần no_default_drops=true" % (fid, spec["id"])
            assert e.get("drop_until"), \
                "%s/%s: drop scopped cần drop_until để không nuốt mục kế" % (
                    fid, spec["id"])
            assert e.get("legal_status", {}).get("blocked") is False
            assert spec["section_id"] == spec["id"], \
                "%s/%s: section_id phải khớp manifest" % (fid, spec["id"])
            cols = spec["columns"]
            assert cols[0]["tag"] is None and cols[0]["input_type"] == "readonly", \
                "%s/%s: cột 0 phải là cột nhãn readonly" % (fid, spec["id"])
            rows = spec["fixed_rows"]
            total_rows = [r for r in rows if r["output_only"]]
            assert len(total_rows) == 1, \
                "%s/%s: đúng 1 dòng tổng output_only, thấy %d" % (
                    fid, spec["id"], len(total_rows))
            assert rows[-1] is total_rows[0] or rows[-1]["output_only"], \
                "%s/%s: dòng tổng phải là dòng CUỐI (theo mẫu)" % (fid, spec["id"])
            # D1.1=A: mọi dòng khác đều là dòng input (không có dòng output ẩn)
            assert len(total_rows) + len(
                [r for r in rows if not r["output_only"]]) == len(rows)


@check("[SPEC] D1.1=A: dong tong output_only + totals sum() dung target_tag")
def t_totals_shape():
    for (fid, sid), want in SECTIONS.items():
        ent = spec_entry(fid, sid)
        spec = entries_of(fid)["specs_by_id"][ent["spec"]]
        assert spec["fixed_rows"][-1]["id"] == want["total_row"], \
            "%s/%s: dòng tổng %r != %r" % (fid, sid,
                                           spec["fixed_rows"][-1]["id"],
                                           want["total_row"])
        assert spec["fixed_rows"][-1]["output_only"] is True
        got_tags = [c["tag"] for c in spec["columns"] if c["tag"]]
        assert got_tags == want["tags"], \
            "%s/%s: tags %r != %r" % (fid, sid, got_tags, want["tags"])
        totals = spec["totals"]
        assert totals, "%s/%s: thiếu totals (D1.1=A computed)" % (fid, sid)
        for t in totals:
            assert t["target_row"] == want["total_row"], \
                "%s/%s: totals target_row %r" % (fid, sid, t["target_row"])
            assert t["target_tag"] in want["tags"], \
                "%s/%s: target_tag %r không phải cột input" % (
                    fid, sid, t["target_tag"])
            # D1.4=A (product) là cơ cấu cổ phần M04/M12 — section này KHÔNG có
            # cột số lượng × mệnh giá nên chỉ được dùng sum()
            assert re.fullmatch(r"sum\([a-z0-9_]+\)", t["formula"]), \
                "%s/%s: formula %r (chỉ sum() đúng contract)" % (
                    fid, sid, t["formula"])
        covered = {t["target_tag"] for t in totals}
        numeric = [t for t in want["tags"] if t != want.get("detail_tag")]
        assert covered == set(numeric), \
            "%s/%s: mọi cột số đều phải có ô tổng computed, thiếu %r" % (
                fid, sid, set(numeric) - covered)


@check("[SPEC] D1.2=B cho M01 + KHONG sub-grid: 'Các tài sản khác' la o mo ta text")
def t_no_subgrid_other_assets():
    for (fid, sid), want in SECTIONS.items():
        ent = spec_entry(fid, sid)
        spec = entries_of(fid)["specs_by_id"][ent["spec"]]
        assert spec["kind"] == "fixed-matrix", \
            "%s/%s: kind %r (không chế component con)" % (fid, sid, spec["kind"])
        assert "loop" not in spec, "%s/%s: fixed-matrix không có loop" % (fid, sid)
        rows = [r["id"] for r in spec["fixed_rows"]]
        assert "tai_san_khac" in rows or sid == "nguon_von_dieule", \
            "%s/%s: thiếu dòng 'Các tài sản khác'" % (fid, sid)
        if want["detail_tag"]:
            col = [c for c in spec["columns"] if c["tag"] == want["detail_tag"]]
            assert len(col) == 1 and col[0]["input_type"] == "text", \
                "%s/%s: ô mô tả phải là 1 cột text (D1.2=B)" % (fid, sid)
            assert not col[0]["computed"] and not col[0]["output_only"], \
                "%s/%s: ô mô tả phải do người dùng nhập" % (fid, sid)
        else:
            assert not any("chi_tiet" in (c["tag"] or "")
                           for c in spec["columns"]), \
                "%s/%s: mẫu KHÔNG có ô mô tả cho 'Vốn khác' -> không tự chế " \
                "cột (D1.2 chưa có quyết định cho M02)" % (fid, sid)
    # PLI_M01: đúng 1 fixed-matrix mới, không đụng 2 section already_complete
    ext = entries_of("PLI_M01")
    fx = [e for e in ext["entries"]
          if ext["specs_by_id"][e["spec"]]["kind"] == "fixed-matrix"]
    assert [ext["specs_by_id"][e["spec"]]["section_id"] for e in fx] == \
        ["tai_san_gop_von_von_dau_tu"], \
        "PLI_M01: chỉ được thêm section missing_grid, thấy %r" % [e for e in fx]


@check("[GAP] nhung gi contract v1 khong khai duoc phai duoc ghi la gap")
def t_documented_gaps():
    for (fid, sid), want in SECTIONS.items():
        ent = spec_entry(fid, sid)
        spec = entries_of(fid)["specs_by_id"][ent["spec"]]
        # (1) cột STT 1..6 in sẵn của mẫu KHÔNG tái lập được ở fixed-matrix v1
        assert not any((c["tag"] or "") == "stt" for c in spec["columns"]), \
            "%s/%s: có cột stt? fixed-matrix v1 không có slot STT" % (fid, sid)
        # (2) % tự tính theo giá trị chưa có formula -> ty_le là ô nhập
        tl = [c for c in spec["columns"] if c["tag"] == "ty_le"][0]
        assert not tl["computed"] and tl["input_type"] == "percent", \
            "%s/%s: ty_le phải là ô nhập percent (v1 chưa có formula %)" % (
                fid, sid)
        # (3) cross_field đối chiếu vốn điều lệ: spec v1 không có slot
        assert "cross_field_rules" not in spec, \
            "%s/%s: slot cross_field_rules chưa thuộc contract v1" % (fid, sid)
        ls = ent.get("legal_status") or {}
        assert ls.get("question"), \
            "%s/%s: gap phải được ghi vào legal_status.question" % (fid, sid)
        assert "gap" in fold(ls["question"]), \
            "%s/%s: legal_status.question phải nêu rõ gap" % (fid, sid)


# ---------------------------------------------------------------------------
# Anchor + drop
# ---------------------------------------------------------------------------
@check("[ANCHOR] anchor dung DUY NHAT 1 paragraph trong blocks + template")
def t_anchor_unique():
    for (fid, sid), want in SECTIONS.items():
        ent = spec_entry(fid, sid)
        folds_b = block_folds(fid)
        hits_b = [i for i, t in enumerate(folds_b)
                  if re.search(ent["anchor"], t)]
        assert len(hits_b) == 1, \
            "%s/%s: anchor khớp blocks %r (cần đúng 1)" % (fid, sid, hits_b)
        assert folds_b[hits_b[0]].startswith(want["anchor_fold"]), \
            "%s/%s: anchor trúng %r thay vì heading evidence %r" % (
                fid, sid, folds_b[hits_b[0]][:60], want["anchor_fold"])
        src = bfm.find_source(fid)
        assert src is not None, "%s: không tìm thấy nguồn docx" % fid
        doc = docx.Document(str(src))
        hits_t = [i for i, p in enumerate(doc.paragraphs)
                  if re.search(ent["anchor"], fold(p.text))]
        assert len(hits_t) == 1, \
            "%s/%s: anchor khớp template %d lần" % (fid, sid, len(hits_t))


@check("[DROP] drop chi trong [anchor, drop_until) va khong nuot muc ke")
def t_drop_scoped_and_boundary():
    for fid in FORM_IDS:
        ext = entries_of(fid)
        src = bfm.find_source(fid)
        doc = docx.Document(str(src))
        raws = [p.text for p in doc.paragraphs]
        folds = [fold(t) for t in raws]
        for e in ext["entries"]:
            spec = ext["specs_by_id"][e["spec"]]
            if spec.get("kind") != "fixed-matrix":
                continue
            start = next(i for i, t in enumerate(folds)
                         if re.search(e["anchor"], t)) + 1
            end = next(i for i in range(start, len(folds))
                       if re.search(e["drop_until"], folds[i]))
            # boundary: paragraph mục kế KHÔNG bị drop pattern chạm tới
            assert not any(re.match(p, folds[end]) for p in e["drop"]), \
                "%s/%s: drop pattern nuốt cả heading mục kế %r" % (
                    fid, spec["id"], raws[end][:60])
            # trong vùng: mọi paragraph KHÔNG RỖNG phải là glue của chính mục
            # này (header/nhãn dòng/STT số) — nếu sót một paragraph có nội dung
            # thì rebuild sẽ để lại text rác dán kèm bảng mới
            missed = [raws[i] for i in range(start, end)
                      if folds[i] and not any(re.match(p, folds[i])
                                              for p in e["drop"])]
            assert not missed, \
                "%s/%s: drop bỏ sót glue trong vùng %r" % (fid, spec["id"],
                                                           missed[:3])
            # ô trống (paragraph rỗng) không bị drop pattern nào nuốt: chúng
            # là khoảng trình bày của mẫu, engine chỉ drop paragraph có text
            blanks = [i for i in range(start, end) if not folds[i]]
            assert blanks, "%s/%s: vùng không có ô trống như evidence" % (
                fid, spec["id"])
            # phần nội dung pháp lý dính kèm nằm NGOÀI vùng phải còn nguyên
            assert re.search(e["drop_until"], folds[end]), \
                "%s/%s: mất mốc drop_until" % (fid, spec["id"])


# ---------------------------------------------------------------------------
# Sandbox rebuild + convert
# ---------------------------------------------------------------------------
def new_sandbox() -> tuple[Path, Path]:
    tmp = Path(tempfile.mkdtemp(prefix="p2c_capital_sandbox_"))
    tpl_dir = tmp / "templates"
    data_dir = tmp / "data"
    tpl_dir.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)
    return tpl_dir, data_dir


def run_generation(fid: str, mau: int, tpl_dir: Path) -> tuple[dict, Path]:
    bfm._form_spec_cache.clear()
    src = bfm.find_source(fid)
    meta, _unmatched = bfm.process_form(src, {"Phu luc I|%d" % mau: "DEMO %s"
                                              % fid}, None)
    tpl = tpl_dir / ("%s.docx" % fid)
    assert tpl.exists(), "%s: generator chưa ghi template" % fid
    return meta, tpl


def fixed_matrix_tokens(blocks: dict, sem: str) -> list[str]:
    """Mọi the input '<sem>.<row>.<tag>' của các bảng trong blocks (blocks lưu
    token dưới dạng runs[].input.key, không phải text '{...}')."""
    toks: list[str] = []
    for b in blocks["blocks"]:
        if b.get("t") != "table":
            continue
        for row in b.get("rows") or []:
            for cell in row.get("cells") or []:
                for r in cell.get("runs") or []:
                    key = ((r.get("input") or {}).get("key"))
                    if key:
                        toks.append(str(key))
    return [t for t in toks if t.startswith(sem + ".")]


@check("[SANDBOX] rebuild PLI_M01: bang 8x4 + glue bien mat + muc ke nguyen ven")
def t_sandbox_m01():
    guards = {p.name: sha256_file(p)
              for p in sorted((ROOT / "app" / "templates").glob("PLI_M0[12].docx"))}
    tpl_dir, _data = new_sandbox()
    saved_tpl, saved_data = bfm.OUT_TPL, bfm.OUT_DATA
    bfm.OUT_TPL, bfm.OUT_DATA = tpl_dir, _data
    bfm._form_spec_cache.clear()
    try:
        want = SECTIONS[("PLI_M01", "tai_san_gop_von_von_dau_tu")]
        meta, tpl = run_generation("PLI_M01", 1, tpl_dir)
        specs = [s for s in meta["section_specs"]
                 if s["kind"] == "fixed-matrix"]
        assert len(specs) == 1, "PLI_M01: 1 section_spec, thấy %d" % len(specs)
        sp = specs[0]
        assert sp["semantic_key"] == want["semantic_key"]
        # đủ 7 dòng pháp lý (kể cả khi trống) + 1 hàng tiêu đề
        assert len(sp["fixed_rows"]) == len(want["rows"])
        assert [r["id"] for r in sp["fixed_rows"]] == want["rows"]
        doc = docx.Document(str(tpl))
        shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
        assert want["table_shape"] in shapes, \
            "PLI_M01: không thấy w:tbl %s, có %r" % (want["table_shape"], shapes)
        texts = [Paragraph(el, doc).text
                 for el in doc.element.body.iterchildren()
                 if el.tag.endswith("}p")]
        folds = [fold(t) for t in texts]
        # glue (header + nhãn dòng rời) phải biến mất — chỉ xét VÙNG mục này
        # (các mục sau vẫn có header 'STT' của bảng thuế, không phải glue)
        end = folds.index(fold(want["next_heading_fold"]))
        region = folds[:end + 1]
        for glue in ("stt", "dong viet nam", "quyen so huu tri tue",
                     "cac tai san khac (ghi ro loai tai san", "tong so",
                     "gia tri von cua tung tai san trong von dau tu"):
            left = [t for t in region if t.startswith(glue)]
            assert not left, "PLI_M01: glue còn sót %r" % left[:2]
        # mục kế (6. Thuế) giữ nguyên kèm token thật
        assert want["next_heading_fold"].rstrip(":") in " ".join(folds), \
            "PLI_M01: mất heading mục kế"
        blob = "".join(texts)
        for tok in want["tail_tokens"]:
            assert tok in blob, "PLI_M01: mất token mục kế %s" % tok
        # blocks qua kiểm định converter (annotate_and_verify đã chạy trong
        # convert_document) + đủ the từng ô
        blocks = dxtb.convert_document(docx.Document(str(tpl)), "PLI_M01",
                                       "DEMO", overrides={}, footnotes={},
                                       section_specs=meta["section_specs"])
        toks = fixed_matrix_tokens(blocks, want["semantic_key"])
        for r in want["rows"]:
            for tag in want["tags"]:
                assert ("%s.%s.%s" % (want["semantic_key"], r, tag)) in toks, \
                    "PLI_M01: thiếu the %s.%s.%s" % (want["semantic_key"], r, tag)
        # ô tổng CŨNG có the (engine tính lại lúc export) nhưng KHÔNG đăng ký
        # thành field form_maps (không có dữ liệu giả ghi đè kết quả tính lại)
        total_key = "%s.%s.gia_tri_vnd" % (want["semantic_key"],
                                           want["total_row"])
        assert total_key in toks
        fields = [f["key"] for f in meta["fields"]]
        assert not [k for k in fields if k.startswith(
            "%s.%s." % (want["semantic_key"], want["total_row"]))], \
            "PLI_M01: dòng tổng bị đăng ký thành field (vi phạm outputOnly)"
        assert ("%s.dong_viet_nam.gia_tri_vnd" % want["semantic_key"]) in fields
        # idempotent: 2 lần chạy cho cùng template
        h1 = sha256_file(tpl)
        run_generation("PLI_M01", 1, tpl_dir)
        assert sha256_file(tpl) == h1, "PLI_M01: rebuild 2 lần lệch nhau"
        # template gốc trong repo KHÔNG bị chạm
        for name, before in guards.items():
            assert sha256_file(ROOT / "app" / "templates" / name) == before, \
                "%s bị ghi ngoài sandbox!" % name
    finally:
        bfm.OUT_TPL, bfm.OUT_DATA = saved_tpl, saved_data
        bfm._form_spec_cache.clear()
        shutil.rmtree(tpl_dir.parent, ignore_errors=True)


@check("[SANDBOX] rebuild PLI_M02: 2 bang moi + N1 giu nguyen + muc 9 nguyen ven")
def t_sandbox_m02():
    tpl_dir, _data = new_sandbox()
    saved_tpl, saved_data = bfm.OUT_TPL, bfm.OUT_DATA
    bfm.OUT_TPL, bfm.OUT_DATA = tpl_dir, _data
    bfm._form_spec_cache.clear()
    try:
        meta, tpl = run_generation("PLI_M02", 2, tpl_dir)
        specs = {s["section_id"]: s for s in meta["section_specs"]}
        assert set(specs) >= {"nddpl_uy_quyen_chu_so_huu", "nguon_von_dieule",
                              "tai_san_gop_von_von_dieu_le"}, \
            "PLI_M02: thiếu section_spec %r" % sorted(specs)
        assert specs["nddpl_uy_quyen_chu_so_huu"]["kind"] == "nested-grid", \
            "PLI_M02: entry N1 (lane people) phải giữ nguyên"
        doc = docx.Document(str(tpl))
        shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
        for sid in ("nguon_von_dieule", "tai_san_gop_von_von_dieu_le"):
            want = SECTIONS[("PLI_M02", sid)]
            assert want["table_shape"] in shapes, \
                "%s: không thấy w:tbl %s, có %r" % (sid, want["table_shape"],
                                                    shapes)
        # header nhiều tầng N1 (14 cột: 2 hàng header + 5 hàng trống + 1 khuôn)
        # không bị hai bảng mới làm lệch
        assert (8, 14) in shapes, "PLI_M02: mất bảng nested-grid N1: %r" % shapes
        texts = [Paragraph(el, doc).text
                 for el in doc.element.body.iterchildren()
                 if el.tag.endswith("}p")]
        folds = [fold(t) for t in texts]
        # glue chỉ được xét trong vùng hai mục (mục 9 vẫn có 'Tổng số lao động')
        glue_end = folds.index("9. nguoi dai dien theo phap luat9:")
        region = folds[:glue_end + 1]
        for glue in ("loai nguon von", "von ngan sach nha nuoc", "von khac",
                     "tong cong", "dong viet nam", "tong so",
                     "cac tai san khac (ghi ro loai tai san"):
            left = [t for t in region if t.startswith(glue)]
            assert not left, "PLI_M02: glue còn sót %r" % left[:2]
        blob = "".join(texts)
        # {#loop} N1 nằm trong Ô BẢNG (không phải paragraph body)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    blob += cell.text
        for tok in ("{von_dieu_le}", "{#ds_nddpl_uy_quyen_chu_so_huu}",
                    "{ndd.ho_ten}"):
            assert tok in blob, "PLI_M02: mất chuỗi pháp lý/token %s" % tok
        blocks = dxtb.convert_document(docx.Document(str(tpl)), "PLI_M02",
                                       "DEMO", overrides={}, footnotes={},
                                       section_specs=meta["section_specs"])
        for sid in ("nguon_von_dieule", "tai_san_gop_von_von_dieu_le"):
            want = SECTIONS[("PLI_M02", sid)]
            toks = fixed_matrix_tokens(blocks, want["semantic_key"])
            for r in want["rows"]:
                for tag in want["tags"]:
                    key = "%s.%s.%s" % (want["semantic_key"], r, tag)
                    assert key in toks, "PLI_M02/%s: thiếu the %s" % (sid, key)
        h1 = sha256_file(tpl)
        run_generation("PLI_M02", 2, tpl_dir)
        assert sha256_file(tpl) == h1, "PLI_M02: rebuild 2 lần lệch nhau"
    finally:
        bfm.OUT_TPL, bfm.OUT_DATA = saved_tpl, saved_data
        bfm._form_spec_cache.clear()
        shutil.rmtree(tpl_dir.parent, ignore_errors=True)


# ---------------------------------------------------------------------------
# Export (D1.1=A)
# ---------------------------------------------------------------------------
@check("[EXPORT] D1.1=A: tong tinh lai tu o du lieu, ghi de stored, conflict blocking")
def t_export_recompute():
    tpl_dir, _data = new_sandbox()
    saved_tpl, saved_data = bfm.OUT_TPL, bfm.OUT_DATA
    bfm.OUT_TPL, bfm.OUT_DATA = tpl_dir, _data
    bfm._form_spec_cache.clear()
    try:
        meta, tpl = run_generation("PLI_M01", 1, tpl_dir)
        specs = meta["section_specs"]
        want = SECTIONS[("PLI_M01", "tai_san_gop_von_von_dau_tu")]
        sem = want["semantic_key"]
        flat = {
            "%s.dong_viet_nam.gia_tri_vnd" % sem: "1.000.000",
            "%s.dong_viet_nam.ty_le" % sem: "60",
            "%s.tai_san_khac.%s" % (sem, "chi_tiet_ngoai_te_tai_san_khac"):
                "Xe hơi phục vụ lãnh đạo, giá trị còn lại 500 triệu",
            "%s.tai_san_khac.gia_tri_vnd" % sem: "500000",
            "%s.tai_san_khac.ty_le" % sem: "40",
            # stored SAI chủ đích: engine phải tính lại, không tin draft
            "%s.tong_so.gia_tri_vnd" % sem: "1",
            "%s.tong_so.ty_le" % sem: "7",
        }
        data = nest(flat)
        report = {r["target"]: r for r in
                  recompute_totals(data, specs, strict=False)}
        g = report["%s.tong_so.gia_tri_vnd" % sem]
        p = report["%s.tong_so.ty_le" % sem]
        assert g["computed"] == "1500000", \
            "PLI_M01: tổng giá trị %r != 1500000" % g["computed"]
        assert p["computed"] == "100", \
            "PLI_M01: tổng tỷ lệ %r != 100" % p["computed"]
        # dòng không kê khai (ngoại tệ/vàng/QSDĐ/QSHTT) không bị tính thành 0
        assert g["had_input"] is True
        # strict + stored != computed -> raise EXPORT_TOTALS_RECOMPUTE_CONFLICT
        try:
            recompute_totals(nest(flat), specs, strict=True)
        except ExportTotalsError as exc:
            codes = {p_.get("code") for p_ in exc.problems}
            assert codes == {"EXPORT_TOTALS_RECOMPUTE_CONFLICT"}, codes
        else:  # pragma: no cover
            raise AssertionError("PLI_M01: stored sai mà không raise conflict")
        # render(): ô tổng trong DOCX xuất ra mang số đã tính lại. Draft thật
        # không bao giờ giữ số ở ô outputOnly (projection) nên render với ô
        # tổng BỎ TRỐNG — cấm tình huống "stored lệch" trên đường export thật.
        export_flat = {k: v for k, v in flat.items()
                       if not k.startswith("%s.%s." % (sem, want["total_row"]))}
        from tools.docx_render import write_rendered
        out = tpl_dir / "PLI_M01_rendered.docx"
        write_rendered(str(tpl), nest(export_flat), str(out), tables_meta=specs)
        rdoc = docx.Document(str(out))
        target = None
        for t in rdoc.tables:
            if len(t.columns) == 4 and len(t.rows) == 8:
                target = t
        assert target is not None, "PLI_M01: mất bảng tài sản trong DOCX render"
        last = [c.text.strip() for c in target.rows[-1].cells]
        assert last == ["Tổng số", "", "1500000", "100"], \
            "PLI_M01: hàng tổng render %r" % last
        # label "Các tài sản khác" giữ nguyên, ô mô tả D1.2=B nằm ở cột chi tiết
        row6 = [c.text.strip() for c in target.rows[6].cells]
        assert fold(row6[0]).startswith("cac tai san khac"), row6
        assert "Xe hơi phục vụ lãnh đạo" in row6[1], \
            "PLI_M01: ô mô tả 'Các tài sản khác' không nhận dữ liệu: %r" % row6
    finally:
        bfm.OUT_TPL, bfm.OUT_DATA = saved_tpl, saved_data
        bfm._form_spec_cache.clear()
        shutil.rmtree(tpl_dir.parent, ignore_errors=True)


@check("[EXPORT] PLI_M02: tong nguon von + tong tai san deu computed outputOnly")
def t_export_m02_totals():
    tpl_dir, _data = new_sandbox()
    saved_tpl, saved_data = bfm.OUT_TPL, bfm.OUT_DATA
    bfm.OUT_TPL, bfm.OUT_DATA = tpl_dir, _data
    bfm._form_spec_cache.clear()
    try:
        meta, _tpl = run_generation("PLI_M02", 2, tpl_dir)
        specs = meta["section_specs"]
        sem_v = "nguon_von_dieu_le"
        sem_t = "co_cau_tai_san_gop_von_dieu_le"
        flat = {
            "%s.von_tu_nhan.so_tien_vnd_va_tuong_duong" % sem_v: "2.000.000.000",
            "%s.von_tu_nhan.ty_le" % sem_v: "100",
            "%s.dong_viet_nam.gia_tri_vnd" % sem_t: "2000000000",
            "%s.dong_viet_nam.ty_le" % sem_t: "100",
        }
        report = {r["target"]: r for r in
                  recompute_totals(nest(flat), specs, strict=False)}
        assert report["%s.tong_cong.so_tien_vnd_va_tuong_duong"
                      % sem_v]["computed"] == "2000000000"
        assert report["%s.tong_cong.ty_le" % sem_v]["computed"] == "100"
        assert report["%s.tong_so.gia_tri_vnd" % sem_t]["computed"] == \
            "2000000000"
        # 'Vốn khác' KHÔNG có cột mô tả (D1.2 chưa có quyết định cho M02)
        spec_v = [s for s in specs if s["section_id"] == "nguon_von_dieule"][0]
        assert [c["tag"] for c in spec_v["columns"] if c["tag"]] == \
            ["so_tien_vnd_va_tuong_duong", "ty_le"]
    finally:
        bfm.OUT_TPL, bfm.OUT_DATA = saved_tpl, saved_data
        bfm._form_spec_cache.clear()
        shutil.rmtree(tpl_dir.parent, ignore_errors=True)


# ---------------------------------------------------------------------------
# BLOCKED / giữ nguyên
# ---------------------------------------------------------------------------
@check("[BLOCKED] PLI_M01: 2 section already_complete van khong co entry rebuild")
def t_blocked_m01_already_mapped():
    raw = json.loads((ROOT / "tools" / "form_specs" / "formation" /
                      "PLI_M01.json").read_text(encoding="utf-8-sig"))
    sids = [(e["spec"].get("section_id") or e["spec"]["id"])
            for e in raw["tables"]]
    for already in ("lua_chon_khu_san_xuat", "thue_chi_tieu_6_1_6_9"):
        assert already not in sids, \
            "PLI_M01: section already_complete %s không được rebuild (đè mất " \
            "chk_khu_* / gd.* / kt.* đã map)" % already
    # entry cũ (list ds_nganh) giữ nguyên
    assert raw["tables"][0]["spec"]["id"] == "ds_nganh"
    assert (raw["tables"][0]["spec"].get("kind") or "list") == "list"
    # blocks hiện trạng vẫn còn binding vùng đã map (chưa regen, không bị đụng)
    blob = (BLOCKS_DIR / "PLI_M01.json").read_text(encoding="utf-8")
    for need in ("chk_khu_cong_nghiep", "chk_khu_che_xuat", "chk_khu_kinh_te",
                 "chk_khu_cong_nghe_cao", "gd.ho_ten", "ap_dung_tu",
                 "von_dau_tu_so"):
        assert re.search(r'"key"\s*:\s*"%s"' % re.escape(need), blob), \
            "blocks PLI_M01 mất binding %s" % need


@check("[BLOCKED] PLI_M02: mo_hinh_to_chuc + entry N1 giu nguyen (count 1 nested)")
def t_blocked_m02_no_packet():
    raw = json.loads((ROOT / "tools" / "form_specs" / "core" /
                      "PLI_M02.json").read_text(encoding="utf-8-sig"))
    entries = raw["tables"]
    ids = [e["spec"]["id"] for e in entries]
    assert "mo_hinh_to_chuc_cong_ty" not in ids, \
        "PLI_M02: mô hình tổ chức công ty chưa có packet quyết định -> không spec"
    nested = [e for e in entries
              if (e["spec"].get("kind") or "list") == "nested-grid"]
    assert len(nested) == 1 and nested[0]["spec"]["id"] == \
        "nddpl_uy_quyen_chu_so_huu", \
        "PLI_M02: entry nested-grid N1 phải giữ NGUYÊN đúng 1 (lane people)"
    assert nested[0]["spec"]["semantic_key"] == "ds_nddpl_uy_quyen_chu_so_huu"
    assert nested[0]["spec"]["columns"][0]["tag"] == "ten_nguoi_dai_dien"
    # blocks hiện trạng (sau regen R3B1) có đúng 1 block loop N1, không bị hai
    # section mới của lane này đụng tới
    loops = [b for b in blocks_json("PLI_M02")["blocks"]
             if b.get("t") == "loop"
             and b.get("loop") == "ds_nddpl_uy_quyen_chu_so_huu"]
    assert len(loops) == 1, "PLI_M02: block loop N1 %d != 1" % len(loops)
    assert [c["tag"] for c in loops[0]["cols"]][0] == "ten_nguoi_dai_dien"


@check("[BLOCKED] D1.2 policy ngoai M01 + M03/M04/M05/M12: khong spec tu doan")
def t_blocked_other_forms():
    # M03/M04/M05 đã có file spec (bảng ngành Wave 2A) nhưng phần vốn/cổ phần
    # của chúng là BLOCKED — không được thêm fixed-matrix hộ owner
    for fid, banned in (("PLI_M03", "nguon_von_dieule"),
                        ("PLI_M04", "co_cau_loai_co_phan"),
                        ("PLI_M05", "nguon_von_dieule")):
        path = ROOT / "tools" / "form_specs" / "formation" / ("%s.json" % fid)
        assert path.exists(), "%s: mất file spec ngành có sẵn" % fid
        raw = json.loads(path.read_text(encoding="utf-8-sig"))
        kinds = [(e["spec"].get("kind") or "list") for e in raw["tables"]]
        assert "fixed-matrix" not in kinds, \
            "%s: xuất hiện fixed-matrix (%s là BLOCKED, fixture F2/F3 chỉ là " \
            "contract runtime của lane khác)" % (fid, banned)
    # PLI_M12 capital: chỉ có entry người đại diện (lane people), không fixed-matrix vốn
    p12 = ROOT / "tools" / "form_specs" / "change" / "PLI_M12.json"
    if p12.exists():
        raw = json.loads(p12.read_text(encoding="utf-8-sig"))
        kinds = {(e["spec"].get("kind") or "list") for e in raw["tables"]}
        assert "fixed-matrix" not in kinds, \
            "PLI_M12: capital (policy ngoài M01) là BLOCKED -> không khai fixed-matrix"
    # 2 file của lane: không thêm form nào khác
    forms = manifest_sections()
    assert "nguon_von_dieule" not in forms["PLI_M01"], \
        "PLI_M01 (DNTN) không có mục nguồn vốn điều lệ — không được bịa section"


@check("[GAP] ghi nhan suite ngoac se do: ID5 + W1 gap guard doi regen Lead")
def t_known_external_red():
    """Không phải mệnh đề fail — ghi nhận hiện trạng cho Lead (luôn PASS)."""
    print("     - tools/test_generator_deep_idempotency.py [ID5] và "
          "tools/test_w1_publication_specs.py [W1-GAP] sẽ đỏ cho tới khi Lead "
          "regen artifact PLI_M01 (spec mới thật, artifact chưa regen) và thu "
          "hẹp guard W1 theo section_id.")


def main() -> int:
    for fn in (t_evidence_manifest, t_rows_and_tags_from_source,
               t_spec_validate, t_totals_shape, t_no_subgrid_other_assets,
               t_documented_gaps, t_anchor_unique, t_drop_scoped_and_boundary,
               t_sandbox_m01, t_sandbox_m02, t_export_recompute,
               t_export_m02_totals, t_blocked_m01_already_mapped,
               t_blocked_m02_no_packet, t_blocked_other_forms,
               t_known_external_red):
        fn()
    ok = sum(1 for p, _ in RESULTS if p)
    fail = sum(1 for p, _ in RESULTS if not p)
    print("-" * 72)
    print("test_p2c_capital_specs: %d PASS / %d FAIL" % (ok, fail))
    for passed, name in RESULTS:
        if not passed:
            print("  -> %s" % name)
    return 1 if fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
