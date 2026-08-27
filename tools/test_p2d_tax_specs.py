#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test_p2d_tax_specs.py — Lane P2D (tax, PLI_M17/M19/M23/M25).

Mệnh đề kiểm chứng cho đợt P2D theo docs/decision_registry.md packet D5
(D5.1=A, D5.3=A, D5.4=A; D5.2=BLOCKED — PLII_M02 thuộc lane khác, không đụng):

  [EVD]     census manifest có evidence cho TẤT CẢ section P2D của 4 mẫu, kèm
            defect danh số M23 đã được ghi nhận.
  [SPEC]    4 file spec của lane qua loader/validator v1; KHÔNG entry kind mới
            (fixed-matrix/nested-grid/output-table) — lý do ở [NO-REBUILD];
            entry ngành `ds_nganh` giữ nguyên hành vi.
  [TAGS]    (cốt lõi) MỌI ô ☐/□ của NHÓM CHỌN-1 và CHỈ TIÊU CỐ ĐỊNH thuộc
            phạm vi D5.1/D5.4 ĐÃ có token riêng: khớp (nhãn trên nguồn) → key,
            đúng THỨ TỰ, và key đó phải (i) là chk block, (ii) nằm trong
            form_maps.fields, (iii) xuất hiện dưới dạng {key} trong template
            đã rebuild. Không thừa, không thiếu, không mồ côi.
  [ANCHOR]  anchor khớp ĐÚNG 1 paragraph (blocks hiện tại + template nguồn),
            drop_until nằm SAU anchor, mọi drop pattern đều khớp >=1 paragraph
            trong khoảng [anchor, drop_until) → boundary không nuốt vùng kế.
  [SRC]     danh số theo TEMPLATE NGUỒN: M17 mục 6, M19 là trang "ĐĂNG KÝ THAY
            ĐỔI THÔNG TIN ĐĂNG KÝ THUẾ", M23 mục 7 (KHÔNG phải 6 như fragment),
            M25 mục 6 → defect pli_m23_heading_numbering xác nhận đúng.
  [D5.3]    chỉ M17 có phần BHXH + dòng "Lưu ý" về ngành nông-lâm-ngư-diêm trả
            lương theo sản phẩm; M19/M23/M25 KHÔNG có → không khai rule.
            Slot cross_field_rules nạp được qua loader và đủ trường.
  [D5.4]    M17: chỉ tiêu "8. Chi nhánh CTCK nước ngoài/CLQNNN tại VN" là mục
            RIÊNG NGOÀI bảng thuế — {chk_ctck} ở paragraph, không trong w:tbl.
  [SLOT]    cross_field_rules (khai entry-level trên entry kind=list, slot duy
            nhất loader chấp nhận) là INERT: template byte-identical + blocks
            identical khi bỏ slot đi.
  [NO-REBUILD] sandbox: nếu khai fixed-matrix cho nhóm BHXH như D5.1 "naive" thì
            form_maps sinh thêm key '<sem>.<row>.<tag>' SONG SONG với chk_bc_*
            (vẫn còn nguyên) → hai nguồn sự thật cho cùng một ô vuông. Suite
            chốt rằng spec thật KHÔNG sinh key nào như vậy.
  [GUARD]   app/templates/** và app/data/** không bị ghi trong suốt suite.

Chạy: python3 -X utf8 tools/test_p2d_tax_specs.py   (từ root repo/worktree)
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import sys
import tempfile
import copy
import contextlib
from pathlib import Path

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import tools.build_form_maps as bfm          # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
MANIFEST_PATH = ROOT / "app" / "data" / "table_specs" / "manifest.json"
BLOCKS_DIR = ROOT / "app" / "data" / "blocks"
FORM_MAPS = ROOT / "app" / "data" / "form_maps.json"
SPEC_DIR = ROOT / "tools" / "form_specs" / "change"

FORMS = ("PLI_M17", "PLI_M19", "PLI_M23", "PLI_M25")

# các section P2D nằm trong phạm vi lane (đối chiếu census manifest)
P2D_SECTIONS = {
    "PLI_M17": ["nganh_nghe_kinh_doanh_chi_nhanh", "thue_chi_tieu_6_1_8",
                "lua_chon_khu_cn_cxc_kt_kccn", "phuong_thuc_dong_bhxh"],
    "PLI_M19": ["nganh_nghe_sau_thay_doi", "thue_thay_doi_A"],
    "PLI_M23": ["nganh_nghe_bo_sung_cap_nhat", "thue_chi_tieu_6_1_6"],
    "PLI_M25": ["nganh_nghe_bo_sung_cap_nhat_ctck", "thue_chi_tieu_6_1_6_ctck"],
}

# Key chỉ tồn tại ở PARAGRAPH level (không thành chk block khi convert) nhưng
# vẫn là một ô ☐ thật trên nguồn và đã được đăng ký trong form_maps.
PARA_ONLY_KEYS = {"PLI_M17": {"chk_ctck"},
                  "PLI_M23": {"chk_de_nghi_cap_gcn"},
                  "PLI_M25": {"chk_de_nghi_cap_gcn"}}

# ---------------------------------------------------------------------------
# Kho evidence D5.1/D5.4, lấy TRỰC TIẾP từ template nguồn (python-docx).
# Mỗi region: (marker_bat_dau, marker_ket_thuc, [(nhan_hoac_None, key), ...])
#   - nhan != None  -> paragraph đó phải CHỨA fold(nhan); đây là evidence nhãn
#   - nhan == None  -> paragraph đó phải chứa ký tự ô vuông □ / ☐
#   - key   == None -> chỉ là mốc nội dung (heading nhóm), không đòi input
# Cặp (nhan, key) được tiêu thụ THEO THỨ TỪ trên xuống, không quay lại.
BOX_RE = re.compile(r"[□☐]")

# fingerprint app/** chụp lúc nạp module — [GUARD] đối chiếu lúc kết thúc
def _fingerprint() -> dict:
    dig = {}
    for sub in ("templates", "data"):
        d = ROOT / "app" / sub
        for p in sorted(d.rglob("*")):
            if p.is_file():
                dig[str(p.relative_to(ROOT))] = hashlib.sha256(
                    p.read_bytes()).hexdigest()
    return dig


REPO_FP_BEFORE = _fingerprint()

TAX_TAIL_GTGT = [
    ("phuong phap tinh thue gtgt", None),
    ("khau tru", "chk_gtgt_khoan_tru"),
    ("truc tiep tren gtgt", "chk_gtgt_truc_tiep_tren_gtgt"),
    ("truc tiep tren doanh so", "chk_gtgt_truc_tiep_doanh_so"),
    ("khong phai nop thue gtgt", "chk_gtgt_khong_phai_nop"),
]
HACH_TOAN = [
    ("hach toan doc lap", "chk_ht_doc_lap"),
    ("co bao cao tai chinh hop nhat", "chk_ht_hop_nhat"),
    ("hach toan phu thuoc", "chk_ht_phu_thuoc"),
]
BOT = [
    ("hoat dong theo du an bot", None),
    (None, "chk_co"),
    (None, "chk_khong"),
]
KHU_CN = [
    ("khu cong nghiep", "chk_khu_cong_nghiep"),
    ("khu che xuat", "chk_khu_che_xuat"),
    ("khu kinh te", "chk_khu_kinh_te"),
    ("khu cong nghe cao", "chk_khu_cong_nghe_cao"),
]

REGIONS = {
    "PLI_M17": [
        ("- chi nhanh/van phong dai dien/dia diem kinh doanh nam trong",
         "3. nganh, nghe kinh doanh", KHU_CN),
        ("6. thong tin dang ky thue", "7. thong tin ve viec dong bao hiem xa hoi",
         HACH_TOAN + BOT + TAX_TAIL_GTGT),
        ("7. thong tin ve viec dong bao hiem xa hoi",
         "8. chi nhanh cong ty chung khoan nuoc ngoai",
         [("phuong thuc dong bao hiem xa hoi", None),
          ("hang thang", "chk_bc_hang_thang"),
          ("03 thang mot lan", "chk_bc_03_thang"),
          ("06 thang mot lan", "chk_bc_06_thang")]),
        ("8. chi nhanh cong ty chung khoan nuoc ngoai", "doanh nghiep cam ket",
         [("danh dau x neu la chi nhanh cong ty chung khoan nuoc ngoai",
           "chk_ctck")]),
    ],
    "PLI_M19": [
        ("- chi nhanh/van phong dai dien/dia diem kinh doanh nam trong",
         "dang ky thay doi nganh, nghe kinh doanh", KHU_CN),
        ("dang ky thay doi thong tin dang ky thue",
         "b. bo sung, cap nhat thong tin dang ky hoat dong",
         HACH_TOAN + BOT),
    ],
    "PLI_M23": [
        ("7. thong tin dang ky thue", "ii. thong tin ve chi nhanh chu quan",
         HACH_TOAN + BOT + TAX_TAIL_GTGT +
         [("de nghi quy co quan cap giay chung nhan dang ky hoat dong",
           "chk_de_nghi_cap_gcn")]),
    ],
    "PLI_M25": [
        ("6. thong tin dang ky thue", "doanh nghiep cam ket",
         HACH_TOAN + BOT + TAX_TAIL_GTGT +
         [("de nghi quy co quan cap giay chung nhan dang ky hoat dong",
           "chk_de_nghi_cap_gcn")]),
    ],
}
# Khu CN chỉ có ở M17/M19; M23/M25 và BHXH chỉ có ở M17 -> negative evidence
NEGATIVE = {
    "PLI_M17": [],
    "PLI_M19": ["bao hiem xa hoi", "phuong phap tinh thue gtgt"],
    "PLI_M23": ["bao hiem xa hoi", "khu cong nghe cao", "nong nghiep"],
    "PLI_M25": ["bao hiem xa hoi", "khu cong nghe cao", "nong nghiep"],
}

# key mà mẫu đó KHÔNG được đăng ký (tương ứng negative ở trên)
NEGATIVE_KEYS = {
    "PLI_M17": [],
    "PLI_M19": ["chk_gtgt_khoan_tru", "chk_gtgt_truc_tiep_tren_gtgt",
                "chk_gtgt_truc_tiep_doanh_so", "chk_gtgt_khong_phai_nop",
                "chk_bc_hang_thang", "chk_bc_03_thang", "chk_bc_06_thang"],
    "PLI_M23": ["chk_bc_hang_thang", "chk_bc_03_thang", "chk_bc_06_thang",
                "chk_khu_cong_nghiep", "chk_khu_che_xuat", "chk_khu_kinh_te",
                "chk_khu_cong_nghe_cao"],
    "PLI_M25": ["chk_bc_hang_thang", "chk_bc_03_thang", "chk_bc_06_thang",
                "chk_khu_cong_nghiep", "chk_khu_che_xuat", "chk_khu_kinh_te",
                "chk_khu_cong_nghe_cao"],
}

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run(*args, **kw):
            try:
                fn(*args, **kw)
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:                        # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def src_paras(form_id: str) -> list[str]:
    """Text thô (chưa fold) các paragraph body của TEMPLATE NGUỒN."""
    doc = docx.Document(str(bfm.find_source(form_id)))
    return [p.text for p in doc.paragraphs]


def folds(form_id: str) -> list[str]:
    return [bfm.fold(re.sub(r"\s+", " ", t)).strip() for t in src_paras(form_id)]


def tpl_paras(form_id: str) -> list[str]:
    doc = docx.Document(str(ROOT / "app" / "templates" / ("%s.docx" % form_id)))
    return [Paragraph(el, doc).text for el in doc.element.body.iterchildren()
            if el.tag.endswith("}p")]


def tpl_tags(form_id: str) -> set[str]:
    return set(re.findall(r"\{([A-Za-z0-9_.]+)\}", "\n".join(tpl_paras(form_id))))


def blocks_json(form_id: str) -> dict:
    return json.loads((BLOCKS_DIR / ("%s.json" % form_id)).read_text(
        encoding="utf-8"))


def form_maps_keys(form_id: str) -> set[str]:
    fm = json.loads(FORM_MAPS.read_text(encoding="utf-8"))
    return {f["key"] for f in fm[form_id]["fields"]}


def manifest_sections() -> dict:
    data = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    return {f["form_id"]: {s["section_id"]: s for s in f["sections"]}
            for f in data["forms"]}


def sandbox_build(form_id: str, raw_spec: dict) -> tuple[bytes, list[str], dict]:
    """Rebuild THẬT trong sandbox tmp (OUT_TPL/OUT_DATA/FORM_SPECS_DIR ghi vào
    temp). Trả về (template bytes, paragraph texts, meta)."""
    tmp = Path(tempfile.mkdtemp(prefix="p2d_sandbox_"))
    (tmp / "t").mkdir(parents=True, exist_ok=True)
    (tmp / "d").mkdir(parents=True, exist_ok=True)
    gd = tmp / "change"
    gd.mkdir(parents=True, exist_ok=True)
    (gd / ("%s.json" % form_id)).write_text(
        json.dumps(raw_spec, ensure_ascii=False, indent=1), encoding="utf-8")
    saved = (bfm.OUT_TPL, bfm.OUT_DATA, bfm.FORM_SPECS_DIR)
    bfm.OUT_TPL, bfm.OUT_DATA, bfm.FORM_SPECS_DIR = tmp / "t", tmp / "d", tmp
    bfm._form_spec_cache.clear()
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            meta, _unmatched = bfm.process_form(
                bfm.find_source(form_id),
                {"Phu luc I|17": "T", "Phu luc I|19": "T",
                 "Phu luc I|23": "T", "Phu luc I|25": "T"}, None)
        tpl = tmp / "t" / ("%s.docx" % form_id)
        assert tpl.exists(), "generator chưa ghi template vào sandbox"
        raw = tpl.read_bytes()
        doc = docx.Document(str(tpl))
        paras = [Paragraph(el, doc).text
                 for el in doc.element.body.iterchildren()
                 if el.tag.endswith("}p")]
        return raw, paras, meta
    finally:
        (bfm.OUT_TPL, bfm.OUT_DATA, bfm.FORM_SPECS_DIR) = saved
        bfm._form_spec_cache.clear()
        shutil.rmtree(tmp, ignore_errors=True)


def load_spec_quiet(form_id: str):
    """load_form_spec mà không in gap_nghiepvu (giữ output suite đọc được)."""
    bfm._form_spec_cache.pop(form_id, None)
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            return bfm.load_form_spec(form_id)
    finally:
        pass


def spec_raw(form_id: str) -> dict:
    return json.loads((SPEC_DIR / ("%s.json" % form_id)).read_text(
        encoding="utf-8-sig"))


def walk_region(fid: str, reg_folds: list[str], start: str, end: str,
                pairs) -> list[str]:
    """Tiêu thụ các cặp (nhan, key) trong khoảng [start, end). Trả về list key
    đã tiêu thụ; raise AssertionError khi thiếu/mismatched."""
    si = ei = None
    for i, t in enumerate(reg_folds):
        if si is None and t.startswith(start):
            si = i
        elif si is not None and t.startswith(end):
            ei = i
            break
    assert si is not None, "%s: không tìm thấy marker đầu %r trên nguồn" % (
        fid, start)
    assert ei is not None, "%s: không tìm thấy marker cuối %r trên nguồn" % (
        fid, end)
    assert si < ei, "%s: khoảng [start,end) rỗng/đảo (%d,%d)" % (fid, si, ei)
    cur = si + 1
    got: list[str] = []
    for label, key in pairs:
        hit = -1
        for j in range(cur, ei):
            t = reg_folds[j]
            if label is None:
                ok = bool(BOX_RE.search(t))
            else:
                ok = label in t
            if ok:
                hit = j
                break
        assert hit >= 0, (
            "%s: trong [%s .. %s] không tìm thấy %r (còn phải khớp %r) — "
            "evidence nguồn lệch so với spec" % (fid, start[:40], end[:40],
                                                 label or "ô ☐/□", key))
        cur = hit + 1
        if key:
            got.append(key)
    return got


# ---------------------------------------------------------------------------
@check("[EVD] census manifest co evidence cho moi section P2D + defect M23")
def t_evidence_manifest():
    secs = manifest_sections()
    for fid, sids in P2D_SECTIONS.items():
        for sid in sids:
            s = secs[fid].get(sid)
            assert s is not None, "%s: thiếu section %r trong manifest" % (fid, sid)
            ev = s.get("source_evidence") or {}
            assert s.get("kind"), "%s/%s thiếu kind" % (fid, sid)
            assert s.get("semantic_key"), "%s/%s thiếu semantic_key" % (fid, sid)
            assert (ev.get("heading") or ev.get("anchor")), \
                "%s/%s thiếu evidence nguồn" % (fid, sid)
            assert "03_chi_nhanh_van_phong_dai_dien_dia_diem" in (ev.get(
                "anchor") or "") or (ev.get("template") or "").endswith(
                "%s.docx" % fid), "%s/%s evidence không trỏ nguồn" % (fid, sid)
    d23 = {d["id"]: d for d in
           next(f for f in json.loads(MANIFEST_PATH.read_text(
               encoding="utf-8"))["forms"]
                if f["form_id"] == "PLI_M23")["defects"]}
    assert "pli_m23_heading_numbering" in d23, \
        "manifest mất defect pli_m23_heading_numbering"
    assert d23["pli_m23_heading_numbering"]["queue"] == "P1-spec-cleanup"


@check("[SPEC] 4 file spec qua loader v1; KHONG entry kind moi; entry nganh giu nguyen")
def t_spec_validate():
    for fid in FORMS:
        raw = spec_raw(fid)
        assert set(raw) <= {"form_id", "group", "note", "gap_nghiepvu",
                            "tables", "section_cond"}, \
            "%s: root key lạ %s" % (fid, sorted(set(raw)))
        assert raw["tables"] and len(raw["tables"]) == 1, \
            "%s: kỳ vọng đúng 1 entry (ds_nganh), thấy %d" % (
                fid, len(raw["tables"]))
        ent = raw["tables"][0]
        assert "kind" not in ent, \
            "%s: KHÔNG được khai kind mới ở đây — xem [NO-REBUILD]" % fid
        ext = load_spec_quiet(fid)
        assert ext and len(ext["entries"]) == 1, "%s: loader nạp %d entry" % (
            fid, len(ext["entries"] or []))
        e = ext["entries"][0]
        assert e["spec"] == "ds_nganh" and e["anchor"], fid
        assert e.get("no_default_drops") is True, \
            "%s: cần no_default_drops=true để không nuốt vùng kế" % fid
        assert e.get("drop_until"), "%s: cần drop_until làm biên giới hạn" % fid
        # khối lượng entry ngành không bị lane này đụng tới
        sp = ext["specs_by_id"]["ds_nganh"]
        assert set(sp) >= {"id", "loop", "cols", "fields"}, \
            "%s: spec ds_nganh mất trường %s" % (
                fid, {"id", "loop", "cols", "fields"} - set(sp))
        assert sp["loop"] == "ds_nganh" and sp["kind"] == "list", fid
        assert [t for _p, t in sp["cols"] if t] == [
            "ten_nganh", "ma_nganh", "la_nganh_chinh"], \
            "%s: cột tag của ds_nganh bị đổi" % fid


@check("[TAGS] moi o □/☐ pham vi D5.1/D5.4 DA co token, dung nhan dung thu tu")
def t_tags_bound():
    for fid in FORMS:
        reg_folds = folds(fid)
        blk = [b for b in blocks_json(fid)["blocks"] if b.get("t") == "chk"]
        blk_keys = [b["key"] for b in blk]
        fm_keys = form_maps_keys(fid)
        tags = tpl_tags(fid)
        para_only = PARA_ONLY_KEYS.get(fid, set())

        want: list[str] = []
        for start, end, pairs in REGIONS[fid]:
            want.extend(walk_region(fid, reg_folds, start, end, pairs))
        assert len(want) == len(set(want)), "%s: key lặp trong kho evidence" % fid

        missing_fm = [k for k in want if k not in fm_keys]
        assert not missing_fm, "%s: form_maps thiếu %s" % (fid, missing_fm)
        missing_tpl = [k for k in want if k not in tags]
        assert not missing_tpl, \
            "%s: template không có {tag} cho %s" % (fid, missing_tpl)
        missing_blk = [k for k in want
                       if k not in para_only and k not in blk_keys]
        assert not missing_blk, "%s: blocks thiếu chk block %s" % (
            fid, missing_blk)
        # không có o □ nao trong pham vi bi mapping THUA (chk block mồ côi)
        orphans = [k for k in blk_keys if k not in want]
        assert not orphans, \
            "%s: chk block không thuộc evidence P2D nào (thừa/đổi key?): %s" % (
                fid, orphans)
        print("     %s: %d ô đã gắn thẻ (%d chk block + %d paragraph-level)"
              % (fid, len(want), len(blk_keys), len(want) - len(blk_keys)))
        # negative evidence: những nội dung mẫu KHÔNG có -> không được có key
        for frag in NEGATIVE[fid]:
            assert not any(frag in t for t in reg_folds), \
                "%s: nguồn đã xuất hiện %r — negative evidence hết hiệu lực" % (
                    fid, frag)
        bad = [k for k in NEGATIVE_KEYS[fid]
               if k in fm_keys or k in tags or k in blk_keys]
        assert not bad, "%s: mẫu không có dòng quy định nhưng lại map %s" % (
            fid, bad)


@check("[ANCHOR] anchor khop dung 1 paragraph; drop_until nam sau anchor; drop an toan")
def t_anchor_boundary():
    for fid in FORMS:
        ext = load_spec_quiet(fid)
        ent = ext["entries"][0]
        bf = [bfm.fold(re.sub(r"\s+", " ", "".join(
            str(r.get("text", "")) for r in (b.get("runs") or [])))).strip()
            for b in blocks_json(fid)["blocks"]]
        hits_b = [i for i, t in enumerate(bf) if re.search(ent["anchor"], t)]
        assert len(hits_b) == 1, "%s: anchor khớp blocks %d lần %r" % (
            fid, len(hits_b), hits_b)
        tf = folds(fid)
        hits_t = [i for i, t in enumerate(tf) if re.search(ent["anchor"], t)]
        assert len(hits_t) == 1, "%s: anchor khớp nguồn %d lần" % (fid, len(hits_t))
        ai = next(i for i, t in enumerate(tf) if re.search(ent["anchor"], t))
        ei = next((i for i, t in enumerate(tf)
                   if i > ai and re.search(ent["drop_until"], t)), None)
        assert ei is not None, \
            "%s: drop_until %r không xuất hiện sau anchor trên nguồn" % (
                fid, ent["drop_until"])
        assert ei > ai, "%s: drop_until đứng trước anchor" % fid
        # MỌI drop pattern phải khớp >=1 paragraph trong [anchor, drop_until)
        for pat in ent.get("drop") or []:
            inside = [t for t in tf[ai:ei] if re.match(pat, t)]
            assert inside, "%s: drop %r không khớp paragraph nào trong vùng" % (
                fid, pat)
        # boundary an toàn: paragraph tại drop_until là heading phần kế và PHẢI
        # sống sót sau rebuild (drop bị chặn trong khoảng, không nuốt vùng kế).
        # So prefix vì rebuild có thể gắn thêm {tag} vào đuôi cùng paragraph.
        boundary_text = next(t for i, t in enumerate(tf) if i == ei)
        prefix = boundary_text[:25]
        rebuilt = tpl_paras(fid)
        assert any(bfm.fold(re.sub(r"\s+", " ", t)).strip().startswith(prefix)
                   for t in rebuilt), \
            "%s: heading biên %r bị nuốt khi rebuild -> drop không bounded" % (
                fid, boundary_text[:50])


@check("[SRC] danh so theo template nguon; M23 la muc 7 chu khong phai 6")
def t_source_numbering():
    want = {
        "PLI_M17": (["6. thong tin dang ky thue"], []),
        "PLI_M19": (["dang ky thay doi thong tin dang ky thue"], []),
        "PLI_M23": (["7. thong tin dang ky thue"],
                    ["6. thong tin dang ky thue"]),
        "PLI_M25": (["6. thong tin dang ky thue"], []),
    }
    for fid, (must, must_not) in want.items():
        tf = folds(fid)
        for m in must:
            assert any(t.startswith(m) for t in tf), \
                "%s: nguồn thiếu heading %r" % (fid, m)
        for m in must_not:
            assert not any(t.startswith(m) for t in tf), \
                "%s: nguồn VẪN còn heading %r — evidence số đã đổi" % (fid, m)
    # ngưỡng ngành 5a/5b của M23 đúng theo nguồn (anchor/drop_until hiện có)
    m23 = folds("PLI_M23")
    assert any(t.startswith("5. nganh, nghe kinh doanh/noi dung hoat dong")
               for t in m23), "M23: mất mục 5"
    assert any(t.startswith("a) nganh, nghe kinh doanh") for t in m23), \
        "M23: mất mục 5a (anchor của entry)"
    assert any(t.startswith("b) noi dung hoat dong") for t in m23), \
        "M23: mất mục 5b (drop_until của entry)"


@check("[D5.3] chi M17 co BHXH + Luu y nong-lam-ngu-diem; slot cross_field_rules nap duoc")
def t_d53_cross_field():
    # M17: dòng quy định thật trên nguồn
    m17 = folds("PLI_M17")
    assert any("phuong thuc dong bao hiem xa hoi" in t
               and "1 trong 3" in t for t in m17), "M17: mất dòng phương thức BHXH"
    assert any(t.startswith("luu y:") for t in m17), "M17: mất 'Lưu ý:'"
    assert any("nong nghiep, lam nghiep, ngu nghiep, diem nghiep" in t
               and "tra luong theo san pham" in t for t in m17), \
        "M17: mất gạch đầu dòng ngành chính nông/lâm/ngư/diên + trả lương theo sản phẩm"
    assert any("danh dau vao phuong thuc dong bao hiem xa hoi hang thang" in t
               for t in m17), "M17: mất gạch đầu dòng ngành khác -> hàng tháng"
    # M19/M23/M25: KHÔNG có phần BHXH -> không khai rule (negative evidence)
    for fid in ("PLI_M19", "PLI_M23", "PLI_M25"):
        assert not any("bao hiem xa hoi" in t for t in folds(fid)), \
            "%s: nguồn đã xuất hiện phần BHXH — phải review lại D5.3 cho mẫu này" % fid
        raw = spec_raw(fid)
        assert not any(e.get("cross_field_rules") for e in raw["tables"]), \
            "%s: khai cross_field_rules trong khi mẫu không có dòng quy định" % fid
    # rule khai ở M17
    raw = spec_raw("PLI_M17")
    rules = raw["tables"][0].get("cross_field_rules") or []
    assert len(rules) == 1, "M17: kỳ vọng đúng 1 rule D5.3, thấy %d" % len(rules)
    r = rules[0]
    for f in ("rule_id", "section_id", "semantic_key", "type", "status",
              "exclusive_group", "allow", "otherwise_allow", "source",
              "violation_code", "severity", "when"):
        assert f in r, "M17: rule D5.3 thiếu trường %r" % f
    assert r["section_id"] == "phuong_thuc_dong_bhxh"
    assert r["semantic_key"] == "phuong_thuc_dong_bhxh_cn"
    assert r["exclusive_group"] == ["chk_bc_hang_thang", "chk_bc_03_thang",
                                    "chk_bc_06_thang"], \
        "M17: exclusive_group phải là 3 phương thức theo đúng thứ tự nguồn"
    assert r["otherwise_allow"] == ["chk_bc_hang_thang"]
    assert r["when"]["status"] == "not_evaluable", \
        "M17: điều kiện ngành chính/trả lương theo sản phẩm không có input -> phải not_evaluable"
    # nạp được qua loader
    ext = load_spec_quiet("PLI_M17")
    assert ext["entries"][0].get("cross_field_rules") == rules, \
        "M17: loader không giữ được slot cross_field_rules"


@check("[D5.4] M17 chi tieu 8 CTCK nuoc ngoai la muc rieng ngoai bang thue")
def t_d54_ctck():
    tf = folds("PLI_M17")
    i8 = next(i for i, t in enumerate(tf)
              if t.startswith("8. chi nhanh cong ty chung khoan nuoc ngoai"))
    i7 = next(i for i, t in enumerate(tf)
              if t.startswith("7. thong tin ve viec dong bao hiem xa hoi"))
    i_cam = next(i for i, t in enumerate(tf) if t.startswith("doanh nghiep cam ket"))
    assert i7 < i8 < i_cam, "M17: thứ tự mục 7 -> 8 -> cam kết bị lệch"
    # {chk_ctck} nằm ở PARAGRAPH, không nằm trong bất kỳ w:tbl nào
    doc = docx.Document(str(ROOT / "app" / "templates" / "PLI_M17.docx"))
    in_tbl = set()
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                in_tbl |= set(re.findall(r"\{([A-Za-z0-9_.]+)\}", c.text))
    assert "chk_ctck" not in in_tbl, \
        "M17: {chk_ctck} đã rơi vào bảng — D5.4 bắt buộc mục riêng ngoài bảng thuế"
    paras = tpl_paras("PLI_M17")
    hits = [i for i, t in enumerate(paras) if "{chk_ctck}" in t]
    assert len(hits) == 1, "M17: {chk_ctck} xuất hiện %d lần" % len(hits)
    before = [t for t in paras[:hits[0]] if "chứng khoán nước ngoài" in t]
    assert before, "M17: {chk_ctck} không đứng sau heading mục 8 CTCK"
    # nằm ngoài vùng bảng thuế (mục 6)
    i6 = next(i for i, t in enumerate(paras)
              if t.strip().startswith("6. Thông tin đăng ký thuế"))
    assert hits[0] > i6, "M17: {chk_ctck} rơi vào trước/trong vùng bảng thuế"
    assert "chk_ctck" in form_maps_keys("PLI_M17"), \
        "M17: chk_ctck chưa được đăng ký trong form_maps"


@check("[SLOT] cross_field_rules la slot INERT (template byte-identical khi bo slot)")
def t_slot_inert():
    base = spec_raw("PLI_M17")
    stripped = copy.deepcopy(base)
    stripped["tables"][0].pop("cross_field_rules", None)
    r_raw, r_paras, r_meta = sandbox_build("PLI_M17", base)
    s_raw, s_paras, s_meta = sandbox_build("PLI_M17", stripped)
    assert r_raw == s_raw, "M17: thêm cross_field_rules làm ĐỔI template!"
    assert r_paras == s_paras, "M17: slot cross_field_rules đổi paragraph"
    assert (r_meta.get("section_specs") is None) == (
        s_meta.get("section_specs") is None), "M17: slot đổi section_specs"
    assert not (r_meta.get("section_specs") or []), \
        "M17: spec thật không được sinh section_specs (xem [NO-REBUILD])"
    # và slot này là duy nhất: khai ở root hoặc trong spec kind mới đều bị chặn
    bad_root = copy.deepcopy(base)
    bad_root["cross_field_rules"] = []
    try:
        sandbox_build("PLI_M17", bad_root)
    except ValueError as exc:
        assert "truong root khong duoc phep" in str(exc), exc
    else:
        raise AssertionError("root cross_field_rules phải bị validator chặn")
    bad_fm = copy.deepcopy(base)
    bad_fm["tables"].append({
        "kind": "fixed-matrix", "pos": "after", "anchor": "^7\\. thong tin",
        "spec": {"id": "demo", "kind": "fixed-matrix", "semantic_key": "demo_k",
                 "columns": [{"header": "c", "tag": "t", "input_type": "checkbox"}],
                 "fixed_rows": [{"id": "r1", "label": "L"}], "totals": [],
                 "cross_field_rules": []}})
    try:
        sandbox_build("PLI_M17", bad_fm)
    except ValueError as exc:
        assert "khong thuoc hop dong v1" in str(exc), exc
    else:
        raise AssertionError("cross_field_rules trong spec fixed-matrix phải bị chặn")


@check("[NO-REBUILD] fixed-matrix naive se double-bind; spec that khong sinh key song song")
def t_no_rebuild():
    naive = spec_raw("PLI_M17")
    naive["tables"].append({
        "kind": "fixed-matrix", "pos": "after",
        "anchor": "^7\\. thong tin ve viec dong bao hiem xa hoi",
        "no_default_drops": True, "drop_until": "^luu y:",
        "spec": {"id": "bhxh_phuong_thuc_dong", "kind": "fixed-matrix",
                 "semantic_key": "phuong_thuc_dong_bhxh_cn",
                 "section_id": "phuong_thuc_dong_bhxh",
                 "columns": [{"header": "Phương thức đóng BHXH", "tag": "chon",
                              "input_type": "checkbox"}],
                 "fixed_rows": [{"id": "hang_thang", "label": "Hàng tháng"},
                                {"id": "ba_thang_mot_lan", "label": "03 tháng một lần"},
                                {"id": "sau_thang_mot_lan", "label": "06 tháng một lần"}],
                 "totals": []}})
    _raw, _paras, meta = sandbox_build("PLI_M17", naive)
    keys = {f["key"] for f in meta["fields"]}
    parallel = {k for k in keys if k.startswith("phuong_thuc_dong_bhxh_cn.")}
    assert len(parallel) == 3, \
        "sandbox mong 3 key song song để chứng minh double-bind, thấy %r" % parallel
    legacy = {k for k in keys if k in ("chk_bc_hang_thang", "chk_bc_03_thang",
                                       "chk_bc_06_thang")}
    assert legacy == {"chk_bc_hang_thang", "chk_bc_03_thang", "chk_bc_06_thang"}, \
        "bản naive phải STILL đăng ký chk_bc_* (đó chính là lý do không được khai)"
    # spec thật: KHÔNG key song song nào
    SEM = ("phuong_thuc_dong_bhxh_cn", "chi_tieu_thue_cn_ddkd",
           "chi_tieu_thue_thay_doi_cn", "chi_tieu_thue_bo_sung_cn",
           "chi_tieu_thue_bo_sung_ctck", "lua_chon_khu_san_xuat",
           "chi_tieu_thue_dntn")
    for fid in FORMS:
        _r, _p, good = sandbox_build(fid, spec_raw(fid))
        gk = {f["key"] for f in good["fields"]}
        parallel = sorted(k for k in gk if "." in k and k.split(".")[0] in SEM)
        assert not parallel, \
            "%s: spec thật sinh key '<semantic_key>.<row>.<tag>' song song: %s" % (
                fid, parallel[:4])
        assert not (good.get("section_specs") or []), \
            "%s: spec thật sinh section_specs (rebuild) — cấm" % fid
        # không mất/thêm key cũ: so với form_maps trong repo
        assert gk == form_maps_keys(fid), \
            "%s: bộ key form_maps sandbox lệch với repo (regen sẽ đụng app/data)" % fid


@check("[GUARD] app/templates va app/data khong bi ghi trong suot suite")
def t_guard_repo_untouched():
    after = _fingerprint()
    assert len(after) > 100, "bảng fingerprint quá nhỏ (%d)" % len(after)
    changed = [k for k in after
               if REPO_FP_BEFORE.get(k) != after[k]]
    added = [k for k in after if k not in REPO_FP_BEFORE]
    removed = [k for k in REPO_FP_BEFORE if k not in after]
    assert not (changed or added or removed), \
        "suite đã ghi ngoài sandbox: changed=%s added=%s removed=%s" % (
            changed[:3], added[:3], removed[:3])


def main() -> int:
    for fn in (t_evidence_manifest, t_spec_validate, t_tags_bound,
               t_anchor_boundary, t_source_numbering, t_d53_cross_field,
               t_d54_ctck, t_slot_inert, t_no_rebuild, t_guard_repo_untouched):
        fn()
    ok = sum(1 for p, _ in _results if p)
    fail = sum(1 for p, _ in _results if not p)
    print("-" * 72)
    print("test_p2d_tax_specs: %d PASS / %d FAIL" % (ok, fail))
    for passed, name in _results:
        if not passed:
            print("  -> %s" % name)
    return 1 if fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
