#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test_phase3_m04_m11.py — Lane P3-C (phase3-p4-p11), python3 thuần.

Mệnh đề kiểm chứng:

  [M11-spec-file-parity] tools/form_specs/formation/PLI_M11.json là migration
            VERBATIM của khai báo inline (DYNAMIC_TABLES 'ds_to_chuc_cshhl' +
            REBUILD_TABLES['PLI_M11']) — file-spec JSON ưu tiên hơn inline qua
            load_form_spec. Sandbox hoá generator GIỐNG pattern
            tools/test_cap_von_specs.py (OUT_TPL/OUT_DATA ghi vào tmp, clear
            _form_spec_cache, restore trong finally) rồi assert:
            (a) load_form_spec('PLI_M11') trả entries TỪ FILE (log
                "[form_specs] ... dung file JSON" + specs_by_id có spec);
            (b) form_maps sandbox PLI_M11: dynamic_tables id
                ds_to_chuc_cshhl loop ds_cshhl, đủ 8 cột (col 0..7), 7 cột tag
                — KHỚP CẤU TRÚC entry baseline app/data/form_maps.json
                (đọc-only);
            (c) template sandbox PLI_M11.docx: đúng số bảng, bảng rebuild 8
                cột header lần lượt khớp (python-docx);
            (d) block input key sinh ra (loop ds_cshhl) giữ nguyên bộ tag
                (so với baseline app/data/blocks/PLI_M11.json).
  [M04-menh-gia-scalar] input scalar 'menh_gia_co_phan' (mệnh giá in như MỘT
            DÒNG RIÊNG theo TT 121/2026 Mẫu 4, không phải cột ma trận) nằm
            trong vùng anchor '7. Thông tin về cổ phần:' của
            app/data/blocks/PLI_M04.json (dùng helpers pure của
            tools.census.table_coverage_audit).
  [M04-chao-ban-shape] spec PLI_M04 section co_phan_duoc_quyen_chao_ban:
            fixed_rows 6, dòng cuối tong_so output_only; cột 0 readonly
            tag null; cột 1 tag so_luong. Runtime contract app_web/ui/
            form-domain.js: cột readonly đang `tag: null` (EVIDENCE cho Lead
            wire tag 'loai_co_phan' sau) + `stt: { mode: 'none' }`.
  [M04-co-cau-stt-generated] contract sectionId 'co_cau_loai_co_phan' có
            `stt: { mode: 'generated' }` (số thứ tự tự sinh theo thứ tự
            export).

Baseline production chỉ được ĐỌC — sandbox không ghi đè app/templates,
app/data. Chạy: python3 -X utf8 tools/test_phase3_m04_m11.py
Exit 0 = toàn PASS; exit 1 = có FAIL.
"""

from __future__ import annotations

import contextlib
import io
import json
import shutil
import sys
import tempfile
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import tools.build_form_maps as bfm                                  # noqa: E402
import tools.docx_to_blocks as dxtb                                  # noqa: E402
from tools.census.table_coverage_audit import (                      # noqa: E402
    anchor_pos_in,
    block_all_text,
    flatten_blocks,
    region_keys_scan,
)

ROOT = Path(__file__).resolve().parent.parent
BASELINE_FORM_MAPS = ROOT / "app" / "data" / "form_maps.json"
BASELINE_TEMPLATE_M11 = ROOT / "app" / "templates" / "PLI_M11.docx"
BASELINE_BLOCKS_M11 = ROOT / "app" / "data" / "blocks" / "PLI_M11.json"
BLOCKS_M04 = ROOT / "app" / "data" / "blocks" / "PLI_M04.json"
SPEC_M04 = ROOT / "tools" / "form_specs" / "formation" / "PLI_M04.json"
SPEC_M11 = ROOT / "tools" / "form_specs" / "formation" / "PLI_M11.json"
FORM_DOMAIN_JS = ROOT / "app_web" / "ui" / "form-domain.js"

# Khai báo inline nguyên bản (DYNAMIC_TABLES 'ds_to_chuc_cshhl' +
# REBUILD_TABLES['PLI_M11'] trong tools/build_form_maps.py @ BASE d41c42f) —
# chuẩn so khớp verbatim của file spec JSON.
INLINE_COLS = [
    ("^stt$", None),
    ("ten to chuc", "ten_to_chuc"),
    ("ma so doanh nghiep|ma so thue", "ma_so_dn"),
    ("^ngay cap", "ngay_cap"),
    ("^noi cap", "noi_cap"),
    ("dia chi", "dia_chi"),
    ("ty le", "ty_le"),
    ("ghi chu", "ghi_chu"),
]
INLINE_DETECT = {"stt", "ten to chuc", "ngay cap"}
INLINE_LOOP = "ds_cshhl"
INLINE_FIELDS = ["ten_to_chuc", "ma_so_dn", "ngay_cap", "noi_cap", "dia_chi",
                 "ty_le"]
INLINE_HEADERS = ["STT", "Tên tổ chức", "Mã số doanh nghiệp/số QĐ thành lập",
                  "Ngày cấp", "Nơi cấp", "Địa chỉ trụ sở chính",
                  "Tỷ lệ sở hữu cổ phần có quyền biểu quyết (%)", "Ghi chú"]
INLINE_TAGS = ["ten_to_chuc", "ma_so_dn", "ngay_cap", "noi_cap", "dia_chi",
               "ty_le", "ghi_chu"]

RESULTS: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run(*args, **kw):
            try:
                fn(*args, **kw)
                RESULTS.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                RESULTS.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:                          # noqa: BLE001
                RESULTS.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


# ---------------------------------------------------------------------------
# Sandbox generator — tái dùng cơ chế tools/test_cap_von_specs.py
# (new_sandbox/run_generation): redirect OUT_TPL/OUT_DATA vào tmp, clear
# _form_spec_cache, restore trong finally. Baseline production chỉ ĐỌC.
# ---------------------------------------------------------------------------
def new_sandbox() -> tuple[Path, Path]:
    tmp = Path(tempfile.mkdtemp(prefix="phase3_laneC_"))
    tpl_dir = tmp / "templates"
    data_dir = tmp / "data"
    tpl_dir.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)
    return tpl_dir, data_dir


def run_generation_m11(tpl_dir: Path) -> dict:
    bfm._form_spec_cache.clear()
    src = bfm.find_source("PLI_M11")
    meta, _unmatched = bfm.process_form(
        src, {"Phu luc I|11": "DEMO PLI_M11"}, None)
    assert (tpl_dir / "PLI_M11.docx").exists(), \
        "generator chưa ghi template sandbox"
    return meta


def baseline_dyn_entry_m11() -> dict:
    fm = json.loads(BASELINE_FORM_MAPS.read_text(encoding="utf-8"))
    entries = [e for e in fm["PLI_M11"]["dynamic_tables"]
               if e["id"] == "ds_to_chuc_cshhl"]
    assert len(entries) == 1, \
        "baseline form_maps phải có đúng 1 entry ds_to_chuc_cshhl"
    return entries[0]


def baseline_fields_keys_m11(prefix: str) -> list[str]:
    fm = json.loads(BASELINE_FORM_MAPS.read_text(encoding="utf-8"))
    return [f["key"] for f in fm["PLI_M11"]["fields"]
            if f["key"].startswith(prefix)]


def baseline_loop_block_m11() -> dict:
    doc = json.loads(BASELINE_BLOCKS_M11.read_text(encoding="utf-8"))
    loops = [b for b in doc["blocks"]
             if b.get("t") == "loop" and b.get("loop") == INLINE_LOOP]
    assert len(loops) == 1, \
        "baseline blocks PLI_M11 phải có đúng 1 block loop ds_cshhl"
    return loops[0]


def spec_entry_m11() -> dict:
    raw = json.loads(SPEC_M11.read_text(encoding="utf-8-sig"))
    tables = [e for e in raw["tables"] if e["spec"]["id"] == "ds_to_chuc_cshhl"]
    assert len(tables) == 1, "file spec phải có đúng 1 entry ds_to_chuc_cshhl"
    return tables[0]


def js_section_slice(src: str, section_id: str) -> str:
    key = "sectionId: '%s'" % section_id
    i = src.index(key)
    j = src.find("sectionId:", i + len(key))
    return src[i:j if j != -1 else len(src)]


# ---------------------------------------------------------------------------
# 1. [M11-spec-file-parity]
# ---------------------------------------------------------------------------
@check("[M11-spec-file-parity] file spec duoc load uu tien, verbatim inline")
def t_m11_spec_file_parity():
    # (a) load_form_spec trả entries TỪ FILE (không phải fallback inline)
    bfm._form_spec_cache.clear()
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            ext = bfm.load_form_spec("PLI_M11")
        log = buf.getvalue()
    finally:
        bfm._form_spec_cache.clear()
    assert ext is not None, "load_form_spec('PLI_M11') trả None (file spec " \
                            "không được nhận diện)"
    assert "ds_to_chuc_cshhl" in ext["specs_by_id"], \
        "specs_by_id thiếu ds_to_chuc_cshhl"
    assert "[form_specs]" in log and "dung file JSON" in log, \
        "log không chứng minh dùng file JSON: %r" % log[:200]

    spec = ext["specs_by_id"]["ds_to_chuc_cshhl"]
    entry = ext["entries"][0]
    # verbatim so với inline
    assert spec["detect"] == INLINE_DETECT, \
        "detect %r != inline %r" % (sorted(spec["detect"]), sorted(INLINE_DETECT))
    assert [(p, t) for p, t in spec["cols"]] == INLINE_COLS, \
        "cols lệch inline: %r" % (spec["cols"],)
    assert spec["loop"] == INLINE_LOOP and spec["fields"] == INLINE_FIELDS, \
        "loop/fields lệch inline"
    assert not spec.get("stt_tag"), \
        "inline không có stt_tag — file spec cũng không được khai"
    assert spec.get("kind", "list") == "list", "kind phải là list (mặc định)"
    assert entry["pos"] == "after" and \
        entry["anchor"] == "danh cho cong ty co phan ke khai", \
        "pos/anchor lệch inline: %r" % (entry,)
    assert entry["headers"] == INLINE_HEADERS, \
        "headers lệch inline: %r" % (entry["headers"],)

    # sandbox generation với file spec + so baseline (đọc-only)
    tpl_dir, _data = new_sandbox()
    saved_tpl, saved_data = bfm.OUT_TPL, bfm.OUT_DATA
    bfm.OUT_TPL, bfm.OUT_DATA = tpl_dir, _data
    try:
        meta = run_generation_m11(tpl_dir)

        # (b) form_maps sandbox: entry ds_to_chuc_cshhl khớp cấu trúc baseline
        dyn = [e for e in meta.get("dynamic_tables") or []
               if e["id"] == "ds_to_chuc_cshhl"]
        assert len(dyn) == 1, "sandbox thiếu dynamic_tables ds_to_chuc_cshhl"
        got = dyn[0]
        assert got["loop"] == INLINE_LOOP, "loop %r != ds_cshhl" % got["loop"]
        assert [c["tag"] for c in got["columns"]] == INLINE_TAGS, \
            "7 cột tag lệch: %r" % (got["columns"],)
        assert [c["col"] for c in got["columns"]] == list(range(1, 8)), \
            "col index phải 1..7 (col 0 STT tag null không đăng ký)"
        assert got == baseline_dyn_entry_m11(), \
            "entry sandbox lệch baseline app/data/form_maps.json:\n got=%r\n" \
            " base=%r" % (got, baseline_dyn_entry_m11())
        raw = spec_entry_m11()
        assert len(raw["spec"]["cols"]) == 8, "spec phải đủ 8 cột (col 0..7)"
        assert raw["spec"]["cols"][0] == ["^stt$", None], \
            "cột 0 phải là STT tag null"

        # (c) template sandbox: đúng số bảng, bảng rebuild 8 cột + headers
        doc = docx.Document(str(tpl_dir / "PLI_M11.docx"))
        base_doc = docx.Document(str(BASELINE_TEMPLATE_M11))
        assert len(doc.tables) == len(base_doc.tables), \
            "số bảng sandbox %d != baseline %d" % (len(doc.tables),
                                                   len(base_doc.tables))
        assert len(doc.tables) >= 1, "sandbox PLI_M11 không có bảng nào"
        tbl = doc.tables[0]
        assert len(tbl.columns) == 8, \
            "bảng rebuild phải 8 cột (col 0..7), thấy %d" % len(tbl.columns)
        hdr = [c.text.strip() for c in tbl.rows[0].cells]
        assert hdr == INLINE_HEADERS, "header lệch: %r" % (hdr,)
        base_hdr = [c.text.strip() for c in base_doc.tables[0].rows[0].cells]
        assert hdr == base_hdr, "header sandbox lệch baseline template"

        # (d) block input key sinh ra (loop ds_cshhl) giữ nguyên bộ tag
        blocks = dxtb.convert_document(
            docx.Document(str(tpl_dir / "PLI_M11.docx")), "PLI_M11",
            "DEMO PLI_M11", overrides={}, footnotes={}, section_specs=[])
        loops = [b for b in blocks["blocks"]
                 if b.get("t") == "loop" and b.get("loop") == INLINE_LOOP]
        assert len(loops) == 1, "blocks sandbox thiếu block loop ds_cshhl"
        got_cols = [{"col": c["col"], "tag": c["tag"]} for c in loops[0]["cols"]]
        assert got_cols == baseline_loop_block_m11()["cols"], \
            "cols block loop lệch baseline blocks/PLI_M11.json: %r" % (got_cols,)
        assert [c["tag"] for c in got_cols] == INLINE_TAGS, \
            "bộ tag block loop lệch: %r" % (got_cols,)
        got_fields = [f["key"] for f in meta["fields"]
                      if f["key"].startswith(INLINE_LOOP + ".")]
        assert got_fields == baseline_fields_keys_m11(INLINE_LOOP + "."), \
            "fields ds_cshhl.* lệch baseline form_maps: %r" % (got_fields,)
        assert got_fields == ["%s.%s" % (INLINE_LOOP, t) for t in
                              INLINE_FIELDS], \
            "fields phải đúng 6 tag của spec fields (không thêm ghi_chu): %r" \
            % (got_fields,)
    finally:
        bfm.OUT_TPL, bfm.OUT_DATA = saved_tpl, saved_data
        bfm._form_spec_cache.clear()
        shutil.rmtree(tpl_dir.parent, ignore_errors=True)


# ---------------------------------------------------------------------------
# 2. [M04-menh-gia-scalar]
# ---------------------------------------------------------------------------
@check("[M04-menh-gia-scalar] menh_gia_co_phan la input dong rieng trong muc 7")
def t_m04_menh_gia_scalar():
    doc = json.loads(BLOCKS_M04.read_text(encoding="utf-8"))
    flat = flatten_blocks(doc)
    texts = [block_all_text(b) for b in flat]
    start = anchor_pos_in(texts, "7. Thông tin về cổ phần:")
    end = anchor_pos_in(texts, "Thông tin về cổ phần được quyền chào bán")
    assert start is not None, "không tìm thấy anchor '7. Thông tin về cổ phần:'"
    assert end is not None, \
        "không tìm thấy anchor kế 'Thông tin về cổ phần được quyền chào bán'"
    assert start < end, "vùng mục 7 lệch: start %d >= end %d" % (start, end)
    inputs, _chks = region_keys_scan(flat, start, end)
    assert "menh_gia_co_phan" in inputs, \
        "input key 'menh_gia_co_phan' (exact) không nằm trong vùng mục 7: %r" \
        % sorted(k for k in inputs if "menh" in k)
    # scalar: key không chấm — KHÔNG phải cột ma trận <sem>.<row>.<tag>
    assert "." not in "menh_gia_co_phan"
    assert not [k for k in inputs if k.endswith(".menh_gia_co_phan")], \
        "mệnh giá không được là cột của ma trận"
    # nhãn dòng trong nguồn phải nằm trong vùng
    assert any("menh gia co phan" in t for t in texts[start:end]), \
        "vùng mục 7 không chứa nhãn 'Mệnh giá cổ phần'"


# ---------------------------------------------------------------------------
# 3. [M04-chao-ban-shape]
# ---------------------------------------------------------------------------
@check("[M04-chao-ban-shape] spec + contract co_phan_duoc_quyen_chao_ban")
def t_m04_chao_ban_shape():
    raw = json.loads(SPEC_M04.read_text(encoding="utf-8-sig"))
    entries = [e for e in raw["tables"]
               if e["spec"]["id"] == "co_phan_duoc_quyen_chao_ban"]
    assert len(entries) == 1, "PLI_M04.json thiếu section co_phan_duoc_quyen_chao_ban"
    spec = entries[0]["spec"]
    rows = spec["fixed_rows"]
    assert len(rows) == 6, "fixed_rows phải 6 phần tử, thấy %d" % len(rows)
    assert rows[-1]["id"] == "tong_so" and rows[-1]["output_only"] is True, \
        "dòng cuối phải tong_so output_only=true: %r" % (rows[-1],)
    col0, col1 = spec["columns"][0], spec["columns"][1]
    assert col0["header"] == "Loại cổ phần được quyền chào bán", \
        "header cột 0 lệch: %r" % col0["header"]
    assert col0["tag"] is None and col0["input_type"] == "readonly", \
        "cột 0 phải tag null + readonly: %r" % (col0,)
    assert col1["tag"] == "so_luong", "cột 1 phải tag so_luong: %r" % (col1,)

    # runtime contract (form-domain.js) — evidence cho Lead wire tag sau
    src = FORM_DOMAIN_JS.read_text(encoding="utf-8")
    sl = js_section_slice(src, "co_phan_duoc_quyen_chao_ban")
    assert "tag: null" in sl, \
        "contract đang không có cột tag: null (Lead đã wire?) — cần cập nhật " \
        "evidence khi Lead wire tag 'loai_co_phan'"
    assert "inputType: 'readonly'" in sl, "cột readonly thiếu inputType readonly"
    assert "Loại cổ phần được quyền chào bán" in sl, "thiếu header cột 0"
    assert "stt: { mode: 'none' }" in sl, \
        "contract phải khai stt: { mode: 'none' }"
    assert "tag: 'so_luong'" in sl, "thiếu cột tag 'so_luong'"


# ---------------------------------------------------------------------------
# 4. [M04-co-cau-stt-generated]
# ---------------------------------------------------------------------------
@check("[M04-co-cau-stt-generated] contract co_cau_loai_co_phan stt generated")
def t_m04_co_cau_stt_generated():
    src = FORM_DOMAIN_JS.read_text(encoding="utf-8")
    sl = js_section_slice(src, "co_cau_loai_co_phan")
    assert "stt: { mode: 'generated' }" in sl, \
        "contract co_cau_loai_co_phan phải khai stt: { mode: 'generated' } " \
        "(số thứ tự tự sinh theo thứ tự export)"
    # phòng chống lệch section: slice này phải là fixed-matrix của PLI_M04
    assert "section_id: 'co_cau_loai_co_phan'" in sl and "PLI_M04" in sl, \
        "slice không phải contract của co_cau_loai_co_phan/PLI_M04"


def main() -> int:
    for fn in (t_m11_spec_file_parity, t_m04_menh_gia_scalar,
               t_m04_chao_ban_shape, t_m04_co_cau_stt_generated):
        fn()
    ok = sum(1 for p, _ in RESULTS if p)
    fail = sum(1 for p, _ in RESULTS if not p)
    print("-" * 72)
    print("test_phase3_m04_m11: %d PASS / %d FAIL" % (ok, fail))
    for passed, name in RESULTS:
        if not passed:
            print("  -> %s" % name)
    return 1 if fail else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
