# -*- coding: utf-8 -*-
"""test_phase5_ui_fixes.py — regression cho các fix bug owner Phase 5
(kiểm bằng UI thật trên app pywebview, evidence trong báo cáo Phase 5):

[F1] Bridge.read_data_json: allowlist cứng {'nganh_vsic.json'}, trả đủ
     794 ngành chính thức; từ chối path lạ và file không cho phép.
[F2] app.js load VSIC Bridge-first: read_data_json('nganh_vsic.json') đứng
     TRƯỚC fetch fallback; console.warn khi danh mục rút gọn (ngăn âm thầm
     dùng mock/vsic_mini.json — 15/794 mã, 9 tên sai, thiếu 455 ngành cấp 4).
[F3] style.css: .loop-fld có underline (#9db3cc) — ô editable bảng lặp phải
     nhìn thấy được; .loop-fld[readonly] nền xám + underline transparent.

Chạy:  python3 -X utf8 tools/test_phase5_ui_fixes.py
Exit 0 = PASS; exit 1 = FAIL.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

APP_JS = ROOT / "app_web" / "ui" / "app.js"
MAIN_PY = ROOT / "app_web" / "main.py"
STYLE = ROOT / "app_web" / "ui" / "style.css"

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run():
            try:
                fn()
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


@check("[F1] Bridge.read_data_json allowlist + trả đủ danh mục chính thức")
def f1_bridge_read_data_json():
    from app_web.main import Bridge
    b = Bridge()
    r = b.read_data_json("nganh_vsic.json")
    assert r.get("ok") is True, "ok=False: %r" % r.get("error")
    nganh = r["data"]["nganh"]
    assert len(nganh) >= 790, "danh mục phải là bản chính thức (794), thấy %d" % len(nganh)
    by_ma = {x["ma"]: x["ten"] for x in nganh}
    assert by_ma.get("0112") == "Trồng ngô và cây lương thực có hạt khác", \
        "0112 phải tên chính thức, thấy %r" % by_ma.get("0112")
    assert by_ma.get("0114") == "Trồng cây mía", \
        "0114 phải tên chính thức, thấy %r" % by_ma.get("0114")
    # từ chối path lạ + file không cho phép
    for bad in ("../secrets.json", "use_cases.json", "blocks/PLI_M01.json",
                "", "nganh_vsic.json\x00.js"):
        r2 = b.read_data_json(bad)
        assert r2.get("ok") is False and r2.get("error") == "FILE_NOT_ALLOWED", \
            "phải FILE_NOT_ALLOWED cho %r, thấy %r" % (bad, r2)


@check("[F2] app.js: VSIC Bridge-first trước fetch, warn danh mục rút gọn")
def f2_app_js_bridge_first():
    src = APP_JS.read_text(encoding="utf-8")
    ib = src.find("read_data_json('nganh_vsic.json')")
    assert ib > 0, "app.js phải gọi Bridge read_data_json('nganh_vsic.json')"
    ihttp = src.find("function vsicFromHttp")
    assert ihttp > ib, "vsicFromHttp (fallback) phải đứng SAU đường Bridge"
    assert "vsicFromBridge" in src and "vsicFromHttp" in src, "thiếu 2 loader"
    # Bridge-first có retry sau pywebviewready
    ipy = src.find("loadVsicList()")
    assert ipy > 0 and "pywebviewready" in src[ipy:ipy + 400], \
        "phải nạp lại qua Bridge sau pywebviewready"
    # warn khi danh mục rút gọn (chống fallback âm thầm)
    assert "RÚT GỌN" in src and "console.warn" in src, \
        "phải console.warn khi state.vsic rút gọn"


@check("[F3] style.css: .loop-fld có underline, [readonly] không underline")
def f3_css_loop_fld_visible():
    css = STYLE.read_text(encoding="utf-8")
    i = css.find(".loop-fld {")
    assert i > 0, "thiếu .loop-fld rule"
    block = css[i:css.find("}", i)]
    assert "border-bottom: 1px solid" in block, \
        ".loop-fld phải có underline nhận diện được: %r" % block
    ir = css.find(".loop-fld[readonly]")
    assert ir > 0, "thiếu .loop-fld[readonly] rule"
    rblock = css[ir:css.find("}", ir)]
    assert "border-bottom-color: transparent" in rblock, \
        "[readonly] phải bỏ underline để tương phản với ô editable: %r" % rblock
    assert "background: #f2f2f2" in rblock, "[readonly] giữ nền xám"


@check("[F4] PLI_M01 ô 'Tôi là' có input chu_hkd.ho_ten xuyên template/blocks/form_maps")
def f4_m01_owner_name_field():
    import docx
    d = docx.Document(str(ROOT / "app" / "templates" / "PLI_M01.docx"))
    for p in d.paragraphs:
        if "Tôi là" in p.text:
            assert "{chu_hkd.ho_ten}" in p._p.xml, \
                "template M01 dòng 'Tôi là' phải có {chu_hkd.ho_ten}, thấy %r" % p.text[:80]
            break
    else:
        raise AssertionError("không tìm thấy dòng 'Tôi là' trong template M01")
    blocks = json.loads((ROOT / "app" / "data" / "blocks" / "PLI_M01.json")
                        .read_text(encoding="utf-8"))
    found = []

    def walk(x):
        if isinstance(x, dict):
            inp = x.get("input")
            if isinstance(inp, dict) and inp.get("key") == "chu_hkd.ho_ten":
                found.append(inp)
            for v in x.values():
                if isinstance(v, (dict, list)):
                    walk(v)
        elif isinstance(x, list):
            for v in x:
                walk(v)

    walk(blocks)
    assert found, "blocks M01 phải có input chu_hkd.ho_ten"
    fm = json.loads((ROOT / "app" / "data" / "form_maps.json").read_text(
        encoding="utf-8"))
    keys = [f.get("key") for f in fm["PLI_M01"]["fields"]]
    assert "chu_hkd.ho_ten" in keys, "form_maps M01 phải đăng ký chu_hkd.ho_ten"
    # nhất quán 3 mẫu anh em cùng dòng 'Tôi là'
    for fid in ("PLI_M02", "PLII_M01", "PLI_M03"):
        d2 = docx.Document(str(ROOT / "app" / "templates" / ("%s.docx" % fid)))
        ok = any("{chu_hkd.ho_ten}" in p._p.xml
                 for p in d2.paragraphs if "Tôi là" in p.text)
        assert ok, "%s phải giữ marker chu_hkd.ho_ten" % fid


def main() -> int:
    f1_bridge_read_data_json()
    f2_app_js_bridge_first()
    f3_css_loop_fld_visible()
    f4_m01_owner_name_field()
    fails = [r for r in _results if not r[0]]
    print("-" * 60)
    print("test_phase5_ui_fixes: %d PASS / %d FAIL"
          % (len(_results) - len(fails), len(fails)))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
