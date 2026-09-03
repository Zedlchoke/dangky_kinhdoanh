# -*- coding: utf-8 -*-
"""Test Phase 5R lane A (metadata/boundary) — 4 defect WLAY-001/006/007a/007b.

Mỗi test pin defect cụ thể trong docs/word_layout_audit_2026-09-03.md:
  W1 WLAY-001 — role=source-indicator ("PL1 - Mau so 9: ...") đã bị UI ẩn nhưng
     còn lộ trong DOCX export ở đúng 44 form. Reproduction trước khi sửa: render
     template PLI_M09 rồi quét w:t — text "PL1 - Mau so 9: ..." vẫn còn. Fix:
     docx_render.render() loại paragraph con trực tiếp của w:body mà
     classify_role() xếp role 'source-indicator'.
  W2 WLAY-006 — PLII_M28 lộ literal ` PAGE   \\* MERGEFORMAT 2` (field code số
     trang của bản nguồn bị converter cũ làm phẳng thành w:t, xác minh byte ở
     file nguồn PL2_Mau_28_*.docx — số "2" stale là số trang của MẪU NGUỒN,
     không phải nội dung pháp lý của biểu mẫu; field gốc thuộc footer nên
     dựng lại fldSimple trong thân tài liệu là SAI bố cục -> quyết định LOẠI
     paragraph, không dựng lại field). Fix: build_form_maps.py drop paragraph
     field-code trang stale khi sinh template.
  W3 WLAY-007a — PLI_M72 variant TNHH_2TV có trang trắng (soffice + pymupdf
     đếm trang có chars < 20). Fix OOXML-đúng có điều kiện, không hard-code
     theo fixture (chi tiết gốc rễ trong docstring test W3).
  W4 WLAY-007b — PLI_M80 kéo banner 'PHỤ LỤC II / BIỂU MẪU SỬ DỤNG TRONG ĐĂNG
     KÝ HỘ KINH DOANH / (Kèm theo Thông tư 68/2025...)' của phụ lục KẾ vào đuôi
     template (contamination bước tách nguồn). Fix: rule boundary chung trong
     build_form_maps.py — dừng tách khi gặp banner PHỤ LỤC khác phụ lục của
     form, chống hồi quy cho mọi form.

Chạy:  python3 -X utf8 tools/test_phase5r_metadata.py
Exit 0 = PASS / 1 = FAIL.
"""
from __future__ import annotations

import io
import json
import re
import shutil
import sys
import unicodedata
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from tools.docx_render import (            # noqa: E402
    LEFTOVER_RE, dummy_data_from_map, render, write_rendered,
)
from tools.layout_roles import classify_role, fold  # noqa: E402

MAPS_PATH = ROOT / "app" / "data" / "form_maps.json"
TEMPLATES = ROOT / "app" / "templates"
OUT_DIR = ROOT / "generated_ho_so" / "_test_phase5r_metadata"

# 44 form lộ source-indicator — docs/word_layout_audit_2026-09-03.md mục 4
WLAY001_FORMS = [
    "PLII_M01", "PLII_M02", "PLII_M03", "PLII_M04", "PLII_M05", "PLII_M06",
    "PLII_M07", "PLII_M08", "PLII_M09", "PLII_M10", "PLII_M11", "PLII_M12",
    "PLII_M13", "PLII_M14", "PLII_M15", "PLII_M16", "PLII_M17", "PLII_M18",
    "PLII_M19",
    "PLI_M01", "PLI_M02", "PLI_M05", "PLI_M06", "PLI_M07", "PLI_M08",
    "PLI_M09", "PLI_M10", "PLI_M11", "PLI_M17", "PLI_M23", "PLI_M25",
    "PLI_M27", "PLI_M28", "PLI_M29", "PLI_M30", "PLI_M31", "PLI_M32",
    "PLI_M33", "PLI_M42", "PLI_M43", "PLI_M45", "PLI_M46", "PLI_M47",
    "PLI_M48",
]

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run(*args):
            try:
                fn(*args)
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def doc_xml(docx_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        return z.read("word/document.xml").decode("utf-8")


def template_xml(fid: str) -> str:
    with zipfile.ZipFile(TEMPLATES / ("%s.docx" % fid)) as z:
        return z.read("word/document.xml").decode("utf-8")


def xml_at(path) -> str:
    """Đọc word/document.xml từ đường dẫn docx tuyệt đối (sandbox/ kết quả)."""
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def wt_texts(xml: str) -> list[str]:
    return re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)


def source_indicator_hits(xml: str) -> list[str]:
    """Chuẩn detector D1 của tools/phase5_visual_gate.py."""
    return [t.strip()[:120] for t in wt_texts(xml)
            if classify_role(fold(t)) == "source-indicator"]


# ---------------------------------------------------------------------------
# W1 — WLAY-001: source-indicator phải bị loại khỏi DOCX, form-number phải GIỮ
# ---------------------------------------------------------------------------

@check("W1 template_list: dung 44 form trong audit co marker body-level")
def t_w1_template_list():
    maps = json.loads(MAPS_PATH.read_text(encoding="utf-8"))
    found = []
    for fid in sorted(maps):
        xml = template_xml(fid)
        body = re.search(r"<w:body>(.*)</w:body>", xml, re.S).group(1)
        body_no_tbl = re.sub(r"<w:tbl[ >].*?</w:tbl>", "", body, flags=re.S)
        if source_indicator_hits(body_no_tbl):
            found.append(fid)
    assert set(found) == set(WLAY001_FORMS), (
        "danh sach form co marker lech audit: chi co %d, thieu %d, du %d"
        % (len(found),
           sorted(set(WLAY001_FORMS) - set(found))[:5],
           sorted(set(found) - set(WLAY001_FORMS))[:5]))


@check("W1 render PLI_M09: khong con source-indicator, GIU 'Mau so 9' phap ly")
def t_w1_render_m09():
    maps = json.loads(MAPS_PATH.read_text(encoding="utf-8"))
    fid = "PLI_M09"
    data = dummy_data_from_map(maps[fid])
    xml = doc_xml(render(TEMPLATES / ("%s.docx" % fid), data))
    hits = source_indicator_hits(xml)
    assert hits == [], ("PLI_M09: con %d source-indicator trong DOCX: %r"
                        % (len(hits), hits[:3]))
    # Dòng pháp lý "Mẫu số 9" (role form-number) phải GIỮ NGUYÊN
    texts = wt_texts(xml)
    assert any(fold(t) == "mau so 9" for t in texts), \
        "mat dong phap ly 'Mau so 9' sau khi loai source-indicator"
    assert any(classify_role(fold(t)) == "form-number" for t in texts), \
        "khong con paragraph nao mang role form-number"


@check("W1 render 44 form audit: 0 source-indicator trong moi DOCX")
def t_w1_render_all_44():
    maps = json.loads(MAPS_PATH.read_text(encoding="utf-8"))
    bad = {}
    for fid in WLAY001_FORMS:
        data = dummy_data_from_map(maps[fid])
        xml = doc_xml(render(TEMPLATES / ("%s.docx" % fid), data))
        hits = source_indicator_hits(xml)
        if hits:
            bad[fid] = hits[:2]
    assert not bad, ("%d/%d form con source-indicator: %r"
                     % (len(bad), len(WLAY001_FORMS),
                        dict(list(bad.items())[:5])))


@check("W1 byte-identical: xml khong co marker thi ham loai tra NGUYEN ven")
def t_w1_byte_identical():
    from tools import docx_render
    body = ("<w:p><w:r><w:t>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>Mẫu số 9</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>PL1 - Mau so 9: Giấy đề nghị</w:t></w:r></w:p>"
            "<w:p/>")
    # chỉ 1 marker -> bị loại, 2 paragraph còn lại + self-closing giữ nguyên
    out = docx_render._drop_source_indicators(
        "<w:body>%s<w:sectPr/></w:body>" % body)
    assert "PL1 - Mau so 9" not in out, "marker khong bi loai"
    assert "Mẫu số 9" in out and "CỘNG HÒA" in out, "mat noi dung khong marker"
    # xml KHÔNG có marker -> trả về nguyên bản (byte-identical, không rebuild)
    clean = ("<w:body><w:p><w:r><w:t>Mẫu số 72</w:t></w:r></w:p>"
             "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>"
             "PL1 - Mau so 72 trong bang phai GIU (ngoai pham vi body-level)"
             "</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body>")
    assert docx_render._drop_source_indicators(clean) == clean, \
        "xml khong marker bi bien doi (mất tính byte-identical)"


@check("W1 byte-identical: 43 form khong marker render giong pipeline cu")
def t_w1_render_rest_byte_identical():
    """Form KHÔNG có source-indicator phải render byte-identical như pipeline
    cũ — bằng chứng: tắt `_drop_source_indicators` (đồng nhất với pipeline
    trước khi thêm bước này) rồi render lại, hai document.xml phải bằng nhau
    từng byte. Vì bước mới là khác biệt DUY NHẤT giữa hai pipeline nên hai
    output bằng nhau <=> không có form nào bị đổi output vô cớ."""
    from tools import docx_render
    maps = json.loads(MAPS_PATH.read_text(encoding="utf-8"))
    checked = 0
    for fid in sorted(maps):
        xml = template_xml(fid)
        body = re.search(r"<w:body>(.*)</w:body>", xml, re.S).group(1)
        body_no_tbl = re.sub(r"<w:tbl[ >].*?</w:tbl>", "", body, flags=re.S)
        if source_indicator_hits(body_no_tbl):
            continue
        tpath = TEMPLATES / ("%s.docx" % fid)
        data = dummy_data_from_map(maps[fid])
        with_new = doc_xml(render(tpath, data))
        orig = docx_render._drop_source_indicators
        docx_render._drop_source_indicators = lambda x: x
        try:
            without_new = doc_xml(render(tpath, data))
        finally:
            docx_render._drop_source_indicators = orig
        assert with_new == without_new, \
            "%s: render doi output dù khong co source-indicator" % fid
        checked += 1
    assert checked == 87 - len(WLAY001_FORMS), \
        "so form kiem tra %d != 87 - 44" % checked


# ---------------------------------------------------------------------------
# W2 — WLAY-006: field code ` PAGE \* MERGEFORMAT 2` phải biến mất khỏi
# template PLII_M28 SINH RA (sandbox regen → /tmp, không đụng app/templates)
# ---------------------------------------------------------------------------

W2_SOURCE = (ROOT / "06_ho_kinh_doanh_va_chuyen_doi" / "01_mau_ho_kinh_doanh"
             / "PL2_Mau_28_thong_bao_ve_viec_ra_soat_thong_tin_dang_ky_ho_"
               "kinh_doanh_cua_ho_kinh_doanh.docx")
W2_SANDBOX = OUT_DIR / "w2_sandbox_regen"


@check("W2 reproduction: ban nguon PL2_Mau_28 chua literal field code")
def t_w2_source_has_defect():
    assert W2_SOURCE.exists(), "khong tim thay file nguon PL2_Mau_28"
    xml = xml_at(W2_SOURCE)
    hits = [t for t in wt_texts(xml) if "MERGEFORMAT" in t]
    assert hits, "ban nguon khong con literal (reproduction het dung luc?)"
    assert any("PAGE" in t for t in hits), hits


@check("W2 sandbox regen PLII_M28: det_field_code tra RONG, noi dung phap ly con")
def t_w2_sandbox_regen_clean():
    from tools import build_form_maps as bfm
    from tools.phase5_visual_gate import det_field_code
    import shutil
    if W2_SANDBOX.exists():
        shutil.rmtree(W2_SANDBOX)
    meta, _unmatched = bfm.process_form(
        W2_SOURCE, bfm.load_csv_titles(), None, out_tpl=W2_SANDBOX)
    assert meta["id"] == "PLII_M28", meta["id"]
    out_docx = W2_SANDBOX / "PLII_M28.docx"
    assert out_docx.exists(), "sandbox regen khong ghi ra template"
    # detector D2 của gate (cùng ngữ pháp với generator drop) phải trả rỗng
    hits = det_field_code(out_docx)
    assert hits == [], "sandbox regen van con field code: %r" % hits[:3]
    # nội dung pháp lý của mẫu phải còn (không nuke cả tài liệu)
    xml = xml_at(out_docx)
    joined = fold(" ".join(wt_texts(xml)))
    assert "mau so 28" in joined, "mat dong 'Mau so 28' phap ly"
    assert "ra soat thong tin" in joined, "mat ten mau phap ly"


@check("W2 unit: drop_flat_field_code_paras chi drop paragraph field code")
def t_w2_unit_drop():
    from docx import Document
    from tools.build_form_maps import drop_flat_field_code_paras
    doc = Document()
    doc.add_paragraph(" PAGE   \\* MERGEFORMAT 2")
    doc.add_paragraph("Giấy chứng nhận đăng ký doanh nghiệp")
    doc.add_paragraph("")
    n = drop_flat_field_code_paras(doc)
    assert n == 1, "phai drop dung 1 paragraph, drop %d" % n
    texts = [p.text for p in doc.paragraphs]
    assert "Giấy chứng nhận đăng ký doanh nghiệp" in texts, "mat chu that"
    assert not any("MERGEFORMAT" in t for t in texts)


# ---------------------------------------------------------------------------
# W3 — WLAY-007a: PLI_M72 variant TNHH_2TV trang trắng.
#
# GỐC RỄ (đã điều tra tại BASE c9d095e + corpus audit
# generated_ho_so/phase5_word_open_tnhh_2tv_20260903_180918/PLI_M72.docx):
# đuôi body sau render của mọi variant M72 là
#   [...nội dung][đường kẻ ký][footnote có điều kiện][paragraph RỖNG][sectPr]
# paragraph rỗng cuối (spacing before=120, kế thừa từ bản nguồn) là con
# TRỰC TIẾP cuối của body ngay trước sectPr. Khi nội dung trước nó lấp vừa
# đầy trang cuối, paragraph rỗng này tự mình tràn sang trang kế -> trang
# trắng 0 ký tự (reproduce: convert corpus docx trên bằng soffice + pymupdf
# → 3 trang, chars/page = [577, 641, 0]).
# FIX OOXML-đúng CÓ ĐIỀU KIỆN tại render: xoá paragraph CUỐI body chỉ khi
# (a) là con trực tiếp cuối trước sectPr cuối, (b) rỗng hoàn toàn (không
# text, không drawing/pict). KHÔNG xoá mù paragraph cuối, KHÔNG hard-code
# page break theo fixture, KHÔNG đụng paragraph cuối BẢNG.
# ---------------------------------------------------------------------------

W3_SANDBOX = OUT_DIR / "w3_sandbox_pdf"


def _render_m72_variant(variant_label: str) -> bytes:
    from tools.test_m72_conditional_export import (
        render_payload, VARIANTS, TEMPLATE_PATH,
    )
    v = next(x for x in VARIANTS if x["label"] == variant_label)
    return render(TEMPLATE_PATH, render_payload(v))


@check("W3 reproduction: corpus audit tnhh_2tv co trang 3 chars=0")
def t_w3_reproduction_corpus():
    """Tái hiện defect từ artifact audit: docx corpus (BASE) convert ra
    trang trắng 0 ký tự — đây là bằng chứng gốc rễ, dùng thay fixture tự chế."""
    corpus = Path("/Users/hoangminhlong/orca/dangky_kinhdoanh/generated_ho_so"
                  "/phase5_word_open_tnhh_2tv_20260903_180918/PLI_M72.docx")
    if not corpus.exists():
        return  # corpus không có ở máy chạy test: bỏ qua phần reproduction
    import subprocess
    W3_SANDBOX.mkdir(parents=True, exist_ok=True)
    dst = W3_SANDBOX / "audit_base.docx"
    shutil.copyfile(corpus, dst)
    subprocess.run(
        ["soffice", "--headless", "--norestore",
         "-env:UserInstallation=file:///tmp/p5r_lo_profile_w3",
         "--convert-to", "pdf", "--outdir", str(W3_SANDBOX), str(dst)],
        capture_output=True, timeout=120)
    import fitz
    pdf = W3_SANDBOX / "audit_base.pdf"
    if not pdf.exists():
        return  # soffice không chạy được ở máy này: bỏ qua, test khác chặn
    doc = fitz.open(str(pdf))
    chars = [len("".join(doc[i].get_text("text").split()))
             for i in range(doc.page_count)]
    assert any(c < 20 for c in chars), \
        ("corpus audit khong con trang trang (%s) — reproduction het thoi "
         "diem dung, hay doi evidence moi" % chars)


@check("W3 xml: khong con paragraph rong dung CUOI body truoc sectPr")
def t_w3_no_trailing_empty_para():
    from tools.docx_render import _split_body_children
    for label in ("DNTN", "TNHH_1TV", "TNHH_2TV", "CTCP", "DHDN"):
        xml = doc_xml(_render_m72_variant(label))
        body = re.search(r"<w:body>(.*)</w:body>", xml, re.S).group(1)
        spans = _split_body_children(body)
        assert spans, "%s: body rong?" % label
        last = spans[-1]
        assert body[last[0]:last[1]].startswith("<w:sectPr"), \
            "%s: span cuối không phải sectPr" % label
        if len(spans) >= 2:
            prev = body[spans[-2][0]:spans[-2][1]]
            texts = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", prev))
            if prev.startswith("<w:p"):
                assert texts.strip(), \
                    ("%s: paragraph rỗng vẫn nằm cuối body trước sectPr "
                     "(gốc rễ trang trắng chưa xử lý)" % label)


@check("W3 pagination: TNHH_2TV khong trang chars<20; 4 variant khac giu page count")
def t_w3_pagination():
    """Đỏ trước khi sửa qua reproduction; sau fix: page count 4 variant
    DNTN/TNHH_1TV/CTCP/DHDN giữ nguyên (2/2/3/2 đo tại BASE), TNHH_2TV
    không còn trang trắng (<20 chars)."""
    import subprocess
    import fitz
    W3_SANDBOX.mkdir(parents=True, exist_ok=True)
    baseline_pages = {"DNTN": 2, "TNHH_1TV": 2, "CTCP": 3, "DHDN": 2}
    for label in ("DNTN", "TNHH_1TV", "TNHH_2TV", "CTCP", "DHDN"):
        docx = W3_SANDBOX / ("M72_%s.docx" % label)
        docx.write_bytes(_render_m72_variant(label))
        subprocess.run(
            ["soffice", "--headless", "--norestore",
             "-env:UserInstallation=file:///tmp/p5r_lo_profile_w3",
             "--convert-to", "pdf", "--outdir", str(W3_SANDBOX), str(docx)],
            capture_output=True, timeout=120)
        pdf = W3_SANDBOX / ("M72_%s.pdf" % label)
        assert pdf.exists(), "%s: soffice khong tao duoc PDF" % label
        doc = fitz.open(str(pdf))
        chars = [len("".join(doc[i].get_text("text").split()))
                 for i in range(doc.page_count)]
        assert not any(c < 20 for c in chars), \
            "%s: con trang trang chars=%s" % (label, chars)
        if label in baseline_pages:
            assert doc.page_count == baseline_pages[label], \
                ("%s: page count %d != baseline %d" %
                 (label, doc.page_count, baseline_pages[label]))


# ---------------------------------------------------------------------------
# W4 — WLAY-007b: banner PHỤ LỤC của phụ lục KẾ lọt vào đuôi PLI_M80.
#
# Reproduction trước khi sửa: file nguồn
# 05_cong_bo_cung_cap_thong_tin/PL1_Mau_80_*.docx (bước tách auto-extract
# vượt ranh giới) chứa 3 đoạn cuối 'PHỤ LỤC II' / 'BIỂU MẪU SỬ DỤNG TRONG
# ĐĂNG KÝ HỘ KINH DOANH' / '(Kèm theo Thông tư số 68/2025...)' — template
# sinh ra in contaminated sang trang cuối. Fix: rule boundary chung
# drop_foreign_annex_tail() trong build_form_maps.process_form — dừng tách
# khi gặp banner PHỤ LỤC <roman khác>; chống hồi quy cho MỌI form bằng
# sandbox regen toàn bộ rồi scan banner lạ.
# ---------------------------------------------------------------------------

W4_SANDBOX = OUT_DIR / "w4_sandbox_regen"


@check("W4 reproduction: nguon PL1_M80 chua banner PHU LUC II")
def t_w4_source_has_banner():
    src = (ROOT / "05_cong_bo_cung_cap_thong_tin" /
           "PL1_Mau_80_giay_de_nghi_thu_hoi_giay_to_do_ke_khai_trong_ho_so_"
           "dang_ky_doanh_nghiep_la_gia_mao.docx")
    assert src.exists(), "khong tim thay nguon PL1_M80"
    xml = xml_at(src)
    assert any("PHỤ LỤC II" in t for t in wt_texts(xml)), \
        "nguon khong con banner (reproduction het dung luc)"


@check("W4 sandbox regen PLI_M80: det_contamination tra RONG, mat banner la")
def t_w4_m80_regen_clean():
    from tools import build_form_maps as bfm
    from tools.phase5_visual_gate import det_contamination
    W4_SANDBOX.mkdir(parents=True, exist_ok=True)
    src = (ROOT / "05_cong_bo_cung_cap_thong_tin" /
           "PL1_Mau_80_giay_de_nghi_thu_hoi_giay_to_do_ke_khai_trong_ho_so_"
           "dang_ky_doanh_nghiep_la_gia_mao.docx")
    meta, _ = bfm.process_form(src, bfm.load_csv_titles(), None,
                               out_tpl=W4_SANDBOX)
    assert meta["id"] == "PLI_M80", meta["id"]
    out = W4_SANDBOX / "PLI_M80.docx"
    assert out.exists()
    hits = det_contamination(out, {"phu_luc": "Phụ lục I"})
    assert hits == [], "PLI_M80 sandbox regen van contamination: %r" % hits
    xml = xml_at(out)
    assert not any("PHỤ LỤC II" in t for t in wt_texts(xml)), \
        "banner 'PHỤ LỤC II' van trong template regen"
    # nội dung pháp lý đuôi mẫu (footnote 3) phải còn
    assert any("Người đề nghị ký và ghi họ tên" in t for t in wt_texts(xml)), \
        "mat footnote phap ly cua M80 sau khi cat duoi"


@check("W4 chong hoi quy: sandbox regen TOAN BO form, 0 banner phu luc khac")
def t_w4_all_forms_scan():
    from tools import build_form_maps as bfm
    W4_SANDBOX.mkdir(parents=True, exist_ok=True)
    titles = bfm.load_csv_titles()
    bad = {}
    seen = 0
    for path in bfm.iter_docx_files():
        meta, _ = bfm.process_form(path, titles, None, out_tpl=W4_SANDBOX)
        seen += 1
        out = W4_SANDBOX / ("%s.docx" % meta["id"])
        xml = xml_at(out)
        own = (meta.get("phu_luc") or "").upper()
        for t in wt_texts(xml):
            m = re.match(r"^\s*PHỤ LỤC\s*(I{1,3}|IV|V|VI{1,3}|IX|X|XI{0,3})\s*$",
                         t, re.IGNORECASE)
            if m and m.group(1).upper() != own.upper():
                bad[meta["id"]] = t.strip()[:60]
                break
    assert seen >= 87, "regen qua it form (%d)" % seen
    assert not bad, ("%d form van chua banner phu luc khac: %r"
                     % (len(bad), dict(list(bad.items())[:5])))


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    tests = [v for k, v in sorted(globals().items())
             if k.startswith("t_") and callable(v)]
    for t in tests:
        t()
    passed = sum(1 for ok, _ in _results if ok)
    failed = len(_results) - passed
    print("-" * 60)
    print("test_phase5r_metadata: %d PASS / %d FAIL" % (passed, failed))
    for ok, name in _results:
        if not ok:
            print(" - FAIL:", name)
    print("Thu muc ket qua:", OUT_DIR)
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
