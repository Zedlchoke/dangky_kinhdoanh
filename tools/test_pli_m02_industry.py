# -*- coding: utf-8 -*-
"""Kiểm tra PHASE 3: bảng ngành nghề ds_nganh trong PLI_M02 (production).

Phạm vi:
1. blocks app/data/blocks/PLI_M02.json có đúng 1 loop "ds_nganh" (chọn theo
   tên loop; form còn loop người/ủy quyền R3B1 thuộc lane khác):
   - hàng tiêu đề 5 cột khớp số ô của body/template;
   - cols zero-based {ma_nganh, ten_nganh, chi_tiet, la_nganh_chinh};
   - cột 0 (STT) KHÔNG tag (STT tự sinh qua biến engine {_stt}).
2. form_maps.json PLI_M02 khai báo dynamic_tables tương ứng (col/tag khớp).
3. Render template bằng tools/docx_render.write_rendered với 3 dòng mẫu:
   - đủ 3 hàng dữ liệu nhân bản, mỗi giá trị nằm đúng ô cột (so text từng cell);
   - STT hiển thị 1/2/3 (tự sinh);
   - không còn sót '{#ds_nganh}', '{/ds_nganh}', tag cột hay {_stt}.

File test thuần Python, exit code 0 = PASS / 1 = FAIL.

Chạy:  python -X utf8 tools/test_pli_m02_industry.py
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document                      # noqa: E402
from docx.oxml.ns import qn                   # noqa: E402
from docx.table import Table                  # noqa: E402
from docx.text.paragraph import Paragraph     # noqa: E402

from tools.docx_render import write_rendered  # noqa: E402

BLOCKS_PATH = ROOT / "app" / "data" / "blocks" / "PLI_M02.json"
FORM_MAPS_PATH = ROOT / "app" / "data" / "form_maps.json"
TEMPLATE_PATH = ROOT / "app" / "templates" / "PLI_M02.docx"
OUT_DIR = ROOT / "generated_ho_so" / "_test_pli_m02_industry"

EXPECTED_HEADERS = ["STT", "Mã ngành", "Tên ngành theo VSIC",
                    "Chi tiết hoạt động", "Ngành chính"]
EXPECTED_COLS = [(1, "ma_nganh"), (2, "ten_nganh"),
                 (3, "chi_tiet"), (4, "la_nganh_chinh")]
LOOP_NAME = "ds_nganh"

SAMPLE_ROWS = [
    {"ma_nganh": "0111", "ten_nganh": "Trồng lúa",
     "chi_tiet": "", "la_nganh_chinh": "X"},
    {"ma_nganh": "4751", "ten_nganh": "Bán lẻ máy vi tính,<phần mềm>",
     "chi_tiet": "Kèm dịch vụ lắp đặt", "la_nganh_chinh": ""},
    {"ma_nganh": "6201", "ten_nganh": "Hoạt động lập trình máy tính &",
     "chi_tiet": "", "la_nganh_chinh": ""},
]

_results: list[tuple[bool, str]] = []


def check(name: str):
    def deco(fn):
        def run():
            try:
                fn()
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:  # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


def fold(text: str) -> str:
    nfkd = unicodedata.normalize("NFD", text.lower())
    out = "".join(c for c in nfkd if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", out.replace("đ", "d")).strip()


def cell_runs_text(cell: dict) -> str:
    return "".join(r.get("text", "") for r in cell["runs"])


def tc_iter(tr):
    for child in tr.iterchildren():
        if child.tag == qn("w:tc"):
            yield child


def tr_cell_texts(tbl, ri: int) -> list[str]:
    """Text từng ô của hàng ri — duyệt trực tiếp w:tc (mỗi ô đúng một lần)."""
    out = []
    for tc in tc_iter(tbl._tbl.tr_lst[ri]):
        txt = "\n".join(
            Paragraph(p_el, tbl.part.document).text
            for p_el in tc.iterchildren()
            if p_el.tag == qn("w:p")
        )
        out.append(txt.strip())
    return out


def body_tables(doc):
    """Các bảng theo thứ tự thân tài liệu (duyệt con trực tiếp body)."""
    return [Table(child, doc) for child in doc.element.body.iterchildren()
            if child.tag == qn("w:tbl")]


def find_industry_table(doc):
    """Tìm bảng có hàng tiêu đề khớp EXPECTED_HEADERS."""
    for ti, tbl in enumerate(doc.tables):
        if len(tbl.rows) and fold(" ".join(tr_cell_texts(tbl, 0))) == \
                fold(" ".join(EXPECTED_HEADERS)):
            return ti, tbl
    return None, None


@check("blocks PLI_M02.json: dung 1 loop ds_nganh (tra theo ten loop), cols zero-based, cot 0 STT khong tag")
def t_blocks_loop_structure():
    data = json.loads(BLOCKS_PATH.read_text(encoding="utf-8"))
    # [Lead sync sau regen R3B1] form có thêm block loop người/ủy quyền
    # (ds_nddpl_uy_quyen_chu_so_huu — lane R3B1 sở hữu). Test ngành này chỉ
    # khóa loop ds_nganh của mình: chọn theo TÊN loop thay vì đếm tổng số
    # block loop trong blocks.
    loops = [b for b in data["blocks"]
             if b.get("t") == "loop" and b.get("loop") == LOOP_NAME]
    assert len(loops) == 1, \
        "kỳ vọng đúng 1 block loop %s, thấy %d" % (LOOP_NAME, len(loops))
    blk = loops[0]
    assert blk["headerRows"] == [0], "headerRows=%r != [0]" % (blk["headerRows"],)

    # header 5 cột, text khớp thiết kế
    hdr_cells = blk["rows"][blk["headerRows"][0]]["cells"]
    assert len(hdr_cells) == 5, "header có %d ô != 5" % len(hdr_cells)
    got_headers = [cell_runs_text(c).strip() for c in hdr_cells]
    assert got_headers == EXPECTED_HEADERS, \
        "header %r != %r" % (got_headers, EXPECTED_HEADERS)

    # mọi hàng trong rows đều 5 ô ("header khớp body" về grid)
    for ri, row in enumerate(blk["rows"]):
        assert len(row["cells"]) == 5, \
            "rows[%d] có %d ô != 5" % (ri, len(row["cells"]))

    # cols zero-based đúng thứ tự; cột 0 không tag
    got_cols = sorted((c["col"], c["tag"]) for c in blk["cols"])
    assert got_cols == EXPECTED_COLS, "cols %r != %r" % (got_cols, EXPECTED_COLS)
    assert all(c["col"] != 0 for c in blk["cols"]), \
        "cột 0 (STT tự sinh) không được có tag trong cols"

    # blankRows hợp lý: >=1 hàng trắng để in; rows + blank + 1 khuôn = tổng hàng
    fm = json.loads(FORM_MAPS_PATH.read_text(encoding="utf-8"))["PLI_M02"]
    dt = next(t for t in fm["dynamic_tables"] if t["loop"] == LOOP_NAME)
    doc = Document(str(TEMPLATE_PATH))
    tables = body_tables(doc)
    tbl = tables[dt["table_index"]]
    n_tr = len(tbl._tbl.tr_lst)
    total = len(blk["rows"]) + blk["blankRows"] + 1
    assert total == n_tr, \
        "rows(%d)+blank(%d)+khuôn(1)=%d != %d hàng template" % (
            len(blk["rows"]), blk["blankRows"], total, n_tr)
    assert blk["blankRows"] >= 1, "cần ít nhất 1 hàng trắng để in"


@check("form_maps PLI_M02: dynamic_tables khai bao dung col/tag cho ds_nganh")
def t_form_maps_metadata():
    fm = json.loads(FORM_MAPS_PATH.read_text(encoding="utf-8"))["PLI_M02"]
    dts = [t for t in fm.get("dynamic_tables") or [] if t["loop"] == LOOP_NAME]
    assert len(dts) == 1, "kỳ vọng 1 dynamic_tables ds_nganh, thấy %d" % len(dts)
    dt = dts[0]
    got_cols = sorted((c["col"], c["tag"]) for c in dt["columns"])
    assert got_cols == EXPECTED_COLS, \
        "columns %r != %r" % (got_cols, EXPECTED_COLS)
    fld_keys = {f["key"] for f in fm.get("fields", [])}
    for _, tag in EXPECTED_COLS:
        assert ("%s.%s" % (LOOP_NAME, tag)) in fld_keys, \
            "thiếu field %s.%s trong form_maps" % (LOOP_NAME, tag)


@check("render DOCX: 3 dong mau vao dung cot, STT tu sinh 1/2/3, khong con tag")
def t_render_three_rows():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / "PLI_M02_ds_nganh.docx"
    write_rendered(TEMPLATE_PATH, {LOOP_NAME: SAMPLE_ROWS}, out_path)

    with zipfile.ZipFile(out_path) as z:
        xml = z.read("word/document.xml").decode("utf-8")

    doc = Document(str(out_path))
    ti, tbl = find_industry_table(doc)
    assert tbl is not None, "không tìm thấy bảng ngành nghề trong file render"
    texts = [tr_cell_texts(tbl, ri) for ri in range(len(tbl.rows))]

    # hàng tiêu đề giữ nguyên
    assert [t.strip() for t in texts[0]] == EXPECTED_HEADERS, \
        "header render %r != %r" % (texts[0], EXPECTED_HEADERS)

    # 3 hàng dữ liệu nằm ở cuối bảng (sau các hàng trắng để in)
    n_data = len(SAMPLE_ROWS)
    data_texts = texts[-n_data:]
    for i, row in enumerate(data_texts):
        want = SAMPLE_ROWS[i]
        assert len(row) == 5, "hàng dữ liệu %d có %d ô != 5" % (i, len(row))
        assert row[0] == str(i + 1), \
            "STT dòng %d = %r, kỳ vọng %r" % (i, row[0], str(i + 1))
        assert row[1] == want["ma_nganh"], \
            "Mã ngành dòng %d = %r != %r" % (i, row[1], want["ma_nganh"])
        assert row[2] == want["ten_nganh"], \
            "Tên ngành dòng %d = %r != %r" % (i, row[2], want["ten_nganh"])
        assert row[3] == want["chi_tiet"], \
            "Chi tiết dòng %d = %r != %r" % (i, row[3], want["chi_tiet"])
        assert row[4] == want["la_nganh_chinh"], \
            "Ngành chính dòng %d = %r != %r" % (i, row[4], want["la_nganh_chinh"])

    # không còn bất kỳ the loop/cột/biến engine nào sót trong XML
    for tok in ("{#%s}" % LOOP_NAME, "{/%s}" % LOOP_NAME,
                "{ma_nganh}", "{ten_nganh}", "{chi_tiet}",
                "{la_nganh_chinh}", "{_stt}"):
        assert tok not in xml, "còn sót token %r trong document.xml" % tok


def main() -> int:
    t_blocks_loop_structure()
    t_form_maps_metadata()
    t_render_three_rows()
    passed = sum(1 for ok, _ in _results if ok)
    failed = len(_results) - passed
    print("-" * 60)
    print("test_pli_m02_industry: %d PASS / %d FAIL" % (passed, failed))
    for ok, name in _results:
        if not ok:
            print(" - FAIL:", name)
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
