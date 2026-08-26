# -*- coding: utf-8 -*-
"""Test occurrence matching cho pipeline rebuild bang (contract P1).

Chay:
    python -m unittest tools.test_table_occurrences -v
    python -X utf8 tools\\test_table_occurrences.py

Phu:
- resolve_regions: 2 heading trung text -> dung vung cho nth=1/nth=2;
  end_markers hoat dong; mac dinh ket thuc tai heading "\\d+." ke tiep.
- Contamination: "Mau so <so>" / heading lon khong thuoc vung -> SKIP + log;
  vung SKIP khong bi xoa paragraph, khong bi chen bang.
- validate_occurrences: loi spec bi chan som (id trung/sai, nth sai).
- DOCX gia lap qua build_form_maps.rebuild_list_tables: 2 bang duoc chen
  voi loop {#occ_a}/{#occ_b}; entry legacy KHONG co "occurrences" giu
  hanh vi cu (1 bang tai anchor, loop theo spec).
"""
from __future__ import annotations

import io
import contextlib
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    from tools import build_form_maps as bfm
    from tools import table_occurrences as occ
except ImportError:  # chay truc tiep tu thu muc tools/
    import build_form_maps as bfm  # type: ignore
    import table_occurrences as occ  # type: ignore

from docx import Document  # noqa: E402


def make_spec_entry(loop="ds_nganh", headers=None, no_default_drops=True):
    """Entry spec giong dang load_form_spec tra ve (spec da chuan hoa)."""
    return {
        "pos": "after",
        "headers": headers or ["STT", "Tên ngành", "Mã ngành",
                               "Ngành, nghề kinh doanh chính"],
        "no_default_drops": no_default_drops,
        "drop": ["^stt$", "^ten nganh\\b", "^ma nganh\\d*$",
                 "^nganh, nghe kinh doanh chinh"],
        "spec": {
            "id": loop,
            "detect": {"stt", "ten nganh", "ma nganh"},
            "stt_tag": True,
            "cols": [("^stt$", None), ("ten nganh", "ten_nganh"),
                     ("ma nganh", "ma_nganh"),
                     ("^nganh chinh$|nganh, nghe kinh doanh chinh",
                      "la_nganh_chinh")],
            "loop": loop,
            "fields": ["ten_nganh", "ma_nganh", "la_nganh_chinh"],
        },
    }


def wrap_ext(entries):
    """Dung ext giong load_form_spec tra ve: entry["spec"] la ID string."""
    specs = {}
    for e in entries:
        s = e["spec"]
        specs[s["id"]] = s
        e["spec"] = s["id"]
    return {"entries": entries, "specs_by_id": specs}


class LogCapture:
    def __init__(self):
        self.lines = []

    def __call__(self, msg):
        self.lines.append(msg)


class TestValidateOccurrences(unittest.TestCase):
    def test_valid_minimal(self):
        out = occ.validate_occurrences(
            [{"occurrence_id": "a", "match_heading": "5. Ngành"}],
            existing_loop_names={"ds_nganh"}, form_id="F")
        self.assertEqual(out[0]["nth"], 1)
        self.assertEqual(out[0]["end_markers"], [])
        self.assertIsNone(out[0]["sync_group"])

    def test_dup_id_rejected(self):
        with self.assertRaises(ValueError):
            occ.validate_occurrences(
                [{"occurrence_id": "a", "match_heading": "x", "nth": 1},
                 {"occurrence_id": "a", "match_heading": "y", "nth": 2}])

    def test_id_collision_with_loop_rejected(self):
        with self.assertRaises(ValueError):
            occ.validate_occurrences(
                [{"occurrence_id": "ds_nganh", "match_heading": "x"}],
                existing_loop_names={"ds_nganh"})

    def test_bad_id_and_nth_rejected(self):
        with self.assertRaises(ValueError):
            occ.validate_occurrences(
                [{"occurrence_id": "khong hop le", "match_heading": "x"}])
        with self.assertRaises(ValueError):
            occ.validate_occurrences(
                [{"occurrence_id": "ok", "match_heading": "x", "nth": 0}])


class TestResolveRegions(unittest.TestCase):
    BODY = [
        "4. Thông tin bổ sung",
        "5. Ngành, nghề kinh doanh",          # head lan 1
        "STT Tên ngành Mã ngành Ngành chính",  # tan du trong vung 1
        "6. Vốn đầu tư",                       # terminator mac dinh
        "5. Ngành, nghề kinh doanh",          # head lan 2
        "STT Tên ngành Mã ngành Ngành chính",  # tan du trong vung 2
        "Ghi chú: các dòng này",
    ]

    def test_nth_1_va_2_dung_vung(self):
        log = LogCapture()
        regions = occ.resolve_regions(
            self.BODY,
            occ.validate_occurrences(
                [{"occurrence_id": "occ_a", "match_heading":
                  "5. Ngành, nghề kinh doanh", "nth": 1},
                 {"occurrence_id": "occ_b", "match_heading":
                  "5. Ngành, nghề kinh doanh", "nth": 2,
                  "end_markers": ["^Ghi chú"]}]),
            log=log)
        by_id = {r["cfg"]["occurrence_id"]: r for r in regions}
        self.assertEqual(by_id["occ_a"]["head"], 1)
        self.assertEqual(by_id["occ_a"]["start"], 2)
        self.assertEqual(by_id["occ_a"]["end"], 3)   # ket thuc tai "6. Vốn"
        self.assertEqual(by_id["occ_b"]["head"], 4)
        self.assertEqual(by_id["occ_b"]["start"], 5)
        # end_markers ["^Ghi chú"] -> ket thuc tai dong 6
        self.assertEqual(by_id["occ_b"]["end"], 6)
        self.assertTrue(all(r["status"] == "ok" for r in regions))
        self.assertTrue(any("occ_a" in ln for ln in log.lines))

    def test_end_markers_custom(self):
        body = ["5. Ngành", "nội dung A", "Ghi chú: x", "sau cung"]
        regions = occ.resolve_regions(
            body,
            occ.validate_occurrences(
                [{"occurrence_id": "o", "match_heading": "5. Ngành",
                  "nth": 1, "end_markers": ["^Ghi chú"]}]),
            log=LogCapture())
        self.assertEqual(regions[0]["end"], 2)

    def test_contamination_mau_so_skip_co_log(self):
        body = list(self.BODY)
        body.insert(5, "Mẫu số 07 - Giấy chứng nhận")  # nhiem vao vung 2
        log = LogCapture()
        regions = occ.resolve_regions(
            body,
            occ.validate_occurrences(
                [{"occurrence_id": "occ_a", "match_heading":
                  "5. Ngành, nghề kinh doanh", "nth": 1},
                 {"occurrence_id": "occ_b", "match_heading":
                  "5. Ngành, nghề kinh doanh", "nth": 2}]),
            log=log)
        by_id = {r["cfg"]["occurrence_id"]: r for r in regions}
        self.assertEqual(by_id["occ_a"]["status"], "ok")
        self.assertEqual(by_id["occ_b"]["status"], "skipped")
        self.assertIn("nhiem mau khac", by_id["occ_b"]["reason"])
        self.assertTrue(any("SKIP" in ln and "occ_b" in ln
                            for ln in log.lines))

    def test_heading_lon_trong_vung_skip(self):
        body = ["5. Ngành", "nội dung",
                "6. Một heading khác không phải end marker",
                "7. Ghi chú"]
        log = LogCapture()
        regions = occ.resolve_regions(
            body,
            occ.validate_occurrences(
                [{"occurrence_id": "o", "match_heading": "5. Ngành",
                  "nth": 1, "end_markers": ["^7\\."]}]),
            log=log)
        self.assertEqual(regions[0]["status"], "skipped")
        self.assertIn("heading lon", regions[0]["reason"])

    def test_thieu_lan_xuat_hien_skip(self):
        log = LogCapture()
        regions = occ.resolve_regions(
            ["5. Ngành"],
            occ.validate_occurrences(
                [{"occurrence_id": "o", "match_heading":
                  "5. Ngành", "nth": 3}]),
            log=log)
        self.assertEqual(regions[0]["status"], "skipped")
        self.assertIn("khong du lan thu 3", regions[0]["reason"])


class TestRebuildDocxFake(unittest.TestCase):
    """DOCX gia lap: 2 heading trung text + noi dung giua."""

    def _build_doc(self, contaminate=False):
        doc = Document()
        doc.add_paragraph("4. Thông tin bổ sung")
        doc.add_paragraph("5. Ngành, nghề kinh doanh")
        # tan du luoi cu: moi tieu de cot 1 paragraph nhu template that
        for txt in ("STT", "Tên ngành", "Mã ngành",
                    "Ngành, nghề kinh doanh chính"):
            doc.add_paragraph(txt)
        doc.add_paragraph("12345678910")
        doc.add_paragraph("6. Vốn đầu tư")
        doc.add_paragraph("5. Ngành, nghề kinh doanh")
        if contaminate:
            doc.add_paragraph("Mẫu số 06b - Công bố nội dung")
        for txt in ("STT", "Tên ngành", "Mã ngành",
                    "Ngành, nghề kinh doanh chính"):
            doc.add_paragraph(txt)
        doc.add_paragraph("6. Vốn đầu tư")
        return doc

    def _entry(self, occurrences):
        e = make_spec_entry()
        e["occurrences"] = occurrences
        e["_occurrences"] = occ.validate_occurrences(
            occurrences, existing_loop_names={"ds_nganh"}, form_id="FAKE")
        return e

    def _body_paras(self, doc):
        from docx.text.paragraph import Paragraph
        return [Paragraph(el, doc).text for el in doc.element.body.iterchildren()
                if el.tag.endswith("}p")]

    def _tables_loop_names(self, doc):
        names = []
        for el in doc.element.body.iterchildren():
            if el.tag.endswith("}tbl"):
                txt = "".join(node.text or "" for node in el.iter()
                              if node.tag.endswith("}t"))
                m_open = [seg for seg in txt.split("{#") if "}" in seg]
                names.append(m_open[0].split("}")[0] if m_open else None)
        return names

    def test_hai_occurrence_hai_loop_rieng(self):
        doc = self._build_doc()
        entry = self._entry([
            {"occurrence_id": "occ_a", "match_heading":
             "5. Ngành, nghề kinh doanh", "nth": 1},
            {"occurrence_id": "occ_b", "match_heading":
             "5. Ngành, nghề kinh doanh", "nth": 2},
        ])
        entry["no_default_drops"] = False  # cho phep default drops scoped
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            bfm.rebuild_list_tables(doc, "FAKE", wrap_ext([entry]))
        loops = self._tables_loop_names(doc)
        self.assertEqual(loops, ["occ_a", "occ_b"])
        paras = [p.strip() for p in self._body_paras(doc)]
        # tan du trong 2 vung da duoc drop scoped (entry drop + default)
        self.assertFalse(any(p == "STT" or p == "Tên ngành"
                             or p == "Mã ngành" for p in paras))
        self.assertFalse(any(p == "12345678910" for p in paras))
        # ...nhung ranh gioi "6. Vốn đầu tư" con nguyen ven (khong nuot)
        self.assertEqual(paras.count("6. Vốn đầu tư"), 2)
        # thuoc tinh tam phai ton tai tren bang de process_form nhan dien
        import json as _json
        marks = [el.get(bfm._OCC_TEMP_ATTR)
                 for el in doc.element.body.iterchildren()
                 if el.tag.endswith("}tbl")]
        infos = [_json.loads(m) for m in marks]
        self.assertEqual([i["occurrence_id"] for i in infos],
                         ["occ_a", "occ_b"])

    def test_contamination_skip_giu_vung_lang_gieng(self):
        doc = self._build_doc(contaminate=True)
        entry = self._entry([
            {"occurrence_id": "occ_a", "match_heading":
             "5. Ngành, nghề kinh doanh", "nth": 1},
            {"occurrence_id": "occ_b", "match_heading":
             "5. Ngành, nghề kinh doanh", "nth": 2},
        ])
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            bfm.rebuild_list_tables(doc, "FAKE", wrap_ext([entry]))
        self.assertEqual(self._tables_loop_names(doc), ["occ_a"])
        out = buf.getvalue()
        self.assertIn("SKIP", out)
        self.assertIn("occ_b", out)
        paras = self._body_paras(doc)
        # vung bi SKIP: tan du GIU NGUYEN (khong nuot chuoi lang gieng)
        self.assertTrue(any(p.strip() == "Tên ngành" for p in paras))
        self.assertTrue(any(p.startswith("Mẫu số 06b") for p in paras))

    def test_legacy_entry_khong_occurrences_giu_hanh_vi_cu(self):
        doc = Document()
        doc.add_paragraph("đầu mẫu")
        doc.add_paragraph("4. Ngành, nghề kinh doanh")   # anchor
        doc.add_paragraph("STT Tên ngành Mã ngành Ngành, nghề kinh doanh chính")
        doc.add_paragraph("cuối mẫu")
        entry = make_spec_entry()
        entry["anchor"] = "^4\\. nganh, nghe kinh doanh"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            bfm.rebuild_list_tables(doc, "FAKE", wrap_ext([entry]))
        loops = self._tables_loop_names(doc)
        self.assertEqual(loops, ["ds_nganh"])
        paras = [p.strip() for p in self._body_paras(doc)]
        self.assertIn("4. Ngành, nghề kinh doanh", paras)
        self.assertIn("đầu mẫu", paras)
        self.assertIn("cuối mẫu", paras)


if __name__ == "__main__":
    unittest.main(verbosity=2)
