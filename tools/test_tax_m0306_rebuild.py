#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test_tax_m0306_rebuild.py — Lane fix-tax-m0306 (Owner P3=A, Class 6).

Owner decision P3=A (docs/decision_registry.md): "rebuild bảng thuế
PLII_M03–M06" — nhóm "số mồ côi/bảng thuế bị làm phẳng". Đánh giá TẠI BASE
5907066 (lane fix-tax-m0306) cho thấy premise Class 6 KHÔNG tái hiện được:

  - Nguồn tách 06_ho_kinh_doanh_va_chuyen_doi/01_mau_ho_kinh_doanh/
    PL2_Mau_03..06_*.docx: 0 w:tbl, KHÔNG có paragraph/dòng nào về "thông
    tin đăng ký thuế" (M03: 26 paragraphs; M04: 20; M05: 21; M06: 36).
  - Nguồn đầy đủ 06_ho_kinh_doanh_va_chuyen_doi/_Bieu_mau_PL_I_II_full.docx
    (bản text bị làm phẳng — 0 w:tbl toàn file): Phụ lục II Mẫu 2 CÓ khối
    "ĐĂNG KÝ THAY ĐỔI NỘI DUNG ĐĂNG KÝ THUẾ" + bảng Chỉ tiêu|cũ|mới
    (P5836-P5841) trong khi Mẫu 3/4/5/6 (P5919-P6017) KHÔNG có dòng thuế
    nào → nếu mẫu II.3-II.6 có bảng thuế thì nguồn này phải lộ dấu hiệu;
    detector của pipeline cũng sẽ thấy. That te: KHÔNG.
  - Census (app/data/table_specs/manifest.json): cả 4 mẫu form_status =
    NO_TABLE, detector.template_w_tbl = 0.
  - Blocks hiện tại: 0 paragraph số standalone (số mồ côi Class 6 = 0).

Theo mệnh đề pháp lý của lane ("KHÔNG bịa hàng/cột; mẫu nguồn không có →
giữ BLOCKED + báo BLOCKED_OWNER"), 4 file
tools/form_specs/household/PLII_M03..06.json được PIN trạng thái BLOCKED
bằng `tables: []` (validator v1 vẫn chạy, generator không chạm document).
Suite này chốt trạng thái đó — mọi mệnh đề đều ĐỎ nếu ai đó "dựng theo suy
đoán" hoặc artifact trôi:

  [SPEC]    4 file spec qua loader/validator v1 của build_form_maps;
            `tables` rỗng, KHÔNG entry nào, gap_nghiepvu + note pin
            BLOCKED_OWNER kèm lý do "nguồn không có bảng thuế" (điều d).
  [SRC]     Nguồn tách thật (find_source): 0 w:tbl và 0 marker thuế fold
            ("thong tin dang ky thue|chi tieu|thong tin dang ky cu|moi").
            Đối chứng có RĂNG: nguồn đầy đủ MẪU II.2 chứa đủ dòng lệnh +
            3 header thuế trong khi MẪU II.3..II.6 không có → detector
            đáng lẽ phải thấy nếu bảng tồn tại.
  [ORPHAN]  (điều b) blocks của cả 4 mẫu: 0 paragraph số standalone
            (^[0-9]+([./][0-9]+)*$ — dạng "4","5.1" của Class 6);
            duy nhất 1 block table = header band 2x2, 0 block loop.
  [SHAPE]   (điều a) template commit: đúng 1 w:tbl 2x2 (header band),
            KHÔNG w:tbl dữ liệu nào; sandbox regen với spec thật cho
            document.xml canonical-identical + form_maps fields trùng khớp
            production (regen là no-op — không drift).
  [RENDER]  (điều c) write_rendered với TOÀN BỘ fields của formMaps điền
            sentinel: 0 token sót ({...}/{#...}/{?...}) trong document.xml
            + header/footer; mỗi sentinel xuất hiện đúng 1 giá trị trong
            văn bản render; KHÔNG có key thuế kiểu '<sem>.<row>.<tag>'
            nào được sinh ra (không bia bảng thuế).
  [NEG]     thay spec bằng entry fixed-matrix BỊA trong sandbox → generator
            trung thực sinh thêm w:tbl nhưng hợp đồng shape của lane phải
            TỪ CHỐI — chứng minh suite có cắn.
  [GUARD]   app/templates/** và app/data/** không bị ghi trong suốt suite.

Chạy: python3 -X utf8 -m tools.test_tax_m0306_rebuild   (từ root repo/worktree)
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
import re
import shutil
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import tools.build_form_maps as bfm          # noqa: E402
import tools.docx_render as dxr              # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SPEC_DIR = ROOT / "tools" / "form_specs" / "household"
BLOCKS_DIR = ROOT / "app" / "data" / "blocks"
TPL_DIR = ROOT / "app" / "templates"
FORM_MAPS = ROOT / "app" / "data" / "form_maps.json"
FULL_DOC = ROOT / "06_ho_kinh_doanh_va_chuyen_doi" / "_Bieu_mau_PL_I_II_full.docx"

FORMS = ["PLII_M03", "PLII_M04", "PLII_M05", "PLII_M06"]

# Marker fold của một bảng/khối thuế PLII (đối chứng Mẫu II.2 trong nguồn đầy
# đủ — nguồn text bị làm phẳng vẫn lộ các dòng này). Nếu Mẫu II.3..II.6 có
# bảng thuế thì ít nhất 1 marker dưới đây phải xuất hiện trong vùng mẫu.
TAX_MARKERS = [
    "dang ky thay doi noi dung dang ky thue",
    "dang ky bo sung, thay doi cac chi tieu dang ky thue nhu sau",
    "chi tieu",
    "thong tin dang ky cu",
    "thong tin dang ky moi",
    "cac chi tieu thong tin dang ky thue",
]

# Số mồ côi Class 6: paragraph chỉ gồm chữ số (cho phép 5.1 / 4/5) — dạng
# STT/số mục của một bảng bị làm phẳng. "1 Trường hợp..." (chú thích nguồn,
# Class 3 glue) hay "(Ký và ghi họ tên)1" KHÔNG khớp pattern này.
ORPHAN_DIGIT_RE = re.compile(r"^\s*\d+(?:[./]\d+)*\s*$")

# Tiêu đề fold để neo đúng mẫu trong nguồn đầy đủ (đoạn PL2 bắt đầu sau TOC
# thứ hai; mỗi "mau so N" xuất hiện ít nhất 2 lần: 1 ở TOC, 1 ở mẫu thật).
FULL_DOC_ANCHORS = {
    2: "giay de nghi dang ky thay doi noi dung dang ky ho kinh doanh",
    3: "dang ky tam ngung kinh doanh/tiep tuc kinh doanh truoc thoi han",
    4: "ve viec cham dut hoat dong ho kinh doanh",
    5: "cap lai giay chung nhan dang ky ho kinh doanh",
    6: "dung thuc hien thu tuc dang ky ho kinh doanh",
}

# Khối thuế Mẫu II.2 trong nguồn đầy đủ (dòng lệnh + 3 header), theo
# PLII_M02.json entry a7_thue_chi_tieu_cu_moi: P5836-P5841.
MAU2_TAX_LINES = [
    "dang ky thay doi noi dung dang ky thue",
    "dang ky bo sung, thay doi cac chi tieu dang ky thue nhu sau",
    "chi tieu",
    "thong tin dang ky cu",
    "thong tin dang ky moi",
    "- chi tieu 5.1: dia chi nhan thong bao thue",
]


_results = []  # (passed, name)


def check(name: str):
    def deco(fn):
        def run(*args, **kw):
            try:
                fn(*args, **kw)
                _results.append((True, name))
                print("PASS %s" % name)
            except AssertionError as exc:
                _results.append((False, "%s: %s" % (name, exc)))
                print("FAIL %s: %s" % (name, exc))
            except Exception as exc:                        # noqa: BLE001
                _results.append((False, "%s: exception %r" % (name, exc)))
                print("FAIL %s: exception %r" % (name, exc))
        return run
    return deco


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def spec_raw(fid: str) -> dict:
    return json.loads((SPEC_DIR / ("%s.json" % fid)).read_text(
        encoding="utf-8-sig"))


def load_spec_quiet(fid: str):
    bfm._form_spec_cache.pop(fid, None)
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            return bfm.load_form_spec(fid)
    finally:
        bfm._form_spec_cache.pop(fid, None)


def block_par_text(b: dict) -> str:
    return "".join(r.get("text", "") for r in b.get("runs", []))


def body_par_texts(path) -> list[str]:
    d = docx.Document(str(path))
    return [p.text for p in d.paragraphs]


def canonical_doc_hash(docx_path) -> str:
    """sha256(word/document.xml dạng cây canonical) — bỏ zip metadata."""
    with zipfile.ZipFile(docx_path) as zf:
        raw = zf.read("word/document.xml")
    def canon(el):
        return {
            "tag": el.tag.split("}")[-1],
            "attrs": sorted((k.split("}")[-1], v) for k, v in el.attrib.items()),
            "text": el.text or "",
            "children": [canon(c) for c in el],
        }
    blob = json.dumps(canon(ET.fromstring(raw)), ensure_ascii=False,
                      sort_keys=True)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()


def sandbox_build(fid: str):
    """process_form THẬT với spec thật của form, ghi ra tmp. Trả về
    (meta, template_path, tmp_dir, saved_generator_state) — caller BẮT BUỘC
    sandbox_release(saved) + dọn tmp."""
    tmp = Path(tempfile.mkdtemp(prefix="m0306_sandbox_"))
    (tmp / "t").mkdir()
    (tmp / "d").mkdir()
    saved = (bfm.OUT_TPL, bfm.OUT_DATA)
    bfm.OUT_TPL, bfm.OUT_DATA = tmp / "t", tmp / "d"
    bfm._form_spec_cache.clear()
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            meta, _unmatched = bfm.process_form(
                bfm.find_source(fid), bfm.load_csv_titles(), None)
        tpl = tmp / "t" / ("%s.docx" % fid)
        assert tpl.exists(), "generator chưa ghi template vào sandbox"
        return meta, tpl, tmp, saved
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        bfm.OUT_TPL, bfm.OUT_DATA = saved
        bfm._form_spec_cache.clear()
        raise


def sandbox_release(state):
    bfm.OUT_TPL, bfm.OUT_DATA = state


def sandbox_build_raw(fid: str, raw_spec: dict):
    """Như sandbox_build nhưng DÙNG SPEC ĐƯỢC CẤP (ghi vào FORM_SPECS_DIR
    sandbox) — phục vụ [NEG]: spec bịa phải đi qua loader thật."""
    tmp = Path(tempfile.mkdtemp(prefix="m0306_sandbox_"))
    (tmp / "t").mkdir()
    (tmp / "d").mkdir()
    gh = tmp / "household"
    gh.mkdir()
    (gh / ("%s.json" % fid)).write_text(
        json.dumps(raw_spec, ensure_ascii=False, indent=1), encoding="utf-8")
    saved = (bfm.OUT_TPL, bfm.OUT_DATA, bfm.FORM_SPECS_DIR)
    bfm.OUT_TPL, bfm.OUT_DATA, bfm.FORM_SPECS_DIR = tmp / "t", tmp / "d", tmp
    bfm._form_spec_cache.clear()
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            meta, _unmatched = bfm.process_form(
                bfm.find_source(fid), bfm.load_csv_titles(), None)
        tpl = tmp / "t" / ("%s.docx" % fid)
        assert tpl.exists(), "generator chưa ghi template vào sandbox"
        return meta, tpl, tmp, saved
    except Exception:
        bfm.OUT_TPL, bfm.OUT_DATA, bfm.FORM_SPECS_DIR = saved
        bfm._form_spec_cache.clear()
        shutil.rmtree(tmp, ignore_errors=True)
        raise


def sandbox_release_3(state):
    bfm.OUT_TPL, bfm.OUT_DATA, bfm.FORM_SPECS_DIR = state


def full_doc_pl2_ranges() -> dict[int, tuple[int, int]]:
    """Khoảng paragraph của từng Mẫu số PLII trong _Bieu_mau_PL_I_II_full.docx.

    Mỗi "Mẫu số N" xuất hiện ở TOC (một dòng) và ở mẫu thật; mẫu thật được
    nhận diện nhờ tiêu đề fold chứa anchor đặc trưng trong 12 paragraph kế.
    Trả về {n: (start, end)} với end = start mẫu kế (riêng mẫu 6 lấy
    start + 60 vì mẫu kế là Mẫu 7 GCN — vẫn trong 1 trang).
    """
    d = docx.Document(str(FULL_DOC))
    folded = [bfm.fold(re.sub(r"\s+", " ", p.text)).strip() for p in d.paragraphs]
    starts: dict[int, int] = {}
    for i, t in enumerate(folded):
        m = re.fullmatch(r"mau so (\d+)", t)
        if not m:
            continue
        n = int(m.group(1))
        if n not in FULL_DOC_ANCHORS:
            continue
        window = " ".join(folded[i + 1: i + 13])
        if FULL_DOC_ANCHORS[n] in window:
            # giữ lần xuất hiện SAU CÙNG (mẫu thật đứng sau TOC)
            starts[n] = i
    missing = sorted(set(FULL_DOC_ANCHORS) - set(starts))
    assert not missing, "không neo được Mẫu %s trong nguồn đầy đủ" % missing
    order = sorted(starts)
    ranges = {}
    for k, n in enumerate(order):
        end = starts[order[k + 1]] if k + 1 < len(order) else starts[n] + 60
        ranges[n] = (starts[n], end)
    return ranges


def nested_data(keys: list[str]) -> dict:
    data: dict = {}
    for k in keys:
        cur = data
        parts = k.split(".")
        for part in parts[:-1]:
            cur = cur.setdefault(part, {})
        cur[parts[-1]] = "VAL::%s" % k
    return data


def xml_text_nodes(docx_bytes: bytes) -> list[str]:
    texts = []
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        for name in zf.namelist():
            if not re.match(r"^word/(document|header\d*|footer\d*)\.xml$", name):
                continue
            root = ET.fromstring(zf.read(name))
            for el in root.iter():
                if el.text:
                    texts.append(el.text)
                if el.tail:
                    texts.append(el.tail)
    return texts


def fingerprint() -> dict[str, str]:
    out = {}
    for base in (TPL_DIR, ROOT / "app" / "data"):
        for p in sorted(base.rglob("*")):
            if p.is_file():
                out[str(p.relative_to(ROOT))] = hashlib.sha256(
                    p.read_bytes()).hexdigest()
    return out


FP_BEFORE = fingerprint()


# ---------------------------------------------------------------------------
# clauses
# ---------------------------------------------------------------------------

@check("[SPEC] 4 spec household PIN BLOCKED: tables rong, qua validator v1")
def t_spec_blocked():
    for fid in FORMS:
        path = SPEC_DIR / ("%s.json" % fid)
        assert path.exists(), "%s: thiếu file spec %s" % (fid, path)
        raw = spec_raw(fid)
        assert raw.get("form_id") == fid, "%s: form_id lệch tên file" % fid
        assert raw.get("group") == "household", "%s: group phải là household" % fid
        assert raw.get("tables") == [], \
            "%s: tables phải rỗng — nguồn không có bảng thuế, cấm khai entry" % fid
        ext = load_spec_quiet(fid)
        assert ext is not None, "%s: load_form_spec trả None (file không đọc được)" % fid
        assert ext["entries"] == [], "%s: loader sinh entry — spec bị khai rebuild" % fid
        assert ext["specs_by_id"] == {}, "%s: specs_by_id phải rỗng" % fid
        note = raw.get("note") or ""
        gap = raw.get("gap_nghiepvu") or ""
        assert "BLOCKED_OWNER" in gap, \
            "%s: gap_nghiepvu phải pin BLOCKED_OWNER (điều d)" % fid
        assert "BLOCKED" in note.upper(), "%s: note phải ghi trạng thái BLOCKED" % fid
        for marker in ("01_mau_ho_kinh_doanh", "_Bieu_mau_PL_I_II_full.docx",
                       "NO_TABLE"):
            assert marker in note, \
                "%s: note thiếu evidence %r" % (fid, marker)


@check("[SRC] nguon tach that: 0 w:tbl + 0 marker thue (doi chung Mau II.2 CO)")
def t_source_no_tax():
    tax_hits = []
    for fid in FORMS:
        src = bfm.find_source(fid)
        assert src.exists(), "%s: không tìm thấy nguồn %s" % (fid, src)
        d = docx.Document(str(src))
        assert len(d.tables) == 0, \
            "%s: nguồn %s có %d w:tbl — lệch đánh giá BASE (phải dựng lại spec)" % (
                fid, src.name, len(d.tables))
        folded = [bfm.fold(re.sub(r"\s+", " ", p.text)).strip()
                  for p in d.paragraphs]
        for marker in TAX_MARKERS:
            for t in folded:
                if marker in t:
                    tax_hits.append((fid, marker, t[:80]))
    assert not tax_hits, \
        "nguồn tách xuất hiện marker thuế (phải dựng bảng, không được pin chặn): %r" % (
            tax_hits[:4])
    # đối chứng CÓ RĂNG: nguồn đầy đủ — Mẫu II.2 chứa khối thuế thật
    assert FULL_DOC.exists(), "thiếu nguồn đầy đủ %s" % FULL_DOC
    full_folded = [bfm.fold(re.sub(r"\s+", " ", p.text)).strip()
                   for p in docx.Document(str(FULL_DOC)).paragraphs]
    ranges = full_doc_pl2_ranges()
    s2, e2 = ranges[2]
    mau2_zone = " || ".join(full_folded[s2:e2])
    for line in MAU2_TAX_LINES:
        assert line in mau2_zone, \
            "Mẫu II.2 nguồn đầy đủ thiếu dòng thuế %r — đối chứng hỏng" % line
    # còn RĂNG hơn: Mẫu II.3..II.6 KHÔNG marker thuế nào trong vùng mẫu thật
    bad = []
    for n in (3, 4, 5, 6):
        s, e = ranges[n]
        zone = " || ".join(full_folded[s:e])
        for marker in TAX_MARKERS:
            if marker in zone:
                bad.append((n, marker))
    assert not bad, \
        "Mẫu PLII nguồn đầy đủ CÓ marker thuế (Class 6 tái hiện — phải rebuild): %r" % bad


@check("[ORPHAN] blocks hien tai: 0 so mo coi, dung 1 table = header band 2x2")
def t_orphan_zero():
    for fid in FORMS:
        data = json.loads((BLOCKS_DIR / ("%s.json" % fid)).read_text(
            encoding="utf-8"))
        blocks = data["blocks"]
        orphans = [(i, block_par_text(b).strip())
                   for i, b in enumerate(blocks)
                   if b["t"] == "p" and ORPHAN_DIGIT_RE.match(block_par_text(b))]
        assert not orphans, \
            "%s: còn số mồ côi Class 6 %r — phải drop trước khi pin" % (fid, orphans)
        tables = [b for b in blocks if b["t"] == "table"]
        loops = [b for b in blocks if b["t"] == "loop"]
        assert not loops, "%s: không được có block loop (nguồn không có lưới)" % fid
        assert len(tables) == 1, \
            "%s: mong đúng 1 block table (header band), thấy %d" % (
                fid, len(tables))
        rows = tables[0]["rows"]
        assert len(rows) == 2 and all(len(r["cells"]) == 2 for r in rows), \
            "%s: header band phải 2x2, thấy %dx%s" % (
                fid, len(rows), [len(r["cells"]) for r in rows])


def assert_one_header_band(docx_path, fid: str) -> None:
    """Hợp đồng shape (điều a): đúng 1 w:tbl = header band 2x2, không bảng
    dữ liệu nào khác. Dùng lại bởi [SHAPE] và [NEG]."""
    d = docx.Document(str(docx_path))
    assert len(d.tables) == 1, \
        "%s: phải đúng 1 w:tbl, thấy %d" % (fid, len(d.tables))
    t = d.tables[0]
    assert len(t.rows) == 2 and len(t.columns) == 2, \
        "%s: w:tbl phải 2x2, thấy %dx%d" % (fid, len(t.rows), len(t.columns))


@check("[SHAPE] template commit + sandbox regen: 1 w:tbl 2x2, khong drift")
def t_template_shape_no_drift():
    prod_maps = json.loads(FORM_MAPS.read_text(encoding="utf-8"))
    for fid in FORMS:
        assert_one_header_band(TPL_DIR / ("%s.docx" % fid), fid)
        state_meta, tpl, tmp, saved = sandbox_build(fid)
        sandbox_release(saved)
        try:
            assert_one_header_band(tpl, fid)
            assert canonical_doc_hash(tpl) == canonical_doc_hash(
                TPL_DIR / ("%s.docx" % fid)), \
                "%s: sandbox regen khác template commit (drift) — specChanged behavior" % fid
            keys_new = sorted(f["key"] for f in state_meta["fields"])
            keys_old = sorted(f["key"] for f in prod_maps[fid]["fields"])
            assert keys_new == keys_old, \
                "%s: bộ key form_maps sandbox lệch production: +%s -%s" % (
                    fid, sorted(set(keys_new) - set(keys_old))[:4],
                    sorted(set(keys_old) - set(keys_new))[:4])
            assert not state_meta.get("section_specs"), \
                "%s: sandbox sinh section_specs — đã rebuild bảng (cấm)" % fid
            assert not state_meta.get("dynamic_tables"), \
                "%s: sandbox sinh dynamic_tables — đã rebuild bảng (cấm)" % fid
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
            bfm._form_spec_cache.clear()


@check("[NEG] thay spec bang entry gia (fixed-matrix) -> hop dong shape TU CHOI")
def t_negative_fabrication_refused():
    """Chứng minh suite có cắn: nếu ai đó dựng bảng thuế theo suy đoán
    (entry fixed-matrix bịa cho PLII_M03), generator sẽ tuân theo spec và
    sinh thêm w:tbl — hợp đồng shape/[SPEC] của lane này phải TỪ CHỐI."""
    fid = "PLII_M03"
    tampered = spec_raw(fid)
    tampered["tables"] = [{
        "kind": "fixed-matrix",
        "pos": "after",
        "anchor": "^1\\. truong hop tam ngung kinh doanh",
        "no_default_drops": True,
        "spec": {
            "id": "thue_bia_suy_doan",
            "kind": "fixed-matrix",
            "semantic_key": "thue_bia_suy_doan",
            "columns": [
                {"header": "Chỉ tiêu", "tag": None, "input_type": "readonly"},
                {"header": "Giá trị", "tag": "gia_tri", "input_type": "text"},
            ],
            "fixed_rows": [{"id": "hang_bia", "label": "- Hàng bịa"}],
            "totals": [],
        },
    }]
    meta, tpl, tmp, saved = sandbox_build_raw(fid, tampered)
    sandbox_release_3(saved)
    try:
        d = docx.Document(str(tpl))
        assert len(d.tables) == 2, \
            "generator phải trung thực với spec bịa (sinh thêm w:tbl), thấy %d" % (
                len(d.tables))
        try:
            assert_one_header_band(tpl, fid)
        except AssertionError:
            pass
        else:
            raise AssertionError(
                "hợp đồng shape đã KHÔNG cắn — spec bịa lọt qua mà không bị phát hiện")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
        bfm._form_spec_cache.clear()


@check("[RENDER] write_rendered: 0 token sot, dung gia tri sentinel, 0 key thue")
def t_render_roundtrip():
    prod_maps = json.loads(FORM_MAPS.read_text(encoding="utf-8"))
    leftover_re = dxr.LEFTOVER_RE
    for fid in FORMS:
        fields = prod_maps[fid]["fields"]
        keys = [f["key"] for f in fields]
        bad_keys = [k for k in keys if ".thue_chi_tieu" in k
                    or k.startswith("thue_chi_tieu")
                    or ".thong_tin_dang_ky_cu" in k
                    or ".thong_tin_dang_ky_moi" in k]
        assert not bad_keys, \
            "%s: form_maps có key bảng thuế bịa %r" % (fid, bad_keys)
        meta, tpl, tmp, saved = sandbox_build(fid)
        sandbox_release(saved)
        try:
            out = tmp / "out" / ("%s_rendered.docx" % fid)
            data = nested_data(keys)
            dxr.write_rendered(tpl, data, out)
            raw = out.read_bytes()
            texts = xml_text_nodes(raw)
            joined = "\n".join(texts)
            leftovers = sorted({m.group(0) for t in texts
                                for m in leftover_re.finditer(t)})
            assert not leftovers, \
                "%s: render còn token sót %r" % (fid, leftovers[:6])
            missing = [k for k in keys if "VAL::%s" % k not in joined]
            assert not missing, \
                "%s: %d/%d sentinel không hiện trong bản render: %r" % (
                    fid, len(missing), len(keys), missing[:5])
            # bảng duy nhất trong bản render vẫn là header band 2x2
            rd = docx.Document(str(out))
            assert len(rd.tables) == 1 and len(rd.tables[0].rows) == 2 \
                and len(rd.tables[0].columns) == 2, \
                "%s: bản render phải giữ đúng header band 2x2" % fid
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
            bfm._form_spec_cache.clear()


@check("[GUARD] app/templates va app/data khong bi ghi trong suot suite")
def t_guard_repo_untouched():
    after = fingerprint()
    changed = [k for k in after if FP_BEFORE.get(k) != after[k]]
    added = [k for k in after if k not in FP_BEFORE]
    removed = [k for k in FP_BEFORE if k not in after]
    assert not (changed or added or removed), \
        "suite đã ghi ngoài sandbox: changed=%s added=%s removed=%s" % (
            changed[:3], added[:3], removed[:3])


def main() -> int:
    for fn in (t_spec_blocked, t_source_no_tax, t_orphan_zero,
               t_template_shape_no_drift, t_render_roundtrip,
               t_negative_fabrication_refused, t_guard_repo_untouched):
        fn()
    ok = sum(1 for p, _ in _results if p)
    fail = sum(1 for p, _ in _results if not p)
    print("-" * 72)
    print("test_tax_m0306_rebuild: %d PASS / %d FAIL" % (ok, fail))
    for passed, name in _results:
        if not passed:
            print("  -> %s" % name)
    return 1 if fail else 0


if __name__ == "__main__":
    sys.exit(main())
